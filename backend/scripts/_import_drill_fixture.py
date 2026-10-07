"""Generate entirely synthetic schema-6 material; never read research originals."""
from pathlib import Path
import hashlib
import json
import sqlite3


LEGACY_SCHEMA = """
CREATE TABLE projects (id TEXT PRIMARY KEY, name TEXT NOT NULL, research_topic TEXT NOT NULL,
 project_type TEXT NOT NULL CHECK(project_type IN ('sci','thesis')), created_at TEXT NOT NULL, updated_at TEXT NOT NULL);
CREATE TABLE files (id TEXT PRIMARY KEY, project_id TEXT NOT NULL REFERENCES projects(id), filename TEXT NOT NULL,
 file_type TEXT NOT NULL, size_bytes INTEGER NOT NULL, sha256 TEXT NOT NULL, uploaded_at TEXT NOT NULL,
 parse_status TEXT NOT NULL CHECK(parse_status IN ('parsed','failed')), error_json TEXT, preview_json TEXT);
CREATE TABLE analysis_setups (file_id TEXT PRIMARY KEY REFERENCES files(id), revision INTEGER NOT NULL, setup_json TEXT NOT NULL);
CREATE TABLE analysis_runs (id TEXT PRIMARY KEY, file_id TEXT NOT NULL REFERENCES files(id), setup_revision INTEGER NOT NULL,
 engine_version TEXT NOT NULL, completed_at TEXT NOT NULL, result_json TEXT NOT NULL, UNIQUE(file_id,setup_revision,engine_version));
CREATE TABLE figures (id TEXT PRIMARY KEY, analysis_run_id TEXT NOT NULL REFERENCES analysis_runs(id),
 renderer_version TEXT NOT NULL, figure_json TEXT NOT NULL, UNIQUE(analysis_run_id,renderer_version));
CREATE TABLE figure_jobs (id TEXT PRIMARY KEY, analysis_run_id TEXT NOT NULL REFERENCES analysis_runs(id),
 renderer_version TEXT NOT NULL, created_at TEXT NOT NULL, job_json TEXT NOT NULL);
CREATE TABLE explanations (id TEXT PRIMARY KEY, analysis_run_id TEXT NOT NULL REFERENCES analysis_runs(id),
 figure_id TEXT NOT NULL REFERENCES figures(id), engine_version TEXT NOT NULL, explanation_json TEXT NOT NULL,
 UNIQUE(figure_id,engine_version));
CREATE TABLE explanation_jobs (id TEXT PRIMARY KEY, analysis_run_id TEXT NOT NULL REFERENCES analysis_runs(id),
 figure_id TEXT NOT NULL REFERENCES figures(id), engine_version TEXT NOT NULL, created_at TEXT NOT NULL, job_json TEXT NOT NULL);
CREATE TABLE reports (id TEXT PRIMARY KEY, analysis_run_id TEXT NOT NULL REFERENCES analysis_runs(id),
 explanation_id TEXT NOT NULL REFERENCES explanations(id), renderer_version TEXT NOT NULL, report_json TEXT NOT NULL,
 input_json TEXT NOT NULL, UNIQUE(explanation_id,renderer_version));
PRAGMA user_version=6;
"""

BUSINESS_SENTINEL = 'SYNTHETIC_PRIVATE_BODY_不得进入报告'
CLOUD_SENTINEL = 'SYNTHETIC_CLOUD_IDENTIFIER_不得进入报告'
TIMESTAMP = '2026-10-05T01:23:45.123456+00:00'


def file_hashes(directory):
    """Byte-level evidence including SQLite itself and every retained orphan asset."""
    return {path.relative_to(directory).as_posix(): hashlib.sha256(path.read_bytes()).hexdigest()
            for path in sorted(Path(directory).rglob('*')) if path.is_file()}


def create_complex_source(directory):
    """Create a new directory, valid assets, old snapshots and unfinished jobs."""
    from openpyxl import Workbook
    from PIL import Image
    from docx import Document
    from app.domain.explanation_content import VERSION

    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=False)
    for folder in ('files', 'figures', 'reports'):
        (directory / folder).mkdir()
    for identifier in ('file-parsed', 'file-failed', 'file-orphan'):
        workbook = Workbook()
        workbook.active.append(['synthetic_group', 'synthetic_value'])
        workbook.active.append(['A', 1])
        workbook.active.append(['B', 2])
        workbook.save(directory / 'files' / (identifier + '.xlsx'))
        workbook.close()
    for revision in (1, 2):
        Image.new('RGB', (64, 64), (revision * 50, 80, 120)).save(directory / 'figures' / f'figure-{revision}.png')
        document = Document()
        document.add_heading('Synthetic legacy report', 0)
        document.add_paragraph(BUSINESS_SENTINEL)
        document.save(directory / 'reports' / f'report-{revision}.docx')

    def metadata(relative):
        content = (directory / relative).read_bytes()
        return {'size_bytes': len(content), 'sha256': hashlib.sha256(content).hexdigest()}

    def encode(value):
        # Deliberate whitespace/Unicode must survive import verbatim.
        return json.dumps(value, ensure_ascii=False, indent=1)

    source = directory / 'paperassist.sqlite3'
    with sqlite3.connect(source) as db:
        db.execute('PRAGMA foreign_keys=ON')
        db.executescript(LEGACY_SCHEMA)
        db.execute('INSERT INTO projects VALUES (?,?,?,?,?,?)',
                   ('project-drill', BUSINESS_SENTINEL, 'synthetic topic', 'sci', TIMESTAMP, TIMESTAMP))
        for identifier, status in [('file-parsed', 'parsed'), ('file-failed', 'failed')]:
            asset = metadata(f'files/{identifier}.xlsx')
            error = encode({'code': 'invalid_workbook', 'message': BUSINESS_SENTINEL}) if status == 'failed' else None
            preview = encode({'synthetic': True, 'sentinel': BUSINESS_SENTINEL}) if status == 'parsed' else None
            db.execute('INSERT INTO files VALUES (?,?,?,?,?,?,?,?,?,?)',
                       (identifier, 'project-drill', identifier + '.xlsx', 'xlsx', asset['size_bytes'],
                        asset['sha256'], TIMESTAMP, status, error, preview))
        db.execute('INSERT INTO analysis_setups VALUES (?,?,?)',
                   ('file-parsed', 3, encode({'revision': 3, 'settings': BUSINESS_SENTINEL})))
        source_sha = metadata('files/file-parsed.xlsx')['sha256']
        for revision in (1, 2):
            run, figure, explanation, report = (f'{kind}-{revision}' for kind in ('run', 'figure', 'explanation', 'report'))
            result = {'id': run, 'file_id': 'file-parsed', 'setup_revision': revision, 'source_sha256': source_sha,
                      'engine': {'id': f'legacy-stat-engine-{revision}'}, 'summary': BUSINESS_SENTINEL}
            db.execute('INSERT INTO analysis_runs VALUES (?,?,?,?,?,?)',
                       (run, 'file-parsed', revision, result['engine']['id'], TIMESTAMP, encode(result)))
            figure_data = {'id': figure, 'analysis_run_id': run, 'file_id': 'file-parsed', 'setup_revision': revision,
                           'source_sha256': source_sha, **metadata(f'figures/{figure}.png')}
            db.execute('INSERT INTO figures VALUES (?,?,?,?)', (figure, run, f'legacy-renderer-{revision}', encode(figure_data)))
            explanation_data = {'id': explanation, 'analysis_run_id': run, 'figure_id': figure,
                                'file_id': 'file-parsed', 'setup_revision': revision,
                                'source_sha256': source_sha, 'figure_sha256': figure_data['sha256'],
                                'sections': [{'text': BUSINESS_SENTINEL}]}
            db.execute('INSERT INTO explanations VALUES (?,?,?,?,?)',
                       (explanation, run, figure, f'legacy-explanation-{revision}', encode(explanation_data)))
            report_data = {'id': report, 'analysis_run_id': run, 'explanation_id': explanation,
                           'source_sha256': source_sha, 'figure_sha256': figure_data['sha256'],
                           **metadata(f'reports/{report}.docx')}
            db.execute('INSERT INTO reports VALUES (?,?,?,?,?,?)',
                       (report, run, explanation, f'legacy-report-{revision}', encode(report_data), encode(explanation_data)))
        for table, rowids in [('figure_jobs', (3, 11, 29)), ('explanation_jobs', (5, 19, 41))]:
            for seq, status in zip(rowids, ('completed', 'running', 'submitting')):
                identifier = ('fj' if table == 'figure_jobs' else 'ej') + f'-{seq}'
                job = {'id': identifier, 'analysis_run_id': 'run-2', 'figure_id': 'figure-2', 'status': status,
                       'created_at': TIMESTAMP, 'response_id': CLOUD_SENTINEL + identifier,
                       'container_id': CLOUD_SENTINEL + '-container', 'payload': {'text': BUSINESS_SENTINEL}}
                if table == 'figure_jobs':
                    db.execute('INSERT INTO figure_jobs(rowid,id,analysis_run_id,renderer_version,created_at,job_json) VALUES (?,?,?,?,?,?)',
                               (seq, identifier, 'run-2', 'synthetic-pending-renderer', TIMESTAMP, encode(job)))
                else:
                    db.execute('INSERT INTO explanation_jobs(rowid,id,analysis_run_id,figure_id,engine_version,created_at,job_json) VALUES (?,?,?,?,?,?,?)',
                               (seq, identifier, 'run-2', 'figure-2', VERSION if status != 'completed' else 'legacy-explanation-2',
                                TIMESTAMP, encode(job)))
    return source, directory
