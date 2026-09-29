"""Prepare a fixed analysis task using full-sheet checks, without running statistics."""

from collections import Counter
from datetime import date, datetime, time, timedelta
from math import isfinite
from typing import Annotated, Literal

from fastapi import APIRouter, Depends, Query
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from pydantic import BaseModel, ConfigDict, Field

from .config import ExcelSettings, get_excel_settings
from .excel import fail, parse_workbook
from .storage import ProjectStore, get_project_store

router = APIRouter(prefix='/api/v1/projects/{project_id}/files/{file_id}', tags=['分析配置'])
Store = Annotated[ProjectStore, Depends(get_project_store)]
Settings = Annotated[ExcelSettings, Depends(get_excel_settings)]
ColumnId = Annotated[str, Field(pattern=r'^[A-Z]{1,3}$')]
MAX_GROUPS = 20


class AnalysisSelection(BaseModel):
    model_config = ConfigDict(extra='forbid', strict=True)
    task_type: Literal['descriptive_boxplot']
    sheet_name: str = Field(min_length=1, max_length=31)
    numeric_column: ColumnId
    group_column: ColumnId | None = None
    unit: str = Field(default='', max_length=80)
    missing_policy: Literal['exclude_selected_missing']


class SaveSelection(AnalysisSelection):
    expected_revision: int = Field(ge=0)


class ColumnProfile(BaseModel):
    id: str
    name: str
    type: Literal['number', 'text', 'boolean', 'datetime', 'formula', 'error', 'mixed', 'empty']
    non_missing_count: int
    missing_count: int
    numeric_count: int
    can_be_numeric: bool
    can_be_group: bool
    issue_cells: list[str]


class SheetProfile(BaseModel):
    sheet_name: str
    row_count: int
    columns: list[ColumnProfile]
    warnings: list[str]


class GroupCount(BaseModel):
    label: str
    count: int


class SelectionCheck(BaseModel):
    ready: bool
    row_count: int
    valid_count: int
    excluded_count: int
    groups: list[GroupCount]
    issues: list[str]
    warnings: list[str]


class SavedSetup(BaseModel):
    file_id: str
    source_sha256: str
    status: Literal['configured']
    revision: int
    updated_at: str
    selection: AnalysisSelection
    check: SelectionCheck


def cell_kind(cell):
    if cell is None:
        return 'missing', None
    value = cell.value
    if value is None or isinstance(value, str) and not value.strip():
        return 'missing', None
    if cell.data_type == 'f':
        return 'formula', None
    if cell.data_type == 'e':
        return 'error', None
    if isinstance(value, bool):
        return 'boolean', value
    if isinstance(value, (datetime, date, time, timedelta)):
        return 'datetime', None
    if isinstance(value, (int, float)):
        return ('number', value) if isfinite(value) else ('error', None)
    return 'text', str(value)


class SheetScan:
    def __init__(self, sheet_name: str, selection: AnalysisSelection | None, collect_values: bool = False):
        self.sheet_name = sheet_name
        self.selection = selection
        self.types: list[Counter] = []
        self.issue_cells: list[list[str]] = []
        self.valid_count = 0
        self.groups = Counter()
        self.values = [] if collect_values else None
        self.grouped_values = {}

    def observe(self, sheet_name, row_index, cells):
        if sheet_name != self.sheet_name or row_index == 1:
            return
        while len(self.types) < len(cells):
            self.types.append(Counter())
            self.issue_cells.append([])
        for index, cell in enumerate(cells):
            kind, _ = cell_kind(cell)
            if kind != 'missing':
                self.types[index][kind] += 1
                if kind != 'number' and len(self.issue_cells[index]) < 5:
                    self.issue_cells[index].append(f'{get_column_letter(index + 1)}{row_index}')
        if self.selection is None:
            return

        def selected(column):
            index = column_index_from_string(column) - 1
            return cell_kind(cells[index] if index < len(cells) else None)

        kind, value = selected(self.selection.numeric_column)
        if kind != 'number':
            return
        if self.values is not None and isinstance(value, int) and abs(value) > 2**53 - 1:
            fail('numeric_precision', f'数值列 {self.selection.numeric_column}{row_index} 的整数超过安全精度范围（±9007199254740991）。请调整单位或核对数据后重新上传，避免计算时丢失精度。')
        if self.selection.group_column is not None:
            kind, group = selected(self.selection.group_column)
            if kind not in ('text', 'number', 'boolean'):
                return
            # Bound category memory/output even for an accidentally selected identifier.
            if group in self.groups or len(self.groups) <= MAX_GROUPS:
                self.groups[group] += 1
                if self.values is not None:
                    self.grouped_values.setdefault(group, []).append(float(value))
        self.valid_count += 1
        if self.values is not None:
            self.values.append(float(value))

    def profile(self, sheet):
        columns = []
        for index, name in enumerate(sheet.columns):
            counts = self.types[index] if index < len(self.types) else Counter()
            kind = next(iter(counts)) if len(counts) == 1 else 'mixed' if counts else 'empty'
            columns.append(ColumnProfile(
                id=get_column_letter(index + 1), name=name, type=kind,
                non_missing_count=sum(counts.values()), missing_count=sheet.row_count - sum(counts.values()),
                numeric_count=counts['number'], can_be_numeric=kind == 'number',
                can_be_group=kind in ('number', 'text', 'boolean'),
                issue_cells=self.issue_cells[index] if index < len(self.issue_cells) else [],
            ))
        return SheetProfile(sheet_name=sheet.name, row_count=sheet.row_count, columns=columns, warnings=sheet.warnings)

    def check(self, profile: SheetProfile) -> SelectionCheck:
        selection = self.selection
        columns = {column.id: column for column in profile.columns}
        if selection.numeric_column not in columns:
            fail('invalid_column', '所选数值列不存在，请重新选择工作表和字段。')
        if selection.group_column is not None and selection.group_column not in columns:
            fail('invalid_column', '所选分组列不存在，请重新选择。')
        if selection.numeric_column == selection.group_column:
            fail('invalid_column', '数值列和分组列不能是同一列。')
        numeric = columns[selection.numeric_column]
        issues, warnings = [], []
        if not numeric.can_be_numeric:
            locations = '，例如 ' + '、'.join(numeric.issue_cells) if numeric.issue_cells else ''
            issues.append('数值列必须包含实际数值，且非缺失值全部为数值；文本数字、布尔、日期、公式或错误不能直接分析' + locations + '。')
        if selection.group_column is not None and not columns[selection.group_column].can_be_group:
            issues.append('分组列应为单一的文本、数值或布尔类型，不能包含公式、日期或混合类型。')
        if self.valid_count == 0:
            issues.append('选中字段没有可用于分析的完整记录，请检查缺失值和字段选择。')
        if len(self.groups) > MAX_GROUPS:
            issues.append(f'有效分组超过 {MAX_GROUPS} 个，请选择分类字段，避免使用样本编号作为分组。')
        groups = [GroupCount(label=('TRUE' if key else 'FALSE') if isinstance(key, bool) else str(key), count=count)
                  for key, count in list(self.groups.items())[:MAX_GROUPS]]
        if self.valid_count and (any(group.count < 2 for group in groups) or not groups and self.valid_count < 2):
            warnings.append('部分分组或全部数据不足 2 条有效记录，后续部分统计量将无法计算。')
        if profile.row_count > self.valid_count:
            warnings.append('采用选中字段的完整记录；数值或分组缺失的行将不参与后续计算，原文件保持不变。')
        return SelectionCheck(ready=not issues, row_count=profile.row_count, valid_count=self.valid_count,
            excluded_count=profile.row_count - self.valid_count, groups=groups, issues=issues, warnings=warnings)


def source(store: ProjectStore, project_id: str, file_id: str):
    record = store.file(project_id, file_id)
    content = store.original(record)
    if record['parse_status'] != 'parsed':
        fail('file_not_parsed', '该文件解析失败，请重新上传有效的 Excel 后配置分析。')
    return record, content


def inspect_selection(store, project_id, file_id, settings, sheet_name, selection=None, *, collect_values=False):
    record, content = source(store, project_id, file_id)
    scan = SheetScan(sheet_name, selection, collect_values)
    preview = parse_workbook(content, record['filename'], settings, observe_row=scan.observe)
    sheet = next((sheet for sheet in preview.sheets if sheet.name == sheet_name), None)
    if sheet is None:
        fail('sheet_not_found', '工作表不存在，请重新选择。')
    return record, scan.profile(sheet), scan


@router.get('/analysis-profile', response_model=SheetProfile)
def get_profile(project_id: str, file_id: str, store: Store, settings: Settings,
                sheet: Annotated[str, Query(min_length=1, max_length=31)]):
    return inspect_selection(store, project_id, file_id, settings, sheet)[1]


@router.post('/analysis-check', response_model=SelectionCheck)
def check_selection(project_id: str, file_id: str, selection: AnalysisSelection, store: Store, settings: Settings):
    _, profile, scan = inspect_selection(store, project_id, file_id, settings, selection.sheet_name, selection)
    return scan.check(profile)


@router.get('/analysis-setup', response_model=SavedSetup | None)
def get_setup(project_id: str, file_id: str, store: Store):
    source(store, project_id, file_id)
    return store.analysis_setup(project_id, file_id)


@router.put('/analysis-setup', response_model=SavedSetup)
def save_setup(project_id: str, file_id: str, request: SaveSelection, store: Store, settings: Settings):
    selection = AnalysisSelection(**request.model_dump(exclude={'expected_revision'}))
    record, profile, scan = inspect_selection(store, project_id, file_id, settings, selection.sheet_name, selection)
    check = scan.check(profile)
    if not check.ready:
        fail('invalid_analysis_selection', ' '.join(check.issues))
    return store.save_analysis_setup(project_id, file_id, request.expected_revision, {
        'file_id': file_id, 'source_sha256': record['sha256'], 'status': 'configured',
        'selection': selection.model_dump(), 'check': check.model_dump(),
    })
