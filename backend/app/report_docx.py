"""Editable report from saved evidence only. Never calls an AI service."""

from io import BytesIO
import re

import docx
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .explanation_content import METHOD, METRICS

VERSION = 'word_analysis_report_v1'
# Spreadsheet/project labels can contain controls that XML 1.0 cannot encode.
ILLEGAL_XML = re.compile('[\x00-\x08\x0b\x0c\x0e-\x1f\ud800-\udfff\ufffe\uffff]')


def clean(value):
    return ILLEGAL_XML.sub('\ufffd', str(value))


def display(value):
    return '无法计算' if value is None else format(value, '.6g') if isinstance(value, float) else str(value)


def paragraph(document, text, style=None):
    return document.add_paragraph(clean(text), style)


def table(document, headers, rows, widths):
    item = document.add_table(rows=1, cols=len(headers))
    item.autofit = False
    borders = OxmlElement('w:tblBorders')
    for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        border = OxmlElement('w:' + side)
        for key, value in {'val': 'single', 'sz': '4', 'color': 'D9D9D9'}.items():
            border.set(qn('w:' + key), value)
        borders.append(border)
    item._tbl.tblPr.append(borders)
    for col, width in zip(item.columns, widths):
        col.width = Cm(width)
    item.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
    for index, values in enumerate([headers, *rows]):
        row = item.rows[0] if index == 0 else item.add_row()
        # A row may grow and, if unusually long, split across pages rather than clip.
        for column, (cell, value, width) in enumerate(zip(row.cells, values, widths)):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            margin = OxmlElement('w:tcMar')
            for side in ('top', 'bottom', 'left', 'right'):
                edge = OxmlElement('w:' + side)
                edge.set(qn('w:w'), '90')
                edge.set(qn('w:type'), 'dxa')
                margin.append(edge)
            props.append(margin)
            if index == 0:
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'EDEDED'); props.append(shading)
            cell.text = clean(value)
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if column == 0 else WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.bold = index == 0
    paragraph(document, '').paragraph_format.space_after = Pt(0)


def build_report(snapshot, report, png):
    project, result, figure, explanation = (snapshot[key] for key in ('project', 'result', 'figure', 'explanation'))
    document = Document()
    page = document.sections[0]
    page.page_width, page.page_height = Cm(21.59), Cm(27.94)
    page.top_margin = page.bottom_margin = page.left_margin = page.right_margin = Cm(2)
    page.header_distance = page.footer_distance = Cm(0.9)
    for name, size in [('Normal', 11), ('Title', 22), ('Subtitle', 12), ('Heading 1', 15), ('Heading 2', 12), ('Caption', 10)]:
        style = document.styles[name]
        style.font.name = 'Arial'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), '宋体')
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.2
    document.core_properties.title = '描述统计分析报告'
    document.core_properties.author = 'PaperAssist System'
    document.core_properties.subject = clean(project['name'])[:255]
    header = page.header.paragraphs[0]
    header.text = 'PaperAssist · 描述统计分析报告'
    header.runs[0].font.size = Pt(9)
    footer = page.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('第 ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    footer.add_run(' 页')

    paragraph(document, '描述统计分析报告', 'Title')
    paragraph(document, project['name'], 'Subtitle')
    paragraph(document, f"配置版本 {result['setup_revision']} · 导出时间 {report['created_at']}（UTC）")
    paragraph(document, '本报告整理已保存的数据检查、真实描述统计、图表与 AI 初稿。解释须由研究者审核后使用；本报告未执行推断检验。')

    document.add_heading('1 数据概况', level=1)
    for label, value in [('项目类型', 'SCI 论文' if project['project_type'] == 'sci' else '毕业论文'),
                         ('研究主题（用户填写）', project['research_topic'] or '未填写'),
                         ('原始文件', result['filename']), ('工作表', result['selection']['sheet_name']),
                         ('数值字段', result['numeric_name'] + '（列 ' + result['selection']['numeric_column'] + '）'),
                         ('分组字段', (result['group_name'] + '（列 ' + result['selection']['group_column'] + '）') if result['group_name'] else '不分组'),
                         ('数值单位', result['selection']['unit'] or '未填写')]:
        paragraph(document, f'{label}：{value}')

    document.add_heading('2 数据质量与处理', level=1)
    check = result['check']
    paragraph(document, f"原始数据 {check['row_count']} 行；有效完整记录 {check['valid_count']} 行；排除 {check['excluded_count']} 行。")
    paragraph(document, '排除规则：仅分析数值字段及所选分组字段均有效的完整记录。未插补缺失值，未自动去重，未删除离群观测；原始文件保持不变。')
    for note in check['issues'] + check['warnings']:
        paragraph(document, note)
    paragraph(document, '质量检查限于已选字段及当前支持的类型、缺失和分组规则，不代表已确认实验设计、样本独立性或数据真实性。')

    document.add_heading('3 分析方法', level=1)
    paragraph(document, METHOD)
    paragraph(document, '统计表及解释最多显示 6 位有效数字；“无法计算”表示样本量不足或数值范围限制，不等同于零。完整数值保存在统计结果记录中。')

    document.add_heading('4 统计结果', level=1)
    groups = [('总体', result['overall'])] + [(item['label'], item['statistics']) for item in result['groups']]
    for title, keys, widths in [('表 1 样本量、均值、离散程度及范围', ['n', 'mean', 'std', 'min', 'max'], [4, 1.2, 2.95, 2.95, 2.95, 2.95]),
                                ('表 2 分位数及四分位距', ['q1', 'median', 'q3', 'iqr'], [4, 3.25, 3.25, 3.25, 3.25])]:
        p = paragraph(document, title, 'Caption'); p.paragraph_format.keep_with_next = True
        table(document, ['范围 / 分组'] + [METRICS[key] for key in keys],
              [[label] + [display(stats[key]) for key in keys] for label, stats in groups], widths)
    for label, stats in groups:
        for note in stats['warnings']:
            paragraph(document, f'{label}：{note}')

    document.add_heading('5 图表与图注', level=1)
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    shape = p.add_run().add_picture(BytesIO(png), width=Cm(17))
    # Cap tall images without changing their aspect ratio.
    if shape.height > Cm(12):
        shape.width = int(shape.width * Cm(12) / shape.height); shape.height = Cm(12)
    shape._inline.docPr.set('descr', clean(figure['title']))
    p = paragraph(document, figure['title'], 'Caption')
    p.paragraph_format.keep_with_next = True
    paragraph(document, figure['caption'])
    paragraph(document, '图片为已保存的 OpenAI Code Interpreter 输出，原图字节保持不变。图表元数据已核对，像素中的文字与数值仍需人工复核。')

    document.add_heading('6 分析解释与论文表述', level=1)
    paragraph(document, '以下六部分原样引用已保存的 AI 初稿，事实引用由系统核对；请结合真实研究设计审阅。')
    for section in explanation['sections']:
        document.add_heading(section['title'], level=2)
        paragraph(document, section['text'])
    document.add_heading('7 适用限制与待补充信息', level=1)
    for limitation in explanation['limitations']:
        paragraph(document, limitation)
    paragraph(document, explanation['verification']['note'])

    document.add_heading('附录 A 解释依据', level=1)
    paragraph(document, '以下按解释段落列出实际引用的事实；与已保存的配置、统计结果及图表元数据对应。')
    for section in explanation['sections']:
        document.add_heading(section['title'], level=2)
        for evidence in section['evidence']:
            paragraph(document, f"{evidence['label']}：{evidence['value']}\n引用键：{evidence['key']}")

    document.add_heading('附录 B 来源与版本', level=1)
    records = [('项目编号', project['id']), ('文件编号', result['file_id']), ('原文件 SHA256', result['source_sha256']),
               ('统计结果编号', result['id']), ('统计完成时间', result['completed_at']),
               ('统计引擎', result['engine']['id']), ('Python', result['engine']['python_version']),
               ('openpyxl', result['engine']['openpyxl_version']), ('图表编号', figure['id']), ('图表保存时间', figure['created_at']),
               ('图片 SHA256', figure['sha256']), ('绘图引擎', figure['engine']['id']),
               ('绘图模型', figure['provenance'].get('model', '未记录')), ('绘图响应编号', figure['provenance'].get('response_id', '未记录')),
               ('解释编号', explanation['id']), ('解释保存时间', explanation['created_at']), ('解释模型', explanation['engine']['model']),
               ('解释引擎 / 提示版本', explanation['engine']['id']), ('解释响应编号', explanation['provenance'].get('response_id', '未记录')),
               ('解释输入 SHA256', explanation['provenance']['input_sha256']), ('报告编号', report['id']),
               ('报告模板', VERSION), ('python-docx', docx.__version__), ('报告输入 SHA256', report['input_sha256'])]
    for label, value in records:
        paragraph(document, f'{label}：{value}')
    paragraph(document, '时间均为带 UTC 时区偏移的保存记录。XML 无法表示的控制字符显示为替代符号“�”，不改变原始数据文件。')
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()
