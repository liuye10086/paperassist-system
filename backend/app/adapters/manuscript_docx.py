"""Local fixed-version fact-draft DOCX; no source fetching or model processing."""
from datetime import datetime, timezone
import hashlib
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.image.image import Image
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor

from app.core.exceptions import StorageError


RENDERER_VERSION = 'manuscript_docx_v5'

SECTION_LABELS = (
    ('title', 'Title', '标题'),
    ('abstract', 'Abstract', '摘要'),
    ('introduction', 'Introduction', '引言'),
    ('results', 'Results', '结果'),
    ('discussion', 'Discussion', '讨论'),
    ('materials_and_methods', 'Materials and methods', '材料与方法'),
    ('references', 'References', '参考文献'),
)


def _font(style, size):
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'), 'SimSun')
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + key), None)


def _document():
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(2.2)
    section.left_margin = section.right_margin = Cm(2.3)
    section.header_distance = section.footer_distance = Cm(1)
    for name, size in (('Normal', 11), ('Title', 18), ('Heading 1', 14),
                       ('Heading 2', 12), ('Heading 3', 11), ('Caption', 10)):
        style = document.styles[name]
        _font(style, size)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
    for name in ('Heading 1', 'Heading 2', 'Heading 3'):
        style = document.styles[name]
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.keep_with_next = True
    document.styles['Title'].font.bold = True
    title_properties = document.styles['Title'].element.get_or_add_pPr()
    for border in list(title_properties.findall(qn('w:pBdr'))):
        title_properties.remove(border)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    page = OxmlElement('w:fldSimple')
    page.set(qn('w:instr'), 'PAGE')
    footer._p.append(page)
    # Stable properties and ZIP timestamps support retrying the same input.
    properties = document.core_properties
    properties.author = 'PaperAssist'
    properties.last_modified_by = 'PaperAssist'
    properties.created = properties.modified = datetime(2000, 1, 1, tzinfo=timezone.utc)
    return document


def _paragraph(document, item, bookmark_number):
    paragraph = document.add_paragraph(item['text'])
    identity = str(item['id'])
    start, end = OxmlElement('w:bookmarkStart'), OxmlElement('w:bookmarkEnd')
    start.set(qn('w:id'), str(bookmark_number))
    start.set(qn('w:name'), 'pa_' + hashlib.sha256(identity.encode()).hexdigest()[:32])
    end.set(qn('w:id'), str(bookmark_number))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def _table(document, item):
    document.add_paragraph(item['title'], style='Caption').paragraph_format.keep_with_next = True
    columns = item['columns']
    if not columns or any(len(row) != len(columns) for row in item['rows']):
        raise StorageError('report_render_failed', '稿件表格结构无效。', 422)
    table = document.add_table(rows=1, cols=len(columns))
    table.style = 'Table Grid'
    table.autofit = False
    for column in table.columns:
        column.width = Cm(16.4 / len(columns))
    header = table.rows[0]
    repeat = OxmlElement('w:tblHeader')
    header._tr.get_or_add_trPr().append(repeat)
    for cell, value in zip(header.cells, columns):
        cell.text = str(value)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for values in item['rows']:
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for cell, value in zip(row.cells, values):
            cell.text = str(value)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def _figure(document, figure, images, en):
    raw = images.get(figure['id'])
    if raw is None:
        raise StorageError('figure_missing', '稿件引用的图片不可用。', 409)
    if hashlib.sha256(raw).hexdigest() != figure['sha256']:
        raise StorageError('figure_changed', '图片与稿件冻结版本不一致。', 409)
    image = Image.from_blob(raw)
    scale = min(6.3 / image.px_width, 6.0 / image.px_height)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.add_run().add_picture(BytesIO(raw), width=Inches(image.px_width * scale),
                                   height=Inches(image.px_height * scale))
    caption = document.add_paragraph(
        ('Figure ' if en else '图 ') + str(figure['number']) + '  ' + figure['title'], style='Caption')
    caption.paragraph_format.keep_with_next = bool(figure['caption'])
    if figure['caption']:
        document.add_paragraph(figure['caption'], style='Caption')
    for panel in figure.get('panels', []):
        document.add_paragraph(panel['label'] + '  ' + panel['description'], style='Caption')


def _full_sections(snapshot):
    sections = snapshot['sections']
    valid = isinstance(sections, list) and len(sections) == len(SECTION_LABELS)
    if valid:
        valid = all(isinstance(section, dict) and section.get('id') == labels[0]
                    and isinstance(section.get('paragraphs'), list)
                    for section, labels in zip(sections, SECTION_LABELS))
    if valid:
        paragraphs = [item for section in sections for item in section['paragraphs']]
        valid = all(isinstance(item, dict) and isinstance(item.get('id'), str) and item['id']
                    and isinstance(item.get('text'), str) for item in paragraphs)
        valid = valid and len({item['id'] for item in paragraphs}) == len(paragraphs)
        title = sections[0]['paragraphs']
        valid = valid and len(title) == 1 and title[0]['text'] == snapshot['title']
    if not valid:
        raise StorageError('report_render_failed', '完整稿件章节或标题结构无效。', 422)
    return sections


def _references(document, snapshot, en, *, required=False):
    for reference in snapshot['references']:
        document.add_paragraph('(' + str(reference['number']) + ') ' + reference['text'])
    if required and not snapshot['references']:
        document.add_paragraph('No verified references have been selected; further material is required.'
                               if en else '尚未选择已核实参考文献，待补充。')


def _full_body(document, snapshot, sections, images, en):
    bookmark_number = 0
    for section, labels in zip(sections[1:], SECTION_LABELS[1:]):
        document.add_heading(labels[1 if en else 2], 1)
        if section['id'] == 'references':
            _references(document, snapshot, en, required=True)
            continue
        for item in section['paragraphs']:
            bookmark_number += 1
            _paragraph(document, item, bookmark_number)
        if section['id'] == 'abstract' and snapshot.get('keywords'):
            document.add_paragraph(('Keywords: ' if en else '关键词：')
                                   + ('; ' if en else '；').join(snapshot['keywords']))
        if section['id'] == 'results':
            # The sections are the writing text; original result paragraphs
            # remain frozen evidence and must not replace a rewritten section.
            for group in snapshot['results_groups']:
                if group['tables'] or group['figures']:
                    document.add_heading(group['title'], 2)
                for table in group['tables']:
                    _table(document, table)
                for figure in group['figures']:
                    _figure(document, figure, images, en)
    if snapshot.get('acknowledgements'):
        document.add_heading('Acknowledgements' if en else '致谢', 1)
        document.add_paragraph(snapshot['acknowledgements'])


def build_manuscript_docx(snapshot, images):
    """Render only the frozen snapshot, retaining unresolved facts and original images."""
    if snapshot.get('schema_version') != 1 or snapshot.get('output_language') not in ('zh-CN', 'en'):
        raise StorageError('report_render_failed', '稿件版本格式不受支持。', 422)
    if snapshot.get('review'):
        report=snapshot['review']['report']
        if not report.get('can_publish') or any(i['severity']=='error' for i in report['issues']):
            raise StorageError('manuscript_review_blocked', '评审候选存在不可豁免的一致性问题，不能导出。', 409)
    if snapshot.get('illustration_check_version') is not None:
        from app.domain.manuscripts.illustrations import check_bindings
        if snapshot['illustration_check_version'] != 1 or check_bindings(snapshot['results_groups'], snapshot['output_language']):
            raise StorageError('report_render_failed', '正文图表引用与冻结编号或分面不一致，请核对后另存版本。', 422)
    sections = _full_sections(snapshot) if 'sections' in snapshot else None
    if sections:
        from app.domain.manuscripts.illustrations import check_references, rendered_groups
        existing={(i['id'],tuple(i.get('values',[]))) for i in snapshot.get('issues',[])}
        unresolved=check_references(rendered_groups(snapshot),snapshot['output_language'])
        if any((i['id'],tuple(i['values'])) not in existing for i in unresolved):
            raise StorageError('report_render_failed', '完整稿正文出现未记录的图表或分面对应问题，请核对后另存版本。', 422)
    document = _document()
    en = snapshot['output_language'] == 'en'
    document.add_paragraph(snapshot['title'], style='Title')
    document.add_paragraph(snapshot['notice'])
    document.sections[0].header.paragraphs[0].text = (
        ('SCI manuscript draft' if en else 'SCI论文初稿') if sections else
        ('Research fact draft' if en else '研究事实稿'))
    bookmark_number = 0
    if sections:
        _full_body(document, snapshot, sections, images, en)
    else:
        document.add_heading('Results' if en else '结果', 1)
        for group in snapshot['results_groups']:
            document.add_heading(group['title'], 2)
            for item in group['paragraphs']:
                bookmark_number += 1
                _paragraph(document, item, bookmark_number)
            for table in group['tables']:
                _table(document, table)
            for figure in group['figures']:
                _figure(document, figure, images, en)
        document.add_heading('Materials and methods' if en else '材料与方法', 1)
        for item in snapshot['methods']:
            bookmark_number += 1
            _paragraph(document, item, bookmark_number)
        if snapshot['references']:
            document.add_heading('References' if en else '参考文献', 1)
            _references(document, snapshot, en)
    if snapshot['issues']:
        document.add_heading('Items requiring review or additional material' if en else '待补充与待核实事项', 1)
        for issue in snapshot['issues']:
            document.add_paragraph(issue['message'])
            if issue.get('values'):
                document.add_paragraph(('Conflicting records: ' if en else '相互矛盾的记录：')
                                       + ' / '.join(issue['values']))
    if snapshot.get('review'):
        review=snapshot['review']
        document.add_heading('Material confirmations and consistency review' if en else '资料确认与一致性复核',1)
        document.add_paragraph('User confirmation is not scientific verification. Original source bindings remain frozen; semantic correctness requires researcher review.'
                               if en else '用户确认不等于科研验证。原始来源绑定保持冻结；正文含义仍需研究人员核实。')
        for record in review['decisions']:
            document.add_paragraph(record['target_id']+' — '+record['status']+': '+record['value'])
            document.add_paragraph(('Reason / evidence: ' if en else '理由／证据：')+record['reason']+' / '+record['evidence'])
        for item in review['metadata']:
            text=item['target_id']+' — '+item['field']+': '+item['value']+'; '+item['evidence']
            if item.get('stale'):
                text+=' — '+('Stale confirmation: applies only to the original frozen fact; new material requires renewed confirmation.'
                             if en else '确认已过期：仅适用于原冻结事实；新资料须重新确认。')
            document.add_paragraph(text)
        for item in review['addendum']:
            document.add_paragraph(item['text'])
        for association in review['relationships']['associations']:
            document.add_paragraph(association['research_question']+' — '+', '.join(association['method_ids']+association['paragraph_ids']+association['figure_ids'])+'; '+association['evidence'])
        for impact in review['report']['impacts']:
            document.add_paragraph(('Impact pending review: ' if en else '影响待核对：')+impact['source_id']+' — '+', '.join(impact['locations']))
        # Preserve every distinct frozen caveat, including material and
        # uncertainty details that the general notice cannot replace. Exact
        # original issues (including their values) were already shown above.
        seen={(item['id'],item['message']) for item in snapshot['issues']}
        for item in review['report']['issues']:
            key=(item['location'],item['message'])
            if key in seen: continue
            seen.add(key)
            document.add_paragraph(item['location']+' — '+item['message'])
    if snapshot['sources']:
        document.add_heading('Selected source materials' if en else '所选来源资料', 1)
        description = document.add_paragraph(
            'The saved manuscript workspace retains the exact source versions. User declarations, image observations and literature findings are separate from computed research results.'
            if en else '工作台中保留了此稿件的确切来源版本。用户声明、图像观察及文献发现与本研究计算结果分别记录。')
        seen = set()
        labels = []
        for source in snapshot['sources']:
            label = source['label']
            if label not in seen:
                labels.append(label)
                seen.add(label)
        materials = document.add_paragraph(('; ' if en else '；').join(labels))
        for paragraph in (description, materials):
            paragraph.paragraph_format.line_spacing = 1
            paragraph.paragraph_format.space_after = Pt(3)
            paragraph.paragraph_format.keep_together = True
            for run in paragraph.runs:
                run.font.size = Pt(10)
        description.paragraph_format.keep_with_next = True
    stream = BytesIO()
    document.save(stream)
    stable = BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as original, ZipFile(stable, 'w', ZIP_DEFLATED) as result:
        for item in original.infolist():
            entry = ZipInfo(item.filename, (1980, 1, 1, 0, 0, 0))
            entry.compress_type = ZIP_DEFLATED
            result.writestr(entry, original.read(item.filename))
    return stable.getvalue()
