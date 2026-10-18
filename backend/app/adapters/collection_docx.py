"""Local bilingual collection Word renderer. Original PNG bytes stay unchanged."""
from io import BytesIO
import hashlib
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED
from docx import Document
from docx.image.image import Image
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from app.core.exceptions import StorageError
from app.adapters.collection_labels import rows


def build_collection_report(snapshot,metadata,pngs):
    if not snapshot['ready']:
        raise StorageError('collection_dependencies_missing','图表或六节解释尚未齐备。',409)
    document=Document()
    en=snapshot['language']=='en'
    prepared_quality=any(source.get('schema_version') in (2,3) for source in snapshot['sources'])
    normal=document.styles['Normal']
    normal.font.name='Calibri'; normal.font.size=Pt(10)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'Microsoft YaHei')
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        style=document.styles[name]
        style.font.name='Calibri'
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        fonts.set(qn('w:eastAsia'),'Microsoft YaHei')
        for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme']:
            fonts.attrib.pop(qn('w:'+key),None)
    document.sections[0].header.paragraphs[0].text='Analysis report' if en else '分析报告'
    document.add_heading(snapshot['title'],0)
    document.add_paragraph(('Export revision: ' if en else '导出版本：')+str(metadata['export_revision']))
    document.add_paragraph('Results retain their own sample sizes and exclusions. No new pooled n, P value or causal conclusion is generated.' if en else '各结果保留各自样本量与排除口径。本报告不新增合计人数、P值或因果结论。')
    for section in snapshot['sections']:
        document.add_heading(section['title'],1)
        if section['key']=='figures':
            for item in snapshot['items']:
                figure=item['figure']; raw=pngs[figure['id']]
                if hashlib.sha256(raw).hexdigest()!=figure['sha256']:
                    raise StorageError('figure_changed','图片内容与冻结来源不一致。',409)
                paragraph=document.add_paragraph()
                paragraph.paragraph_format.keep_with_next=True
                image=Image.from_blob(raw)
                scale=min(6/image.px_width,6.5/image.px_height)
                paragraph.add_run().add_picture(BytesIO(raw),width=Inches(image.px_width*scale),height=Inches(image.px_height*scale))
                caption=document.add_paragraph(('Figure ' if en else '图 ')+str(item['figure_number'])+'  '+figure['title'])
                caption.paragraph_format.keep_with_next=True
                document.add_paragraph(item['caption'])
                if 'presentation' in figure:
                    from app.adapters.figure_description import description_lines
                    for line in description_lines(figure['presentation'],en):document.add_paragraph(line)
        elif section['key']=='explanations':
            for item in snapshot['items']:
                document.add_heading(('Figure ' if en else '图 ')+str(item['figure_number'])+'  '+item['figure']['title'],2)
                for entry in item['explanation']['sections']:
                    document.add_heading(entry['title'],3)
                    document.add_paragraph(entry['text'])
                    for fact in entry.get('evidence',[]):
                        document.add_paragraph(str(fact.get('label',fact.get('key','')))+': '+str(fact.get('value','')))
                for limitation in item['explanation'].get('limitations',[]): document.add_paragraph(limitation)
        else:
            for entry in section['content']:
                if section['key']=='processing' and isinstance(entry,dict) and entry.get('schema_version') in (2,3):
                    from app.domain.dataset_processing import quality_lines
                    for line in quality_lines(entry,snapshot['language']):document.add_paragraph(line)
                elif isinstance(entry,str): document.add_paragraph(entry)
                elif 'text' in entry:
                    document.add_paragraph(('Figure ' if en else '图 ')+str(entry['figure_number'])+': '+entry['text'])
                    for limitation in entry.get('limitations',[]): document.add_paragraph(limitation)
                else:
                    table=document.add_table(rows=0,cols=2); table.style='Light Shading Accent 1'
                    for key,value in rows(entry,en,prepared_quality=prepared_quality):
                        row=table.add_row()
                        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
                        cells=row.cells
                        cells[0].text=str(key)
                        cells[1].text=value
                    document.add_paragraph()
    document.add_heading('Appendix — Source references' if en else '附录 来源记录',1)
    document.add_paragraph('These identifiers link the saved results, images and explanations. The SHA-256 digest identifies the original PNG bytes.' if en else '以下标识用于追溯已保存的统计结果、图表和解释。SHA-256摘要用于核对原PNG字节。')
    document.add_paragraph(('Report ID: ' if en else '报告记录ID：')+metadata['id'])
    for item in snapshot['items']:
        document.add_heading(('Figure ' if en else '图 ')+str(item['figure_number'])+'  '+item['figure']['title'],2)
        for labels,value in [
            (('统计结果ID：','Statistical result ID: '),item['result']['id']),
            (('图表配置版本ID：','Figure configuration version ID: '),item['configuration']['id']),
            (('图表记录ID：','Figure ID: '),item['figure']['id']),
            (('解释记录ID：','Explanation ID: '),item['explanation']['id']),
            (('原PNG SHA-256：','Original PNG SHA-256: '),item['figure']['sha256'])]:
            document.add_paragraph(labels[int(en)]+value)
    stream=BytesIO(); document.save(stream)
    stable=BytesIO()
    with ZipFile(BytesIO(stream.getvalue())) as source, ZipFile(stable,'w',ZIP_DEFLATED) as target:
        for entry in source.infolist():
            info=ZipInfo(entry.filename,(1980,1,1,0,0,0)); info.compress_type=ZIP_DEFLATED
            target.writestr(info,source.read(entry.filename))
    return stable.getvalue()
