"""Sparse original XML evidence, with openpyxl used only for display types/styles."""
from datetime import date, datetime, time, timedelta
from io import BytesIO
import json
from posixpath import normpath
from zipfile import ZipFile
from defusedxml.ElementTree import fromstring
from fastapi import HTTPException
from openpyxl import load_workbook
from openpyxl.utils.cell import get_column_letter
from app.adapters.excel import inspect_archive
from app.adapters.excel_process import _error
from app.core.config import ExcelSettings
from app.core.exceptions import StorageError
from app.domain.datasets.contracts import ExcelStructure, SheetStructure, StructureCell, StructureFormula, StructureField, a1, rectangle

NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
REL = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
MAX_OUTPUT_BYTES = 4194304


def structure_envelope_bytes(structure):
    return json.dumps({'structure': structure.model_dump()}, ensure_ascii=False, allow_nan=False).encode('utf-8')


def _range(positions):
    if not positions:
        return None
    low = (min(r for r,c in positions), min(c for r,c in positions))
    high = (max(r for r,c in positions), max(c for r,c in positions))
    return f'{get_column_letter(low[1])}{low[0]}:{get_column_letter(high[1])}{high[0]}'


def extract_excel_structure(content: bytes, filename: str, settings: ExcelSettings) -> ExcelStructure:
    book = None
    try:
        if not content:
            raise _error('empty_file')
        if len(content) > settings.max_upload_bytes:
            raise _error('file_too_large')
        stream = BytesIO(content)
        inspect_archive(stream, settings)
        stream.seek(0)
        book = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
        sheets = []
        scanned = 0
        display_scanned = 0
        with ZipFile(BytesIO(content)) as archive:
            workbook = fromstring(archive.read('xl/workbook.xml'))
            relations = {r.attrib['Id']:r.attrib for r in fromstring(archive.read('xl/_rels/workbook.xml.rels'))}
            for descriptor in workbook.iter(NS+'sheet'):
                relation = relations[descriptor.attrib[REL+'id']]
                if not relation['Type'].endswith('/worksheet'):
                    continue
                target = relation['Target']
                path = normpath(target.lstrip('/') if target.startswith('/') else 'xl/'+target)
                xml = fromstring(archive.read(path))
                worksheet = book[descriptor.attrib['name']]
                worksheet.reset_dimensions()
                # Read once, avoiding random read-only coordinate scans.
                display = {}
                for display_row in worksheet.iter_rows():
                    display_scanned += len(display_row)
                    if display_scanned > settings.max_cells:
                        raise _error('workbook_too_large')
                    for display_cell in display_row:
                        if getattr(display_cell,'coordinate',None):
                            display[display_cell.coordinate] = display_cell
                cells = []
                seen = set()
                previous_row = 0
                for row in xml.findall(NS+'sheetData/'+NS+'row'):
                    row_number = int(row.attrib.get('r', previous_row+1))
                    if row_number <= previous_row or row_number > settings.max_rows:
                        raise ValueError('Invalid row')
                    previous_row = row_number
                    previous_column = 0
                    for node in row.findall(NS+'c'):
                        coordinate = node.attrib.get('r',f'{get_column_letter(previous_column+1)}{row_number}')
                        r,c = a1(coordinate)
                        if r != row_number or c <= previous_column or coordinate in seen:
                            raise ValueError('Contradictory coordinate')
                        previous_column = c
                        seen.add(coordinate)
                        scanned += 1
                        if r > settings.max_rows or c > settings.max_columns or scanned > settings.max_cells:
                            raise _error('workbook_too_large')
                        raw_type = node.attrib.get('t')
                        if raw_type not in (None,'n','s','str','inlineStr','b','e','d'):
                            raise ValueError('Unsupported XML cell type')
                        vnode = node.find(NS+'v')
                        raw_value = vnode.text if vnode is not None else None
                        fnode = node.find(NS+'f')
                        cell = display.get(coordinate)
                        number_format = getattr(cell,'number_format','General')
                        value = getattr(cell,'value',None)
                        formula = None
                        if fnode is not None:
                            kind = 'formula'
                            value = '='+(fnode.text or '')
                            formula = StructureFormula(text=fnode.text, attributes=dict(fnode.attrib),
                                cache_state='missing' if vnode is None else 'empty' if not vnode.text else 'present',
                                cached_type=raw_type, cached_value=raw_value)
                        elif raw_type == 'inlineStr':
                            kind = 'text'
                            inline = node.find(NS+'is')
                            value = '' if inline is None else ''.join(
                                t.text or '' for part in inline
                                for t in ([part] if part.tag == NS+'t' else part.findall(NS+'t') if part.tag == NS+'r' else []))
                        elif raw_type in ('s','str'):
                            kind = 'text'
                            value = value if raw_type == 's' else (raw_value or '')
                        elif raw_type == 'b':
                            if raw_value not in ('0','1'):
                                raise ValueError('Invalid boolean')
                            kind,value = 'boolean',raw_value == '1'
                        elif raw_type == 'e':
                            kind,value = 'error',raw_value
                        elif isinstance(value,(datetime,date,time,timedelta)):
                            kind = 'datetime' if isinstance(value,datetime) else 'date' if isinstance(value,date) else 'time' if isinstance(value,time) else 'duration'
                            value = value.isoformat() if not isinstance(value,timedelta) else str(value)
                        elif raw_value is not None:
                            kind,value = 'number',raw_value
                        else:
                            kind,value = 'blank',None
                        cells.append(StructureCell(row=r,column=c,coordinate=coordinate,kind=kind,value=value,
                            raw_type=raw_type,raw_value=raw_value,number_format=number_format,formula=formula))
                merged = [n.attrib['ref'] for n in xml.findall(NS+'mergeCells/'+NS+'mergeCell')]
                positions = [(c.row,c.column) for c in cells]
                span = positions.copy()
                for merged_range in merged:
                    bounds = rectangle(merged_range)
                    if bounds[1][0] > settings.max_rows or bounds[1][1] > settings.max_columns:
                        raise _error('workbook_too_large')
                    span.extend(bounds)
                row_count = max((r for r,c in span),default=0)
                width = max((c for r,c in span),default=0)
                if row_count*width > settings.max_cells:
                    raise _error('workbook_too_large')
                headers = {c.column:c.value for c in cells if c.row == 1}
                fields = [StructureField(column=c,letter=get_column_letter(c),header_value=headers.get(c),
                    observed_types=sorted({cell.kind for cell in cells if cell.column == c and cell.row > 1})) for c in range(1,width+1)]
                nonempty = [(c.row,c.column) for c in cells if c.value not in (None,'')]
                rows = {r for r,c in nonempty}
                warnings = []
                if worksheet.sheet_state != 'visible': warnings.append('structure_hidden_sheet')
                if any(f.header_value in (None,'') for f in fields): warnings.append('structure_blank_headers')
                names = [f.header_value for f in fields if f.header_value not in (None,'')]
                if len(set(names)) != len(names): warnings.append('structure_duplicate_headers')
                if merged: warnings.append('structure_merged_cells')
                if any(c.formula for c in cells): warnings.append('structure_formula_unverified')
                if not nonempty: warnings.append('structure_empty_sheet')
                if fields: warnings.append('structure_first_row_candidate')
                if any(isinstance(c.value,str) and c.value and not c.value.strip() for c in cells): warnings.append('structure_whitespace_values')
                if any(c.kind == 'number' for c in cells): warnings.append('structure_precision_preserved')
                sheets.append(SheetStructure(index=len(sheets),name=worksheet.title,state=worksheet.sheet_state,
                    row_count=row_count,column_count=width,nonempty_row_count=len(rows),nonempty_data_row_count=len(rows-{1}),
                    stored_range=_range(positions),content_range=_range(nonempty),merged_ranges=merged,fields=fields,cells=cells,warnings=warnings))
        result = ExcelStructure(filename=filename,epoch='1904' if book.epoch.year == 1904 else '1900',sheets=sheets)
        if len(structure_envelope_bytes(result)) > MAX_OUTPUT_BYTES:
            raise _error('parse_output_invalid')
        return result
    except StorageError:
        raise
    except HTTPException as exc:
        code = exc.detail.get('code') if isinstance(exc.detail,dict) else 'parse_failed'
        raise _error(code if code in ('invalid_workbook','workbook_too_large') else 'parse_failed') from None
    except MemoryError:
        raise
    except Exception:
        raise _error('parse_failed') from None
    finally:
        if book is not None:
            book.close()
