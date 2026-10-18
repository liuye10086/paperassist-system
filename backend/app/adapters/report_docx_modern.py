"""Bilingual local reports from immutable saved artifacts only."""
from io import BytesIO
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.adapters.report_docx import clean, paragraph, table
from app.domain.artifact_contracts import artifact_language
from app.domain.correlation_facts import method_summary, reason_text


def build_report(snapshot, report, png):
    project,result,figure,explanation=(snapshot[key] for key in ('project','result','figure','explanation'))
    language=explanation['language']
    def t(zh,en): return zh if language=='zh-CN' else en
    def display(value): return t('无法计算','Unavailable') if value is None else format(value,'.6g') if isinstance(value,float) else str(value)
    correlation=result.get('result_type')=='correlation'
    title=t('相关分析报告','Correlation analysis report') if correlation else t('描述统计分析报告','Descriptive statistics report')
    document=Document()
    page=document.sections[0]
    page.page_width,page.page_height=Cm(21.59),Cm(27.94)
    page.top_margin=page.bottom_margin=page.left_margin=page.right_margin=Cm(2)
    page.header_distance=page.footer_distance=Cm(0.9)
    for name,size in [('Normal',11),('Title',22),('Subtitle',12),('Heading 1',15),('Heading 2',12),('Caption',10)]:
        style=document.styles[name]; style.font.name='Arial'; style.font.size=Pt(size)
        style.font.color.rgb=RGBColor(0,0,0)
        style._element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'),'宋体')
        style.paragraph_format.space_after=Pt(6); style.paragraph_format.line_spacing=1.15
    document.core_properties.title=title; document.core_properties.author='PaperAssist System'
    document.core_properties.subject=clean(project['name'])[:255]
    page.header.paragraphs[0].text='PaperAssist · '+title
    footer=page.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(t('第 ','Page ')); field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    if language=='zh-CN': footer.add_run(' 页')
    paragraph(document,title,'Title'); paragraph(document,project['name'],'Subtitle')
    paragraph(document,t('配置版本','Configuration revision')+f' {result["setup_revision"]} · '+t('导出时间','Exported')+f' {report["created_at"]} (UTC)')
    paragraph(document,t('本报告记录完整行统计、已核验图表及六节解释，研究者需审核解释及适用前提。',
        'This report records complete-case statistics, verified figure metadata and six-section interpretation. Researchers must review the wording and applicable assumptions.'))
    document.add_heading(t('1 数据概况','1 Data overview'),level=1)
    for label,value in [(t('原始文件','Source file'),result['filename']), (t('工作表','Worksheet'),result['selection']['sheet_name']),
                        (t('研究主题 用户填写','Research topic entered by user'),project['research_topic'] or t('未填写','Not specified'))]:
        paragraph(document,label+': '+str(value))
    source=result.get('source')
    if source:
        paragraph(document,t('数据版本','Dataset version')+f': {source["dataset_version"]}')
        if source.get('parent_dataset_version') is not None: paragraph(document,t('原始数据版本','Parent dataset version')+f': {source["parent_dataset_version"]}')
    if correlation:
        for variable in result['variables']:
            paragraph(document,t('变量','Variable')+f': {variable["name"]} ({variable["column"]}); '+t('单位','Unit')+': '+(variable['unit'] or t('未填写','Not specified')))
    else:
        paragraph(document,t('数值字段','Numeric variable')+': '+result['numeric_name'])
        paragraph(document,t('单位','Unit')+': '+(result['selection']['unit'] or t('未填写','Not specified')))
        paragraph(document,t('分组字段','Grouping variable')+': '+(result['group_name'] or t('不分组','No grouping')))
    plan=result.get('analysis_context',{}).get('plan',{})
    research=plan.get('research_request')
    if research:
        document.add_heading(t('用户研究声明','User research declarations'),level=2)
        for key,zh,en in [('question','研究问题','Research question'),('population','研究人群或对象','Study population or subjects'),
                         ('study_design','研究设计','Study design'),('observation_unit','观察单位','Observation unit'),
                         ('relationship_description','对应关系说明','Relationship description'),
                         ('output_requirements','输出要求','Output requirements')]:
            paragraph(document,t(zh,en)+': '+(research.get(key) or t('未提供','Not provided')))
        structures={'independent':t('独立观察','Independent observations'),'paired':t('配对观察','Paired observations'),
                    'repeated':t('重复测量','Repeated measurements'),'unknown':t('未明确','Not specified')}
        paragraph(document,t('观察结构','Observation structure')+': '+structures.get(research.get('observation_structure'),t('未明确','Not specified')))
        paragraph(document,t('上述文字按用户确认的研究声明保留，程序不证明研究设计、样本来源或测量真实性。',
            'These declarations are retained from the user-confirmed research request; the program does not establish design validity, sample provenance or measurement authenticity.'))
    assumptions=result['selection'].get('assumptions')
    if assumptions:
        document.add_heading(t('确认的分析前提','Confirmed analysis assumptions'),level=2)
        def declaration(value): return t('已声明','Declared') if value is True else t('未确认','Not confirmed')
        paragraph(document,t('独立观察声明','Independent-observation declaration')+': '+declaration(assumptions['independent_observations']))
        forms={'linear':t('线性','Linear'),'monotonic':t('单调','Monotonic'),'unknown':t('未明确','Not specified')}
        paragraph(document,t('关系形态','Relationship form')+': '+forms.get(assumptions['relationship_form'],t('未明确','Not specified')))
        paragraph(document,t('推断前提确认','Inferential-assumption confirmation')+': '+declaration(assumptions['inference_assumptions_confirmed']))
    document.add_heading(t('2 数据质量与处理','2 Data quality and processing'),level=1)
    check=result['check']
    paragraph(document,t(f'总行数 {check["row_count"]}；共同完整记录 {check["valid_count"]}；排除 {check["excluded_count"]} 行。',
        f'Source rows: {check["row_count"]}; common complete records: {check["valid_count"]}; excluded rows: {check["excluded_count"]}.'))
    from app.domain.dataset_processing import append_quality
    if not append_quality(document, snapshot, language):
        paragraph(document,t('仅纳入所选变量均有效的完整行，保留原始行对应。未插补、去重或删除离群值；原始数据保留。',
            'Only rows complete across selected variables were included, preserving original row correspondence. No imputation, deduplication or outlier removal was performed. Original data are preserved.'))
    audit=result.get('row_audit')
    if audit:
        paragraph(document,t('纳入原始行号 前二十行','Included original row numbers first twenty')+': '+', '.join(str(row) for row in audit['included_rows'][:20]))
        for excluded in audit['excluded_rows'][:20]:
            paragraph(document,t('排除原始行','Excluded original row')+f' {excluded["row"]}: '+', '.join(excluded['coordinates'])+'; '+t('所选变量缺失','selected-variable missingness'))
        paragraph(document,t('本报告仅展示前二十条纳入及排除坐标，完整坐标审计保存在统计结果中。',
            'This report shows the first twenty included and excluded coordinates; the complete coordinate audit is preserved in the saved result.'))
    if language=='zh-CN':
        for note in check['issues']+check['warnings']: paragraph(document,note)
    document.add_heading(t('3 分析方法','3 Analysis methods'),level=1)
    if correlation: paragraph(document,method_summary(result,language))
    else:
        from app.domain.descriptive_facts import METHOD as EN_METHOD
        from app.domain.explanation_content import METHOD
        from app.domain.dataset_processing import analysis_stage_text
        paragraph(document,analysis_stage_text(METHOD if language=='zh-CN' else EN_METHOD,result))
    paragraph(document,t('表格采用最多 6 位有效数字；无法计算时显示原因，保存结果保留完整精度。',
        'Tables use at most 6 significant digits. Unavailable statistics retain their reason; saved results preserve full precision.'))
    document.add_heading(t('4 统计结果','4 Statistical results'),level=1)
    if correlation:
        stats=result['statistics']
        if stats['mode']=='pair':
            p=stats['p_value']; ci=stats['confidence_interval']
            rows=[(t('相关系数','Correlation coefficient'),display(stats['coefficient'])),
                  (t('有效 n','Effective n'),display(stats['n'])),
                  (t('双侧 P 值','Two-sided P value'),display(p['value']) if p['value'] is not None else reason_text(p['reason'],language)),
                  (t('95% 置信区间','95% confidence interval'),f'{display(ci["low"])} '+t('至','to')+f' {display(ci["high"])}' if ci['low'] is not None else reason_text(ci['reason'],language))]
            table(document,[t('指标','Metric'),t('结果','Result')],rows,[6,11])
            if ci['low'] is None: paragraph(document,reason_text(ci['reason'],language))
        else:
            paragraph(document,t('探索矩阵仅显示系数及共同有效 n，不报告 P 值、区间、显著性星号或多重检验结论。',
                'This exploratory matrix reports coefficients and common effective n only; no P values, intervals, significance stars or multiple-testing inference.'))
            paragraph(document,t('共同有效 n','Common effective n')+': '+display(stats['n']))
            names=[variable['name'] for variable in result['variables']]
            widths=[4]+[13/len(names)]*len(names)
            table(document,[t('变量','Variable'),*names],[[name,*[display(value) for value in values]] for name,values in zip(names,stats['matrix'])],widths)
        warning_text={'small_sample':t('小样本结果可能不稳定','Small-sample results may be unstable'),
            'monte_carlo_p_value':t('P 值使用带固定种子的 Monte Carlo 排列估计','P value uses Monte Carlo permutation with a recorded fixed seed'),
            'asymptotic_p_value':t('P 值使用大样本近似','P value uses a large-sample approximation'),
            'exploratory_matrix':t('矩阵仅作探索','The matrix is exploratory'),
            'correlation_not_causation':t('相关不等于因果','Correlation does not establish causation')}
        for warning in stats.get('warnings',[]): paragraph(document,t('统计提示','Statistical note')+': '+warning_text.get(warning,str(warning)))
    else:
        from app.domain.explanation_content import METRICS as ZH_METRICS
        from app.domain.descriptive_facts import METRICS as EN_METRICS
        metrics=ZH_METRICS if language=='zh-CN' else EN_METRICS
        groups=[(t('总体','Overall'),result['overall'])]+[(item['label'],item['statistics']) for item in result['groups']]
        for keys in [('n','mean','std','min','max'),('q1','median','q3','iqr')]:
            table(document,[t('范围或分组','Scope or group'),*[metrics[key] for key in keys]],
                [[label,*[display(stats[key]) for key in keys]] for label,stats in groups],[4]+[13/len(keys)]*len(keys))
    document.add_heading(t('5 图表与图注','5 Figure and caption'),level=1)
    p=document.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next=True
    shape=p.add_run().add_picture(BytesIO(png),width=Cm(17))
    if shape.height>Cm(12): shape.width=int(shape.width*Cm(12)/shape.height); shape.height=Cm(12)
    shape._inline.docPr.set('descr',clean(figure['title']))
    paragraph(document,figure['title'],'Caption'); paragraph(document,figure['caption'])
    paragraph(document,t('图像使用已保存 PNG 原字节。独立云端结构化数值已与本地结果核对；图片像素仍需人工确认。',
        'The saved PNG bytes are embedded unchanged. Independent cloud values and effective n were checked against local results; image pixels still require human review.'))
    document.add_heading(t('6 分析解释与论文表述','6 Interpretation and manuscript wording'),level=1)
    for section in explanation['sections']:
        document.add_heading(section['title'],level=2); paragraph(document,section['text'])
    document.add_heading(t('7 适用限制','7 Applicability and limitations'),level=1)
    for limitation in explanation['limitations']: paragraph(document,limitation)
    paragraph(document,explanation['verification']['note'])
    document.add_heading(t('附录 A 解释依据','Appendix A Interpretation evidence'),level=1)
    for section in explanation['sections']:
        document.add_heading(section['title'],level=2)
        for evidence in section['evidence']: paragraph(document,f'{evidence["label"]}: {evidence["value"]}\n'+t('引用键','Reference key')+': '+evidence['key'])
    document.add_heading(t('附录 B 来源与版本','Appendix B Sources and versions'),level=1)
    records=[('Project ID',project['id']),('File ID',result['file_id']),('Source SHA256',result['source_sha256']),
        ('Analysis ID',result['id']),('Analysis engine',result['engine']['id']),('Figure ID',figure['id']),
        ('PNG SHA256',figure['sha256']),('Plot engine',figure['engine']['id']),('Explanation ID',explanation['id']),
        ('Explanation engine',explanation['engine']['id']),('Explanation input SHA256',explanation['provenance']['input_sha256']),
        ('Report ID',report['id']),('Report renderer',report['engine']['id']),('Report input SHA256',report['input_sha256'])]
    if result.get('analysis_plan'): records.append(('Plan SHA256',result['analysis_plan']['sha256']))
    for key,label in [('id','Plan ID'),('decision_id','Plan decision ID'),('plan_revision','Plan revision'),
                      ('workflow_version','Plan workflow'),('rules_version','Plan rules')]:
        if key in plan: records.append((label,plan[key]))
    records.extend((key,value) for key,value in result['engine'].items() if key!='id')
    source_labels={'Project ID':'项目编号','File ID':'文件编号','Source SHA256':'原文件 SHA256','Analysis ID':'统计结果编号',
        'Analysis engine':'统计引擎','Figure ID':'图表编号','PNG SHA256':'图片 SHA256','Plot engine':'绘图引擎',
        'Explanation ID':'解释编号','Explanation engine':'解释引擎','Explanation input SHA256':'解释输入 SHA256',
        'Report ID':'报告编号','Report renderer':'报告模板','Report input SHA256':'报告输入 SHA256','Plan SHA256':'计划 SHA256',
        'Plan ID':'计划编号','Plan decision ID':'计划决定编号','Plan revision':'计划代次','Plan workflow':'计划工作流','Plan rules':'计划规则'}
    for label,value in records: paragraph(document,(source_labels.get(label,label) if language=='zh-CN' else label)+': '+str(value))
    paragraph(document,t('时间使用 UTC 保存记录。XML 不支持的控制字符显示为替代符号。',
        'Times use saved UTC records. XML-incompatible controls are displayed as replacement characters.'))
    output=BytesIO(); document.save(output); return output.getvalue()
