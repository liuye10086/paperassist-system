"""School formatting over the common manuscript paragraphs, tables and images."""
from copy import deepcopy
from hashlib import sha256
from io import BytesIO
import json
import re

from docx.enum.section import WD_SECTION_START
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from lxml import etree as E

from app.core.exceptions import StorageError
from app.domain.document_layout import LayoutOptions
from app.domain.manuscripts.school_export import MANUAL_LABELS, validate_school_export
from .document_layout.package import NS, encode, parse, read_package, write_package

SCHOOL_RENDERER_VERSION='school-docx-v1'
STYLE_NAMES={'body':'Normal','heading1':'Heading 1','heading2':'Heading 2','heading3':'Heading 3','title_zh':'Title','title_en':'School English Title','abstract_heading_zh':'School Chinese Abstract','abstract_heading_en':'School English Abstract','keywords_zh':'School Chinese Keywords','keywords_en':'School English Keywords','toc_heading':'School Contents','references_heading':'School References','acknowledgements_heading':'School Acknowledgements','figure_caption':'School Figure Caption','table_caption':'School Table Caption'}


def _prop(parent,name):
    node=parent.find(qn('w:'+name))
    if node is None:
        node=OxmlElement('w:'+name)
        if name=='pPr':parent.insert(0,node)
        else:parent.append(node)
    return node


def _style(document,category,plan):
    name=STYLE_NAMES.get(category,'Normal')
    style=document.styles[name] if name in document.styles else document.styles.add_style(name,WD_STYLE_TYPE.PARAGRAPH)
    spec=plan['styles'].get(category,plan['styles'].get('body',{}))
    if name!='Normal':style.base_style=document.styles['Normal']
    from .manuscript_docx import _font
    _font(style,spec.get('font_size_pt',12))
    style.font.name=spec.get('font_latin','Times New Roman')
    fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
    fonts.set(qn('w:eastAsia'),spec.get('font_east_asia','宋体'))
    fonts.set(qn('w:ascii'),spec.get('font_latin','Times New Roman'))
    fonts.set(qn('w:hAnsi'),spec.get('font_latin','Times New Roman'))
    style.paragraph_format.space_before=Pt(0);style.paragraph_format.space_after=Pt(0)
    alignment={'left':WD_ALIGN_PARAGRAPH.LEFT,'center':WD_ALIGN_PARAGRAPH.CENTER,'right':WD_ALIGN_PARAGRAPH.RIGHT,'both':WD_ALIGN_PARAGRAPH.JUSTIFY,'justify':WD_ALIGN_PARAGRAPH.JUSTIFY}
    style.paragraph_format.alignment=alignment.get(spec.get('alignment'),WD_ALIGN_PARAGRAPH.LEFT)
    ppr=style.element.get_or_add_pPr();spacing=_prop(ppr,'spacing')
    spacing.set(qn('w:line'),str(spec.get('line_spacing','440')))
    spacing.set(qn('w:lineRule'),spec.get('line_rule','exact'))
    for key,attribute in [('space_before_lines','beforeLines'),('space_after_lines','afterLines')]:
        if key in spec:spacing.set(qn('w:'+attribute),str(spec[key]))
    _prop(ppr,'ind').set(qn('w:firstLineChars'),str(spec.get('first_line_chars','0')))
    if category not in ('body','keywords_zh','keywords_en'):
        style.paragraph_format.keep_with_next=True
    if category not in ('heading1','heading2','heading3'):
        # These front-matter headings must not enter the body TOC through an
        # inherited built-in Heading outline level.
        _prop(ppr,'outlineLvl').set(qn('w:val'),'9')
    return style


def _field(paragraph,instruction):
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),instruction)
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';run.append(text);field.append(run);paragraph._p.append(field)


def _section_settings(document,section,spec):
    page=spec['page'];section.page_width=Cm(page['width_cm']);section.page_height=Cm(page['height_cm'])
    for name in ('top','bottom','left','right'):setattr(section,name+'_margin',Cm(page['margin_'+name+'_cm']))
    section.different_first_page_header_footer=bool(spec.get('different_first_page'))
    if spec.get('different_odd_even_pages'):
        document.settings.odd_and_even_pages_header_footer=True
    number=spec.get('page_number',{});pg=_prop(section._sectPr,'pgNumType')
    pg.set(qn('w:fmt'),number.get('format') or 'decimal')
    pg.attrib.pop(qn('w:start'),None)
    if number.get('start') is not None:pg.set(qn('w:start'),str(number['start']))
    starts={'continuous':WD_SECTION_START.CONTINUOUS,'evenPage':WD_SECTION_START.EVEN_PAGE,'oddPage':WD_SECTION_START.ODD_PAGE,'nextPage':WD_SECTION_START.NEW_PAGE}
    section.start_type=starts.get(spec.get('break_type'),WD_SECTION_START.NEW_PAGE)
    for kind,objects in [('headers',{'default':section.header,'first':section.first_page_header,'even':section.even_page_header}),('footers',{'default':section.footer,'first':section.first_page_footer,'even':section.even_page_footer})]:
        entries={v['type']:v for v in spec.get(kind,[])}
        for typ,target in objects.items():
            original=entries.get(typ)
            linked=bool(original and original.get('linked_to_previous'))
            # A new document's first section has no predecessor.
            target.is_linked_to_previous=linked and len(document.sections)>1
            if target.is_linked_to_previous:continue
            for child in list(target._element):target._element.remove(child)
            paragraph=target.add_paragraph();paragraph.alignment=WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent=Pt(0)
            if original:
                if original.get('text'):paragraph.add_run(original['text'])
                for instruction in dict.fromkeys(i.strip() for i in original.get('fields',[])):
                    # Only passive page fields are copied. No instruction text
                    # in a template can create links or executable fields.
                    if re.fullmatch(r'(?:PAGE|NUMPAGES|SECTIONPAGES)(?:\s+\\\*\s+MERGEFORMAT)?',instruction,re.I):_field(paragraph,instruction)
                for run in paragraph.runs:
                    run.font.name='Times New Roman';run.font.size=Pt(9)
                    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),'宋体')


def _rewrite_labels(text,plan):
    for category,pattern in [('tables',r'(?P<prefix>Table\s+|表\s*)(?P<number>\d+)(?!\d|\.\d)'),('figures',r'(?P<prefix>Figure\s+|Fig\.\s*|图\s*)(?P<number>\d+)(?!\d|\.\d)')]:
        mapping=plan['label_map'][category]
        text=re.sub(pattern,lambda m:m['prefix']+mapping.get(m['number'],m['number']),text,flags=re.I)
    return text


class Footnotes:
    def __init__(self,document,snapshot,enabled):
        self.document=document;self.enabled=enabled;self.entries=[]
        self.references={str(r['number']):r for r in snapshot['references']}

    def paragraph(self,item,bookmark,plan):
        from .manuscript_docx import _paragraph
        supplied=deepcopy(item);supplied['text']=_rewrite_labels(item['text'],plan)
        _paragraph(self.document,supplied,bookmark);paragraph=self.document.paragraphs[-1]
        if not self.enabled:return paragraph
        allowed={b.get('project_reference_id') for b in item.get('citations',[])}
        if not allowed:return paragraph
        marker=re.compile(r'\[(\d+(?:,\s*\d+)*)\]|\((\d+(?:,\s*\d+)*)\)')
        matches=[]
        for match in marker.finditer(supplied['text']):
            numbers=[n.strip() for n in (match[1] or match[2]).split(',')]
            if all(n in self.references and self.references[n].get('project_reference_id') in allowed for n in numbers):matches.append((match,numbers))
        if not matches:return paragraph
        for node in list(paragraph._p):
            if node.tag==qn('w:r'):paragraph._p.remove(node)
        end=paragraph._p.find(qn('w:bookmarkEnd'))
        paragraph._p.remove(end)
        cursor=0
        for match,numbers in matches:
            paragraph.add_run(supplied['text'][cursor:match.start()])
            note_id=len(self.entries)+1
            self.entries.append((note_id,'\n'.join('['+n+'] '+self.references[n]['text'] for n in numbers)))
            run=paragraph.add_run();run.font.superscript=True
            ref=OxmlElement('w:footnoteReference');ref.set(qn('w:id'),str(note_id));run._r.append(ref)
            cursor=match.end()
        paragraph.add_run(supplied['text'][cursor:])
        paragraph._p.append(end)
        return paragraph

    def finish(self,content):
        if not self.entries:return content
        parts=read_package(content);notes=OxmlElement('w:footnotes')
        for identity,typ,tag in [(-1,'separator','separator'),(0,'continuationSeparator','continuationSeparator')]:
            note=OxmlElement('w:footnote');note.set(qn('w:id'),str(identity));note.set(qn('w:type'),typ)
            p=OxmlElement('w:p');r=OxmlElement('w:r');r.append(OxmlElement('w:'+tag));p.append(r);note.append(p);notes.append(note)
        for identity,text in self.entries:
            note=OxmlElement('w:footnote');note.set(qn('w:id'),str(identity));p=OxmlElement('w:p')
            run=OxmlElement('w:r');run.append(OxmlElement('w:footnoteRef'));p.append(run)
            run=OxmlElement('w:r');rpr=OxmlElement('w:rPr');fonts=OxmlElement('w:rFonts')
            fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman');fonts.set(qn('w:eastAsia'),'宋体');rpr.append(fonts)
            size=OxmlElement('w:sz');size.set(qn('w:val'),'18');rpr.append(size);run.append(rpr)
            t=OxmlElement('w:t');t.text=' '+text;run.append(t);p.append(run);note.append(p);notes.append(note)
        parts['word/footnotes.xml']=encode(notes)
        rels=parse(parts['word/_rels/document.xml.rels'])
        namespace='http://schemas.openxmlformats.org/package/2006/relationships'
        E.SubElement(rels,'{'+namespace+'}Relationship',Id='rIdSchoolFootnotes',Type='http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes',Target='footnotes.xml')
        parts['word/_rels/document.xml.rels']=encode(rels)
        types=parse(parts['[Content_Types].xml']);namespace='http://schemas.openxmlformats.org/package/2006/content-types'
        E.SubElement(types,'{'+namespace+'}Override',PartName='/word/footnotes.xml',ContentType='application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml')
        parts['[Content_Types].xml']=encode(types)
        settings=parse(parts['word/settings.xml']);fp=_prop(settings,'footnotePr');_prop(fp,'numRestart').set(qn('w:val'),'eachPage')
        parts['word/settings.xml']=encode(settings)
        root=parse(parts['word/document.xml'])
        for section in root.xpath('//w:sectPr',namespaces=NS):
            fp=section.find(qn('w:footnotePr'))
            if fp is None:
                fp=OxmlElement('w:footnotePr');index=0
                while index<len(section) and section[index].tag in (qn('w:headerReference'),qn('w:footerReference')):index+=1
                section.insert(index,fp)
            _prop(fp,'numRestart').set(qn('w:val'),'eachPage')
        parts['word/document.xml']=encode(root)
        return write_package(parts)


def _add_toc(content,plan,language):
    from .document_layout.engine import Editor
    editor=Editor(content,LayoutOptions(toc={'depth':plan['toc']['depth']}),language)
    pending=editor.finish();parts=read_package(pending);root=parse(parts['word/document.xml']);body=root.find(qn('w:body'))
    marker=root.xpath('//w:bookmarkStart[@w:name="PA_SCHOOL_TOC"]',namespaces=NS)[0].getparent()
    title,node=body[0],body[1];body.remove(title);body.remove(node)
    ppr=_prop(title,'pPr');_prop(ppr,'pStyle').set(qn('w:val'),STYLE_NAMES['toc_heading'].replace(' ',''))
    index=body.index(marker);body.insert(index,title);body.insert(index+1,node);body.remove(marker)
    parts['word/document.xml']=encode(root)
    return write_package(parts)


class _SchoolPagination2:
    """Consume a real source section's restart only on its first use."""
    def __init__(self,document):
        self.document=document;self.seen=set();self.active_spec=None

    def apply(self,section,spec):
        applied=deepcopy(spec)
        if spec['id'] in self.seen:
            applied['page_number']=dict(applied.get('page_number',{}),start=None,mode='continue_or_default')
        _section_settings(self.document,section,applied)
        self.seen.add(spec['id']);self.active_spec=spec

    def new_section(self,spec):
        self.apply(self.document.add_section(WD_SECTION_START.NEW_PAGE),spec)


def build_school_docx(snapshot,images,plan,*,source_map=None):
    from .manuscript_docx import _document,_table,_figure,_save_document
    validate_school_export(plan,snapshot)
    if not plan['report']['can_export']:raise StorageError('manuscript_review_blocked','学校导出存在一致性硬错误。',409)
    document=_document()
    language_nodes={}
    mapping=source_map is not None or plan.get('language_mapping_version')==1
    def mark(paragraph,chapter=None,identity=False):
        if not mapping:return
        node=getattr(paragraph,'_p',paragraph)
        language_nodes[node]=dict(chapter_id=chapter['id'] if chapter else None,target_language=chapter['output_language'] if chapter else snapshot['output_language'],protected_identity=identity,scope='reference_metadata' if chapter and chapter['purpose']=='references' else 'narrative')
    def capture(chapter=None):
        if not mapping:return
        for part in document.part.package.parts:
            if str(part.partname)!='' and (str(part.partname)=='/word/document.xml' or re.fullmatch(r'/word/(?:header|footer)\d+\.xml',str(part.partname))):
                for node in part.element.iter(qn('w:p')):
                    if node not in language_nodes:mark(node,chapter)
    # Preserve supplied soft breaks and the source paragraph bookmark, while
    # keeping short lines from stretching under the school's justified body.
    _prop(_prop(document.settings.element,'compat'),'doNotExpandShiftReturn')
    for category in STYLE_NAMES:_style(document,category,plan)
    sections={s['chapter_id']:s for s in plan['sections']}
    pagination=_SchoolPagination2(document) if plan.get('pagination_version')==2 else None
    if pagination:pagination.apply(document.sections[0],sections['__cover__'])
    else:_section_settings(document,document.sections[0],sections['__cover__'])
    front=plan.get('front_matter')
    if front is None:
        paragraph=document.add_paragraph(snapshot['thesis_context']['profile']['school'],style='Title')
        mark(paragraph,identity=True)
        document.add_paragraph(snapshot['title'],style='Title')
        document.add_paragraph('School thesis engineering draft; verify synthetic materials and missing details.' if mapping and snapshot['output_language']=='en' else '学校论文工程稿  合成资料及缺失信息须核对')
        fields=plan['options'].get('manual_fields',{});profile=snapshot['thesis_context']['profile']
        for key,label in MANUAL_LABELS.items():
            paragraph=document.add_paragraph(label+'：'+(fields.get(key) or profile.get(key) or '待填写'))
            mark(paragraph,identity=True)
        document.add_paragraph('Only determined cover rules are applied; formal cover, identity and signatures require explicit user material.' if mapping and snapshot['output_language']=='en' else '封面仅应用可确定规则；原模板未确定的正式封面、身份和签名须人工补齐。')
    else:
        document.add_paragraph(snapshot['title'],style='Title')
        _front_fields(document,plan,None,sections['__cover__']['id'],mark=mark,section_settings=pagination.apply if pagination else None)
    capture()
    manual={i['chapter_id']:i['text'] for i in plan['options'].get('manual_chapters',[])}
    keywords={i['chapter_id']:i['values'] for i in plan['options'].get('keywords',[])}
    notes=Footnotes(document,snapshot,plan.get('citation_footnotes',False))
    bookmark=0;toc_inserted=False;previous_spec=None
    for chapter_index,chapter in enumerate(snapshot['sections']):
        role=chapter['purpose'];en=chapter['output_language']=='en'
        toc_just_inserted=False
        if plan['toc']['enabled'] and not toc_inserted and role not in ('manual','abstract','references','acknowledgements'):
            if pagination:pagination.new_section(sections['__toc__'])
            else:
                section=document.add_section(WD_SECTION_START.NEW_PAGE);_section_settings(document,section,sections['__toc__'])
            paragraph=document.add_paragraph();start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),'500000');start.set(qn('w:name'),'PA_SCHOOL_TOC')
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'500000');paragraph._p.append(start);paragraph._p.append(end);toc_inserted=True
            toc_just_inserted=True
        current_spec=sections[chapter['id']]
        chapter_page_break=False
        if pagination:
            if pagination.active_spec['id']!=current_spec['id']:pagination.new_section(current_spec)
            else:chapter_page_break=chapter['level']==1 or chapter_index==0 or toc_just_inserted
        elif chapter['level']==1 or previous_spec is None or current_spec['id']!=previous_spec['id']:
            section=document.add_section(WD_SECTION_START.NEW_PAGE);_section_settings(document,section,current_spec)
        previous_spec=current_spec
        if role=='abstract':category='abstract_heading_en' if en else 'abstract_heading_zh'
        elif role=='references':category='references_heading'
        elif role=='acknowledgements':category='acknowledgements_heading'
        elif role=='manual':category='abstract_heading_zh'
        else:category='heading'+str(min(3,chapter['level']))
        heading=document.add_paragraph(chapter['title'],style=STYLE_NAMES[category])
        if chapter_page_break:heading.paragraph_format.page_break_before=True
        # References/acknowledgements enter the TOC; front matter and manual
        # declarations retain their school heading appearance without outline.
        if role in ('references','acknowledgements'):_prop(heading._p.get_or_add_pPr(),'outlineLvl').set(qn('w:val'),str(min(3,chapter['level'])-1))
        if role=='references':
            for reference in snapshot['references']:document.add_paragraph('['+str(reference['number'])+'] '+reference['text'])
            if not snapshot['references']:document.add_paragraph('No verified references selected.' if en else '尚未选择已核实参考文献，待补充。')
            if front is not None:
                last=_front_fields(document,plan,chapter['id'],current_spec['id'],mark=lambda p,identity=False:mark(p,chapter,identity),section_settings=pagination.apply if pagination else None)
                previous_spec=plan['front_matter_sections'][last]
            capture(chapter)
            continue
        if role=='abstract':
            title=plan['options'].get('english_title') or (snapshot['title'] if snapshot['output_language']=='en' else '英文题名待填写') if en else snapshot['title']
            document.add_paragraph(title,style=STYLE_NAMES['title_en' if en else 'title_zh'])
        paragraphs=[dict(id='manual:'+chapter['id']+':'+str(index),text=line,citations=[],source_ids=[]) for index,line in enumerate(manual[chapter['id']].splitlines())] if role=='manual' and manual.get(chapter['id']) else chapter['paragraphs']
        for item in paragraphs:
            bookmark+=1;notes.paragraph(item,bookmark,plan)
        if front is not None:
            last=_front_fields(document,plan,chapter['id'],current_spec['id'],mark=lambda p,identity=False:mark(p,chapter,identity),section_settings=pagination.apply if pagination else None)
            previous_spec=plan['front_matter_sections'][last]
        if role=='abstract':
            values=keywords.get(chapter['id'],snapshot.get('keywords',[]) if chapter['output_language']==snapshot['output_language'] else [])
            paragraph=document.add_paragraph(style=STYLE_NAMES['keywords_en' if en else 'keywords_zh'])
            paragraph.add_run('Keywords: ' if en else '关键词：').bold=True
            paragraph.add_run(('; ' if en else '；').join(values) if values else ('Pending user input' if en else '待用户提供'))
        for group in snapshot['results_groups']:
            if group['id'] not in chapter.get('result_group_ids',[]):continue
            for table in group['tables']:
                item=deepcopy(table)
                label=plan['label_map']['tables'].get(str(table['number']),str(table['number']))
                title=re.sub(r'^(?:Table\s+|表\s*)'+re.escape(str(table['number']))+r'\.?\s*','',item['title'],flags=re.I)
                item['title']=('Table ' if en else '表 ')+label+'  '+title
                before=len(document.paragraphs);_table(document,item)
                document.paragraphs[before].style=STYLE_NAMES['table_caption']
                for row in document.tables[-1].rows:
                    for cell in row.cells:
                        for paragraph in cell.paragraphs:
                            paragraph.paragraph_format.line_spacing=1;paragraph.paragraph_format.space_after=Pt(0)
                            _prop(paragraph._p.get_or_add_pPr(),'ind').set(qn('w:firstLineChars'),'0')
                            for run in paragraph.runs:run.font.size=Pt(10.5)
            for figure in group['figures']:
                item=deepcopy(figure);item['number']=plan['label_map']['figures'].get(str(figure['number']),str(figure['number']))
                before=len(document.paragraphs);_figure(document,item,images,en)
                # Inline drawings need their actual height. A fixed body-text
                # line height clips the bitmap in Word, WPS and LibreOffice.
                document.paragraphs[before].paragraph_format.line_spacing=1
                for paragraph in document.paragraphs[before+1:]:paragraph.style=STYLE_NAMES['figure_caption']
        capture(chapter)
    if plan['toc']['enabled'] and not toc_inserted:
        # All remaining eligible headings are references/acknowledgements.
        raise StorageError('school_export_input_invalid','当前章节没有可放置目录的正文入口。',422)
    if mapping:
        for index,(node,metadata) in enumerate(language_nodes.items()):
            start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(700000+index));start.set(qn('w:name'),'PA_LANGUAGE_'+str(index))
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(700000+index));node.append(start);node.append(end)
    content=notes.finish(_save_document(document))
    content=_add_toc(content,plan,snapshot['output_language']) if plan['toc']['enabled'] else content
    if source_map is not None:
        from io import BytesIO
        from zipfile import ZipFile
        from lxml import etree
        with ZipFile(BytesIO(content)) as archive:
            by_name={'PA_LANGUAGE_'+str(i):m for i,m in enumerate(language_nodes.values())}
            for name in sorted(archive.namelist()):
                if name not in ('word/document.xml','word/footnotes.xml') and not re.fullmatch(r'word/(?:header|footer)\d+\.xml',name):continue
                root=etree.fromstring(archive.read(name))
                for index,node in enumerate(root.iter(qn('w:p'))):
                    text=''.join(t.text or '' for t in node.iter(qn('w:t')))
                    if not text:continue
                    bookmark=next((b.get(qn('w:name')) for b in node.iter(qn('w:bookmarkStart')) if b.get(qn('w:name')) in by_name),None)
                    if bookmark:metadata=by_name[bookmark]
                    elif name=='word/footnotes.xml':
                        reference=next((c for c in snapshot['sections'] if c['purpose']=='references'),None)
                        metadata=dict(chapter_id=reference['id'] if reference else None,target_language=reference['output_language'] if reference else snapshot['output_language'],protected_identity=False,scope='reference_metadata')
                    else:
                        # The only text introduced after capture is the renderer's
                        # own TOC heading/cache. Never infer chapter ownership by text.
                        style=node.find(qn('w:pPr')+'/'+qn('w:pStyle'))
                        fields=[f.get(qn('w:instr'),'') for f in node.iter(qn('w:fldSimple'))]
                        if not (style is not None and style.get(qn('w:val'))==STYLE_NAMES['toc_heading'].replace(' ','') or any('TOC' in f for f in fields)):raise ValueError('Unmapped school export text')
                        metadata=dict(chapter_id=None,target_language=snapshot['output_language'],protected_identity=False,generated_field=True,scope='label')
                    source_map.append(dict(location=name+':paragraph:'+str(index),text=text,**metadata,**({'source_bookmark':bookmark} if bookmark else {})))
    return content


def final_school_source_map(pending,content,source_map,language):
    """Rebase exact paragraph ownership after field-cache insertion by anchors.

    TOC labels belong to the explicitly linked heading. Untouched auxiliary
    parts retain their exact XML positions; ordinary text is never guessed.
    """
    from io import BytesIO
    from zipfile import ZipFile
    from lxml import etree
    import json
    by_mark={r['source_bookmark']:r for r in source_map if r.get('source_bookmark')}
    by_location={r['location']:r for r in source_map}
    result=[]
    with ZipFile(BytesIO(pending)) as before,ZipFile(BytesIO(content)) as after:
        targets={}
        if 'customXml/paperassist-layout.xml' in after.namelist():
            manifest=json.loads(etree.fromstring(after.read('customXml/paperassist-layout.xml')).text)
            root=etree.fromstring(after.read('word/document.xml'))
            for target in manifest['toc_targets']:
                nodes=root.xpath('//w:bookmarkStart[@w:name=$name]',namespaces={'w':qn('w:p').split('}')[0][1:]},name=target['bookmark'])
                if len(nodes)!=1:raise ValueError('School TOC target is not unique')
                paragraph=nodes[0]
                while paragraph is not None and paragraph.tag!=qn('w:p'):paragraph=paragraph.getparent()
                marks=[b.get(qn('w:name')) for b in paragraph.iter(qn('w:bookmarkStart')) if b.get(qn('w:name')) in by_mark] if paragraph is not None else []
                if len(marks)!=1:raise ValueError('School TOC target has no exact source chapter')
                targets[target['bookmark']]=by_mark[marks[0]]
        for name in sorted(after.namelist()):
            if name not in ('word/document.xml','word/footnotes.xml') and not re.fullmatch(r'word/(?:header|footer)\d+\.xml',name):continue
            raw=after.read(name);root=etree.fromstring(raw)
            for index,node in enumerate(root.iter(qn('w:p'))):
                text=''.join(t.text or '' for t in node.iter(qn('w:t')))
                if not text:continue
                location=name+':paragraph:'+str(index)
                marks=[b.get(qn('w:name')) for b in node.iter(qn('w:bookmarkStart')) if b.get(qn('w:name')) in by_mark]
                links=[h.get(qn('w:anchor')) for h in node.iter(qn('w:hyperlink')) if h.get(qn('w:anchor')) in targets]
                if len(marks)==1:metadata=by_mark[marks[0]]
                elif not marks and len(links)==1:
                    metadata=dict(targets[links[0]],protected_identity=False,generated_field=True,scope='label',toc_target=links[0])
                    metadata.pop('source_bookmark',None)
                elif name!='word/document.xml' and name in before.namelist() and raw==before.read(name) and location in by_location:metadata=by_location[location]
                else:
                    style=node.find(qn('w:pPr')+'/'+qn('w:pStyle'))
                    if marks or links or style is None or style.get(qn('w:val'))!=STYLE_NAMES['toc_heading'].replace(' ',''):raise ValueError('Final school export text is not mapped')
                    metadata=dict(chapter_id=None,target_language=language,protected_identity=False,generated_field=True,scope='label')
                result.append(dict({k:v for k,v in metadata.items() if k not in ('location','text')},location=location,text=text))
    return result


def complete_school_fields(content,plan):
    if not plan['toc']['enabled']:return content,dict(status='not_requested')
    from .pdf_renderer import update_layout_fields,renderer_identity
    from .document_layout import merge_field_cache
    cache=update_layout_fields(content)
    result=merge_field_cache(content,cache)
    return result,dict(status='updated',software='LibreOffice',renderer_identity=renderer_identity(),toc_entries=len(cache['toc_lines']),field_count=len(cache['fields']),source_sha256=sha256(content).hexdigest(),page_count=None,word_verified=False)


def _front_fields(document,plan,chapter_id,section_id,mark=None,section_settings=None):
    """Apply only the confirmed paragraph source's finite passive attributes."""
    for field in plan['front_matter']['fields']:
        if field['chapter_id']!=chapter_id:continue
        if field['section_id']!=section_id:
            section=document.add_section(WD_SECTION_START.NEW_PAGE)
            if section_settings:section_settings(section,plan['front_matter_sections'][field['section_id']])
            else:_section_settings(document,section,plan['front_matter_sections'][field['section_id']])
            section_id=field['section_id']
        value=plan['options'].get('field_values',{}).get(field['id'],'').strip()
        paragraph=document.add_paragraph(field['label']+'：'+(value or '待填写'))
        if mark:mark(paragraph,identity=field['kind'] in ('identity','signature','date'))
        spec=plan['front_matter_sources'][field['style_source_id']]['style'] if field['style_source_id'] else {}
        ppr=paragraph._p.get_or_add_pPr()
        if 'alignment' in spec:_prop(ppr,'jc').set(qn('w:val'),spec['alignment'])
        for key,attribute in [('space_before','before'),('space_after','after'),('line_spacing','line'),('line_rule','lineRule')]:
            if key in spec:_prop(ppr,'spacing').set(qn('w:'+attribute),str(spec[key]))
        for run in paragraph.runs:
            if 'font_size_pt' in spec:run.font.size=Pt(spec['font_size_pt'])
            fonts=run._r.get_or_add_rPr().get_or_add_rFonts()
            for key,attr in [('font_east_asia','eastAsia'),('font_latin','ascii'),('font_latin','hAnsi')]:
                if key in spec:fonts.set(qn('w:'+attr),spec[key])
    return section_id
