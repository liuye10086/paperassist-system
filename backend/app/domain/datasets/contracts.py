from typing import Literal
import re
from decimal import Decimal
from pydantic import BaseModel, ConfigDict, Field, model_validator
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter

STRUCTURE_PARSER_VERSION = 'openpyxl-3.1.5-structure-v1'
CellKind = Literal['blank','text','number','boolean','date','datetime','time','duration','error','formula']
RawCellType = Literal['n','s','str','inlineStr','b','e','d']
WarningCode = Literal['structure_hidden_sheet','structure_blank_headers','structure_duplicate_headers','structure_merged_cells','structure_formula_unverified','structure_empty_sheet','structure_first_row_candidate','structure_whitespace_values','structure_precision_preserved']


class StrictModel(BaseModel):
    model_config = ConfigDict(extra='forbid', strict=True)


def a1(value):
    if not re.fullmatch(r'[A-Z]+[1-9][0-9]*', value):
        raise ValueError('Invalid coordinate')
    row, column = coordinate_to_tuple(value)
    if row > 100000 or column > 256:
        raise ValueError('Coordinate exceeds limits')
    return row, column


def rectangle(value):
    if value is None:
        return None
    parts = value.split(':')
    if len(parts) not in (1,2):
        raise ValueError('Invalid range')
    start, end = a1(parts[0]), a1(parts[-1])
    if start[0] > end[0] or start[1] > end[1]:
        raise ValueError('Invalid range')
    return start, end


class StructureFormula(StrictModel):
    text: str | None
    attributes: dict[str,str]
    cache_state: Literal['missing','empty','present']
    cached_type: RawCellType | None
    cached_value: str | None
    cache_validity: Literal['unverified'] = 'unverified'

    @model_validator(mode='after')
    def coherent(self):
        if (self.cache_state == 'present') != (self.cached_value not in (None,'')):
            raise ValueError('Invalid cache evidence')
        return self


class StructureCell(StrictModel):
    row: int = Field(ge=1,le=100000)
    column: int = Field(ge=1,le=256)
    coordinate: str
    kind: CellKind
    value: str | bool | None
    raw_type: RawCellType | None
    raw_value: str | None
    number_format: str
    formula: StructureFormula | None

    @model_validator(mode='after')
    def coherent(self):
        if a1(self.coordinate) != (self.row,self.column) or (self.kind == 'formula') != (self.formula is not None):
            raise ValueError('Invalid cell identity')
        if self.kind == 'blank' and self.value is not None:
            raise ValueError('Invalid blank')
        if self.kind == 'boolean' and type(self.value) is not bool:
            raise ValueError('Invalid boolean')
        if self.kind not in ('blank','boolean') and type(self.value) is not str:
            raise ValueError('Invalid display value')
        if self.kind == 'number' and (not re.fullmatch(r'[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?',self.value) or not Decimal(self.value).is_finite()):
            raise ValueError('Invalid number')
        if self.formula and self.value != '='+(self.formula.text or ''):
            raise ValueError('Invalid formula text')
        if self.formula and (self.formula.cached_type != self.raw_type or self.formula.cached_value != self.raw_value):
            raise ValueError('Invalid formula cache evidence')
        if self.kind == 'number' and self.value != self.raw_value:
            raise ValueError('Invalid numeric evidence')
        if self.kind != 'formula':
            allowed_types = {
                'blank': (None, 'n'), 'text': ('s', 'str', 'inlineStr'),
                'number': (None, 'n'), 'boolean': ('b',), 'error': ('e',),
                'date': (None, 'n', 'd'), 'datetime': (None, 'n', 'd'),
                'time': (None, 'n', 'd'), 'duration': (None, 'n', 'd'),
            }
            if self.raw_type not in allowed_types[self.kind]:
                raise ValueError('Invalid raw type evidence')
            if self.kind == 'boolean' and (self.raw_value not in ('0', '1') or self.value != (self.raw_value == '1')):
                raise ValueError('Invalid boolean evidence')
            if self.kind == 'blank' and self.raw_value not in (None, ''):
                raise ValueError('Invalid blank evidence')
            if self.kind == 'error' and self.value != self.raw_value:
                raise ValueError('Invalid error evidence')
            if self.kind == 'text' and self.raw_type == 'str' and self.value != (self.raw_value or ''):
                raise ValueError('Invalid text evidence')
            if self.kind == 'text' and self.raw_type == 'inlineStr' and self.raw_value is not None:
                raise ValueError('Invalid inline text evidence')
        return self


class StructureUnit(StrictModel):
    status: Literal['unconfirmed'] = 'unconfirmed'
    value: None = None


class StructureField(StrictModel):
    column: int = Field(ge=1,le=256)
    letter: str
    header_value: str | bool | None
    observed_types: list[CellKind]
    unit: StructureUnit = Field(default_factory=StructureUnit)

    @model_validator(mode='after')
    def coherent(self):
        if self.letter != get_column_letter(self.column) or len(set(self.observed_types)) != len(self.observed_types):
            raise ValueError('Invalid field')
        return self


class SheetStructure(StrictModel):
    index: int = Field(ge=0,le=49)
    name: str
    state: Literal['visible','hidden','veryHidden']
    row_count: int = Field(ge=0,le=100000)
    column_count: int = Field(ge=0,le=256)
    nonempty_row_count: int = Field(ge=0,le=100000)
    nonempty_data_row_count: int = Field(ge=0,le=100000)
    stored_range: str | None
    content_range: str | None
    merged_ranges: list[str]
    fields: list[StructureField]
    cells: list[StructureCell]
    warnings: list[WarningCode]

    @model_validator(mode='after')
    def coherent(self):
        for value in [self.stored_range,self.content_range,*self.merged_ranges]:
            bounds = rectangle(value)
            if bounds and (bounds[1][0] > self.row_count or bounds[1][1] > self.column_count):
                raise ValueError('Range exceeds sheet')
        positions = [(c.row,c.column) for c in self.cells]
        if positions != sorted(set(positions)) or any(r > self.row_count or c > self.column_count for r,c in positions):
            raise ValueError('Invalid cell ordering')
        if [f.column for f in self.fields] != list(range(1,self.column_count+1)):
            raise ValueError('Invalid fields')
        nonempty = {c.row for c in self.cells if c.value not in (None,'')}
        if self.nonempty_row_count != len(nonempty) or self.nonempty_data_row_count != len(nonempty-{1}):
            raise ValueError('Invalid row statistics')
        def extent(points):
            if not points: return None
            return f'{get_column_letter(min(c for r,c in points))}{min(r for r,c in points)}:{get_column_letter(max(c for r,c in points))}{max(r for r,c in points)}'
        content = [(c.row,c.column) for c in self.cells if c.value not in (None,'')]
        if self.stored_range != extent(positions) or self.content_range != extent(content):
            raise ValueError('Invalid range evidence')
        span = positions.copy()
        for merged in self.merged_ranges:
            span.extend(rectangle(merged))
        if self.row_count != max((r for r,c in span),default=0) or self.column_count != max((c for r,c in span),default=0):
            raise ValueError('Invalid extent')
        for field in self.fields:
            header = next((c.value for c in self.cells if c.row == 1 and c.column == field.column),None)
            observed = {c.kind for c in self.cells if c.row > 1 and c.column == field.column}
            if field.header_value != header or set(field.observed_types) != observed:
                raise ValueError('Invalid field evidence')
        return self


class ExcelStructure(StrictModel):
    schema_version: Literal[1] = 1
    filename: str
    epoch: Literal['1900','1904']
    sheets: list[SheetStructure]

    @model_validator(mode='before')
    @classmethod
    def strict_schema(cls, value):
        if isinstance(value,dict) and 'schema_version' in value and type(value['schema_version']) is not int:
            raise ValueError('Invalid schema version')
        return value

    @model_validator(mode='after')
    def coherent(self):
        if type(self.schema_version) is not int:
            raise ValueError('Invalid schema version')
        if not 1 <= len(self.sheets) <= 50 or [s.index for s in self.sheets] != list(range(len(self.sheets))) or len({s.name for s in self.sheets}) != len(self.sheets):
            raise ValueError('Invalid sheets')
        return self
