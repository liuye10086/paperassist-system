"""Sparse scan: work proportional to stored cells, never rows × columns."""
from collections import Counter, defaultdict
from decimal import Decimal, localcontext
from fractions import Fraction
from math import isfinite
import statistics

from openpyxl.utils.cell import get_column_letter
from app.domain.dataset_analysis import binary64, decimal_label
from app.domain.dataset_preparation.contracts import PreparedSheet
from app.domain.datasets.contracts import rectangle
from app.domain.dataset_profile.contracts import PAGE_SIZE
from app.domain.dataset_profile.temporal import candidate_kind, temporal_key
from fastapi import HTTPException

METRICS = ('mean', 'std', 'min', 'max', 'q1', 'median', 'q3', 'iqr')
FIELD_ISSUES = ('missing', 'outlier', 'mixed_type', 'error', 'formula', 'time_candidate',
    'time_invalid', 'time_ambiguous', 'numeric_precision')


def missing(cell):
    return cell is None or cell.kind == 'blank' or (cell.kind == 'text' and not cell.value.strip())


def typed_key(cell):
    return (cell.kind, Decimal(cell.value) if cell.kind == 'number' else cell.value)


def item(row, column, cell=None, *, first=None, size=None):
    return dict(row=row, column=column, coordinate=f'{get_column_letter(column)}{row}' if column else None,
        kind=cell.kind if cell else None, value=cell.value if cell else None, group_first_row=first, group_size=size)


def fence_label(value):
    # Fractions built from binary64 have power-of-two denominators. This
    # precision covers the exact decimal expansion including subnormal values.
    with localcontext() as context:
        context.prec = 2000
        return decimal_label(Decimal(value.numerator) / Decimal(value.denominator))


def numeric_profile(cells):
    if not cells:
        return None, [], []
    values, precision = [], []
    for cell in cells:
        exact = Decimal(cell.value)
        try:
            values.append(binary64(exact, get_column_letter(cell.column), cell.row, cell.value))
        except (HTTPException, ArithmeticError, ValueError, OverflowError):
            precision.append(cell)
    result = dict(status='unavailable' if precision else 'ok', n=len(cells),
        **{metric: None for metric in METRICS}, lower_fence=None, upper_fence=None,
        outlier_count=0, warnings=[])
    if precision:
        result['warnings'] = ['numeric_precision', 'numeric_unavailable']
        return result, [], precision
    ordered = sorted(values)
    def quartile(quarter):
        index, remainder = divmod((len(ordered) - 1) * quarter, 4)
        return ((4 - remainder) * Fraction(ordered[index]) + remainder * Fraction(ordered[min(index + 1, len(ordered) - 1)])) / 4
    q1, median, q3 = (quartile(q) for q in (1, 2, 3))
    lower, upper = q1 - (q3 - q1) * Fraction(3, 2), q3 + (q3 - q1) * Fraction(3, 2)
    def metric(name, calculate):
        try:
            value = float(calculate())
            if isfinite(value):
                result[name] = value
                return
        except (ArithmeticError, ValueError):
            pass
        if 'metric_unavailable' not in result['warnings']:
            result['warnings'].append('metric_unavailable')
    for name, value in (('q1', q1), ('median', median), ('q3', q3), ('iqr', q3-q1)):
        metric(name, lambda value=value: value)
    metric('mean', lambda: statistics.mean(ordered))
    if len(ordered) < 2:
        result['warnings'].append('std_insufficient')
    else:
        metric('std', lambda: statistics.stdev(ordered))
    result.update(min=ordered[0], max=ordered[-1], lower_fence=fence_label(lower), upper_fence=fence_label(upper))
    outliers = [cell for cell, value in zip(cells, values) if value < lower or value > upper]
    result['outlier_count'] = len(outliers)
    return result, outliers, precision


class ProfileScan:
    def __init__(self, sheet):
        self.sheet = sheet
        prepared = isinstance(sheet, PreparedSheet)
        if prepared:
            (self.start, left), (self.end, right) = rectangle(sheet.data_range)
            self.data_range = sheet.data_range
            fields = [(f.column, f.name, f.unit.model_dump()) for f in sheet.fields]
        else:
            content = [c for c in sheet.cells if c.value not in (None, '')]
            self.start, self.end = 2, max((c.row for c in content), default=0)
            left, right = 1, max((c.column for c in content), default=0)
            self.data_range = f'A2:{get_column_letter(right)}{self.end}' if self.end >= 2 and right else None
            headers = {c.column: c.value for c in sheet.cells if c.row == 1}
            fields = [(c, str(headers[c]) if headers.get(c) not in (None, '') else f'未命名列 {get_column_letter(c)}',
                dict(status='unconfirmed', value=None)) for c in range(left, right + 1)]
        self.record_count = max(0, self.end - self.start + 1)
        self.columns = {c: {} for c, _, _ in fields}
        rows = defaultdict(list)
        for cell in sheet.cells:
            if self.start <= cell.row <= self.end and cell.column in self.columns:
                self.columns[cell.column][cell.row] = cell
                if not missing(cell):
                    rows[cell.row].append((cell.column, typed_key(cell)))
        self.blank_row_count = self.record_count - len(rows)
        groups = defaultdict(list)
        # Omitting missing cells is equivalent to comparing the whole selected
        # rectangle with one missing sentinel. Blank rows stay a separate count.
        for row, entries in rows.items():
            groups[tuple(sorted(entries))].append(row)
        self.duplicate_members = sorted((sorted(members) for members in groups.values() if len(members) > 1), key=lambda members: members[0])
        self.duplicate_groups = len(self.duplicate_members)
        self.duplicate_rows = sum(map(len, self.duplicate_members))
        self.duplicate_extra_rows = self.duplicate_rows - self.duplicate_groups
        self.fields, self.issue_cells = [], {}
        for column, name, unit in fields:
            cells = sorted((c for c in self.columns[column].values() if not missing(c)), key=lambda c: c.row)
            counts = Counter(c.kind for c in cells)
            numeric, outliers, precision = numeric_profile([c for c in cells if c.kind == 'number'])
            category_counts = Counter(typed_key(c) for c in cells if c.kind in ('text', 'boolean'))
            top = category_counts.most_common(10)
            categories = dict(distinct_count=len(category_counts), top=[dict(kind=k, value=v, count=n) for (k, v), n in top],
                other_count=sum(category_counts.values()) - sum(n for _, n in top)) if category_counts else None
            temporal = []
            for kind in ('date', 'datetime', 'time', 'duration'):
                values = [c.value for c in cells if c.kind == kind]
                if values:
                    ordered = sorted(values, key=lambda value: temporal_key(kind, value))
                    temporal.append(dict(kind=kind, count=len(values), min=ordered[0], max=ordered[-1]))
            issues = {kind: [] for kind in FIELD_ISSUES if kind != 'missing'}
            issues.update(outlier=outliers, numeric_precision=precision,
                mixed_type=cells if len(counts) > 1 else [],
                error=[c for c in cells if c.kind == 'error'], formula=[c for c in cells if c.kind == 'formula'])
            for cell in cells:
                candidate = candidate_kind(cell.value) if cell.kind == 'text' else None
                if candidate:
                    issues[candidate].append(cell)
            self.issue_cells[column] = issues
            missing_count = self.record_count - len(cells)
            self.fields.append(dict(column=column, letter=get_column_letter(column), name=name, unit=unit,
                type_counts=dict(counts), missing_count=missing_count, non_missing_count=len(cells),
                distinct_count=len({typed_key(c) for c in cells}), numeric=numeric, categories=categories, temporal=temporal,
                time_candidates=dict(valid_count=len(issues['time_candidate']), invalid_count=len(issues['time_invalid']),
                    ambiguous_count=len(issues['time_ambiguous'])), issue_counts={
                    'missing': missing_count, **{kind: len(members) for kind, members in issues.items()}}))

    def issues(self, kind, column, page):
        offset = (page - 1) * PAGE_SIZE
        if kind == 'duplicate':
            items = []
            seen = 0
            for group in self.duplicate_members:
                for row in group:
                    if offset <= seen < offset + PAGE_SIZE:
                        items.append(item(row, None, first=group[0], size=len(group)))
                    seen += 1
            return self.duplicate_rows, items
        if kind == 'missing':
            cells = self.columns[column]
            total = self.record_count - sum(not missing(cell) for cell in cells.values())
            items, seen = [], 0
            # Only the requested column is scanned, and only one page allocated.
            for row in range(self.start, self.end + 1):
                if missing(cells.get(row)):
                    if offset <= seen < offset + PAGE_SIZE:
                        items.append(item(row, column, cells.get(row)))
                    seen += 1
                    if len(items) == PAGE_SIZE:
                        break
            return total, items
        cells = self.issue_cells[column][kind]
        return len(cells), [item(c.row, column, c) for c in cells[offset:offset + PAGE_SIZE]]
