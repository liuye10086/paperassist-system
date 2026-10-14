"""Frozen contracts for one rectangular, single-header preparation recipe."""
from decimal import Decimal
from collections import Counter
import re
from typing import Literal
from openpyxl.utils.cell import get_column_letter
from pydantic import Field, model_validator
from app.domain.analysis_sources import DatasetAnalysisSource
from app.domain.datasets.contracts import CellKind, StrictModel, a1, rectangle

TRANSFORM_VERSION = 'excel_prepare_v1'
WORKFLOW_VERSION = 'dataset_preparation_v1'
MAX_OBJECT_BYTES = 4194304
MAX_CHANGES = 20
Operation = Literal['trim_text', 'missing_marker']


def canonical_range(value: str) -> str:
    start, end = rectangle(value)
    return f'{get_column_letter(start[1])}{start[0]}:{get_column_letter(end[1])}{end[0]}'


def validate_value(kind, value):
    if kind == 'blank':
        if value is not None:
            raise ValueError('Invalid blank')
    elif kind == 'boolean':
        if type(value) is not bool:
            raise ValueError('Invalid boolean')
    elif type(value) is not str:
        raise ValueError('Invalid display value')
    elif kind == 'number' and (not re.fullmatch(r'[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]+)?', value) or not Decimal(value).is_finite()):
        raise ValueError('Invalid number')
    elif kind == 'formula' and not value.startswith('='):
        raise ValueError('Invalid formula display')


class PreparationUnit(StrictModel):
    status: Literal['unconfirmed', 'confirmed']
    value: str | None = Field(max_length=80)

    @model_validator(mode='after')
    def coherent(self):
        if (self.status == 'unconfirmed') != (self.value is None):
            raise ValueError('Unit requires explicit confirmation')
        return self


class PreparationField(StrictModel):
    column: int = Field(ge=1, le=256)
    name: str = Field(min_length=1, max_length=160)
    unit: PreparationUnit
    trim_text: bool = False
    missing_markers: list[str] = Field(default_factory=list, max_length=20)

    @model_validator(mode='after')
    def coherent(self):
        if any(not 1 <= len(marker) <= 160 for marker in self.missing_markers) or len(set(self.missing_markers)) != len(self.missing_markers):
            raise ValueError('Invalid missing markers')
        return self


class PreparationRecipe(StrictModel):
    header_row: int = Field(ge=1, le=100000)
    data_range: str
    fields: list[PreparationField] = Field(min_length=1, max_length=256)

    @model_validator(mode='after')
    def coherent(self):
        start, end = rectangle(self.data_range)
        if self.data_range != canonical_range(self.data_range) or self.header_row >= start[0]:
            raise ValueError('Invalid header or canonical data range')
        if [f.column for f in self.fields] != list(range(start[1], end[1] + 1)):
            raise ValueError('Fields must be contiguous original columns')
        return self


class PreparedCell(StrictModel):
    row: int = Field(ge=1, le=100000)
    column: int = Field(ge=1, le=256)
    coordinate: str
    kind: CellKind
    value: str | bool | None

    @model_validator(mode='after')
    def coherent(self):
        if a1(self.coordinate) != (self.row, self.column):
            raise ValueError('Invalid cell identity')
        validate_value(self.kind, self.value)
        return self


def intersects(first, second):
    a, b = rectangle(first)
    c, d = rectangle(second)
    return a[0] <= d[0] and c[0] <= b[0] and a[1] <= d[1] and c[1] <= b[1]


class PreparedSheet(StrictModel):
    index: int = Field(ge=0, le=49)
    name: str = Field(min_length=1, max_length=31)
    data_range: str
    fields: list[PreparationField] = Field(min_length=1, max_length=256)
    cells: list[PreparedCell]
    merged_ranges: list[str]

    @property
    def row_count(self):
        start, end = rectangle(self.data_range)
        return end[0] - start[0] + 1

    @property
    def column_count(self):
        return len(self.fields)

    @model_validator(mode='after')
    def coherent(self):
        start, end = rectangle(self.data_range)
        if self.data_range != canonical_range(self.data_range) or [f.column for f in self.fields] != list(range(start[1], end[1] + 1)):
            raise ValueError('Invalid prepared rectangle')
        previous = (0, 0)
        for cell in self.cells:
            position = (cell.row, cell.column)
            if position <= previous or not (start[0] <= cell.row <= end[0] and start[1] <= cell.column <= end[1]):
                raise ValueError('Invalid prepared cell ordering or bounds')
            previous = position
        if len(set(self.merged_ranges)) != len(self.merged_ranges) or any(not intersects(area, self.data_range) for area in self.merged_ranges):
            raise ValueError('Invalid merged area evidence')
        return self


class CellChange(StrictModel):
    row: int = Field(ge=1, le=100000)
    column: int = Field(ge=1, le=256)
    coordinate: str
    before: dict[str, str | bool | None]
    after: dict[str, str | bool | None]
    operations: list[Operation] = Field(min_length=1, max_length=2)

    @model_validator(mode='after')
    def coherent(self):
        if a1(self.coordinate) != (self.row, self.column):
            raise ValueError('Invalid change identity')
        for evidence in (self.before, self.after):
            if set(evidence) != {'kind', 'value'} or evidence['kind'] not in ('blank', 'text'):
                raise ValueError('Only text operations are allowed')
            validate_value(evidence['kind'], evidence['value'])
        if self.before['kind'] != 'text' or self.before == self.after or self.operations not in (['trim_text'], ['missing_marker'], ['trim_text', 'missing_marker']):
            raise ValueError('Invalid change operations')
        return self


class PreparationColumnSummary(StrictModel):
    column: int = Field(ge=1, le=256)
    name: str = Field(min_length=1, max_length=160)
    unit: PreparationUnit
    observed_types: list[CellKind]
    missing_count: int = Field(ge=0, le=100000)
    non_missing_count: int = Field(ge=0, le=100000)
    trimmed_cell_count: int = Field(ge=0, le=100000)
    missing_marker_cell_count: int = Field(ge=0, le=100000)

    @model_validator(mode='after')
    def coherent(self):
        if len(set(self.observed_types)) != len(self.observed_types):
            raise ValueError('Duplicate observed types')
        return self


class PreparationIssue(StrictModel):
    code: Literal['formula_unverified', 'error_values', 'mixed_types', 'unconfirmed_unit', 'merged_cells']
    column: int | None = Field(ge=1, le=256)
    count: int = Field(ge=0)


class PreparationSummary(StrictModel):
    record_count: int = Field(ge=1, le=100000)
    changed_cell_count: int = Field(ge=0)
    changed_row_count: int = Field(ge=0, le=100000)
    trimmed_cell_count: int = Field(ge=0)
    missing_marker_cell_count: int = Field(ge=0)
    columns: list[PreparationColumnSummary] = Field(min_length=1, max_length=256)
    unresolved: list[PreparationIssue]
    changes_truncated: bool

    @model_validator(mode='after')
    def coherent(self):
        if self.changed_row_count > min(self.record_count, self.changed_cell_count):
            raise ValueError('Invalid changed rows')
        if not max(self.trimmed_cell_count, self.missing_marker_cell_count) <= self.changed_cell_count <= self.trimmed_cell_count + self.missing_marker_cell_count:
            raise ValueError('Invalid operation counts')
        if self.changes_truncated != (self.changed_cell_count > MAX_CHANGES):
            raise ValueError('Invalid changes truncation')
        if self.trimmed_cell_count != sum(c.trimmed_cell_count for c in self.columns) or self.missing_marker_cell_count != sum(c.missing_marker_cell_count for c in self.columns):
            raise ValueError('Invalid column operation totals')
        if any(c.missing_count + c.non_missing_count != self.record_count for c in self.columns):
            raise ValueError('Invalid column row counts')
        return self


class PreparedDatasetObject(StrictModel):
    schema_version: Literal[1] = 1
    parent_source: DatasetAnalysisSource
    transform_version: Literal['excel_prepare_v1'] = TRANSFORM_VERSION
    recipe: PreparationRecipe
    sheet: PreparedSheet
    summary: PreparationSummary
    changes: list[CellChange] = Field(max_length=MAX_CHANGES)

    @model_validator(mode='before')
    @classmethod
    def strict_schema(cls, value):
        if isinstance(value, dict) and 'schema_version' in value and type(value['schema_version']) is not int:
            raise ValueError('Invalid schema version')
        return value

    @model_validator(mode='after')
    def coherent(self):
        if self.parent_source.kind != 'dataset_version' or (self.parent_source.sheet_index, self.parent_source.sheet_name) != (self.sheet.index, self.sheet.name):
            raise ValueError('Invalid raw parent identity')
        if self.sheet.data_range != self.recipe.data_range or self.sheet.fields != self.recipe.fields:
            raise ValueError('Recipe differs from effective sheet')
        if self.summary.record_count != self.sheet.row_count or [(c.column, c.name, c.unit) for c in self.summary.columns] != [(f.column, f.name, f.unit) for f in self.recipe.fields]:
            raise ValueError('Summary differs from recipe')
        if len(self.changes) != min(MAX_CHANGES, self.summary.changed_cell_count):
            raise ValueError('Invalid changes preview length')
        effective, observed = {}, {field.column: Counter() for field in self.recipe.fields}
        for cell in self.sheet.cells:
            effective[cell.row, cell.column] = cell
            if cell.kind != 'blank' and not (cell.kind == 'text' and not cell.value.strip()):
                observed[cell.column][cell.kind] += 1
        expected_issues = []
        for column, field in zip(self.summary.columns, self.recipe.fields):
            counts = observed[field.column]
            if column.non_missing_count != sum(counts.values()) or set(column.observed_types) != set(counts):
                raise ValueError('Summary differs from effective cells')
            if (not field.trim_text and column.trimmed_cell_count) or (not field.missing_markers and column.missing_marker_cell_count):
                raise ValueError('Summary operations differ from recipe')
            for code, kind in (('formula_unverified', 'formula'), ('error_values', 'error')):
                if counts[kind]:
                    expected_issues.append((code, field.column, counts[kind]))
            if len(counts) > 1:
                expected_issues.append(('mixed_types', field.column, sum(counts.values())))
            if field.unit.status == 'unconfirmed':
                expected_issues.append(('unconfirmed_unit', field.column, 1))
        if self.sheet.merged_ranges:
            expected_issues.append(('merged_cells', None, len(self.sheet.merged_ranges)))
        if [(i.code, i.column, i.count) for i in self.summary.unresolved] != expected_issues:
            raise ValueError('Issues differ from effective data')
        if self.summary.changed_cell_count > len(self.sheet.cells):
            raise ValueError('Changed cells exceed effective stored cells')
        previous = (0, 0)
        fields = {f.column: f for f in self.recipe.fields}
        for change in self.changes:
            position = (change.row, change.column)
            cell = effective.get(position)
            if position <= previous or cell is None or change.after != {'kind': cell.kind, 'value': cell.value}:
                raise ValueError('Invalid change evidence')
            field = fields[change.column]
            value = change.before['value']
            operations = []
            if field.trim_text and value.strip() != value:
                value = value.strip()
                operations.append('trim_text')
            kind = 'text'
            if value in field.missing_markers:
                kind, value = 'blank', None
                operations.append('missing_marker')
            if operations != change.operations or change.after != {'kind': kind, 'value': value}:
                raise ValueError('Change does not follow approved text rules')
            previous = position
        return self
