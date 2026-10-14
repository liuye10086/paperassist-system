"""One pure sparse-cell transform; never reads Excel or infers cleaning rules."""
from collections import Counter
from openpyxl.utils.cell import get_column_letter
from app.core.exceptions import StorageError
from app.domain.analysis_sources import DatasetAnalysisSource
from app.domain.datasets.contracts import SheetStructure, rectangle
from app.domain.dataset_preparation.contracts import (
    MAX_CHANGES, CellChange, PreparationColumnSummary, PreparationField,
    PreparationIssue, PreparationRecipe, PreparationSummary, PreparationUnit,
    PreparedCell, PreparedDatasetObject, PreparedSheet, canonical_range, intersects,
)


def invalid():
    return StorageError('preparation_invalid', '整理规则无效，请检查表头行、数据区域、字段名称和单位确认。', 422)


def _within_sheet(sheet, recipe):
    start, end = rectangle(recipe.data_range)
    if recipe.header_row > sheet.row_count or end[0] > sheet.row_count or end[1] > sheet.column_count:
        raise invalid()
    return start, end


def validate_recipe(sheet: SheetStructure, recipe: PreparationRecipe) -> PreparationRecipe:
    """Validate an input recipe at create/revise without running the transform."""
    try:
        validated = PreparationRecipe.model_validate(recipe.model_dump() if isinstance(recipe, PreparationRecipe) else recipe)
        _within_sheet(sheet, validated)
        return validated
    except (ValueError, TypeError, AttributeError) as exc:
        raise invalid() from exc


def build_recipe(sheet: SheetStructure, header_row=1, data_range=None) -> PreparationRecipe:
    """Suggest content columns and header+1..last content row for explicit review."""
    try:
        if type(header_row) is not int or not 1 <= header_row <= min(100000, sheet.row_count):
            raise ValueError('Invalid header row')
        if data_range is None:
            content = [c for c in sheet.cells if c.row > header_row and c.value not in (None, '')]
            if not content:
                raise ValueError('No data after header')
            data_range = f'{get_column_letter(min(c.column for c in content))}{header_row + 1}:{get_column_letter(max(c.column for c in content))}{max(c.row for c in content)}'
        data_range = canonical_range(data_range)
        start, end = rectangle(data_range)
        headers = {c.column: c.value for c in sheet.cells if c.row == header_row}
        fields = [PreparationField(column=column,
            # This is an editable suggestion; the full original header stays in the raw version.
            name=str(headers[column])[:160] if headers.get(column) not in (None, '') else f'未命名列 {get_column_letter(column)}',
            unit=PreparationUnit(status='unconfirmed', value=None)) for column in range(start[1], end[1] + 1)]
        recipe = PreparationRecipe(header_row=header_row, data_range=data_range, fields=fields)
        return validate_recipe(sheet, recipe)
    except (ValueError, TypeError, AttributeError) as exc:
        raise invalid() from exc


def prepare_dataset(sheet: SheetStructure, parent_source, recipe: PreparationRecipe) -> PreparedDatasetObject:
    """Apply only per-column strip and exact text markers in the selected rectangle."""
    try:
        # Copy/revalidate mutable request models once; never share references upstream.
        recipe = validate_recipe(sheet, recipe)
        parent_source = DatasetAnalysisSource.model_validate(parent_source.model_dump() if isinstance(parent_source, DatasetAnalysisSource) else parent_source)
        if parent_source.kind != 'dataset_version' or (parent_source.sheet_index, parent_source.sheet_name) != (sheet.index, sheet.name):
            raise ValueError('Raw parent and sheet differ')
        start, end = _within_sheet(sheet, recipe)
        record_count = end[0] - start[0] + 1
        fields = {f.column: f for f in recipe.fields}
        types = {f.column: Counter() for f in recipe.fields}
        nonmissing = Counter()
        trims, markers = Counter(), Counter()
        changes, cells, changed_rows = [], [], set()
        changed_count = 0
        for raw in sheet.cells:
            if not (start[0] <= raw.row <= end[0] and start[1] <= raw.column <= end[1]):
                continue
            field = fields[raw.column]
            kind, value = raw.kind, raw.value
            operations = []
            if kind == 'text':
                if field.trim_text and value.strip() != value:
                    value = value.strip()
                    operations.append('trim_text')
                    trims[raw.column] += 1
                if value in field.missing_markers:
                    kind, value = 'blank', None
                    operations.append('missing_marker')
                    markers[raw.column] += 1
            cell = PreparedCell(row=raw.row, column=raw.column, coordinate=raw.coordinate, kind=kind, value=value)
            cells.append(cell)
            # Whitespace-only text stays text in the object but is analytically missing.
            if kind != 'blank' and not (kind == 'text' and not value.strip()):
                nonmissing[raw.column] += 1
                types[raw.column][kind] += 1
            if operations:
                changed_count += 1
                changed_rows.add(raw.row)
                if len(changes) < MAX_CHANGES:
                    changes.append(CellChange(row=raw.row, column=raw.column, coordinate=raw.coordinate,
                        before={'kind': raw.kind, 'value': raw.value}, after={'kind': kind, 'value': value}, operations=operations))
        merges = [area for area in sheet.merged_ranges if intersects(area, recipe.data_range)]
        columns, unresolved = [], []
        for field in recipe.fields:
            column_types = types[field.column]
            columns.append(PreparationColumnSummary(column=field.column, name=field.name, unit=field.unit,
                observed_types=sorted(column_types), missing_count=record_count - nonmissing[field.column],
                non_missing_count=nonmissing[field.column], trimmed_cell_count=trims[field.column],
                missing_marker_cell_count=markers[field.column]))
            for code, kind in (('formula_unverified', 'formula'), ('error_values', 'error')):
                if column_types[kind]:
                    unresolved.append(PreparationIssue(code=code, column=field.column, count=column_types[kind]))
            if len(column_types) > 1:
                unresolved.append(PreparationIssue(code='mixed_types', column=field.column, count=nonmissing[field.column]))
            if field.unit.status == 'unconfirmed':
                unresolved.append(PreparationIssue(code='unconfirmed_unit', column=field.column, count=1))
        if merges:
            unresolved.append(PreparationIssue(code='merged_cells', column=None, count=len(merges)))
        return PreparedDatasetObject(parent_source=parent_source, recipe=recipe,
            sheet=PreparedSheet(index=sheet.index, name=sheet.name, data_range=recipe.data_range,
                fields=recipe.fields, cells=cells, merged_ranges=merges),
            summary=PreparationSummary(record_count=record_count, changed_cell_count=changed_count,
                changed_row_count=len(changed_rows), trimmed_cell_count=sum(trims.values()),
                missing_marker_cell_count=sum(markers.values()), columns=columns, unresolved=unresolved,
                changes_truncated=changed_count > MAX_CHANGES), changes=changes)
    except (ValueError, TypeError, AttributeError) as exc:
        raise invalid() from exc
