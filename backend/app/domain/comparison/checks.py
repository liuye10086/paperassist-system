"""Full prepared-sheet typed categories, ID uniqueness and original-row audit."""
from decimal import Decimal
from fastapi import HTTPException
from openpyxl.utils.cell import column_index_from_string, get_column_letter
from app.core.exceptions import StorageError
from app.domain.comparison.contracts import canonical_number, selected_columns
from app.domain.dataset_analysis import binary64
from app.domain.dataset_profile.scan import missing
from app.domain.analysis_plan.profile import PlanProfileScan as ProfileScan
from app.domain.datasets.contracts import rectangle
from app.domain.dataset_rows import source_rows
from app.domain.dataset_preparation.contracts import PreparedSheet


def typed_identity(cell, *, id_role=False):
    if missing(cell): return None
    if cell.kind=='text':
        if not id_role and len(cell.value)>500:
            raise StorageError('group_type','分组标签最长500字符，请在新的整理版本中核对。',422)
        return ('text',cell.value)
    if cell.kind=='number':
        try:return ('number',canonical_number(cell.value))
        except ValueError:
            raise StorageError('pair_id_type' if id_role else 'group_type','数值分组或配对ID超出可安全表示的标签范围。',422) from None
    if cell.kind=='boolean': return ('boolean','true' if cell.value else 'false')
    raise StorageError('pair_id_type' if id_role else 'group_type','配对ID或组字段只能包含文本、数值、布尔或缺失值。',422)


def friendly_label(identity):
    kind,value=identity
    return ('TRUE' if value=='true' else 'FALSE') if kind=='boolean' else value


def group_options(sheet, dataset_id, sheet_name, group_column):
    if not isinstance(sheet,PreparedSheet) or sheet.name!=sheet_name:
        raise StorageError('sheet_not_found','请选择已确认整理数据及其工作表。',422)
    if group_column not in {get_column_letter(f.column) for f in sheet.fields}:
        raise StorageError('invalid_column','分组字段不存在。',422)
    cells={(c.row,c.column):c for c in sheet.cells};start,end=rectangle(sheet.data_range)
    categories={};missing_count=0
    for row in source_rows(sheet):
        identity=typed_identity(cells.get((row,column_index_from_string(group_column))))
        if identity is None:missing_count+=1
        else:categories[identity]=categories.get(identity,0)+1
    return dict(dataset_id=dataset_id,sheet_name=sheet_name,group_column=group_column,row_count=sheet.row_count,
        missing_count=missing_count,options=[dict(kind=k,value=v,label=friendly_label((k,v)),count=n) for (k,v),n in categories.items()])


def selected_records(sheet, selection):
    if not isinstance(sheet,PreparedSheet) or sheet.name!=selection.sheet_name:
        raise StorageError('sheet_not_found','比较分析须使用冻结的已确认整理工作表。',422)
    fields={get_column_letter(f.column):f for f in sheet.fields}
    if any(c not in fields for c in selected_columns(selection)):
        raise StorageError('invalid_column','所选比较字段不存在。',422)
    numeric=[selection.before_column,selection.after_column] if selection.design=='paired' else [selection.numeric_column]
    for letter in numeric:
        unit=fields[letter].unit
        if unit.status=='confirmed' and unit.value!=selection.unit:
            raise StorageError('unit_conflict','测量字段单位与声明不一致；不会自动换算数值。',422)
    if selection.design=='paired' and fields[numeric[0]].unit!=fields[numeric[1]].unit:
        raise StorageError('unit_conflict','配对前后字段单位须可比且一致，请重新核对整理单位。',422)
    cells={(c.row,c.column):c for c in sheet.cells};start,end=rectangle(sheet.data_range)
    def cell(row,letter):return cells.get((row,column_index_from_string(letter)))
    if selection.design=='paired':
        identifiers=set()
        for row in source_rows(sheet):
            identity=typed_identity(cell(row,selection.pair_id_column),id_role=True)
            if identity is None:raise StorageError('pair_id_missing',f'配对ID缺失：{selection.pair_id_column}{row}；所有来源行都需要ID。',422)
            if identity in identifiers:raise StorageError('pair_id_duplicate',f'配对ID重复：{selection.pair_id_column}{row}；不会自动去重。',422)
            identifiers.add(identity)
    identities=[(g.kind,g.value) for g in selection.groups]
    groups=[[],[]] if selection.design=='paired' else [[] for _ in identities]
    group_rows=[[] for _ in groups];included=[];excluded=[];observed=set()
    for row in source_rows(sheet):
        absent=[];group_index=0
        if selection.design!='paired':
            identity=typed_identity(cell(row,selection.group_column))
            if identity is None:absent.append(f'{selection.group_column}{row}')
            elif identity not in identities:
                excluded.append(dict(row=row,coordinates=[f'{selection.group_column}{row}'],reason='group_not_selected'));continue
            else:group_index=identities.index(identity);observed.add(identity)
        values=[]
        for letter in numeric:
            measured=cell(row,letter)
            if missing(measured):absent.append(f'{letter}{row}')
            elif measured.kind!='number':raise StorageError('numeric_type',f'比较测量字段须是实际数值或缺失值：{letter}{row}。',422)
            else:
                try:values.append(binary64(Decimal(measured.value),letter,row,measured.value))
                except HTTPException as exc:raise StorageError('numeric_precision','所选测量值超出安全精度或表示范围。',422) from exc
        if absent:excluded.append(dict(row=row,coordinates=absent,reason='selected_missing'))
        else:
            included.append(row)
            if selection.design=='paired':
                for index,value in enumerate(values):groups[index].append(value);group_rows[index].append(row)
            else:groups[group_index].append(values[0]);group_rows[group_index].append(row)
    if selection.design!='paired' and any(identity not in observed for identity in identities):
        raise StorageError('groups_missing','所选参与组在冻结工作表中不存在。',422)
    variables=[dict(column=c,name=fields[c].name,unit=selection.unit) for c in numeric]
    labels=[fields[c].name for c in numeric] if selection.design=='paired' else [friendly_label(i) for i in identities]
    audit=dict(included_rows=included,excluded_rows=excluded,missing_policy=selection.missing_policy,group_rows=group_rows,pair_id_column=selection.pair_id_column)
    return groups,variables,audit,labels


def build_comparison_check(sheet, selection, research):
    from app.domain.analysis_plan.contracts import ModernPlanCheck
    from app.domain.analysis_plan.research import missing_research
    from app.domain.comparison.kernel import validate_comparison_inputs, ComparisonError
    issues=[];selected=[]
    def issue(code,field,count,message,coordinates=(),resolution=None):
        issues.append(dict(id=f'{code}:{field or "all"}',code=code,field=field,count=count,
            severity='warning' if resolution else 'blocking',message=message,coordinates=list(coordinates)[:5],
            allowed_resolutions=[resolution] if resolution else []))
    profile=ProfileScan(sheet,selected_columns=selected_columns(selection));lookup={f['letter']:f for f in profile.fields}
    numeric=[selection.before_column,selection.after_column] if selection.design=='paired' else [selection.numeric_column]
    for letter in selected_columns(selection):
        field=lookup.get(letter)
        if field is None:continue
        metrics=field['numeric'] or {};unit=field['unit'];label=selection.unit if letter in numeric else unit['value'] or ''
        selected.append(dict(letter=letter,name=field['name'],type_counts=field['type_counts'],missing_count=field['missing_count'],
            non_missing_count=field['non_missing_count'],prepared_unit=unit,declared_unit=dict(status='label' if label else 'unknown',value=label or None),
            observed_min=metrics.get('min'),observed_max=metrics.get('max')))
        if letter in numeric:
            if unit['status']=='unconfirmed':issue('unit_unconfirmed',letter,1,'单位标签须由用户确认；不自动换算。',resolution='declare_label' if label else 'unknown')
            count,items=profile.issues('outlier',field['column'],1)
            if count:issue('outlier',letter,count,'IQR异常候选保持纳入，不自动删除。',[i['coordinate'] for i in items],'retain')
    n=excluded_count=0;group_counts=[]
    try:
        groups,_,audit,labels=selected_records(sheet,selection)
        n=len(audit['included_rows']);excluded_count=len(audit['excluded_rows'])
        group_counts=[dict(label=label,count=len(values)) for label,values in zip(labels,groups)]
        validate_comparison_inputs(groups,method=selection.method,posthoc=selection.posthoc)
        if min(map(len,groups))<10:issue('small_sample',None,n,'样本量较小，推断可能不稳定。',resolution='accept_partial_metrics')
        for reason,resolution,message in [('selected_missing','exclude_selected_missing','缺失测量按完整记录或完整对排除，保留原来源。'),('group_not_selected','retain','仅采用确认的参与组；其余组记录明确排除，原数据保留。')]:
            excluded=[item for item in audit['excluded_rows'] if item['reason']==reason]
            if excluded:issue(reason,None,len(excluded),message,[c for i in excluded for c in i['coordinates']],resolution)
    except (StorageError,ComparisonError) as exc:issue(exc.code,None,n,getattr(exc,'message','比较字段检查未通过。'))
    missing_fields=missing_research(research)
    if research['observation_structure']!=('paired' if selection.design=='paired' else 'independent') and 'observation_structure' not in missing_fields:missing_fields.append('observation_structure')
    assumptions=selection.assumptions
    if not assumptions.independent_units:missing_fields.append('independent_units')
    required='normal' if selection.method in ('welch_t','paired_t','welch_anova') else 'symmetric_differences' if selection.method=='wilcoxon' else 'rank'
    if assumptions.distribution!=required:missing_fields.append('distribution')
    if not assumptions.inference_confirmed:missing_fields.append('inference_confirmed')
    blocking=any(i['severity']=='blocking' for i in issues)
    check=dict(ready=not blocking,row_count=sheet.row_count,valid_count=n,excluded_count=excluded_count,groups=group_counts,
        issues=[i['message'] for i in issues if i['severity']=='blocking'],warnings=[i['message'] for i in issues if i['severity']=='warning'])
    return ModernPlanCheck(rules_version='analysis_plan_rules_v3',data_range=sheet.data_range,row_count=sheet.row_count,
        selected_fields=selected,selection_check=check,issues=issues,research_missing=missing_fields,
        limitations=['研究设计及分布前提由用户确认，程序检查不能证明独立性、正态性或差值对称性。',
            ('统计阶段采用当前整理保留行的完整记录或完整对，不新增插补、去重或自动剔除。' if getattr(sheet,'active_source_rows',None) is not None else '完整记录或完整对；不插补、不去重、不自动剔除异常候选。'),
            '差值方向为第一组减第二组；配对为前测减后测。',
            '两两95%区间是普通区间，未作同时校正；全部预声明组对采用Holm校正。']).model_dump(mode='json')
