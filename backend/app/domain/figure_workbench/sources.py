"""Reconstruct immutable prepared populations and typed projection fields."""
from decimal import Decimal
from copy import deepcopy
from app.domain.regression.bounded import compute_bounded
from app.domain.regression.common import NumericError
from openpyxl.utils.cell import get_column_letter,column_index_from_string
from app.core.exceptions import StorageError
from app.domain.datasets.contracts import rectangle
from app.domain.dataset_rows import source_rows
from app.domain.dataset_analysis import binary64,load_dataset
from app.domain.figure_workbench.kernel import compute_projection,ProjectionError
from app.domain.figure_workbench.material import legacy_material,projection_material
from app.domain.figure_workbench.contracts import FigureConfig
from app.domain.external_processing import digest


def load_source(store,project_id,file_id,run_id):
    result=store.analysis_result_by_id(project_id,file_id,run_id)
    from app.domain.analysis_plan.service import validate_plan_reference
    plan=validate_plan_reference(store,project_id,file_id,result)
    if not plan or result.get('source',{}).get('kind')!='prepared_dataset_version':
        raise StorageError('task_source_conflict','新图表需要来源完整的已确认分析计划和整理版本。',409)
    _,source,sheet=load_dataset(store,project_id,file_id,result['source']['dataset_id'],expected_source=result['source'])
    return result,sheet


def fields_for(sheet):
    return [{'column':get_column_letter(field.column),'name':field.name,'unit':field.unit.value or '', 'unit_status':field.unit.status,
        'type':','.join(sorted({c.kind for c in sheet.cells if c.column==field.column and c.kind!='blank'})) or 'blank'} for field in sheet.fields]


def projection_rows(sheet,config,*,included_rows=None):
    p=config['projection'] or {};kind=config['figure_kind'];fields={get_column_letter(f.column):f for f in sheet.fields}
    required={'line':['x_column','y_column'],'bar':['category_column']+(['y_column'] if p.get('bar_stat')=='mean' else []),'violin':['y_column'],'scatter':['group_column']}[kind]
    if any(not p.get(key) for key in required): raise StorageError('figure_projection_invalid','请补充图表必需字段。',422,params={'reason':'missing_fields'})
    chosen={role:p.get(role+'_column') for role in ('x','y','group','category') if p.get(role+'_column')}
    if any(letter not in fields for letter in chosen.values()): raise StorageError('invalid_column','投影字段不在已确认整理范围内。',422)
    cells={(c.row,c.column):c for c in sheet.cells};start,end=rectangle(sheet.data_range);rows=[]
    if end[0]-start[0]+1>10000: raise StorageError('plot_input_limit','图表投影最多支持一万条纳入记录。',422)
    for row in source_rows(sheet):
        if included_rows is not None and row not in included_rows:continue
        item={'row':row}
        for role,letter in chosen.items():
            cell=cells.get((row,column_index_from_string(letter)))
            if cell is None or cell.kind=='blank' or cell.kind=='text' and not cell.value.strip():value=None
            elif role=='y' or role=='x' and p.get('x_type','numeric')=='numeric':
                if cell.kind!='number': raise StorageError('numeric_type','数值字段包含非数值类型，不会自动转换。',422)
                value=binary64(Decimal(cell.value),letter,row,cell.value)
            elif role=='x' and p.get('x_type')=='datetime':
                if cell.kind not in ('date','datetime'):raise StorageError('task_input_invalid','时间X必须为实际日期/日期时间单元格。',422)
                value=cell.value
            else:
                if cell.kind not in ('text','number','boolean'):
                    if role=='group' and p.get('missing_group')=='exclude':item['_group_invalid']=True;item[role]=None;continue
                    raise StorageError('figure_projection_invalid','分类字段包含非法类型。',422,params={'reason':'invalid_group' if role=='group' else 'invalid_category'})
                value=binary64(Decimal(cell.value),letter,row,cell.value) if cell.kind=='number' else cell.value
            item[role]=value
        rows.append(item)
    return rows


def calculate(store,project_id,file_id,result,sheet,config,*,frozen=None,saved_expected=None,figure_contract=1):
    config=FigureConfig.model_validate(config).model_dump()
    from app.domain.artifact_contracts import artifact_language
    if config['language']!=artifact_language(result):raise StorageError('task_source_conflict','图表语言必须与已确认结果一致。',409)
    if config['figure_kind'] in ('line','bar','violin'):
        population,population_excluded=population_rows(sheet,result)
        rows=projection_rows(sheet,config,included_rows=population)
        try: computed=deepcopy(frozen) if frozen is not None else compute_bounded('figure_projection',rows,{'figure_kind':config['figure_kind'],'projection':config['projection']})
        except NumericError as exc:raise StorageError(exc.code,str(exc),422,params={'reason':getattr(exc,'reason','invalid_selection')}) from exc
        computed['audit'].update(source_n=sheet.row_count,population_excluded_n=len(population_excluded),population_excluded=population_excluded)
        payload,expected=projection_material(config,rows,computed,{f['column']:f for f in fields_for(sheet)},figure_contract=figure_contract)
    else:
        payload,expected=legacy_material(store,project_id,file_id,result,sheet,config,saved_expected=saved_expected,figure_contract=figure_contract)
        audit=result.get('row_audit',{'included_rows':[],'excluded_rows':[]})
        computed={'schema_version':1,'tool_version':'figure_projection_v1','method':{'kind':config['figure_kind'],'basis':'original_statistical_valid_rows'},
            'audit':{'population_n':result['check']['row_count'],'valid_n':result['check']['valid_count'],'excluded_n':result['check']['excluded_count'],'missing_y_n':0,'included_rows':audit.get('included_rows',[]),'excluded':audit.get('excluded_rows',[])},'series':[]}
        if config['figure_kind']=='scatter' and config['projection']:
            rows=projection_rows(sheet,config,included_rows=set(audit['included_rows']))
            if len(rows)!=result['statistics']['n'] or any(r.get('group') is None for r in rows):raise StorageError('task_input_invalid','散点显示分组必须覆盖原相关的每条有效记录，不允许删点。',422)
            groups=[r['group'] for r in rows]
            if len(set((type(g).__name__,str(g)) for g in groups))>20:raise StorageError('plot_input_limit','显示分组超过二十组。',422)
            payload['display_groups']=groups;expected['display_groups']=groups
    return computed,payload,expected

def population_rows(sheet,result):
    """Confirmed region and participating groups; statistical missingness is not population."""
    from app.domain.comparison.checks import typed_identity
    start,end=rectangle(sheet.data_range)
    area=result.get('analysis_context',{}).get('plan',{}).get('check',{}).get('data_range',sheet.data_range)
    low,high=rectangle(area)
    if low[0]<start[0] or high[0]>end[0] or low[1]<start[1] or high[1]>end[1]:raise source_conflict_population()
    selection=result['selection'];group_column=selection.get('group_column');groups=selection.get('groups',[])
    identities={(g['kind'],g['value']) for g in groups}
    cells={(c.row,c.column):c for c in sheet.cells};included=set();excluded=[]
    for row in source_rows(sheet):
        reason=None
        if not low[0]<=row<=high[0]:reason='outside_confirmed_range'
        elif group_column and identities:
            identity=typed_identity(cells.get((row,column_index_from_string(group_column))))
            if identity is None:reason='population_group_missing'
            elif identity not in identities:reason='population_group_not_selected'
        if reason:excluded.append({'row':row,'reason':reason})
        else:included.add(row)
    return included,excluded


def source_conflict_population():
    return StorageError('task_source_conflict','投影人群超出已确认整理范围。',409)
