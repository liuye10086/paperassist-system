"""Analysis over immutable raw or prepared datasets, without reopening Excel."""
from collections import Counter
from decimal import Decimal
from math import isfinite
from types import SimpleNamespace
import re
from openpyxl.utils.cell import get_column_letter, column_index_from_string
from app.adapters.excel import fail
from app.core.exceptions import StorageError
from app.domain.analysis_sources import raw_dataset_source, prepared_dataset_source, source_dict
from app.domain.datasets.service import DatasetService
from app.domain.analysis import SheetScan, ColumnProfile, SheetProfile, MAX_GROUPS
from app.domain.dataset_preparation.contracts import PreparedSheet


def load_dataset(store, project_id, file_id, dataset_id, *, expected_source=None):
    record = store.file(project_id, file_id)
    service = DatasetService(store)
    try:
        row, _, sheet = service.dataset(project_id, file_id, dataset_id)
    except StorageError as exc:
        if exc.code != 'dataset_missing' or exc.status != 404:
            raise
        from app.domain.dataset_preparation.service import DatasetPreparationService
        row, _, content = DatasetPreparationService(store).load_version(project_id, file_id, dataset_id)
        source = prepared_dataset_source(row, content.parent_source.model_dump())
        sheet = content.sheet
    else:
        source = raw_dataset_source(row, sheet, file_id)
    if record['sha256'] != source['source_sha256'] or (expected_source is not None and source_dict(expected_source) != source):
        raise StorageError('source_conflict', '分析数据来源已变化，请重新选择并保存数据版本。', 409)
    return record, source, sheet


def kind_value(cell):
    if cell is None or cell.kind == 'blank' or (cell.kind == 'text' and not cell.value.strip()):
        return 'missing', None
    if cell.kind in ('date', 'datetime', 'time', 'duration'):
        return 'datetime', None
    if cell.kind == 'number':
        return 'number', Decimal(cell.value)
    return cell.kind, cell.value


def binary64(value, column, row, lexical):
    if re.fullmatch(r'[+-]?[0-9]+', lexical) and value.copy_abs() > 2**53 - 1:
        fail('numeric_precision', f'数值列 {column}{row} 的整数超过安全精度范围（±9007199254740991）。', params={'column': column, 'row': row})
    number = float(value)
    if not isfinite(number) or (number == 0 and value != 0):
        fail('numeric_precision', f'数值列 {column}{row} 超出可表示的数值范围。', params={'column': column, 'row': row})
    return number


def decimal_label(value):
    if value == 0:
        return '0'
    # Do not use normalize(): it rounds through the current decimal context.
    sign, digits, exponent = value.as_tuple()
    digits = list(digits)
    while len(digits) > 1 and digits[-1] == 0:
        digits.pop()
        exponent += 1
    canonical = Decimal((sign, tuple(digits), exponent))
    return format(canonical, 'f') if -20 <= canonical.adjusted() <= 20 else str(canonical)


class DatasetScan(SheetScan):
    def observe_dataset(self, row_index, cells):
        while len(self.types) <= max(cells, default=-1):
            self.types.append(Counter())
            self.issue_cells.append([])
        for index, cell in cells.items():
            kind, value = kind_value(cell)
            if kind != 'missing':
                self.types[index][kind] += 1
                if kind != 'number' and len(self.issue_cells[index]) < 5:
                    self.issue_cells[index].append(f'{get_column_letter(index + 1)}{row_index}')
        if self.selection is None:
            return
        def selected(column):
            index = column_index_from_string(column) - 1
            return kind_value(cells.get(index))
        kind, value = selected(self.selection.numeric_column)
        if kind != 'number':
            return
        # Validate representability on check/save as well as run.
        numeric_index = column_index_from_string(self.selection.numeric_column) - 1
        number = binary64(value, self.selection.numeric_column, row_index, cells[numeric_index].value)
        if self.selection.group_column is not None:
            kind, group = selected(self.selection.group_column)
            if kind not in ('number', 'text', 'boolean'):
                return
            group = decimal_label(group) if kind == 'number' else group
            if group in self.groups or len(self.groups) <= MAX_GROUPS:
                self.groups[group] += 1
                if self.values is not None:
                    self.grouped_values.setdefault(group, []).append(number)
        self.valid_count += 1
        if self.values is not None:
            self.values.append(number)


def scan_dataset(sheet, selection=None, collect_values=False):
    if selection is not None and selection.sheet_name != sheet.name:
        fail('sheet_not_found', '所选数据版本与工作表不一致，请重新选择。')
    if isinstance(sheet, PreparedSheet):
        return scan_prepared(sheet, selection, collect_values)
    content = [c for c in sheet.cells if c.value not in (None, '')]
    last_row = max((c.row for c in content), default=0)
    width = max((c.column for c in content), default=0)
    cells = {(c.row, c.column): c for c in sheet.cells}
    headers = [cells.get((1, column)) for column in range(1, width + 1)]
    columns = [str(c.value) if c is not None and c.value not in (None, '') else f'未命名列 {get_column_letter(i)}'
        for i, c in enumerate(headers, 1)]
    warnings = ['第一行作为字段名；分析使用第2行至最后有值行的完整记录，不自动清洗；普通小数采用 binary64，可能存在浮点误差。']
    if any(c is None or c.value in (None, '') for c in headers):
        warnings.append('存在空表头，已用 Excel 列字母标记；未自动填充数据。')
    if len(set(columns)) < len(columns):
        warnings.append('存在重复列名，已按原始列顺序保留。')
    if sheet.merged_ranges:
        warnings.append('存在合并单元格，覆盖格不自动补值。')
    scan = DatasetScan(sheet.name, selection, collect_values)
    rows = {}
    for cell in sheet.cells:
        if 1 < cell.row <= last_row and cell.column <= width:
            rows.setdefault(cell.row, {})[cell.column - 1] = cell
    for row, stored_cells in rows.items():
        scan.observe_dataset(row, stored_cells)
    profile = scan.profile(SimpleNamespace(name=sheet.name, row_count=max(0, last_row-1), columns=columns, warnings=warnings))
    return profile, scan


def scan_prepared(sheet, selection, collect_values):
    scan = DatasetScan(sheet.name, selection, collect_values)
    rows = {}
    for cell in sheet.cells:
        rows.setdefault(cell.row, {})[cell.column - 1] = cell
    for row, cells in rows.items():
        scan.observe_dataset(row, cells)
    columns = []
    for field in sheet.fields:
        index = field.column - 1
        counts = scan.types[index] if index < len(scan.types) else Counter()
        kind = next(iter(counts)) if len(counts) == 1 else 'mixed' if counts else 'empty'
        columns.append(ColumnProfile(id=get_column_letter(field.column), name=field.name, type=kind,
            non_missing_count=sum(counts.values()), missing_count=sheet.row_count-sum(counts.values()),
            numeric_count=counts['number'], can_be_numeric=kind=='number', can_be_group=kind in ('number','text','boolean'),
            issue_cells=scan.issue_cells[index] if index < len(scan.issue_cells) else [], unit=field.unit))
    warnings = [f'分析使用已确认整理版本的 {sheet.data_range} 区域，保留区域内空行；普通小数采用 binary64，可能存在浮点误差。']
    if sheet.merged_ranges:
        warnings.append('存在合并单元格，覆盖格未自动补值。')
    if any(f.unit.status == 'unconfirmed' for f in sheet.fields):
        warnings.append('部分字段单位尚未确认；整理确认不代表数据真实性已核验。')
    return SheetProfile(sheet_name=sheet.name,row_count=sheet.row_count,columns=columns,warnings=warnings), scan
