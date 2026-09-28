from datetime import date, datetime, time, timedelta
from io import BytesIO
from math import isfinite
from posixpath import normpath
from typing import Annotated
from zipfile import BadZipFile, ZipFile

from defusedxml.ElementTree import fromstring, iterparse
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter
from pydantic import BaseModel

from .config import ExcelSettings, get_excel_settings

PREVIEW_ROWS = 20
CellValue = str | int | float | bool | None
router = APIRouter(prefix="/api/v1/excel", tags=["Excel 预览"])


class SheetPreview(BaseModel):
    name: str
    columns: list[str]
    row_count: int
    column_count: int
    preview_rows: list[list[CellValue]]
    warnings: list[str]


class WorkbookPreview(BaseModel):
    filename: str
    sheets: list[SheetPreview]


class PreviewConfig(BaseModel):
    max_upload_bytes: int
    preview_row_limit: int = PREVIEW_ROWS


def fail(code: str, message: str, status: int = 422):
    raise HTTPException(status_code=status, detail={"code": code, "message": message})


def too_large():
    fail("workbook_too_large", "工作簿超过解析规模限制，请减少工作表、行列或拆分文件后重试。")


def inspect_archive(stream: BytesIO, settings: ExcelSettings) -> bool:
    """Check decompressed size and XML coordinates before allocating cell rows."""
    try:
        archive = ZipFile(stream)
    except BadZipFile:
        fail("invalid_workbook", "文件不是有效的 .xlsx 工作簿，可能已损坏或加密。请在 Excel 中另存为 .xlsx。")
    with archive:
        entries = archive.infolist()
        names = set(archive.namelist())
        if not {"[Content_Types].xml", "xl/workbook.xml"}.issubset(names):
            fail("invalid_workbook", "文件内容不是有效的 .xlsx 工作簿，请勿只修改扩展名。")
        if len(entries) > 10000 or sum(entry.file_size for entry in entries) > settings.max_uncompressed_bytes:
            too_large()
        if any(entry.flag_bits & 1 for entry in entries):
            fail("invalid_workbook", "暂不支持加密工作簿，请先移除密码。")
        if (any(name.lower().endswith("vbaproject.bin") for name in names)
                or b"macroEnabled" in archive.read("[Content_Types].xml")):
            fail("invalid_workbook", "暂不支持带宏工作簿，请另存为不含宏的 .xlsx。")
        relationships = {
            element.attrib["Id"]: element.attrib
            for element in fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        }
        sheet_parts = []
        workbook_xml = fromstring(archive.read("xl/workbook.xml"))
        for sheet in workbook_xml.iter("{http://schemas.openxmlformats.org/spreadsheetml/2006/main}sheet"):
            relation_id = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
            relation = relationships[relation_id]
            relationship_prefix = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/"
            if relation["Type"] not in {relationship_prefix + "worksheet", relationship_prefix + "chartsheet"}:
                # openpyxl otherwise treats unknown non-chart types as worksheets.
                raise ValueError("Unsupported sheet relationship type")
            if relation.get("TargetMode") == "External":
                raise ValueError("External sheet is unsupported")
            target = relation["Target"]
            part = normpath(target.lstrip("/") if target.startswith("/") else "xl/" + target)
            if part not in names:
                raise ValueError("Missing sheet part")
            if relation["Type"].endswith("/worksheet"):
                sheet_parts.append(part)
        if len(sheet_parts) > settings.max_sheets:
            too_large()
        has_merges = False
        xml_cells = 0
        for part in sheet_parts:
            row_index = column_index = 0
            with archive.open(part) as sheet_xml:
                for event, element in iterparse(sheet_xml, events=("start", "end")):
                    tag = element.tag.rsplit("}", 1)[-1]
                    if event == "start" and tag == "row":
                        row_index = int(element.attrib.get("r", row_index + 1))
                        column_index = 0
                        if row_index < 1:
                            raise ValueError("Invalid row coordinate")
                        if row_index > settings.max_rows:
                            too_large()
                    if event == "start" and tag == "c":
                        xml_cells += 1
                        coordinate = element.attrib.get("r")
                        row, column = coordinate_to_tuple(coordinate) if coordinate else (row_index, column_index + 1)
                        column_index = column
                        if row < 1 or column < 1:
                            raise ValueError("Invalid cell coordinate")
                        if (row > settings.max_rows or column > settings.max_columns
                                or xml_cells > settings.max_cells):
                            too_large()
                    if tag == "mergeCell":
                        has_merges = True
                    if event == "end":
                        element.clear()
        return has_merges


def json_value(value) -> CellValue:
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    if isinstance(value, float) and not isfinite(value):
        return str(value)
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    # Array formulas have a text attribute; never execute a formula.
    return str(getattr(value, "text", value))


def parse_workbook(content: bytes, filename: str, settings: ExcelSettings) -> WorkbookPreview:
    workbook = None
    try:
        with BytesIO(content) as stream:
            has_merges = inspect_archive(stream, settings)
            stream.seek(0)
            workbook = load_workbook(stream, read_only=True, data_only=False, keep_links=False)
            if not workbook.worksheets:
                fail("invalid_workbook", "工作簿中没有可读取的数据工作表。")
            if len(workbook.worksheets) > settings.max_sheets:
                too_large()
            sheets = []
            scanned_cells = 0
            for sheet in workbook.worksheets:
                # Ignore incorrect dimensions written by some spreadsheet producers.
                sheet.reset_dimensions()
                headers = []
                preview = []
                last_row = width = 0
                has_formula = False
                for row_index, cells in enumerate(sheet.iter_rows(), start=1):
                    scanned_cells += len(cells)
                    if (row_index > settings.max_rows or len(cells) > settings.max_columns
                            or scanned_cells > settings.max_cells):
                        too_large()
                    values = [json_value(cell.value) for cell in cells]
                    has_formula |= any(cell.data_type == "f" for cell in cells)
                    used = [index for index, value in enumerate(values, start=1) if value is not None and value != ""]
                    if used:
                        last_row = row_index
                        width = max(width, used[-1])
                    if row_index == 1:
                        headers = values
                    elif row_index <= PREVIEW_ROWS + 1:
                        preview.append(values)
                row_count = max(0, last_row - 1)
                headers = (headers + [None] * width)[:width]
                notes = []
                columns = [str(value) if value is not None and value != "" else f"未命名列 {get_column_letter(index)}"
                           for index, value in enumerate(headers, start=1)]
                if any(value is None or value == "" for value in headers):
                    notes.append("存在空表头，已用 Excel 列字母标记；未自动填充数据。")
                if len(set(columns)) < len(columns):
                    notes.append("存在重复列名，已按原始列顺序保留。")
                if not width:
                    notes.append("这是空表，没有可预览的数据。")
                elif not row_count:
                    notes.append("此工作表只有表头，没有数据行。")
                if has_formula:
                    notes.append("包含公式：仅展示公式原文，未执行计算；如需数值，请上传经核对的数值副本。")
                if has_merges:
                    notes.append("工作簿包含合并单元格，预览仅保留实际存储值，不自动填充或还原合并布局。")
                if sheet.sheet_state != "visible":
                    notes.append("这是隐藏工作表，预览仍包含其数据。")
                sheets.append(SheetPreview(
                    name=sheet.title, columns=columns, row_count=row_count, column_count=width,
                    preview_rows=[(row + [None] * width)[:width] for row in preview[:row_count]],
                    warnings=notes,
                ))
            return WorkbookPreview(filename=filename, sheets=sheets)
    except HTTPException:
        raise
    except Exception:
        # Third-party parser exceptions must not expose file content or host paths.
        fail("parse_failed", "Excel 解析失败，文件可能已损坏或包含不支持的内容。请用 Excel 打开并另存为 .xlsx 后重试。")
    finally:
        if workbook is not None:
            workbook.close()


@router.get("/config", response_model=PreviewConfig)
def preview_config(settings: Annotated[ExcelSettings, Depends(get_excel_settings)]):
    return PreviewConfig(max_upload_bytes=settings.max_upload_bytes)


@router.post("/preview", response_model=WorkbookPreview)
def preview_excel(
    settings: Annotated[ExcelSettings, Depends(get_excel_settings)],
    file: Annotated[UploadFile | None, File(description="一个 .xlsx 工作簿")] = None,
):
    if file is None:
        fail("missing_file", "请选择一个 .xlsx 文件后上传。")
    try:
        filename = (file.filename or "").replace("\\", "/").rsplit("/", 1)[-1]
        if not filename.lower().endswith(".xlsx"):
            fail("unsupported_file_type", "仅支持 .xlsx 文件，暂不支持 .xls、.xlsm 或其他格式。", 415)
        content = file.file.read(settings.max_upload_bytes + 1)
        if len(content) > settings.max_upload_bytes:
            fail("file_too_large", f"文件超过大小限制（{settings.max_upload_bytes} 字节），请缩小文件后重试。", 413)
        if not content:
            fail("empty_file", "上传的文件为空，请选择包含工作簿内容的 .xlsx 文件。")
        return parse_workbook(content, filename, settings)
    finally:
        file.file.close()
