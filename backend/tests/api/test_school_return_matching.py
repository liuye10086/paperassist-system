from copy import deepcopy
from io import BytesIO
from docx import Document
from app.adapters.manuscript_docx import _paragraph
from app.domain.manuscripts.school_return import extract,apply_return
from zipfile import ZipFile,ZIP_DEFLATED
from lxml import etree


def package_change(content,mutate):
    with ZipFile(BytesIO(content)) as package:parts={n:package.read(n) for n in package.namelist()}
    mutate(parts);stream=BytesIO()
    with ZipFile(stream,'w',ZIP_DEFLATED) as package:
        for name,data in parts.items():package.writestr(name,data)
    return stream.getvalue()


def snapshot():
    def section(identity,title):return dict(id=identity,title=title,purpose='other',paragraphs=[dict(id='p:'+identity,text='合成普通文字。',origin='user_provided',source_ids=[],citations=[],numeric_bindings=[])])
    return dict(sections=[section('a','同名'),section('b','同名')])


def raw(doc):
    stream=BytesIO();doc.save(stream);return stream.getvalue()


def test_bookmark_priority_duplicate_title_and_no_bookmark_ambiguity():
    base=snapshot();doc=Document();doc.add_heading('同名',1);_paragraph(doc,base['sections'][1]['paragraphs'][0],1)
    sections=extract(raw(doc),'return.docx',base)['sections']
    assert sections[0]['suggested_chapter_id']=='b' and sections[0]['match_method']=='bookmark'
    plain=Document();plain.add_heading('同名',1);plain.add_paragraph('回传正文。')
    ambiguous=extract(raw(plain),'plain.docx',base)['sections'][0]
    assert ambiguous['match_method']=='ambiguous' and ambiguous['suggested_chapter_id'] is None


def test_unknown_front_matter_and_changed_bound_claim_not_dropped():
    base=snapshot();doc=Document();doc.add_paragraph('未分配额外内容。');doc.add_heading('新章',1);doc.add_paragraph('另一段。')
    result=extract(raw(doc),'return.docx',base)
    assert result['sections'][0]['paragraphs']==['未分配额外内容。']
    base['sections']=base['sections'][:1];base['sections'][0]['paragraphs'][0]['citations']=[dict(review_id='fixed')]
    incoming=[dict(id='new',paragraphs=['改变论述。'],paragraph_items=[dict(text='改变论述。',original_id='p:a')])]
    candidate,_,issues=apply_return(base,incoming,[dict(input_section_id='new',chapter_id='a')])
    assert any(p['code']=='bound_claim_changed' for p in issues)
    assert candidate['sections'][0]['paragraphs'][0]['citations']==[]


def test_fixed_export_decorations_and_unchanged_images_allowed():
    base=snapshot();base['sections']=base['sections'][:1]
    doc=Document();doc.add_paragraph('封面人工信息');doc.sections[0].header.paragraphs[0].text='原页眉'
    doc.add_heading('同名',1);_paragraph(doc,base['sections'][0]['paragraphs'][0],1)
    baseline=raw(doc);doc.paragraphs[-1].runs[0].text+='附加普通说明。'
    result=extract(raw(doc),'return.docx',base,baseline)
    assert not any(i['severity']=='error' for i in result['issues']),result['issues']
    assert len(result['sections'])==1
    assert result['retained_nonbody'][0]['text']=='封面人工信息'


def test_unchanged_exported_citation_rendering_restores_exact_original_support():
    base=snapshot();base['sections']=base['sections'][:1]
    base['sections'][0]['paragraphs'][0]['text']='已核实合成论述。[1]'
    base['sections'][0]['paragraphs'][0]['citations']=[dict(review_id='fixed')]
    doc=Document();doc.add_heading('同名',1)
    displayed=deepcopy(base['sections'][0]['paragraphs'][0]);displayed['text']='已核实合成论述。'
    _paragraph(doc,displayed,1);baseline=raw(doc)
    result=extract(baseline,'return.docx',base,baseline)
    candidate,differences,issues=apply_return(base,result['sections'],[dict(input_section_id=result['sections'][0]['id'],chapter_id='a')])
    assert not issues and not differences
    assert candidate==base


def test_changed_footnote_target_or_field_does_not_restore_old_support():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    base=snapshot();base['sections']=base['sections'][:1]
    base['sections'][0]['paragraphs'][0].update(text='原已核实论述。[1]',citations=[dict(review_id='fixed')])
    for marker in ('footnoteReference','endnoteReference','instrText'):
        doc=Document();doc.add_heading('同名',1);_paragraph(doc,dict(base['sections'][0]['paragraphs'][0],text='原已核实论述。'),1)
        node=OxmlElement('w:'+marker)
        if marker=='instrText':node.text=' REF old_target '
        else:node.set(qn('w:id'),'1')
        doc.paragraphs[-1].add_run()._r.append(node);baseline=raw(doc)
        def mutate(parts):
            root=etree.fromstring(parts['word/document.xml']);found=root.find('.//'+qn('w:'+marker))
            if marker=='instrText':found.text=' REF changed_target '
            else:found.set(qn('w:id'),'2')
            parts['word/document.xml']=etree.tostring(root)
        result=extract(package_change(baseline,mutate),'return.docx',base,baseline)
        assert any(p['severity']=='error' and p['code']=='citation_changed' for p in result['issues']),marker
        assert result['sections'][0]['paragraphs']!=['原已核实论述。[1]']


def test_no_bookmark_title_fallback_cannot_hide_changed_note_target():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    base=snapshot();base['sections']=base['sections'][:1]
    base['sections'][0]['paragraphs'][0]['citations']=[dict(review_id='fixed')]
    doc=Document();doc.add_heading('同名',1);p=doc.add_paragraph(base['sections'][0]['paragraphs'][0]['text'])
    node=OxmlElement('w:footnoteReference');node.set(qn('w:id'),'1');p.add_run()._r.append(node)
    baseline=raw(doc);node.set(qn('w:id'),'2')
    result=extract(raw(doc),'return.docx',base,baseline)
    assert any(i['code']=='citation_changed' and i['severity']=='error' for i in result['issues'])


def test_inserted_prose_before_unchanged_caption_uses_stable_chapter_match():
    base=snapshot();base['sections']=base['sections'][:1]
    doc=Document();doc.add_heading('同名',1);_paragraph(doc,base['sections'][0]['paragraphs'][0],1)
    caption=doc.add_paragraph('Figure 1. Frozen supplied caption')
    baseline=raw(doc);caption.insert_paragraph_before('新增普通段落。')
    result=extract(raw(doc),'return.docx',base,baseline)
    assert not any(i['severity']=='error' for i in result['issues']),result['issues']
    assert result['sections'][0]['paragraphs']==['合成普通文字。','新增普通段落。']
    caption.runs[0].text='Figure 1. Changed scientific caption'
    changed=extract(raw(doc),'return.docx',base,baseline)
    assert any(i['code']=='nonbody_changed' and i['severity']=='error' for i in changed['issues'])


def test_actual_word_saveas_only_changes_requested_prose():
    import json
    from pathlib import Path
    import pytest
    folder=Path(__file__).resolve().parents[3]/'script/.runtime/stage07-batch04'
    if not (folder/'return-main.docx').exists():pytest.skip('Local actual Word evidence is not checked into Git')
    base=json.loads((folder/'migration-adopted.json').read_text(encoding='utf-8-sig'))['versions'][0]['snapshot']
    baseline=(folder/'migration-v4.docx').read_bytes();returned=(folder/'return-main.docx').read_bytes()
    result=extract(returned,'return-main.docx',base,baseline)
    assert not [i for i in result['issues'] if i['severity']=='error'],result['issues']
    mappings=[dict(input_section_id=s['id'],chapter_id=s['suggested_chapter_id']) for s in result['sections']]
    candidate,differences,issues=apply_return(base,result['sections'],mappings)
    assert not issues,issues
    assert len(differences)==1 and differences[0]['kind']=='text_changed'
    assert differences[0]['after']==differences[0]['before'].replace('工程操作说明','回传补充说明',1)
    for part,tag,attribute in [('word/footnotes.xml','t',None),('word/document.xml','footnoteReference','id')]:
        def mutate(parts):
            from docx.oxml.ns import qn
            root=etree.fromstring(parts[part]);node=root.find('.//'+qn('w:'+tag))
            if attribute:node.set(qn('w:'+attribute),'999')
            else:node.text+=' changed citation content'
            parts[part]=etree.tostring(root)
        changed=extract(package_change(returned,mutate),'changed.docx',base,baseline)
        assert any(i['severity']=='error' and i['code']=='citation_changed' for i in changed['issues'])


def test_word_saveas_field_cache_and_part_serialization_equivalence():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    base=snapshot()
    doc=Document();doc.sections[0].header.paragraphs[0].text='Frozen school header'
    footer=doc.sections[0].footer.paragraphs[0]
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),' PAGE ')
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';run.append(text);field.append(run);footer._p.append(field)
    def field_run(parent,kind=None,instr=None):
        run=OxmlElement('w:r');node=OxmlElement('w:fldChar' if kind else 'w:instrText')
        if kind:node.set(qn('w:fldCharType'),kind)
        else:node.text=instr
        run.append(node);parent.append(run)
    start=doc.add_paragraph();field_run(start._p,'begin');field_run(start._p,instr=' TOC \\o "1-3" \\h \\z \\u ');field_run(start._p,'separate')
    entry=doc.add_paragraph();link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'PATOCfixed')
    for value in ('同名',None,'1'):
        run=OxmlElement('w:r');node=OxmlElement('w:t' if value else 'w:tab')
        if value:node.text=value
        run.append(node);link.append(run)
    entry._p.append(link);field_run(entry._p,'end')
    heading=doc.add_heading('同名',1);mark=OxmlElement('w:bookmarkStart');mark.set(qn('w:id'),'900');mark.set(qn('w:name'),'PATOCfixed');heading._p.append(mark)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),'900');heading._p.append(end)
    stable=OxmlElement('w:bookmarkStart');stable.set(qn('w:id'),'901');stable.set(qn('w:name'),'PA00000001');heading._p.append(stable)
    stable_end=OxmlElement('w:bookmarkEnd');stable_end.set(qn('w:id'),'901');heading._p.append(stable_end)
    _paragraph(doc,base['sections'][0]['paragraphs'][0],1)
    other=doc.add_heading('同名',1)
    for identity,name in [('902','PATOCother'),('903','PA00000002')]:
        mark=OxmlElement('w:bookmarkStart');mark.set(qn('w:id'),identity);mark.set(qn('w:name'),name);other._p.append(mark)
        end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),identity);other._p.append(end)
    _paragraph(doc,base['sections'][1]['paragraphs'][0],2)
    baseline=raw(doc)
    def saveas(parts):
        # Equivalent field representation, page cache refresh and run formatting.
        footer=etree.fromstring(parts['word/footer1.xml']);p=footer.find('.//'+qn('w:p'));p.remove(p.find(qn('w:fldSimple')))
        field_run(p,'begin');field_run(p,instr=' PAGE  \\* MERGEFORMAT ');field_run(p,'separate')
        run=OxmlElement('w:r');node=OxmlElement('w:t');node.text='9';run.append(node);p.append(run);field_run(p,'end')
        parts['word/footer1.xml']=etree.tostring(footer)
        header=etree.fromstring(parts.pop('word/header1.xml'));header.find('.//'+qn('w:p')).set(qn('w:rsidR'),'12345')
        parts['word/header9.xml']=etree.tostring(header)
        rels=etree.fromstring(parts['word/_rels/document.xml.rels'])
        for rel in rels:
            if rel.get('Target')=='header1.xml':rel.set('Target','header9.xml')
        parts['word/_rels/document.xml.rels']=etree.tostring(rels)
        root=etree.fromstring(parts['word/document.xml']);paras=root.findall('.//'+qn('w:p'))
        link=paras[1].find(qn('w:hyperlink'));paras[1].remove(link);paras[0].append(link)
        last=link[-1];link.remove(last);field_run(link,'begin');field_run(link,instr=' PAGEREF PATOCfixed \\h ');field_run(link,'separate');last[0].text='9';link.append(last);field_run(link,'end')
        parts['word/document.xml']=etree.tostring(root)
    returned=package_change(baseline,saveas)
    result=extract(returned,'return.docx',base,baseline)
    assert not [i for i in result['issues'] if i['severity']=='error'],result['issues']
    assert result['sections'][0]['paragraphs']==['合成普通文字。']
    assert any(i['code']=='export_cache_normalized' for i in result['issues'])
    for mutation in ('header','toc_title','toc_target','toc_unknown','toc_same_title_target'):
        def change(parts):
            name='word/header9.xml' if mutation=='header' else 'word/document.xml'
            root=etree.fromstring(parts[name]);link=root.find('.//'+qn('w:hyperlink'))
            if mutation=='header':root.find('.//'+qn('w:t')).text='Changed header'
            elif mutation=='toc_title':link.find('.//'+qn('w:t')).text='Changed title'
            elif mutation=='toc_target':link.set(qn('w:anchor'),'missing')
            elif mutation=='toc_same_title_target':
                link.set(qn('w:anchor'),'PATOCother')
                link.find('.//'+qn('w:instrText')).text=' PAGEREF PATOCother \\h '
            else:link.append(OxmlElement('w:object'))
            parts[name]=etree.tostring(root)
        changed=extract(package_change(returned,change),'changed.docx',base,baseline)
        assert any(i['severity']=='error' for i in changed['issues']),mutation
