"""Focused synthetic pipeline tests; mock receipts are never real CPU evidence."""
from tests.api.test_analysis import client
from tests.api.test_review_rounds import make_round


def upload_in_project(instance, project, rows):
    from io import BytesIO
    from openpyxl import Workbook
    book = Workbook()
    book.active.title = '研究 & 数据'
    for row in rows:
        book.active.append(row)
    raw = BytesIO()
    book.save(raw)
    book.close()
    response = instance.post(f'/api/v1/projects/{project}/files', files={'file': ('synthetic-old-data.xlsx', raw.getvalue())})
    assert response.status_code == 201, response.text
    file = response.json()['file']
    return f'/api/v1/projects/{project}/files/{file["id"]}', file


def mock_local_summary(values, **kwargs):
    """Synthetic pipeline receipt only; actual Linux calculations belong to root smoke."""
    from app.domain.fixed_statistics_runtime import record_execution
    from app.domain.statistics_execution.service import RESOURCE_CONTRACT
    record_execution(dict(contract_version='fixed_statistics_v2', sandbox_version=RESOURCE_CONTRACT['sandbox_version'],
        input_bytes=30, output_bytes=300, timeout_seconds=30, elapsed_seconds=0.01, cache_hit=False,
        stdout_policy='bounded_protocol', stderr_policy='discard', input_policy='frozen_projection', output_policy='bounded_pipe',
        memory_hard_limit=1073741824, isolation='linux_landlock_seccomp', cpu_hard_limit_seconds=30, cpu_affinity_count=1,
        child_cpu_seconds=0.01, child_peak_rss_bytes=1024, error_code=None))
    return dict(n=len(values), mean=5.0, std=2.0, min=1.0, q1=3.0, median=5.0, q3=7.0, max=9.0, iqr=4.0, warnings=[])


def queued_analysis(client, monkeypatch):
    from tests.api.test_manuscript_review_v2_bridge import published
    from tests.helpers.analysis_plan import confirmed_sample
    from app.domain.tasks.execution import run_task
    from app.domain.dataset_profile.scan import numeric_profile
    # Isolated pipeline substitute for the Linux-only plan profile boundary.
    monkeypatch.setattr('app.domain.analysis_plan.profile.bounded_numeric_profile', numeric_profile)
    file_path, file, _, _ = confirmed_sample(client)
    project = file_path.split('/')[4]
    with monkeypatch.context() as fixture:
        fixture.setattr('tests.api.test_descriptive.add_file', lambda instance, rows: upload_in_project(instance, project, rows))
        root, base, _, _ = published(client)
    url = root.rsplit('/manuscripts/', 1)[0] + '/reviews'
    response = client.post(url + '/rounds', json=dict(title='Synthetic analysis revision', project_type='sci', synthetic=True,
        manuscript=dict(id=root.rsplit('/', 1)[1], revision=base['revision'], version_id=base['id']),
        sources=[dict(role='reviewer', label='Reviewer 1', format='text', text='1. Please update statistics.\n2. Clarify the methods.')]),
        headers={'Idempotency-Key': 'science-round'})
    assert response.status_code == 202, response.text
    assert run_task(response.json()['task']['id'], 1)['status'] == 'succeeded'
    round = client.get(url + '/rounds/' + response.json()['id']).json()
    path = url + '/rounds/' + round['id']
    body = dict(plan_revision=round['revision'], item_id=round['items'][0]['id'], file_id=file['id'], expected_setup_revision=1)
    response = client.post(path + '/statistics-tasks', json=body, headers={'Idempotency-Key': 'science-statistics'})
    assert response.status_code == 202, response.text
    return url, path, round, base, file_path, response.json(), body


def adoption_body(round, base, link):
    return dict(plan_revision=round['revision'], item_id=round['items'][0]['id'], task_link_id=link['id'],
                supersedes_source_ids=[base['snapshot']['sources'][0]['id']],
                affected_locations=[dict(kind='paragraph', id='results:1'), dict(kind='paragraph', id='abstract:1')],
                reason='明确采用新确认数据的统计结果；原分析仅保留历史。')


def test_confirmed_local_statistics_queue_and_receipt_adoption_public_producer(client, monkeypatch):
    from app.domain.tasks.execution import run_task
    url, path, round, base, file_path, link, request = queued_analysis(client, monkeypatch)
    assert link['task']['task_type'] == 'statistics_compute'
    assert link['task']['status'] == 'queued' and not link['result']['available']
    assert client.post(path + '/statistics-tasks', json=request, headers={'Idempotency-Key': 'science-statistics'}).json()['id'] == link['id']
    body = adoption_body(round, base, link)
    assert client.post(path + '/analysis-adoptions', json=body, headers={'Idempotency-Key': 'too-early'}).status_code == 409
    monkeypatch.setattr('app.domain.descriptive.summarize', mock_local_summary)
    succeeded = run_task(link['task']['id'], 1)
    assert succeeded['status'] == 'succeeded', succeeded
    response = client.post(path + '/analysis-adoptions', json=body, headers={'Idempotency-Key': 'adopt-analysis'})
    assert response.status_code == 201, response.text
    adopted = response.json()
    assert adopted['source']['id'] == succeeded['result_analysis_run_id']
    assert adopted['input_snapshot']['execution_receipt'] == succeeded['execution_receipt']
    assert adopted['pending'] and {x['section_id'] for x in adopted['impact']} == {'abstract', 'results', 'discussion', 'captions'}
    assert client.post(path + '/analysis-adoptions', json=body, headers={'Idempotency-Key': 'adopt-analysis'}).json()['id'] == adopted['id']
    assert client.post(path + '/analysis-adoptions', json=body, headers={'Idempotency-Key': 'adopt-same-result'}).status_code == 409
    from tests.api.test_review_rounds import edits
    unchanged_revision = round['items'][0]['revision']
    changes = edits(round)
    changes[1]['plan'] += ' Preserve the existing method and clarify its wording.'
    saved_plan = client.post(path + '/plans', json=dict(expected_revision=round['revision'], items=changes), headers={'Idempotency-Key': 'unrelated-method-plan'})
    assert saved_plan.status_code == 201, saved_plan.text
    round = saved_plan.json()
    assert round['items'][0]['revision'] == unchanged_revision
    context = client.get(path + '/revision-context').json()
    assert adopted['source']['id'] in {e['id'] for e in context['grounded_plans'][0]['evidence']}, 'New plan must read the actually adopted result'
    facts = {f['path']: f for f in context['analysis_facts']}
    assert '/overall/n' in facts, 'Unrelated item changes must retain current receipt-backed analysis facts'
    assert facts['/overall/n']['value'] != '8'
    bindings = [dict(fact_id=facts[p]['id'], value=facts[p]['value']) for p in ('/overall/n', '/overall/mean')]
    text = 'n = ' + bindings[0]['value'] + '; mean = ' + bindings[1]['value']
    operations = [dict(item_id=round['items'][0]['id'], kind='replace_paragraph', location_id=pid,
        text=text, evidence_ids=[adopted['source']['id']], numeric_bindings=bindings, reason='采用新的逐值统计。') for pid in ('results:1', 'abstract:1')]
    data = dict(plan_revision=round['revision'], expected_current_revision=context['expected_current_revision'],
                analysis_adoption_ids=[adopted['id']], operations=operations)
    preview = client.post(path + '/revisions/preview', json=data)
    assert preview.status_code == 200, preview.text
    assert preview.json()['can_adopt'], preview.json()
    adopted_run = client.post(path + '/revisions', json=data | dict(preview_sha256=preview.json()['preview_sha256']), headers={'Idempotency-Key': 'publish-new-results'})
    assert adopted_run.status_code == 202, adopted_run.text
    run = adopted_run.json()
    assert run_task(run['task']['id'], 1)['status'] == 'succeeded'
    result = client.get(url + '/revisions/' + run['id']).json()['result']
    assert run_task(result['output']['task_id'], 1)['status'] == 'succeeded'
    manuscript = url.rsplit('/reviews', 1)[0] + '/manuscripts/' + result['output']['manuscript_id']
    output = client.get(manuscript + '/versions/' + str(result['output']['revision']))
    assert output.status_code == 200, output.text
    snapshot = output.json()['snapshot']
    assert any(s['id'] == adopted['source']['id'] for s in snapshot['sources'])
    assert snapshot['review_analysis_provenance'][0]['adoption_id'] == adopted['id']
    assert client.get(manuscript + '/versions/' + str(base['revision'])).json()['snapshot'] == base['snapshot']
    checked = client.get(manuscript + '/reviews/context')
    assert checked.status_code == 200, checked.text
    assert not [i for i in checked.json()['issues'] if i['severity'] == 'error'], checked.json()['issues']


def test_science_links_public_task_with_frozen_item_and_live_status(client):
    url, round, _, base = make_round(client)
    path = url + '/rounds/' + round['id']
    response = client.get(path + '/science')
    assert response.status_code == 200, response.text
    assert response.json()['task_links'] == []
    body = dict(plan_revision=round['revision'], item_id=round['items'][0]['id'],
                task_id=base['task_id'], purpose='writing')
    response = client.post(path + '/task-links', json=body, headers={'Idempotency-Key': 'science-link'})
    assert response.status_code == 201, response.text
    link = response.json()
    assert link['task']['id'] == base['task_id']
    assert link['task']['status'] == 'succeeded'
    assert link['input_snapshot']['manuscript_version_id'] == base['id']
    assert link['input_snapshot']['task_input_digest']
    assert link['result'].get('artifacts', {}).get('manuscript', {}).get('version_id') == base['id']
    assert client.post(path + '/task-links', json=body, headers={'Idempotency-Key': 'science-link'}).status_code == 200
    current = client.get(path).json()
    assert current['items'][0]['status'] != 'completed'
    assert len(client.get(path + '/science').json()['task_links']) == 1
    from app.maintenance import review_science_audit as audit
    assert hasattr(audit, 'audit_review_science'), 'Frozen analysis receipt and task-link audit is missing'
    from app.db.database import database_connection
    with database_connection() as db:
        assert all(check['invalid'] == 0 for check in audit.audit_review_science(db))
    writing_body = body | dict(task_id=base['snapshot']['writing_provenance']['request_id'])
    writing = client.post(path + '/task-links', json=writing_body, headers={'Idempotency-Key': 'science-writing-candidate'})
    assert writing.status_code == 201, writing.text
    assert writing.json()['result'].get('artifacts', {}).get('writing', {}).get('candidate_available') is True


def test_task_link_rejects_foreign_project_wrong_type_item_and_stale_plan(client):
    from tests.api.test_projects import create_project
    url, round, _, base = make_round(client)
    path = url + '/rounds/' + round['id'] + '/task-links'
    body = dict(plan_revision=round['revision'], item_id=round['items'][0]['id'], task_id=base['task_id'], purpose='writing')
    other = create_project(client)
    foreign = f'/api/v1/projects/{other["id"]}/reviews/rounds/{round["id"]}/task-links'
    assert client.post(foreign, json=body, headers={'Idempotency-Key': 'foreign-project'}).status_code == 404
    assert client.post(path, json=body | dict(item_id='foreign-item'), headers={'Idempotency-Key': 'foreign-item'}).status_code == 422
    assert client.post(path, json=body | dict(purpose='statistics'), headers={'Idempotency-Key': 'wrong-task-type'}).status_code == 422
    assert client.post(path, json=body | dict(plan_revision=round['revision']+1), headers={'Idempotency-Key': 'stale-plan'}).status_code == 409


def test_adoption_rejects_fake_receipt_foreign_source_and_stale_setup(client, monkeypatch):
    from copy import deepcopy
    from psycopg.types.json import Jsonb
    from app.db.database import database_connection
    from app.domain.tasks.execution import run_task
    import json
    url, path, round, base, file_path, link, request = queued_analysis(client, monkeypatch)
    monkeypatch.setattr('app.domain.descriptive.summarize', mock_local_summary)
    task = run_task(link['task']['id'], 1)
    assert task['status'] == 'succeeded', task
    body = adoption_body(round, base, link)
    wrong = client.post(path + '/analysis-adoptions', json=body | dict(supersedes_source_ids=['foreign-analysis']), headers={'Idempotency-Key': 'wrong-source'})
    assert wrong.status_code == 422, wrong.text
    fake = deepcopy(task['execution_receipt'])
    fake['result_sha256'] = '0'*64
    with database_connection(write=True) as db:
        db.execute('UPDATE tasks SET execution_receipt=%s WHERE id=%s', (Jsonb(fake), task['id']))
    rejected = client.post(path + '/analysis-adoptions', json=body, headers={'Idempotency-Key': 'fake-receipt'})
    assert rejected.status_code == 409, rejected.text
    with database_connection(write=True) as db:
        db.execute('UPDATE tasks SET execution_receipt=%s WHERE id=%s', (Jsonb(task['execution_receipt']), task['id']))
        saved = db.execute('SELECT setup_json FROM analysis_setups WHERE file_id=%s', (request['file_id'],)).fetchone()
        setup = json.loads(saved['setup_json'])
        setup['revision'] += 1
        db.execute('UPDATE analysis_setups SET revision=%s,setup_json=%s WHERE file_id=%s', (setup['revision'], json.dumps(setup), request['file_id']))
    rejected = client.post(path + '/analysis-adoptions', json=body, headers={'Idempotency-Key': 'stale-setup'})
    assert rejected.status_code == 409, rejected.text
    assert client.get(path + '/science').json()['analysis_adoptions'] == []


def test_same_file_reanalysis_can_publish_only_explicitly_superseded_old_setup(client, monkeypatch):
    from tests.api.test_dataset_analysis import version
    from tests.integration.test_dataset_preparation import run, review, decide, decision
    from tests.helpers.legacy_plans import seed_legacy_plan
    from tests.helpers.analysis_plan import run_plan, review_plan, decide_plan, plan_decision
    from app.domain.dataset_profile.scan import numeric_profile
    from app.domain.tasks.execution import run_task
    monkeypatch.setattr('app.domain.analysis_plan.profile.bounded_numeric_profile', numeric_profile)
    monkeypatch.setattr('app.domain.descriptive.summarize', mock_local_summary)
    url, round, _, base = make_round(client, '1. Please update statistics.')
    source = base['snapshot']['sources'][0]
    file_path = url.rsplit('/reviews', 1)[0] + '/files/' + source['file_id']
    old_setup = client.get(file_path + '/analysis-setup').json()
    raw = version(client, file_path)
    defaults = client.get(file_path + '/dataset-versions/' + raw['id'] + '/preparation-defaults', params=dict(header_row=1, data_range='A2:B5')).json()
    created = client.post(file_path + '/dataset-preparations', json=dict(dataset_id=raw['id'], recipe=defaults['recipe']), headers={'Idempotency-Key': 'same-file-preparation'})
    assert created.status_code == 202, created.text
    task = created.json()
    assert run(task)['status'] == 'waiting_confirmation'
    confirmed = decide(client, task, decision(review(client, task)), 'same-file-confirm').json()
    assert run(confirmed)['status'] == 'succeeded'
    prepared = review(client, task)['result']
    planned = seed_legacy_plan(client, file_path + '/analysis-plans', json=dict(dataset_id=prepared['id'], expected_setup_revision=old_setup['revision'], selection=old_setup['selection']), headers={'Idempotency-Key': 'same-file-plan'})
    assert planned.status_code == 202, planned.text
    plan_task = planned.json()
    assert run_plan(plan_task)['status'] == 'waiting_confirmation'
    accepted = decide_plan(client, plan_task, plan_decision(review_plan(client, plan_task)), 'same-file-plan-confirm').json()
    assert run_plan(accepted)['status'] == 'succeeded'
    path = url + '/rounds/' + round['id']
    queued = client.post(path + '/statistics-tasks', json=dict(plan_revision=round['revision'], item_id=round['items'][0]['id'], file_id=source['file_id'], expected_setup_revision=old_setup['revision']+1), headers={'Idempotency-Key': 'same-file-statistics'})
    assert queued.status_code == 202, queued.text
    link = queued.json()
    assert run_task(link['task']['id'], 1)['status'] == 'succeeded'
    saved = client.post(path + '/analysis-adoptions', json=adoption_body(round, base, link), headers={'Idempotency-Key': 'same-file-adopt'})
    assert saved.status_code == 201, saved.text
    adopted = saved.json()
    context = client.get(path + '/revision-context').json()
    facts = {f['path']: f for f in context['analysis_facts']}
    bindings = [dict(fact_id=facts[p]['id'], value=facts[p]['value']) for p in ('/overall/n', '/overall/mean')]
    operations = [dict(item_id=round['items'][0]['id'], kind='replace_paragraph', location_id=pid, text='n = '+bindings[0]['value']+'; mean = '+bindings[1]['value'], numeric_bindings=bindings, evidence_ids=[adopted['source']['id']], reason='Same file explicitly reanalysed under new confirmed setup.') for pid in ('results:1', 'abstract:1')]
    body = dict(plan_revision=round['revision'], expected_current_revision=context['expected_current_revision'], analysis_adoption_ids=[adopted['id']], operations=operations)
    preview = client.post(path + '/revisions/preview', json=body)
    assert preview.status_code == 200 and preview.json()['can_adopt'], preview.text
    saved = client.post(path + '/revisions', json=body | dict(preview_sha256=preview.json()['preview_sha256']), headers={'Idempotency-Key': 'same-file-publish'})
    assert saved.status_code == 202, saved.text
    result = run_task(saved.json()['task']['id'], 1)
    assert result['status'] == 'succeeded', result
    output = client.get(url + '/revisions/' + saved.json()['id']).json()['result']['output']
    assert run_task(output['task_id'], 1)['status'] == 'succeeded'
    public = client.get(url.rsplit('/reviews', 1)[0] + '/manuscripts/' + output['manuscript_id'] + '/versions/' + str(output['revision']))
    assert public.status_code == 200, public.text


def test_explicit_new_material_resumes_same_round_without_false_completion(client):
    from io import BytesIO
    from docx import Document
    from tests.api.test_review_rounds import edits, upload_parsed
    from app.domain.tasks.execution import run_task
    url, round, _, base = make_round(client, '1. Clarify the method.\n2. Change correlation to causation.')
    path = url + '/rounds/' + round['id']
    items = edits(round)
    items[0]['needs'] = [dict(kind='material', description='需提供现有方法原始材料。')]
    saved = client.post(path + '/plans', json=dict(expected_revision=round['revision'], items=items), headers={'Idempotency-Key': 'await-material'})
    assert saved.status_code == 201, saved.text
    round = saved.json()
    decided = client.post(path + '/decisions', json=dict(expected_revision=round['revision'], decisions=[dict(item_id=round['items'][1]['id'], action='not_adopt', reason='缺乏因果依据，明确保留相关。')]), headers={'Idempotency-Key': 'unaffected-decision'})
    assert decided.status_code == 201, decided.text
    round = decided.json()
    original_decision = round['items'][1]['decision']
    original_wait = round['wait']['id']
    task = round['task']
    resume = dict(operation='resume', wait_id=original_wait, expected_task_revision=task['revision'], input_version=task['input_version'])
    resumed = client.post('/api/v1/tasks/' + task['id'] + '/resume', json=resume, headers={'Idempotency-Key': 'material-still-missing'})
    assert resumed.status_code == 202, resumed.text
    assert run_task(task['id'], resumed.json()['revision'])['status'] == 'waiting_input'
    waiting = client.get(path).json()
    document = Document()
    document.add_paragraph('Synthetic method material: the original method and limitations remain unchanged. No experiment was performed.')
    raw = BytesIO()
    document.save(raw)
    material = upload_parsed(client, round['project_id'], raw.getvalue(), 'synthetic-method-supplement.docx')
    items = edits(waiting)
    items[0].update(needs=[], evidence_ids=[material['document_version_id']], plan='依据新固定原材料核对现有方法，仍须另交具体改文；不声称实验完成。')
    supplied = client.post(path + '/plans', json=dict(expected_revision=waiting['revision'], items=items,
        materials=[dict(kind='document_version', file_id=material['file_id'], id=material['document_version_id'])]), headers={'Idempotency-Key': 'supply-fixed-material'})
    assert supplied.status_code == 201, supplied.text
    current = supplied.json()
    assert current['items'][1]['decision'] == original_decision
    assert current['items'][0]['status'] != 'completed'
    assert current['versions'][0]['items'][0]['needs'] == []
    assert any(v['items'][0]['needs'] for v in current['versions'])
    task = current['task']
    resumed = client.post('/api/v1/tasks/' + task['id'] + '/resume', json=dict(operation='resume', wait_id=current['wait']['id'], expected_task_revision=task['revision'], input_version=task['input_version']), headers={'Idempotency-Key': 'resume-with-material'})
    assert resumed.status_code == 202, resumed.text
    assert run_task(task['id'], resumed.json()['revision'])['status'] == 'succeeded'
    final = client.get(path).json()
    assert final['id'] == round['id'] and final['round_number'] == round['round_number']
    assert final['items'][0]['status'] == 'pending'
    assert final['items'][1]['status'] == 'not_adopted'


def test_existing_synthetic_search_link_uses_public_support_and_citation_flow(client):
    from tests.api.test_literature import create
    from tests.api.test_literature_citations import binding, draft_body, paragraph, CitationNcbi
    from app.domain.literature.execution import run_literature_task
    url, round, _, _ = make_round(client, '1. Add a reference to the introduction.')
    path = url + '/rounds/' + round['id']
    task = create(client, round['project_id']).json()
    completed = run_literature_task(task['id'], task['revision'], client=CitationNcbi(3))
    assert completed['status'] == 'succeeded'
    saved = client.post(path + '/task-links', json=dict(plan_revision=round['revision'], item_id=round['items'][0]['id'], task_id=task['id'], purpose='retrieval'), headers={'Idempotency-Key': 'existing-synthetic-search'})
    assert saved.status_code == 201, saved.text
    result = saved.json()['result']
    assert result['artifacts']['literature']['search_id'] == task['id']
    assert result['artifacts']['literature']['retrieved_count'] == 3
    assert result['artifacts']['literature']['support_review_required'] is True
    assert client.get(result['public_result_url']).status_code == 200
    lit = url.rsplit('/reviews', 1)[0] + '/literature'
    references = client.get(result['artifacts']['literature']['references_url']).json()['items']
    reference = next(r for r in references if r['record']['abstract'])
    claim, support, _, _ = binding(client, lit, task, reference)
    cited = client.post(lit + '/citations/drafts', json=draft_body(paragraph(claim, support)), headers={'Idempotency-Key': 'existing-stage05-citation'})
    assert cited.status_code == 201, cited.text
    assert cited.json()['snapshot']['output']['references']
    current = client.get(path).json()
    assert current['items'][0]['status'] == 'waiting_material'
    assert not current['actual_revisions']


def test_same_opinion_plan_edit_can_explicitly_reuse_frozen_analysis_adoption(client, monkeypatch):
    from tests.api.test_review_rounds import edits
    from app.domain.tasks.execution import run_task
    from app.db.database import database_connection
    url, path, round, base, _, link, _ = queued_analysis(client, monkeypatch)
    monkeypatch.setattr('app.domain.descriptive.summarize', mock_local_summary)
    assert run_task(link['task']['id'], 1)['status'] == 'succeeded'
    saved = client.post(path + '/analysis-adoptions', json=adoption_body(round, base, link), headers={'Idempotency-Key': 'reusable-analysis'})
    assert saved.status_code == 201, saved.text
    adopted = saved.json()
    frozen = adopted['input_snapshot']
    original_item_revision = round['items'][0]['revision']
    changes = edits(round)
    changes[0].update(plan='复用已计算结果；先补充当前方法核对材料并确认具体处理。', needs=[dict(kind='analysis', description='当前计划尚待明确是否复用此实际分析。'), dict(kind='material', description='当前方法补料尚未确认。')], flags=['ambiguous'])
    updated = client.post(path + '/plans', json=dict(expected_revision=round['revision'], items=changes), headers={'Idempotency-Key': 'edit-adopted-opinion'})
    assert updated.status_code == 201, updated.text
    round = updated.json()
    assert round['items'][0]['revision'] == original_item_revision + 1
    context = client.get(path + '/revision-context').json()
    facts = {f['path']: f for f in context['analysis_facts']}
    assert '/overall/n' in facts, 'Explicit reuse of this same opinion must retain receipt-backed facts after its plan changes'
    assert adopted['source']['id'] in {e['id'] for e in context['grounded_plans'][0]['evidence']}
    bindings = [dict(fact_id=facts[p]['id'], value=facts[p]['value']) for p in ('/overall/n', '/overall/mean')]
    text = 'n = ' + bindings[0]['value'] + '; mean = ' + bindings[1]['value']
    operations = [dict(item_id=round['items'][0]['id'], kind='replace_paragraph', location_id=pid, text=text,
        numeric_bindings=bindings, evidence_ids=[adopted['source']['id']], reason='明确在当前计划中复用同一实际分析采用。') for pid in ('abstract:1', 'results:1')]
    body = dict(plan_revision=round['revision'], expected_current_revision=context['expected_current_revision'], analysis_adoption_ids=[adopted['id']], operations=operations)
    pending = client.post(path + '/revisions/preview', json=body)
    assert pending.status_code == 200, pending.text
    assert not pending.json()['can_adopt'] and any('当前方法补料' in i['reason'] for i in pending.json()['unexecuted'])
    assert any('当前计划尚待明确' in i['reason'] for i in pending.json()['unexecuted'])
    changes = edits(round)
    changes[0]['needs'] = []
    supplied = client.post(path + '/plans', json=dict(expected_revision=round['revision'], items=changes), headers={'Idempotency-Key': 'reuse-plan-material-checked'})
    assert supplied.status_code == 201, supplied.text
    round = supplied.json()
    body['plan_revision'] = round['revision']
    confirmation = client.post(path + '/revisions/preview', json=body)
    assert confirmation.status_code == 200, confirmation.text
    assert not confirmation.json()['can_adopt'] and any('重要决定' in i['reason'] for i in confirmation.json()['unexecuted'])
    decided = client.post(path + '/decisions', json=dict(expected_revision=round['revision'], decisions=[dict(item_id=round['items'][0]['id'], action='accept', reason='明确确认当前计划复用此已计算结果，不新增因果或实验结论。')]), headers={'Idempotency-Key': 'reuse-current-decision'})
    assert decided.status_code == 201, decided.text
    current = decided.json()
    assert client.post(path + '/revisions/preview', json=body).status_code == 409
    body['plan_revision'] = current['revision']
    ready = client.post(path + '/revisions/preview', json=body)
    assert ready.status_code == 200 and ready.json()['can_adopt'], ready.text
    revision = client.post(path + '/revisions', json=body | dict(preview_sha256=ready.json()['preview_sha256']), headers={'Idempotency-Key': 'publish-reused-analysis'})
    assert revision.status_code == 202, revision.text
    assert run_task(revision.json()['task']['id'], 1)['status'] == 'succeeded'
    output = client.get(url + '/revisions/' + revision.json()['id']).json()['result']['output']
    assert run_task(output['task_id'], 1)['status'] == 'succeeded'
    public = client.get(url.rsplit('/reviews', 1)[0] + '/manuscripts/' + output['manuscript_id'] + '/versions/' + str(output['revision']))
    assert public.status_code == 200, public.text
    assert client.get(path + '/science').json()['analysis_adoptions'][0]['input_snapshot'] == frozen
    with database_connection() as db:
        assert db.execute('SELECT count(*) AS n FROM review_analysis_adoptions WHERE round_id=%s', (round['id'],)).fetchone()['n'] == 1
