from datetime import datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile
import re

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook
from openpyxl.styles import Font

from tests.helpers.auth import login_test_client

from app.main import app


@pytest.fixture
def client():
    with TestClient(app) as test_client:
        yield login_test_client(test_client)


def workbook_bytes():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "研究数据"
    sheet.append(["编号", "值", "日期", "是否入组", "缺失"])
    for index in range(25):
        sheet.append([f"S{index:02}", index, datetime(2026, 9, 28), False, None])
    workbook.create_sheet("空表")
    workbook.create_sheet("只有表头").append(["组别", "计数"])
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def upload(client, content=None, name="研究.xlsx"):
    return client.post("/api/v1/excel/preview", files={
        "file": (name, workbook_bytes() if content is None else content,
                 "application/octet-stream"),
    })


def test_health_is_preserved(client):
    assert client.get("/api/v1/health").json() == {
        "status": "ok", "service": "paperassist-system",
    }


def test_reads_all_sheets_and_only_twenty_rows(client):
    response = upload(client)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["filename"] == "研究.xlsx"
    assert [sheet["name"] for sheet in data["sheets"]] == ["研究数据", "空表", "只有表头"]
    sheet = data["sheets"][0]
    assert sheet["columns"] == ["编号", "值", "日期", "是否入组", "缺失"]
    assert (sheet["row_count"], sheet["column_count"]) == (25, 5)
    assert len(sheet["preview_rows"]) == 20
    assert sheet["preview_rows"][0] == ["S00", 0, "2026-09-28T00:00:00", False, None]
    assert sheet["preview_rows"][-1][0] == "S19"


def test_empty_and_header_only_sheets(client):
    response = upload(client)
    assert response.status_code == 200
    empty, headers = response.json()["sheets"][1:]
    assert (empty["row_count"], empty["column_count"]) == (0, 0)
    assert empty["preview_rows"] == []
    assert any("空表" in warning for warning in empty["warnings"])
    assert (headers["row_count"], headers["column_count"]) == (0, 2)
    assert any("表头" in warning for warning in headers["warnings"])


def test_preserves_gaps_duplicate_headers_and_formulas(client):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["重复", "重复", None])
    sheet.append([0, False, "=1+1"])
    sheet.append([None, None, None])
    sheet.append([None, "末行", 7])
    sheet["J50"].font = Font(bold=True)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    response = upload(client, stream.getvalue())
    assert response.status_code == 200
    result = response.json()["sheets"][0]
    assert result["columns"] == ["重复", "重复", "未命名列 C"]
    assert (result["row_count"], result["column_count"]) == (3, 3)
    assert result["preview_rows"] == [[0, False, "=1+1"], [None, None, None], [None, "末行", 7]]
    assert any("公式" in warning for warning in result["warnings"])
    assert any("重复" in warning for warning in result["warnings"])


@pytest.mark.parametrize("name", ["data.xls", "data.xlsm", "data.csv", "data.xlsx.exe"])
def test_rejects_unsupported_extension(client, name):
    response = upload(client, name=name)
    assert response.status_code == 415
    assert ".xlsx" in response.json()["detail"]["message"]


@pytest.mark.parametrize("content, code", [(b"", "empty_file"), (b"not excel", "invalid_workbook")])
def test_invalid_content_has_clear_error(client, content, code):
    response = upload(client, content)
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == code


def test_missing_file(client):
    response = client.post("/api/v1/excel/preview")
    assert response.status_code == 422
    assert "文件" in response.json()["detail"]["message"]


def test_configurable_file_limit(client, monkeypatch):
    monkeypatch.setenv("EXCEL_MAX_UPLOAD_BYTES", "1024")
    assert client.get("/api/v1/excel/config").json()["max_upload_bytes"] == 1024
    response = upload(client, b"x" * 1025)
    assert response.status_code == 413
    assert response.json()["detail"]["code"] == "file_too_large"


def test_request_limit_also_applies_without_content_length(client, monkeypatch):
    monkeypatch.setenv("EXCEL_MAX_UPLOAD_BYTES", "10")
    response = client.post("/api/v1/excel/preview", content=iter([b"x" * 70000]))
    assert response.status_code == 413


def test_uppercase_extension(client):
    assert upload(client, name="DATA.XLSX").status_code == 200


def test_zip_without_workbook_is_rejected(client):
    stream = BytesIO()
    with ZipFile(stream, "w", ZIP_DEFLATED) as archive:
        archive.writestr("other.txt", "not an Excel workbook")
    response = upload(client, stream.getvalue())
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "invalid_workbook"


@pytest.mark.parametrize("variable, limit", [
    ("EXCEL_MAX_UNCOMPRESSED_BYTES", "100"),
    ("EXCEL_MAX_SHEETS", "2"),
    ("EXCEL_MAX_ROWS", "10"),
    ("EXCEL_MAX_COLUMNS", "4"),
    ("EXCEL_MAX_CELLS", "50"),
])
def test_parse_resource_limits(client, monkeypatch, variable, limit):
    monkeypatch.setenv(variable, limit)
    response = upload(client)
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "workbook_too_large"


def test_damaged_sheet_returns_parse_error(client):
    stream = BytesIO()
    with ZipFile(BytesIO(workbook_bytes())) as source, ZipFile(stream, "w") as target:
        for entry in source.infolist():
            content = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                content = b"<broken>"
            target.writestr(entry.filename, content)
    response = upload(client, stream.getvalue())
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "parse_failed"


def rewrite_workbook(transform):
    stream = BytesIO()
    with ZipFile(BytesIO(workbook_bytes())) as source, ZipFile(stream, "w", ZIP_DEFLATED) as target:
        for entry in source.infolist():
            name, content = transform(entry.filename, source.read(entry.filename))
            if name is not None:
                target.writestr(name, content)
    return stream.getvalue()


def test_optional_cell_and_row_coordinates(client):
    def transform(name, content):
        if name.startswith("xl/worksheets/"):
            content = re.sub(rb'(<(?:c|row)) r="[^"]+"', rb'\1', content)
        return name, content
    response = upload(client, rewrite_workbook(transform))
    assert response.status_code == 200
    sheet = response.json()["sheets"][0]
    assert (sheet["row_count"], sheet["column_count"]) == (25, 5)


def test_duplicate_xml_cells_cannot_bypass_budget(client, monkeypatch):
    monkeypatch.setenv("EXCEL_MAX_CELLS", "200")
    def transform(name, content):
        if name == "xl/worksheets/sheet1.xml":
            content = re.sub(rb'<sheetData>.*?</sheetData>',
                b'<sheetData><row r="1">' + b'<c r="A1" t="n"><v>1</v></c>' * 201 + b'</row></sheetData>', content)
        return name, content
    response = upload(client, rewrite_workbook(transform))
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "workbook_too_large"


def test_unknown_sheet_relationship_cannot_skip_preflight(client):
    def transform(name, content):
        if name == "xl/_rels/workbook.xml.rels":
            content = content.replace(b'/relationships/worksheet', b'/relationships/other')
        return name, content
    response = upload(client, rewrite_workbook(transform))
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "parse_failed"


def test_sheet_part_can_live_outside_standard_directory(client, monkeypatch):
    monkeypatch.setenv("EXCEL_MAX_CELLS", "200")
    def transform(name, content):
        if name == "xl/_rels/workbook.xml.rels":
            content = content.replace(b'/xl/worksheets/sheet1.xml', b'/xl/custom/data.xml')
        if name == "xl/worksheets/sheet1.xml":
            name = "xl/custom/data.xml"
            content = re.sub(rb'<sheetData>.*?</sheetData>',
                b'<sheetData><row r="1">' + b'<c r="A1" t="n"><v>1</v></c>' * 201 + b'</row></sheetData>', content)
        return name, content
    response = upload(client, rewrite_workbook(transform))
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "workbook_too_large"


def test_missing_sheet_part_does_not_silently_drop_a_sheet(client):
    response = upload(client, rewrite_workbook(lambda name, content:
        (None if name == "xl/worksheets/sheet1.xml" else name, content)))
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "parse_failed"


def test_exact_file_size_limit_is_accepted(client, monkeypatch):
    content = workbook_bytes()
    monkeypatch.setenv("EXCEL_MAX_UPLOAD_BYTES", str(len(content)))
    assert upload(client, content).status_code == 200


def test_misreported_dimensions_do_not_truncate_data(client):
    def transform(name, content):
        if name == "xl/worksheets/sheet1.xml":
            content = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1:A1"', content)
        return name, content
    response = upload(client, rewrite_workbook(transform))
    assert response.status_code == 200
    assert response.json()["sheets"][0]["row_count"] == 25


def test_merged_cells_preserve_missing_values(client):
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["分组", "记录"])
    sheet.append(["A", 1])
    sheet.append([None, 2])
    sheet.merge_cells("A2:A3")
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    response = upload(client, stream.getvalue())
    assert response.status_code == 200
    sheet = response.json()["sheets"][0]
    assert sheet["preview_rows"] == [["A", 1], [None, 2]]
    assert any("合并" in message for message in sheet["warnings"])


def test_entirely_empty_workbook(client):
    workbook = Workbook()
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    response = upload(client, stream.getvalue())
    assert response.status_code == 200
    assert response.json()["sheets"][0]["column_count"] == 0


def test_renamed_macro_workbook_is_rejected(client):
    def transform(name, content):
        if name == "[Content_Types].xml":
            content = content.replace(b'spreadsheetml.sheet.main+xml', b'spreadsheetml.sheet.macroEnabled.main+xml')
        return name, content
    response = upload(client, rewrite_workbook(transform))
    assert response.status_code == 422
    assert response.json()["detail"]["code"] == "invalid_workbook"
