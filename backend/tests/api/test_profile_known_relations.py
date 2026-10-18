"""Real saved-version API evidence for independently defined known relations."""
from tests.helpers.statistics import statistics_post
from hashlib import sha256
from io import BytesIO
from time import perf_counter

from openpyxl import load_workbook
import pytest

from tests.api.test_analysis import client  # noqa: F401
from tests.api.test_dataset_analysis import version, forbid_original
from tests.browser.task_browser_support import block_real_providers
from tests.helpers.profile_known_relations import workbook_bytes, expected_profile, sample_rows
from app.core.config import get_data_dir
from app.db.database import database_connection


def snapshot():
    with database_connection() as db:
        names = [r['tablename'] for r in db.execute(
            "SELECT tablename FROM pg_tables WHERE schemaname=current_schema() ORDER BY tablename").fetchall()]
        excluded = {'users', 'sessions', 'auth_login_attempts', 'session_revocations', 'password_recovery_codes'}
        rows = {name: sorted([dict(r) for r in db.execute('SELECT * FROM "' + name + '"').fetchall()], key=repr)
                for name in names if name not in excluded}
    assets = {str(p.relative_to(get_data_dir())): p.read_bytes() for p in get_data_dir().rglob('*') if p.is_file()}
    return rows, assets


def test_known_relations_saved_profile_and_reads_preserve_all_business_state(client, monkeypatch):
    attempts = block_real_providers(monkeypatch)
    oracle = expected_profile()
    oracle['numeric'][0]['n'] = -1
    assert expected_profile()['numeric'][0]['n'] == 24
    oracle = expected_profile()
    content = workbook_bytes()
    from datetime import datetime
    from datetime import timezone
    from types import SimpleNamespace
    class LaterClock(datetime):
        @staticmethod
        def now(tz=None):
            return datetime(2030, 12, 31, 23, 59, 59, tzinfo=tz)
    with monkeypatch.context() as clock:
        clock.setattr('openpyxl.writer.excel.datetime', SimpleNamespace(datetime=LaterClock, timezone=timezone))
        assert workbook_bytes() == content
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ['已知关系']
    assert list(book.active.values)[1:] == [tuple(row) for row in sample_rows()]
    assert all(cell.data_type != 'f' for row in book.active for cell in row)
    book.close()
    project = client.post('/api/v1/projects', json=dict(name='已知关系专项', research_topic='合成数据', project_type='sci')).json()
    upload = client.post(f'/api/v1/projects/{project["id"]}/files', files={'file': ('known-relations.xlsx', content)})
    assert upload.status_code == 201, upload.text
    record = upload.json()['file']
    base = f'/api/v1/projects/{project["id"]}/files/{record["id"]}'
    started = perf_counter()
    raw = version(client, base)
    elapsed = perf_counter() - started
    selection = dict(task_type='descriptive_boxplot', sheet_name='已知关系', numeric_column='A', group_column=None,
                     unit='未确认', missing_policy='exclude_selected_missing', dataset_id=raw['id'], expected_revision=0)
    saved = client.put(base + '/analysis-setup', json=selection)
    assert saved.status_code == 200, saved.text
    run = statistics_post(client, base + '/analysis-runs', json={'expected_revision': 1})
    assert run.status_code == 200, run.text
    assert run.json()['overall']['n'] == 24
    assert run.json()['overall']['mean'] == pytest.approx(343/24)
    before = snapshot()
    monkeypatch.setattr('openpyxl.load_workbook', forbid_original)
    url = base + '/dataset-versions/' + raw['id']
    for _ in range(2):
        response = client.get(url + '/profile')
        assert response.status_code == 200, response.text
        profile = response.json()
        for key in ('record_count', 'blank_row_count', 'data_range', 'duplicate_groups', 'duplicate_rows', 'duplicate_extra_rows'):
            assert profile[key] == oracle[key]
        assert profile['source']['dataset_id'] == raw['id']
        assert profile['sheet_name'] == oracle['sheet']
        assert [(f['column'], f['letter'], f['name']) for f in profile['fields']] == [
            (1, 'A', '合成X'), (2, 'B', '同向Y'), (3, 'C', '反向Z'),
            (4, 'D', '固定值'), (5, 'E', '日期文本')]
        for index, field in enumerate(profile['fields']):
            assert field['issue_counts'] == {
                'missing': 1 if index < 3 else 0,
                'outlier': 1 if index < 3 else 0,
                'mixed_type': 0, 'error': 0, 'formula': 0, 'numeric_precision': 0,
                'time_candidate': 23 if index == 4 else 0,
                'time_invalid': 1 if index == 4 else 0,
                'time_ambiguous': 1 if index == 4 else 0,
            }
            assert field['unit'] == {'status': 'unconfirmed', 'value': None}
            assert field['missing_count'] == (1 if index < 3 else 0)
            assert field['type_counts'] == ({'number': 24 if index < 3 else 25} if index < 4 else {'text': 25})
            if index < 4:
                for key, expected in oracle['numeric'][index].items():
                    assert field['numeric'][key] == (pytest.approx(expected) if isinstance(expected, float) else expected)
            else:
                assert field['time_candidates'] == oracle['time_candidates']
                assert field['temporal'] == []
            for kind, total in field['issue_counts'].items():
                items = []
                for page in range(1, max(1, (total + 19)//20) + 1):
                    issue = client.get(url + '/profile/issues', params=dict(kind=kind, column=index+1, page=page))
                    assert issue.status_code == 200, issue.text
                    assert issue.json()['total'] == total
                    items.extend(issue.json()['items'])
                assert len(items) == total
                expected_coords = {'missing': oracle['missing'], 'outlier': oracle['outliers'],
                                   'time_invalid': oracle['time_invalid'], 'time_ambiguous': oracle['time_ambiguous']}.get(kind)
                if expected_coords is not None:
                    assert [i['coordinate'] for i in items] == [c for c in expected_coords if c[0] == field['letter']]
                if kind == 'time_candidate' and index == 4:
                    assert [i['coordinate'] for i in items] == [f'E{r}' for r in list(range(2,25))]
        duplicates = client.get(url + '/profile/issues?kind=duplicate').json()
        assert duplicates['total'] == 3
        assert [i['row'] for i in duplicates['items']] == [7, 22, 23]
        assert all(i['group_first_row'] == 7 and i['group_size'] == 3 for i in duplicates['items'])
        paged = []
        for start in (2, 22):
            page = client.get(url + '/rows', params=dict(start_row=start)).json()
            paged.extend(page['rows'])
        assert [r['row'] for r in paged] == list(range(2,27))
        actual = []
        for row in paged:
            values = {c['column']: c['value'] for c in row['cells']}
            actual.append([int(values[c]) if c < 5 and values.get(c) is not None else values.get(c) for c in range(1,6)])
        assert actual == sample_rows()
        pairs = [r for r in actual if r[0] is not None]
        assert len(pairs) == 24 and all(y == 3*x+7 and z == 100-2*x for x,y,z,*_ in pairs)
        download = client.get(url + '/download')
        assert download.status_code == 200 and download.json()['sheet']['row_count'] == 26
        original = client.get(base + '/download')
        assert original.status_code == 200 and original.content == content
        assert client.get(base + '/analysis-result').json()['result'] == run.json()
        assert snapshot() == before
    assert attempts == []
    for table in ('model_calls', 'budget_reservations', 'usage_events'):
        assert before[0][table] == []
    forbidden = {'correlation', 'correlations', 'p_value', 'causality', 'causal', 'pearson', 'spearman'}
    def check_keys(value):
        if isinstance(value, dict):
            assert not forbidden.intersection(value)
            for child in value.values():
                check_keys(child)
        elif isinstance(value, list):
            for child in value:
                check_keys(child)
    check_keys(profile)
    print(f'known_relations bytes={len(content)} sha256={sha256(content).hexdigest()} saved_structure_seconds={elapsed:.6f}')
