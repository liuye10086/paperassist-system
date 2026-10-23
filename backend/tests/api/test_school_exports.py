"""School export intents remain distinct from frozen manuscript facts."""
from copy import deepcopy
import pytest
from tests.api.test_analysis import client  # noqa: F401
from tests.api.test_school_templates import setup as school_setup,analysis as base_analysis
from tests.api.test_literature_evidence import post
from app.domain.tasks.execution import run_task
from app.db.database import database_connection


def published(client):
    from app.domain.school_templates.execution import run_school_template_task
    path,template=school_setup(client)
    item=post(client,path+'/school-templates',template).json()
    def school_analysis(*args):
        analysis=base_analysis()
        analysis['sections']=[dict(id='section-1',page=dict(width_cm=21,height_cm=29.7,margin_top_cm=2.5,margin_bottom_cm=2.5,margin_left_cm=2.5,margin_right_cm=2.5),
            source=dict(part='word/document.xml',path='/w:document/w:body/w:p[100]/w:pPr/w:sectPr',text='synthetic page'),
            page_number=dict(format='decimal',start='1',mode='restart'),break_type=None,different_first_page=False,different_odd_even_pages=False,headers=[],footers=[])]
        return analysis
    assert run_school_template_task(item['task']['id'],1,parser=school_analysis)['status']=='succeeded'
    item=client.get(path+'/school-templates').json()['items'][0]
    chapters=[dict(id='school-intro',title='绪论',level=1),dict(id='school-methods',title='研究方法',level=1),dict(id='school-results',title='结果',level=1)]
    confirmation=post(client,path+'/school-templates/'+item['id']+'/confirm',dict(analysis_sha256=item['analysis_sha256'],decisions=[dict(conflict_id='conflict-1',option_id='text-2')],chapters=chapters)).json()
    profile=post(client,path+'/thesis-template-binding',dict(confirmation_id=confirmation['id'],expected_revision=0)).json()
    request=dict(title='合成学校导出',output_language='zh-CN',sources=[],thesis=dict(confirmation_id=confirmation['id'],profile_revision=profile['revision'],chapter_mappings=[dict(chapter_id=c['id'],purpose=p,output_language='project',result_group_ids=[]) for c,p in zip(chapters,['introduction','materials_and_methods','results'])]))
    root=post(client,path+'/manuscripts',request,'school').json()
    assert run_task(root['versions'][0]['task_id'],1)['status']=='succeeded'
    base=path+'/manuscripts/'+root['id']
    return path,base,template,root


def test_school_export_preview_and_idempotent_sealed_intent(client):
    path,base,_,root=published(client)
    endpoint=base+'/versions/1/exports'
    preview=post(client,endpoint+'/preview',{})
    assert preview.status_code==200,preview.text
    plan=preview.json()['school']
    assert preview.json()['schema_version']==1 and plan['confirmation_id']==root['versions'][0]['snapshot']['thesis_context']['confirmation_id']
    created=post(client,endpoint,{},'export')
    assert created.status_code==201,created.text
    assert post(client,endpoint,{'school':{}},'export').status_code==200
    assert post(client,endpoint,{'school':{'manual_fields':{'author':'changed'}}},'export').status_code==409
    view=created.json()
    assert view['school']['plan']==plan
    with database_connection() as db:
        row=db.execute('SELECT * FROM manuscript_exports WHERE id=%s',(view['id'],)).fetchone()
        assert row['school_input_json']['plan']==plan and row['school_input_sha256']
    assert client.get(base+'/versions/1').json()['snapshot']==root['versions'][0]['snapshot']


def test_school_export_uses_frozen_profile_and_rejects_bad_manual_contract(client):
    path,base,_,root=published(client)
    profile=client.get(path+'/thesis-profile').json()
    fields={k:profile[k] for k in ('school','department','major','degree_type','education_level','applicable_version')}
    assert client.put(path+'/thesis-profile',json=dict(expected_revision=profile['revision'],**(fields|{'major':'changed'}))).status_code==200
    endpoint=base+'/versions/1/exports'
    assert post(client,endpoint,{},'frozen-export').status_code==201
    assert post(client,endpoint,{'school':{'manual_fields':{'invented':'value'}}},'bad').status_code==422
    assert post(client,endpoint+'/preview',{'school':{'manual_chapters':[{'chapter_id':'school-intro','text':'wrong purpose'}]}}).status_code==422
    assert post(client,endpoint+'/preview',{'school':{'keywords':[{'chapter_id':'school-intro','values':['bad']} ]}}).status_code==422


def test_sci_export_retains_empty_request_and_no_school_metadata(client):
    from tests.api.test_manuscripts import setup,body
    path,_,_=setup(client)
    root=post(client,path,body(),'sci').json()
    run_task(root['versions'][0]['task_id'],1)
    endpoint=path+'/'+root['id']+'/versions/1/exports'
    assert post(client,endpoint,{'school':{}},'bad').status_code==422
    created=post(client,endpoint,{},'old').json()
    assert 'school' not in created
    assert run_task(created['task_id'],1)['status']=='succeeded'
    with database_connection() as db:
        row=db.execute('SELECT * FROM manuscript_exports WHERE id=%s',(created['id'],)).fetchone()
        assert row['school_input_json'] is None and row['school_input_sha256'] is None
        assert 'school' not in row['metadata_json']


def test_school_word_receipt_pdf_failure_and_intent_tampering(client,monkeypatch):
    from app.core.exceptions import StorageError
    from psycopg.types.json import Jsonb
    path,base,_,root=published(client)
    endpoint=base+'/versions/1/exports'
    # Renderer body/receipt are tested independently; use a real editable DOCX here.
    from docx import Document
    from io import BytesIO
    stream=BytesIO();doc=Document();doc.add_paragraph('Synthetic school export');doc.save(stream)
    monkeypatch.setattr('app.adapters.manuscript_docx.build_manuscript_docx',lambda *a,**k:stream.getvalue())
    receipt={'status':'updated','renderer_manifest':{'contract':'test-v1'},'word_verified':False,'wps_verified':False}
    monkeypatch.setattr('app.adapters.school_docx.complete_school_fields',lambda raw,plan:(raw,receipt))
    created=post(client,endpoint,{},'word').json()
    def future_compile(*args):raise AssertionError('Old export must not recompile current rules')
    monkeypatch.setattr('app.domain.manuscripts.school_export.compile_school_export',future_compile)
    assert run_task(created['task_id'],1)['status']=='succeeded'
    done=client.get(base+'/exports/'+created['id']).json()
    assert done['school']['field_update']==receipt
    assert client.get(done['download_url']).status_code==200
    monkeypatch.setattr('app.domain.pdf_exports.service.renderer_identity',lambda:{'contract':'test-v1'})
    pdf=client.post(path+'/pdf-exports',json=dict(source_kind='manuscript_export',source_id=created['id'],request_key='pdf')).json()
    def fail(raw):raise StorageError('pdf_conversion_failed','synthetic failure',422)
    monkeypatch.setattr('app.domain.pdf_exports.execution.render_pdf',fail)
    assert run_task(pdf['task_id'],1)['status']=='failed'
    assert client.get(done['download_url']).status_code==200
    with database_connection(write=True) as db:
        row=db.execute('SELECT school_input_json FROM manuscript_exports WHERE id=%s',(created['id'],)).fetchone()
        changed=deepcopy(row['school_input_json']);changed['plan']['toc']['depth']=1
        db.execute('UPDATE manuscript_exports SET school_input_json=%s WHERE id=%s',(Jsonb(changed),created['id']))
    assert client.get(done['download_url']).status_code==409
    from app.maintenance.manuscript_audit import audit_manuscripts
    with database_connection() as db:
        assert any(check['table']=='manuscript_exports' and check['invalid'] for check in audit_manuscripts(db))


def test_school_field_failure_does_not_publish_or_expose_download(client,monkeypatch):
    from app.core.exceptions import StorageError
    _,base,_,_=published(client)
    from docx import Document
    from io import BytesIO
    stream=BytesIO();Document().save(stream)
    monkeypatch.setattr('app.adapters.manuscript_docx.build_manuscript_docx',lambda *a,**k:stream.getvalue())
    def fail(*a):raise StorageError('layout_field_update_failed','synthetic failure',422)
    monkeypatch.setattr('app.adapters.school_docx.complete_school_fields',fail)
    created=post(client,base+'/versions/1/exports',{},'failed').json()
    assert run_task(created['task_id'],1)['status']=='failed'
    done=client.get(base+'/exports/'+created['id']).json()
    assert done['error_code']=='layout_field_update_failed' and done['download_url'] is None
    assert done['school']['field_update']['status']=='pending'


def test_school_export_owner_scope_source_change_and_audit(client):
    from tests.integration.test_project_ownership import account_client
    from app.maintenance.manuscript_audit import audit_manuscripts
    from psycopg.types.json import Jsonb
    _,base,template,_=published(client)
    endpoint=base+'/versions/1/exports'
    saved=post(client,endpoint,{},'school').json()
    with account_client('school-export-other@example.test')[0] as other:
        assert post(other,endpoint+'/preview',{}).status_code==404
        assert post(other,endpoint,{},'foreign').status_code==404
        assert other.get(base+'/exports/'+saved['id']).status_code==404
    with database_connection() as db:
        assert not any(check['invalid'] for check in audit_manuscripts(db))
    with database_connection(write=True) as db:
        db.execute('UPDATE files SET sha256=%s WHERE id=%s',('a'*64,template['template_file_id']))
    assert post(client,endpoint+'/preview',{}).status_code==409
    assert run_task(saved['task_id'],1)['status']=='failed'
    failed=client.get(base+'/exports/'+saved['id'])
    assert failed.status_code==200 and failed.json()['status']=='failed'
    assert failed.json()['error_code']=='task_source_conflict'
