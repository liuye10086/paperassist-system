from copy import deepcopy
from io import BytesIO
import pytest
from docx import Document

def fixed_result():
    change=dict(item_id='i1',kind='replace_paragraph',before='Association was observed.',after='An association was observed; causality is not established.',reason='Retain existing limitations.',location=dict(kind='paragraph',id='p1',section_id='discussion',section_title='Discussion'),evidence=[dict(id='e1',sha256='a'*64,content={'text':'existing limitation'})])
    return dict(actual_changes=[change],items=[dict(id='i1',source_id='r1',source_label='Reviewer 1',number=1,original_number='1.',original_text='Clarify limitation',translation=dict(language='zh-CN',text='说明局限'),status='completed',actual_changes=[change]),dict(id='i2',source_id='r2',source_label='Reviewer 2',number=1,original_text='Add an independent experiment',status='waiting_material',uncompleted_reason='No new experiment is available.')],output=dict(manuscript_id='m',version_id='v2',revision=2,task_id='producer'),synthetic=True)

def test_reply_is_grounded_and_language_does_not_overwrite_original():
    from app.domain.review_rounds.documents import build_snapshot
    result=fixed_result();before=deepcopy(result)
    en=build_snapshot(dict(id='rev',round_id='round',snapshot_json=dict(project_type='sci',output_language='zh-CN'),result_json=result,result_sha256='hash'),'reply','en')
    zh=build_snapshot(dict(id='rev',round_id='round',snapshot_json=dict(project_type='sci',output_language='zh-CN'),result_json=result,result_sha256='hash'),'reply','zh-CN')
    assert en['language']=='en' and zh['language']=='zh-CN'
    assert 'not completed' in en['items'][1]['text'].lower()
    assert 'No new experiment' in en['items'][1]['text']
    assert en['items'][0]['actual_changes'][0]['after']==result['actual_changes'][0]['after']
    assert en['items'][0]['original_text']=='Clarify limitation'
    assert zh['items'][0]['translation']==result['items'][0]['translation']
    assert result==before
    assert all('page' not in location and 'line' not in location for item in en['items'] for location in item['locations'])

def test_completed_without_actual_change_cannot_fabricate_reply():
    from app.domain.review_rounds.documents import build_snapshot
    from app.core.exceptions import StorageError
    result=fixed_result();result['actual_changes']=[];result['items'][0]['actual_changes']=[]
    with pytest.raises(StorageError):
        build_snapshot(dict(id='rev',round_id='round',snapshot_json=dict(project_type='sci'),result_json=result,result_sha256='hash'),'reply','en')

def test_independent_docx_preserves_sources_diff_and_pending_reason():
    from app.domain.review_rounds.documents import build_snapshot
    from app.adapters.review_docx import render_review_document
    snap=build_snapshot(dict(id='rev',round_id='round',snapshot_json=dict(project_type='sci'),result_json=fixed_result(),result_sha256='hash'),'comparison','en')
    raw=render_review_document(snap)
    assert raw==render_review_document(snap)
    document=Document(BytesIO(raw));text='\n'.join(p.text for p in document.paragraphs)
    assert 'Reviewer 1' in text and 'Reviewer 2' in text
    assert 'Association was observed.' in text and 'causality is not established' in text
    assert 'No new experiment is available.' in text
    assert 'p1' in text and 'existing limitation' in text
from tests.api.test_analysis import client
from tests.api.test_review_revisions import context_and_request

def completed_revision(client):
    from app.domain.tasks.execution import run_task
    url,round,base,context,body=context_and_request(client)
    path=url+'/rounds/'+round['id']+'/revisions'
    preview=client.post(path+'/preview',json=body);assert preview.status_code==200,preview.text
    body['preview_sha256']=preview.json()['preview_sha256']
    created=client.post(path,json=body,headers={'Idempotency-Key':'document-adopt'});assert created.status_code==202,created.text
    run=created.json();assert run_task(run['task']['id'],1)['status']=='succeeded'
    run=client.get(url+'/revisions/'+run['id']).json()
    early=client.post(url+'/revisions/'+run['id']+'/documents',json=dict(kind='reply',language='en'),headers={'Idempotency-Key':'too-early'})
    assert early.status_code==409,early.text
    assert run_task(run['result']['output']['task_id'],1)['status']=='succeeded'
    return url,client.get(url+'/revisions/'+run['id']).json()

def create_document(client,url,run,kind='reply',language='en',key='reply'):
    response=client.post(url+'/revisions/'+run['id']+'/documents',json=dict(kind=kind,language=language),headers={'Idempotency-Key':key})
    assert response.status_code==201,response.text
    return response.json()

def test_versions_confirmation_languages_and_owner_are_independent(client):
    url,run=completed_revision(client)
    doc=create_document(client,url,run)
    assert doc['snapshot']['items'][0]['status']=='completed'
    assert doc['snapshot']['items'][0]['source_label']=='Reviewer 1'
    assert doc['snapshot']['items'][0]['locations'][0]['paragraph_number']>=1
    assert doc['snapshot']['items'][1]['status']=='waiting_material'
    assert doc['snapshot']['output']==run['result']['output']
    assert client.post(url+'/revisions/'+run['id']+'/documents',json=dict(kind='reply',language='en'),headers={'Idempotency-Key':'reply'}).json()==doc
    assert client.post(url+'/revisions/'+run['id']+'/documents',json=dict(kind='comparison',language='en'),headers={'Idempotency-Key':'reply'}).status_code==409
    confirm=url+'/documents/'+doc['id']+'/confirm'
    confirmed=client.post(confirm,json=dict(expected_version=doc['version']),headers={'Idempotency-Key':'confirmed'})
    assert confirmed.status_code==200,confirmed.text
    assert confirmed.json()['confirmed_at']
    request=dict(expected_version=doc['version'],language='zh-CN',items=[dict(item_id=i['item_id'],text='人工明确说明：'+i['text']) for i in doc['snapshot']['items']])
    version=client.post(url+'/documents/'+doc['id']+'/versions',json=request,headers={'Idempotency-Key':'edit'})
    assert version.status_code==201,version.text
    version=version.json();assert version['version']==2 and version['language']=='zh-CN' and version['confirmed_at'] is None
    assert version['snapshot']['items'][0]['original_text']==doc['snapshot']['items'][0]['original_text']
    assert version['snapshot']['items'][0]['actual_changes']==doc['snapshot']['items'][0]['actual_changes']
    assert client.get(url+'/documents/'+doc['id']).json()['confirmed_at']
    assert client.post(url+'/documents/'+doc['id']+'/versions',json=request,headers={'Idempotency-Key':'stale'}).status_code==409
    comparison=create_document(client,url,run,'comparison','zh-CN','comparison')
    from tests.integration.test_project_ownership import account_client
    with account_client('document-foreign@example.test')[0] as other:
        for suffix in ('/documents/'+doc['id'],'/revisions/'+run['id']+'/documents'):
            assert other.get(url+suffix).status_code==404
        assert other.post(confirm,json=dict(expected_version=1),headers={'Idempotency-Key':'foreign-confirm'}).status_code==404
    assert client.get(url+'/revisions/'+run['id']+'/documents').status_code==200

def test_export_pdf_fixed_sha_download_and_failure_preserves_success(client,monkeypatch):
    from types import SimpleNamespace
    from hashlib import sha256
    from app.domain.tasks.execution import run_task
    url,run=completed_revision(client);doc=create_document(client,url,run)
    path=url+'/documents/'+doc['id']+'/exports'
    response=client.post(path,json={},headers={'Idempotency-Key':'export'})
    assert response.status_code==202,response.text
    export=response.json();assert client.post(path,json={},headers={'Idempotency-Key':'export'}).json()==export
    assert client.get(url+'/document-exports/'+export['id']+'/download').status_code==409
    assert run_task(export['task_id'],1)['status']=='succeeded'
    assert run_task(export['task_id'],1) is None
    export=client.get(url+'/document-exports/'+export['id']).json()
    raw=client.get(export['download_url']);assert raw.status_code==200 and raw.content.startswith(b'PK')
    assert sha256(raw.content).hexdigest()==export['sha256']
    task=client.get('/api/v1/tasks/'+export['task_id']).json();assert task['input_version']['document_id']==doc['id']
    assert task['input_version']['revision_id']==run['id'] and task['input_version']['round_id']==run['round_id']
    workspace=client.get('/api/v1/tasks/'+export['task_id']+'/workspace');assert workspace.status_code==200,workspace.text
    failed_doc=create_document(client,url,run,'comparison','en','comparison-for-failure')
    failed=client.post(url+'/documents/'+failed_doc['id']+'/exports',json={},headers={'Idempotency-Key':'failed-export'}).json()
    from app.adapters.review_docx import render_review_document
    def fail(_):raise OSError('forced isolated export failure')
    monkeypatch.setattr('app.domain.review_rounds.document_execution.render_review_document',fail)
    assert run_task(failed['task_id'],1)['status']=='failed'
    assert client.get(export['download_url']).content==raw.content
    assert client.get(url+'/revisions/'+run['id']).json()['result']['output']==run['result']['output']
    task=client.get('/api/v1/tasks/'+failed['task_id']).json()
    retry=client.post('/api/v1/tasks/'+failed['task_id']+'/retry',json=dict(expected_revision=task['revision']))
    assert retry.status_code==202,retry.text
    assert client.post('/api/v1/tasks/'+failed['task_id']+'/retry',json=dict(expected_revision=task['revision'])).status_code==409
    monkeypatch.setattr('app.domain.review_rounds.document_execution.render_review_document',render_review_document)
    assert run_task(failed['task_id'],retry.json()['revision'])['status']=='succeeded'
    assert client.get(export['download_url']).content==raw.content
    monkeypatch.setattr('app.domain.pdf_exports.service.renderer_identity',lambda:{'contract':'test-v1'})
    monkeypatch.setattr('app.domain.pdf_exports.execution.renderer_identity',lambda:{'contract':'test-v1'})
    monkeypatch.setattr('app.domain.pdf_exports.execution.render_pdf',lambda data:SimpleNamespace(content=b'%PDF-1.7\nsynthetic test only',manifest={'contract':'test-v1'}))
    pdfpath=url.removesuffix('/reviews')+'/pdf-exports'
    pdf=client.post(pdfpath,json=dict(source_kind='review_document_export',source_id=export['id'],request_key='pdf'))
    assert pdf.status_code==201,pdf.text
    pdf=pdf.json();assert pdf['source_sha256']==export['sha256']
    pdf_task=run_task(pdf['task_id'],1)
    assert pdf_task['status']=='succeeded',pdf_task
    complete=client.get(pdfpath+'/'+pdf['id']).json();assert client.get(complete['download_url']).content.startswith(b'%PDF-')
    assert client.get(pdfpath,params=dict(source_kind='review_document_export',source_id=export['id'])).json()['items'][0]['source_sha256']==export['sha256']
    from tests.integration.test_project_ownership import account_client
    with account_client('export-foreign@example.test')[0] as other:
        assert other.get(export['download_url']).status_code==404
        assert other.get(complete['download_url']).status_code==404

def test_expired_lease_recovery_cannot_publish_stale_document(client):
    from app.adapters.task_execution_store import TaskExecutionStore
    from app.db.database import database_connection
    from app.domain.tasks.execution import run_task
    url,run=completed_revision(client);doc=create_document(client,url,run)
    export=client.post(url+'/documents/'+doc['id']+'/exports',json={},headers={'Idempotency-Key':'lease-export'}).json()
    execution=TaskExecutionStore();old=execution.claim(export['task_id'],1)
    assert old is not None and execution.recover_expired()==0
    with database_connection(write=True) as db:
        db.execute("UPDATE task_attempts SET lease_expires_at=clock_timestamp()-interval '1 second' WHERE task_id=%s",(export['task_id'],))
    assert execution.recover_expired()==1 and not execution.heartbeat(old)
    assert run_task(export['task_id'],1) is None
    task=client.get('/api/v1/tasks/'+export['task_id']).json()
    assert run_task(export['task_id'],task['revision'])['status']=='succeeded'
    exported=client.get(url+'/document-exports/'+export['id']).json()
    assert client.get(exported['download_url']).content.startswith(b'PK')

def test_school_comparison_without_journal_and_no_sci_reply():
    from app.domain.review_rounds.documents import build_snapshot
    from app.core.exceptions import StorageError
    row=dict(id='rev',round_id='round',snapshot_json=dict(project_type='thesis'),result_json=fixed_result(),result_sha256='hash')
    assert build_snapshot(row,'comparison','zh-CN')['project_type']=='thesis'
    with pytest.raises(StorageError):build_snapshot(row,'reply','en')

def test_structured_changes_are_readable_without_exposing_internal_objects():
    from app.domain.review_rounds.documents import build_snapshot
    from app.adapters.review_docx import render_review_document
    result=fixed_result();result['actual_changes']=[]
    metadata=dict(title='Synthetic source title',year='2024',authors=[dict(collective='Synthetic author group')])
    citation=dict(metadata=metadata,review=dict(status='supported'),metadata_source=dict(receipt='PRIVATE_RECEIPT_MARKER'),project_reference_id='opaque-reference-uuid')
    for kind,before,after in [('bind_citation',[],[citation]),('reorder_paragraphs',['opaque-a','opaque-b','opaque-c'],['opaque-c','opaque-a','opaque-b'])]:
        change=dict(item_id='i1',kind=kind,before=before,after=after,reason='Existing fixed evidence.',location=dict(kind='section',id='methods',section_title='Methods'),evidence=[dict(id='source-test',kind='document_version',sha256='b'*64,content=dict(text='Fixed evidence excerpt.',object_key='PRIVATE_OBJECT_KEY',receipt=dict(secret='PRIVATE_RECEIPT_MARKER')))])
        result['actual_changes'].append(change)
    before=deepcopy(result)
    snap=build_snapshot(dict(id='rev',round_id='round',snapshot_json=dict(project_type='sci'),result_json=result,result_sha256='hash'),'reply','en')
    initial=snap['items'][0]['text']
    assert 'Synthetic source title' in initial and '2024' in initial and 'supported' in initial
    assert '3 → 1 → 2' in initial
    raw=render_review_document(snap);text='\n'.join(p.text for p in Document(BytesIO(raw)).paragraphs)
    assert 'Synthetic author group' in text and 'Fixed evidence excerpt.' in text
    for value in ('PRIVATE_RECEIPT_MARKER','PRIVATE_OBJECT_KEY','opaque-a','opaque-reference-uuid',"'metadata'",'"object_key"'):
        assert value not in text
    assert result==before
    assert snap['items'][0]['actual_changes'][0]['after'][0]['metadata_source']['receipt']=='PRIVATE_RECEIPT_MARKER'

def test_partial_change_text_preserves_actual_work_and_unfinished_reason():
    from app.domain.review_rounds.documents import build_snapshot
    result=fixed_result();result['items'][0].update(status='waiting_material',uncompleted_reason='Independent experiment remains unavailable.')
    snap=build_snapshot(dict(id='rev',round_id='round',snapshot_json=dict(project_type='sci'),result_json=result,result_sha256='hash'),'reply','en')
    assert 'partial' in snap['items'][0]['text'].lower()
    assert result['actual_changes'][0]['after'] in snap['items'][0]['text']
    assert 'Independent experiment remains unavailable.' in snap['items'][0]['text']
