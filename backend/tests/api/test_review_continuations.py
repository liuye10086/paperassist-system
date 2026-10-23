"""Bounded user returns, exact next inputs and user-only progress."""
from copy import deepcopy
from io import BytesIO

from docx import Document

from tests.api.test_analysis import client
from tests.api.test_review_rounds import upload_parsed


def successful_revision(client):
    from tests.api.test_review_revisions import context_and_request
    from app.domain.tasks.execution import run_task
    url, round, base, context, body = context_and_request(client, '1. Clarify the method.')
    path = url+'/rounds/'+round['id']+'/revisions'
    preview = client.post(path+'/preview', json=body).json()
    response = client.post(path, json=body|dict(preview_sha256=preview['preview_sha256']),
                           headers={'Idempotency-Key': 'return-parent'})
    assert response.status_code == 202, response.text
    revision = response.json()
    assert run_task(revision['task']['id'], 1)['status'] == 'succeeded'
    output = client.get(url+'/revisions/'+revision['id']).json()['result']['output']
    assert run_task(output['task_id'], 1)['status'] == 'succeeded'
    root = url.rsplit('/reviews', 1)[0]+'/manuscripts/'+output['manuscript_id']
    parent = client.get(root+'/versions/'+str(output['revision'])).json()
    return url, round, base, revision, root, parent


def register_return(client, url, round, root, parent, revision_id=None, *, mutate=None, suffix='clean'):
    from app.domain.tasks.execution import run_task
    options=dict(school=dict(include_toc=False)) if parent['snapshot'].get('thesis_context') else {}
    response = client.post(root+'/versions/'+str(parent['revision'])+'/exports', json=options,
                           headers={'Idempotency-Key': 'return-export-'+suffix})
    assert response.status_code == 201, response.text
    export = response.json()
    task=run_task(export['task_id'], 1)
    assert task['status'] == 'succeeded', task
    download = client.get(root+'/exports/'+export['id']+'/download')
    assert download.status_code == 200, download.text
    doc = Document(BytesIO(download.content))
    body_text={p['text'] for s in parent['snapshot'].get('sections',[]) for p in s['paragraphs']}
    paragraph = next(p for p in doc.paragraphs if 'The existing method is retained.' in p.text) if revision_id else next(p for p in doc.paragraphs if p.text.strip() and p.text in body_text)
    if mutate:
        mutate(doc, paragraph)
    else:
        paragraph.runs[0].text += ' Additional ordinary wording from the user.' if revision_id else ' 补充普通用户表述。'
    stream = BytesIO(); doc.save(stream)
    source = upload_parsed(client, round['project_id'], stream.getvalue(), 'synthetic-return-'+suffix+'.docx')
    body = dict(manuscript=dict(id=root.rsplit('/', 1)[1], revision=parent['revision'], version_id=parent['id']),
                revision_id=revision_id, file_id=source['file_id'], document_version_id=source['document_version_id'],
                base_export_id=export['id'], synthetic=True)
    response = client.post(url+'/rounds/'+round['id']+'/returns', json=body,
                           headers={'Idempotency-Key': 'register-return-'+suffix})
    assert response.status_code == 201, response.text
    return response.json(), body


def next_request(round, parent, *, return_id=None):
    return dict(title='Explicit next round', project_type=round['project_type'], synthetic=True,
                manuscript=dict(id=round['manuscript']['id'], revision=parent['revision'], version_id=parent['id']),
                previous_round_id=round['id'], user_return_id=return_id, materials=[],
                school_confirmation_id=round['constraints']['school_confirmation_id'],
                journal_target_id=round['constraints']['journal_target_id'],
                sources=[dict(role='advisor' if round['project_type']=='thesis' else 'reviewer', label='Next reviewer', format='text', text='请澄清下一轮普通表述。')])


def test_sci_user_return_public_producer_and_explicit_next_round(client):
    from app.domain.tasks.execution import run_task
    url, round, old, revision, root, parent = successful_revision(client)
    response = client.get(url+'/rounds/'+round['id']+'/continuation-context')
    assert response.status_code == 200, response.text
    assert {v['version_id'] for v in response.json()['eligible_inputs']} >= {old['id'], parent['id']}
    saved, body = register_return(client, url, round, root, parent, revision['id'])
    assert saved['can_adopt'], saved['issues']
    assert saved['source']['file_id'] == body['file_id'] and saved['synthetic']
    assert any(d['kind']=='text_changed' for d in saved['differences'])
    adopt = dict(mappings=saved['mappings'], expected_current_revision=parent['revision'], preview_sha256=saved['preview_sha256'])
    response = client.post(url+'/returns/'+saved['id']+'/adopt', json=adopt, headers={'Idempotency-Key': 'adopt-return'})
    assert response.status_code == 201, response.text
    output = response.json()['output']; assert output['status']=='pending'
    assert client.post(url+'/returns/'+saved['id']+'/adopt', json=adopt, headers={'Idempotency-Key':'adopt-return'}).status_code == 200
    selected = dict(id=output['version_id'], revision=output['revision'])
    payload = next_request(round, selected, return_id=saved['id'])
    assert client.post(url+'/rounds', json=payload, headers={'Idempotency-Key':'premature-next'}).status_code == 409
    assert run_task(output['task_id'], 1)['status'] == 'succeeded'
    public = client.get(root+'/versions/'+str(output['revision']))
    assert public.status_code == 200, public.text
    assert public.json()['parent_version_id'] == parent['id']
    assert public.json()['snapshot']['user_return_provenance']['return_id'] == saved['id']
    assert client.get(saved['original_download_url']).status_code == 200
    assert client.get(root+'/versions/'+str(parent['revision'])).json()['snapshot'] == parent['snapshot']
    response = client.post(url+'/rounds', json=payload, headers={'Idempotency-Key':'next-return'})
    assert response.status_code == 202, response.text
    next_round = response.json()
    assert next_round['continuation']['user_return_id'] == saved['id']
    assert next_round['manuscript']['version_id'] == output['version_id']
    older = client.post(url+'/rounds', json=next_request(round, old), headers={'Idempotency-Key':'explicit-old'})
    assert older.status_code == 202, older.text
    assert older.json()['manuscript']['version_id'] == old['id']
    previous = client.get(url+'/rounds/'+round['id']).json()
    assert next_round['id'] in previous['next_round_ids'] and previous['actual_revisions']
    run_task(next_round['task']['id'],1)
    continued=client.get(url+'/rounds/'+next_round['id']).json()
    context=client.get(url+'/rounds/'+next_round['id']+'/revision-context').json()
    location=next(p for p in context['locations'] if p['kind']=='paragraph' and p['section_id']=='materials_and_methods')
    request=dict(plan_revision=continued['revision'],expected_current_revision=context['expected_current_revision'],operations=[dict(
        item_id=continued['items'][0]['id'],kind='replace_paragraph',location_id=location['id'],text=location['text']+' Further next-round wording.',
        evidence_ids=[context['evidence'][0]['id']],reason='Continue from the explicitly selected user manuscript.')])
    path=url+'/rounds/'+continued['id']+'/revisions'
    preview=client.post(path+'/preview',json=request)
    assert preview.status_code==200,preview.text
    assert preview.json()['can_adopt'],[i for i in preview.json()['public_issues'] if i['severity']=='error']
    saved_revision=client.post(path,json=request|dict(preview_sha256=preview.json()['preview_sha256']),headers={'Idempotency-Key':'after-user-return'})
    assert saved_revision.status_code==202,saved_revision.text
    assert run_task(saved_revision.json()['task']['id'],1)['status']=='succeeded'
    next_output=client.get(url+'/revisions/'+saved_revision.json()['id']).json()['result']['output']
    assert next_output['parent_version_id']==output['version_id']
    assert run_task(next_output['task_id'],1)['status']=='succeeded'
    assert client.get(root+'/versions/'+str(next_output['revision'])).status_code==200
    from app.db.database import database_connection
    from app.maintenance.review_continuation_audit import audit_review_continuations
    with database_connection() as db:
        assert all(c['invalid']==0 for c in audit_review_continuations(db))


def test_return_preview_cas_scope_ambiguity_and_protected_content(client):
    from tests.api.test_projects import create_project
    url, round, old, revision, root, parent = successful_revision(client)
    saved, body = register_return(client, url, round, root, parent, revision['id'])
    preview = client.post(url+'/returns/'+saved['id']+'/preview', json=dict(mappings=[]))
    assert preview.status_code == 200 and not preview.json()['can_adopt']
    adopt = dict(mappings=saved['mappings'], expected_current_revision=parent['revision']-1, preview_sha256=saved['preview_sha256'])
    assert client.post(url+'/returns/'+saved['id']+'/adopt', json=adopt, headers={'Idempotency-Key':'stale-return'}).status_code == 409
    adopt['expected_current_revision'] = parent['revision']; adopt['preview_sha256'] = '0'*64
    assert client.post(url+'/returns/'+saved['id']+'/adopt', json=adopt, headers={'Idempotency-Key':'changed-preview'}).status_code == 409
    other = create_project(client)
    assert client.get(f"/api/v1/projects/{other['id']}/reviews/returns/{saved['id']}").status_code == 404
    def protected(doc, paragraph):
        paragraph.runs[0].text += ' Ordinary user wording.'
        doc.sections[0].header.paragraphs[0].text = 'Changed protected header'
    blocked, _ = register_return(client, url, round, root, parent, revision['id'], mutate=protected, suffix='protected')
    assert not blocked['can_adopt'] and any(p['code']=='citation_changed' for p in blocked['issues'])
    assert client.get(blocked['original_download_url']).status_code == 200


def test_school_return_and_next_round_need_no_journal(client):
    from tests.api.test_school_exports import published
    from app.domain.tasks.execution import run_task
    path, root, _, manuscript = published(client); parent = manuscript['versions'][0]
    url = path+'/reviews'
    body = dict(title='学校合成回传', project_type='thesis', synthetic=True, manuscript=dict(id=manuscript['id'],revision=parent['revision'],version_id=parent['id']),
                school_confirmation_id=parent['snapshot']['thesis_context']['confirmation_id'], materials=[], sources=[dict(role='advisor',label='导师',format='text',text='补充普通说明。')])
    created = client.post(url+'/rounds',json=body,headers={'Idempotency-Key':'school-return-round'})
    assert created.status_code==202,created.text
    run_task(created.json()['task']['id'],1)
    round = client.get(url+'/rounds/'+created.json()['id']).json()
    saved, _ = register_return(client,url,round,root,parent)
    assert saved['can_adopt'],saved['issues']
    response = client.post(url+'/returns/'+saved['id']+'/adopt',json=dict(mappings=saved['mappings'],expected_current_revision=1,preview_sha256=saved['preview_sha256']),headers={'Idempotency-Key':'school-return-adopt'})
    assert response.status_code==201,response.text
    output = response.json()['output']; assert run_task(output['task_id'],1)['status']=='succeeded'
    selected = client.get(root+'/versions/'+str(output['revision'])).json()
    assert selected['snapshot']['thesis_context']==parent['snapshot']['thesis_context']
    response = client.post(url+'/rounds',json=next_request(round,selected,return_id=saved['id']),headers={'Idempotency-Key':'school-next'})
    assert response.status_code==202,response.text
    assert response.json()['constraints']['journal_target_id'] is None


def test_user_progress_corrections_exact_refs_and_task_read_isolation(client):
    from app.domain.tasks.execution import run_task
    url, round, old, revision, root, parent = successful_revision(client)
    assert client.get(url+'/progress').json()['items']==[]
    body = dict(status='revision',event_date='2026-10-23',note='合成用户独立记录',synthetic=True,round_id=round['id'],manuscript=dict(id=round['manuscript']['id'],revision=parent['revision'],version_id=parent['id']))
    response = client.post(url+'/progress',json=body,headers={'Idempotency-Key':'user-progress'})
    assert response.status_code==201,response.text
    event=response.json(); assert event['source_kind']=='project' and event['synthetic']
    assert client.post(url+'/progress',json=body,headers={'Idempotency-Key':'user-progress'}).status_code==200
    correction = body|dict(status='not_recorded',note='合成纠正记录',synthetic=False,corrects_event_id=event['id'])
    response = client.post(url+'/progress',json=correction,headers={'Idempotency-Key':'correct-progress'})
    assert response.status_code==201,response.text
    corrected=response.json(); assert corrected['synthetic']
    assert client.post(url+'/progress',json=correction,headers={'Idempotency-Key':'duplicate-correction'}).status_code==409
    client.get(url+'/rounds/'+round['id']); client.get('/api/v1/tasks/'+revision['task']['id']+'/workspace')
    client.get(root+'/versions/'+str(parent['revision']))
    listing=client.get(url+'/progress').json()
    assert len(listing['items'])==2 and listing['current']['id']==corrected['id']
    invalid=body|dict(export_id='foreign-export')
    assert client.post(url+'/progress',json=invalid,headers={'Idempotency-Key':'foreign-export'}).status_code in (404,422)


def test_sci_target_progress_reuses_existing_submission_truth(client):
    from tests.api.test_journal_adaptation import adaptation_ready
    root,_,base,fixed,_,_=adaptation_ready(client)
    url=root.rsplit('/manuscripts/',1)[0]+'/reviews'
    body=dict(status='submitted',event_date='2026-10-23',note='synthetic user event',synthetic=True,target_id=fixed['target_id'])
    response=client.post(url+'/progress',json=body,headers={'Idempotency-Key':'target-user-progress'})
    assert response.status_code==201,response.text
    assert response.json()['source_kind']=='journal' and response.json()['synthetic']
    from app.db.database import database_connection
    with database_connection() as db:
        assert db.execute('SELECT count(*) n FROM journal_submission_events').fetchone()['n']==1
        assert db.execute('SELECT count(*) n FROM project_review_progress_events').fetchone()['n']==0
    listing=client.get(url+'/progress').json()
    assert listing['journal_current_by_target'][fixed['target_id']]['id']==response.json()['id']


def test_existing_journal_progress_normalizes_optional_review_refs(client):
    from tests.api.test_journal_adaptation import adaptation_ready
    root,_,_,fixed,_,_=adaptation_ready(client)
    project=root.rsplit('/manuscripts/',1)[0]
    response=client.post(project+'/journals/submission-events',json=dict(target_id=fixed['target_id'],status='revision',event_date='2026-10-23',note='Legacy synthetic journal event'),headers={'Idempotency-Key':'legacy-journal-progress'})
    assert response.status_code==201,response.text
    saved=response.json()
    response=client.get(project+'/reviews/progress')
    assert response.status_code==200,response.text
    event=next(e for e in response.json()['items'] if e['id']==saved['id'])
    assert event['source_kind']=='journal' and event['synthetic']
    assert event['round_id'] is None and event['manuscript'] is None
    from app.db.database import database_connection
    with database_connection() as db:
        assert db.execute('SELECT count(*) n FROM project_review_progress_events').fetchone()['n']==0


def test_changed_return_claim_drops_support_but_exact_statement_keeps_it():
    from app.domain.review_rounds.returns import apply_supported_return
    def paragraph(identity,text):
        return dict(id=identity,text=text,origin='user_provided',source_ids=[],citations=[dict(review_id='fixed')],numeric_bindings=[],claim_bindings=['fixed'])
    base=dict(output_language='en',sections=[dict(id='a',title='A',purpose='other',paragraphs=[paragraph('same','Exact statement.'),paragraph('changed','Old statement.')])])
    incoming=[dict(id='incoming',paragraphs=['Exact statement.','Reworded statement.'],paragraph_items=[dict(text='Exact statement.',original_id='same'),dict(text='Reworded statement.',original_id='changed')])]
    candidate,_,issues=apply_supported_return(base,incoming,[dict(input_section_id='incoming',chapter_id='a')])
    assert candidate['sections'][0]['paragraphs'][0]==base['sections'][0]['paragraphs'][0]
    changed=candidate['sections'][0]['paragraphs'][1]
    assert changed['citations']==[] and not changed.get('claim_bindings')
    assert any(i['code']=='bound_claim_changed' for i in issues)
    assert base['sections'][0]['paragraphs'][1]['text']=='Old statement.'


def test_unadopted_return_original_is_retained_by_review_history(client):
    from tests.api.test_review_rounds import make_round
    url,round,_,base=make_round(client)
    doc=Document();doc.add_heading('Materials and methods',1);doc.add_paragraph('Synthetic incomplete return retained for matching.')
    stream=BytesIO();doc.save(stream)
    source=upload_parsed(client,round['project_id'],stream.getvalue(),'synthetic-unadopted-return.docx')
    body=dict(manuscript=dict(id=round['manuscript']['id'],revision=base['revision'],version_id=base['id']),
              file_id=source['file_id'],document_version_id=source['document_version_id'],synthetic=True)
    response=client.post(url+'/rounds/'+round['id']+'/returns',json=body,headers={'Idempotency-Key':'unadopted-return'})
    assert response.status_code==201,response.text
    assert response.json()['output'] is None and not response.json()['can_adopt']
    endpoint=f"/api/v1/projects/{round['project_id']}/files/{source['file_id']}"
    assessment=client.get(endpoint+'/lifecycle/assessment')
    assert assessment.status_code==200,assessment.text
    assert 'review_history_retained' in {r['code'] for r in assessment.json()['reasons']}
    assert client.get(endpoint+'/download').status_code==200


def test_new_material_invalidates_only_changed_item_decisions(client):
    from tests.api.test_review_rounds import make_round,edits
    url,round,_,_=make_round(client,'1. Clarify the method.\n2. Clarify the wording.')
    path=url+'/rounds/'+round['id']
    response=client.post(path+'/decisions',json=dict(expected_revision=round['revision'],decisions=[
        dict(item_id=round['items'][0]['id'],action='accept',reason='Use existing material.'),
        dict(item_id=round['items'][1]['id'],action='not_adopt',reason='Retain this independent original expression.'),
    ]),headers={'Idempotency-Key':'before-material'})
    assert response.status_code==201,response.text
    saved=response.json();independent=saved['items'][1]['decision']
    doc=Document();doc.add_paragraph('Synthetic newly supplied existing method material.')
    stream=BytesIO();doc.save(stream)
    source=upload_parsed(client,round['project_id'],stream.getvalue(),'synthetic-new-material.docx')
    items=edits(saved);items[0]['evidence_ids'].append(source['document_version_id'])
    response=client.post(path+'/plans',json=dict(expected_revision=saved['revision'],items=items,
                         materials=[dict(kind='document_version',file_id=source['file_id'],id=source['document_version_id'])]),
                         headers={'Idempotency-Key':'material-only-for-first'})
    assert response.status_code==201,response.text
    updated=response.json()
    assert updated['items'][0]['decision'] is None
    assert updated['items'][1]['decision']==independent


def test_explicit_older_input_retains_synthetic_previous_round_marker(client):
    from tests.api.test_school_exports import published
    path,_,_,manuscript=published(client);base=manuscript['versions'][0]
    url=path+'/reviews'
    body=dict(title='Synthetic school history',project_type='thesis',synthetic=True,
              manuscript=dict(id=manuscript['id'],revision=base['revision'],version_id=base['id']),
              school_confirmation_id=base['snapshot']['thesis_context']['confirmation_id'],
              sources=[dict(role='advisor',label='Advisor',format='text',text='保留明确合成来源。')])
    response=client.post(url+'/rounds',json=body,headers={'Idempotency-Key':'synthetic-history'})
    assert response.status_code==202,response.text
    previous=response.json()
    response=client.post(url+'/rounds',json=next_request(previous,base)|dict(synthetic=False),headers={'Idempotency-Key':'history-next'})
    assert response.status_code==202,response.text
    assert response.json()['synthetic'] is True


def test_school_same_file_reanalysis_publishes_only_receipt_backed_supersession(client,monkeypatch):
    from tests.api.test_thesis_manuscripts import bound
    from tests.api.test_review_science import upload_in_project,mock_local_summary
    from tests.api.test_dataset_analysis import version
    from tests.integration.test_dataset_preparation import run,review,decide,decision
    from tests.helpers.legacy_plans import seed_legacy_plan
    from tests.helpers.analysis_plan import run_plan,review_plan,decide_plan,plan_decision
    from tests.api.test_analysis import selection
    from app.adapters.storage import get_project_store
    from app.domain.tasks.execution import run_task
    from app.domain.dataset_profile.scan import numeric_profile
    from app.domain.manuscripts.service import ManuscriptService
    from app.domain.report_collections.service import scoped_store
    from app.db.database import database_connection
    from app.core.exceptions import StorageError
    import pytest

    monkeypatch.setattr('app.domain.analysis_plan.profile.bounded_numeric_profile',numeric_profile)
    monkeypatch.setattr('app.domain.descriptive.summarize',mock_local_summary)
    project_path,_,request=bound(client)
    file_path,file=upload_in_project(client,project_path.rsplit('/',1)[1],[['Value','Group']]+[[3,'A'] for _ in range(8)])
    response=client.put(file_path+'/analysis-setup',json=selection(expected_revision=0))
    assert response.status_code==200,response.text
    old_setup=response.json()
    summary=dict(n=8,mean=3.0,std=0.0,min=3.0,q1=3.0,median=3.0,q3=3.0,max=3.0,iqr=0.0,warnings=[])
    old=get_project_store().save_analysis_run(project_path.rsplit('/',1)[1],file['id'],dict(
        file_id=file['id'],filename=file['filename'],source_sha256=file['sha256'],setup_revision=old_setup['revision'],
        selection=old_setup['selection'],check=old_setup['check'],numeric_name='Value',group_name='Group',
        engine=dict(id='synthetic_fixture',python_version='synthetic_fixture',std_ddof=1,quantile_method='linear_inclusive'),
        started_at=old_setup['updated_at'],completed_at=old_setup['updated_at'],overall=summary,groups=[dict(label='A',statistics=summary)]))
    request.update(sources=[dict(kind='analysis_run',id=old['id'],file_id=file['id'],role='computed_results')],
                   results_groups=[dict(id='school-data',title='合成结果',analysis_run_ids=[old['id']])])
    request['thesis']['chapter_mappings'][-1]['result_group_ids']=['school-data']
    response=client.post(project_path+'/manuscripts',json=request,headers={'Idempotency-Key':'school-same-file-base'})
    assert response.status_code==201,response.text
    manuscript=response.json();base=manuscript['versions'][0]
    assert run_task(base['task_id'],1)['status']=='succeeded'
    url=project_path+'/reviews'
    response=client.post(url+'/rounds',json=dict(title='合成学校同文件重分析',project_type='thesis',synthetic=True,
        manuscript=dict(id=manuscript['id'],revision=base['revision'],version_id=base['id']),school_confirmation_id=request['thesis']['confirmation_id'],
        sources=[dict(role='advisor',label='导师',format='text',text='请更新统计结果。')]),headers={'Idempotency-Key':'school-same-file-round'})
    assert response.status_code==202,response.text
    assert run_task(response.json()['task']['id'],1)['status']=='succeeded'
    path=url+'/rounds/'+response.json()['id'];round=client.get(path).json()
    raw=version(client,file_path)
    defaults=client.get(file_path+'/dataset-versions/'+raw['id']+'/preparation-defaults',params=dict(header_row=1,data_range='A2:B5')).json()
    response=client.post(file_path+'/dataset-preparations',json=dict(dataset_id=raw['id'],recipe=defaults['recipe']),headers={'Idempotency-Key':'school-same-file-data'})
    assert response.status_code==202,response.text
    task=response.json();assert run(task)['status']=='waiting_confirmation'
    approved=decide(client,task,decision(review(client,task)),'school-same-file-data-confirm').json()
    assert run(approved)['status']=='succeeded'
    prepared=review(client,task)['result']
    response=seed_legacy_plan(client,file_path+'/analysis-plans',json=dict(dataset_id=prepared['id'],expected_setup_revision=old_setup['revision'],selection=old_setup['selection']),headers={'Idempotency-Key':'school-same-file-plan'})
    assert response.status_code==202,response.text
    plan=response.json();assert run_plan(plan)['status']=='waiting_confirmation'
    accepted=decide_plan(client,plan,plan_decision(review_plan(client,plan)),'school-same-file-plan-confirm').json()
    assert run_plan(accepted)['status']=='succeeded'
    response=client.post(path+'/statistics-tasks',json=dict(plan_revision=round['revision'],item_id=round['items'][0]['id'],file_id=file['id'],expected_setup_revision=old_setup['revision']+1),headers={'Idempotency-Key':'school-same-file-statistics'})
    assert response.status_code==202,response.text
    link=response.json();assert run_task(link['task']['id'],1)['status']=='succeeded'
    context=client.get(path+'/revision-context').json()
    paragraphs={p['id']:p for s in base['snapshot']['sections'] for p in s['paragraphs']}
    affected=[dict(kind='paragraph',id=p['id']) for p in context['locations'] if p['kind']=='paragraph' and old['id'] in paragraphs[p['id']].get('source_ids',[])]
    response=client.post(path+'/analysis-adoptions',json=dict(plan_revision=round['revision'],item_id=round['items'][0]['id'],task_link_id=link['id'],supersedes_source_ids=[old['id']],affected_locations=affected,reason='明确采用同一文件新确认范围和计划的结果。'),headers={'Idempotency-Key':'school-same-file-adopt'})
    assert response.status_code==201,response.text
    adopted=response.json();context=client.get(path+'/revision-context').json()
    facts={f['path']:f for f in context['analysis_facts']}
    bindings=[dict(fact_id=facts[p]['id'],value=facts[p]['value']) for p in ('/overall/n','/overall/mean')]
    operations=[dict(item_id=round['items'][0]['id'],kind='replace_paragraph',location_id=p['id'],
        text='n = '+bindings[0]['value']+'; mean = '+bindings[1]['value']+('；见表 1。' if paragraphs[p['id']].get('cross_references') else ''),numeric_bindings=bindings,evidence_ids=[adopted['source']['id']],reason='逐值采用同文件新的计算结果。') for p in affected]
    body=dict(plan_revision=round['revision'],expected_current_revision=context['expected_current_revision'],analysis_adoption_ids=[adopted['id']],operations=operations)
    preview=client.post(path+'/revisions/preview',json=body)
    assert preview.status_code==200,preview.text
    assert preview.json()['can_adopt'],[(i.get('category'),i.get('location'),i.get('message')) for i in [*preview.json()['public_issues'],*preview.json()['constraint_issues']] if i['severity']=='error']
    response=client.post(path+'/revisions',json=body|dict(preview_sha256=preview.json()['preview_sha256']),headers={'Idempotency-Key':'school-same-file-publish'})
    assert response.status_code==202,response.text
    result=run_task(response.json()['task']['id'],1)
    assert result['status']=='succeeded',result
    output=client.get(url+'/revisions/'+response.json()['id']).json()['result']['output']
    task=run_task(output['task_id'],1)
    assert task['status']=='succeeded',task
    root=project_path+'/manuscripts/'+manuscript['id']
    public=client.get(root+'/versions/'+str(output['revision']))
    assert public.status_code==200,public.text
    assert {s['id'] for s in public.json()['snapshot']['sources']}=={old['id'],adopted['source']['id']}
    assert public.json()['snapshot']['thesis_context']==base['snapshot']['thesis_context']
    assert client.get(root+'/versions/'+str(base['revision'])).json()['snapshot']==base['snapshot']
    # The stale original alone has no receipt-backed supersession capability.
    with database_connection() as db:
        owner=db.execute('SELECT user_id FROM manuscript_versions WHERE id=%s',(base['id'],)).fetchone()['user_id']
        service=ManuscriptService(scoped_store(owner,db.config))
        with pytest.raises(StorageError):service.verify_task(db,service.store.require_task(db,base['task_id']),current=True)
    profile=client.get(project_path+'/thesis-profile').json()
    fields={k:profile[k] for k in ('school','department','major','degree_type','education_level','applicable_version')}
    response=client.put(project_path+'/thesis-profile',json=dict(expected_revision=profile['revision'],**(fields|dict(major='Changed independent school configuration'))))
    assert response.status_code==200,response.text
    with database_connection() as db:
        service=ManuscriptService(scoped_store(owner,db.config))
        with pytest.raises(StorageError):service.verify_task(db,service.store.require_task(db,output['task_id']),current=True)
