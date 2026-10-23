"""Fourth batch: explicit school changes preserve fixed parents and sources."""
from copy import deepcopy
from io import BytesIO
from docx import Document
from tests.api.test_analysis import client  # noqa: F401
from tests.api.test_thesis_manuscripts import bound
from tests.api.test_school_templates import post, analysis
from app.domain.tasks.execution import run_task
from app.db.database import database_connection


def base(client):
    path, template, request = bound(client)
    root = post(client, path+'/manuscripts', request, 'base').json()
    assert run_task(root['versions'][0]['task_id'], 1)['status']=='succeeded'
    return path, template, request, root, path+'/manuscripts/'+root['id']


def updated(client,path,template,request,reverse=False):
    from app.domain.school_templates.execution import run_school_template_task
    old=client.get(path+'/school-templates').json()['items'][0]
    response=post(client,path+'/school-templates',template|dict(name='工程更新',previous_version_id=old['id']),'update')
    assert response.status_code==202,response.text
    item=response.json();assert item['previous_version_id']==old['id']
    assert run_school_template_task(item['task']['id'],1,parser=analysis)['status']=='succeeded'
    item=client.get(path+'/school-templates').json()['items'][0]
    chapters=deepcopy(old['confirmation']['chapters']);chapters[0]['title']='更新绪论'
    if reverse:chapters.reverse()
    confirmation=post(client,path+'/school-templates/'+item['id']+'/confirm',dict(analysis_sha256=item['analysis_sha256'],decisions=[dict(conflict_id='conflict-1',option_id='text-2')],chapters=chapters),'confirm-update')
    assert confirmation.status_code==201,confirmation.text
    return confirmation.json()


def test_migrate_old_fixed_parent_and_replay_preserves_original(client):
    path,template,request,root,endpoint=base(client)
    confirmation=updated(client,path,template,request,reverse=True)
    assert post(client,path+'/thesis-template-binding',dict(confirmation_id=confirmation['id'],expected_revision=1)).status_code==200
    url=endpoint+'/versions/1/school-changes'
    body=dict(kind='template_migration',target_confirmation_id=confirmation['id'])
    response=post(client,url,body,'change');assert response.status_code==201,response.text
    change=response.json();assert change['can_adopt']
    assert any(d['kind']=='template_changed' and d['before']=='绪论' and d['after']=='更新绪论' for d in change['differences'])
    assert any(d['kind']=='template_changed' and d['location']=='chapter_order' and '绪论' in d['before'] and '更新绪论' in d['after'] for d in change['differences'])
    assert post(client,url,body,'change').json()==change
    assert len(client.get(endpoint+'/school-changes').json()['items'])==1
    args=dict(mappings=change['mappings'],expected_revision=1,preview_sha256=change['preview_sha256'])
    adopted=post(client,endpoint+'/school-changes/'+change['id']+'/adopt',args,'adopt')
    assert adopted.status_code==201,adopted.text
    version=adopted.json()['versions'][0]
    assert version['parent_version_id']==root['versions'][0]['id']
    assert version['snapshot']['thesis_context']['confirmation_id']==confirmation['id']
    assert 'school_change_provenance' in version['snapshot']
    assert run_task(version['task_id'],1)['status']=='succeeded'
    assert post(client,endpoint+'/school-changes/'+change['id']+'/adopt',args,'adopt').status_code==200
    second=post(client,url,body,'branch').json()
    args=dict(mappings=second['mappings'],expected_revision=2,preview_sha256=second['preview_sha256'])
    branch=post(client,endpoint+'/school-changes/'+second['id']+'/adopt',args,'branch-adopt')
    assert branch.status_code==201,branch.text
    assert branch.json()['versions'][0]['parent_version_id']==root['versions'][0]['id']
    assert run_task(branch.json()['versions'][0]['task_id'],1)['status']=='succeeded'
    assert client.get(endpoint+'/versions/1').json()['snapshot']==root['versions'][0]['snapshot']


def upload_return(client,path,snapshot,edit='补充合成方法说明。'):
    doc=Document()
    for section in snapshot['sections']:
        doc.add_heading(section['title'],1)
        for p in section['paragraphs']:doc.add_paragraph(p['text']+(edit if section['purpose']=='materials_and_methods' else ''))
    stream=BytesIO();doc.save(stream)
    return upload_raw(client,path,stream.getvalue())


def upload_raw(client,path,content):
    from tests.api.test_file_batches import register
    project=dict(id=path.rsplit('/',1)[1])
    upload,batch,item,raw=register(client,project,content,'return.docx',key='return-'+__import__('uuid').uuid4().hex)
    result=client.post(upload+'/'+batch['id']+'/items/'+item['client_id'],files={'file':(item['filename'],raw)}).json()
    with database_connection() as db:
        task=db.execute("SELECT t.* FROM tasks t JOIN file_parse_results r ON r.task_id=t.id WHERE r.file_id=%s",(result['file_id'],)).fetchone()
    from app.domain.tasks.file_parse_execution import run_parse_file_task
    from app.adapters.docx_parser import parse_docx
    assert run_parse_file_task(task['id'],1,parser=parse_docx)['status']=='succeeded'
    with database_connection() as db:parsed=db.execute('SELECT id FROM file_parse_results WHERE file_id=%s',(result['file_id'],)).fetchone()
    return dict(kind='docx_return',file_id=result['file_id'],document_version_id=parsed['id'])


def test_return_edit_adopts_while_bound_profile_changes_and_numbers_block(client):
    path,template,request,root,endpoint=base(client)
    body=upload_return(client,path,root['versions'][0]['snapshot'])
    confirmation=updated(client,path,template,request)
    post(client,path+'/thesis-template-binding',dict(confirmation_id=confirmation['id'],expected_revision=1))
    response=post(client,endpoint+'/versions/1/school-changes',body,'returned')
    assert response.status_code==201,response.text
    change=response.json();assert change['can_adopt'],change['issues']
    assert change['original_download_url'] and any(d['kind']=='text_changed' for d in change['differences'])
    adopted=post(client,endpoint+'/school-changes/'+change['id']+'/adopt',dict(mappings=change['mappings'],expected_revision=1,preview_sha256=change['preview_sha256']),'return-adopt')
    assert adopted.status_code==201,adopted.text
    version=adopted.json()['versions'][0]
    assert version['snapshot']['thesis_context']==root['versions'][0]['snapshot']['thesis_context']
    assert run_task(version['task_id'],1)['status']=='succeeded'
    changed=upload_return(client,path,root['versions'][0]['snapshot'],'数值99 mg。')
    change=post(client,endpoint+'/versions/1/school-changes',changed,'numbers').json()
    assert not change['can_adopt'] and any(d['kind']=='numeric_changed' for d in change['differences'])
    assert client.get(endpoint).json()['current_revision']==2


def test_return_mapping_permission_and_source_fence(client):
    from tests.integration.test_project_ownership import account_client
    path,_,_,root,endpoint=base(client)
    body=upload_return(client,path,root['versions'][0]['snapshot'])
    change=post(client,endpoint+'/versions/1/school-changes',body,'returned').json()
    url=endpoint+'/school-changes/'+change['id']
    missing=post(client,url+'/preview',dict(mappings=[]))
    assert missing.status_code==200 and not missing.json()['can_adopt']
    with account_client('other-changes@example.test')[0] as other:assert other.get(url).status_code==404
    with database_connection(write=True) as db:db.execute('UPDATE files SET sha256=%s WHERE id=%s',('a'*64,body['file_id']))
    latest=client.get(url);assert latest.status_code==200,latest.text
    assert not latest.json()['can_adopt']
    assert post(client,url+'/adopt',dict(mappings=change['mappings'],expected_revision=1,preview_sha256=change['preview_sha256']),'changed-source').status_code==409


def test_real_local_docx_export_return_and_frozen_historical_validation(client,monkeypatch):
    from tests.api.test_school_exports import published
    path,endpoint,_,root=published(client)
    exported=post(client,endpoint+'/versions/1/exports',dict(school=dict(include_toc=False)),'local-word').json()
    assert run_task(exported['task_id'],1)['status']=='succeeded'
    raw=client.get(endpoint+'/exports/'+exported['id']+'/download').content
    doc=Document(BytesIO(raw));paragraph=next(p for p in doc.paragraphs if p.text.startswith('本章正文待补充'))
    paragraph.runs[0].text+='补充可追溯的合成说明。'
    stream=BytesIO();doc.save(stream)
    source=upload_raw(client,path,stream.getvalue())|dict(base_export_id=exported['id'])
    response=post(client,endpoint+'/versions/1/school-changes',source,'word-return')
    assert response.status_code==201,response.text
    change=response.json();assert change['can_adopt'],change['issues']
    response=post(client,endpoint+'/school-changes/'+change['id']+'/adopt',dict(mappings=change['mappings'],expected_revision=1,preview_sha256=change['preview_sha256']),'word-adopt')
    assert response.status_code==201,response.text
    version=response.json()['versions'][0]
    assert run_task(version['task_id'],1)['status']=='succeeded'
    def future(*args,**kwargs):raise AssertionError('Historical validation must not recompile returned body')
    monkeypatch.setattr('app.domain.manuscripts.school_changes.SchoolChangeService.candidate',future)
    assert client.get(endpoint+'/versions/2').json()['snapshot']==version['snapshot']
    assert client.get(endpoint+'/exports/'+exported['id']+'/download').content==raw


def test_reviewed_parent_change_keeps_shared_review_history_contract(client):
    import json,subprocess,shutil
    from pathlib import Path
    path,template,request,root,endpoint=base(client)
    review=post(client,endpoint+'/reviews',dict(base_revision=1),'review').json()
    reviewed=post(client,endpoint+'/reviews/'+review['id']+'/publish',dict(expected_revision=1),'publish-review')
    assert reviewed.status_code==201,reviewed.text
    parent=reviewed.json()['versions'][0]
    assert run_task(parent['task_id'],1)['status']=='succeeded'
    confirmation=updated(client,path,template,request)
    change=post(client,endpoint+'/versions/2/school-changes',dict(kind='template_migration',target_confirmation_id=confirmation['id']),'review-migration').json()
    adopted=post(client,endpoint+'/school-changes/'+change['id']+'/adopt',dict(mappings=change['mappings'],expected_revision=2,preview_sha256=change['preview_sha256']),'review-migrate-adopt')
    assert adopted.status_code==201,adopted.text
    snapshot_=adopted.json()['versions'][0]['snapshot'];history=snapshot_['review_history'][-1]
    assert history==dict(provenance=parent['snapshot']['review_provenance'],review=parent['snapshot']['review'],applies_to_current_text=False,status='historical_requires_recheck')
    frontend=Path(__file__).resolve().parents[3]/'frontend'
    if shutil.which('node') and (frontend/'node_modules/esbuild').exists():
        script="const{buildSync}=require('esbuild');const fs=require('fs');const value=JSON.parse(fs.readFileSync(0,'utf8'));const code=buildSync({stdin:{contents:\"export {validSnapshot} from './src/features/manuscripts/manuscriptTypes.ts'\",resolveDir:process.cwd()},bundle:true,platform:'node',format:'cjs',write:false}).outputFiles[0].text;const m={exports:{}};new Function('module','exports',code)(m,m.exports);if(!m.exports.validSnapshot(value,value.project_id))process.exit(1);"
        checked=subprocess.run(['node','-e',script],input=json.dumps(snapshot_),text=True,capture_output=True,cwd=frontend)
        assert checked.returncode==0,checked.stderr
