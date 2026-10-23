from tests.api.test_analysis import client
from tests.api.test_manuscript_review_v2_bridge import published

def test_review_round_fixed_input_and_public_worker(client):
    root,base,_,_=published(client)
    project=root.split('/')[4];url=f'/api/v1/projects/{project}/reviews'
    result=client.get(url+'/inputs')
    assert result.status_code==200,result.text
    payload=dict(title='离线返修',project_type='sci',manuscript=dict(id=root.split('/')[-1],revision=base['revision'],version_id=base['id']),synthetic=True,sources=[dict(role='reviewer',label='Reviewer 1',format='text',text='1. Clarify the method.\n2. Add experiment to confirm conclusion.')],materials=[])
    response=client.post(url+'/rounds',json=payload,headers={'Idempotency-Key':'round'})
    assert response.status_code==202,response.text
    round=response.json()
    from app.domain.review_rounds.execution import run_review_round_task
    assert run_review_round_task(round['task']['id'],1)['status']=='waiting_input'
    round=client.get(url+'/rounds/'+round['id']).json()
    assert len(round['items'])==2
    assert round['items'][1]['status']=='waiting_material'
    assert all(i['status']!='completed' for i in round['items'])
    summary=client.get(f'/api/v1/projects/{project}/summary').json()['future']['revision']
    assert summary['state']=='available' and summary['current_round']['id']==round['id']
    assert len(summary['opinions'])==2 and summary['pending_materials']
    assert client.post(url+'/rounds',json=payload,headers={'Idempotency-Key':'round'}).status_code==200
    workspace=client.get('/api/v1/tasks/'+round['task']['id']).json()
    assert workspace['input_version']['base_version_id']==base['id']
    from app.db.database import database_connection
    from app.maintenance.review_round_audit import audit_review_rounds
    with database_connection() as db:assert audit_review_rounds(db)[0]['invalid']==0
    import pytest
    from app.core.exceptions import StorageError
    from sqlalchemy.exc import DBAPIError
    with pytest.raises(DBAPIError,match='Immutable review round history'):
        with database_connection(write=True) as db:db.execute('UPDATE review_rounds SET round_number=99 WHERE id=%s',(round['id'],))

def make_round(client,text='1. Please clarify this comment.'):
    root,base,_,_=published(client);project=root.split('/')[4];url=f'/api/v1/projects/{project}/reviews'
    body=dict(title='Synthetic review round',project_type='sci',manuscript=dict(id=root.split('/')[-1],revision=base['revision'],version_id=base['id']),synthetic=False,sources=[dict(role='reviewer',label='Reviewer 1',format='text',text=text)],materials=[])
    response=client.post(url+'/rounds',json=body,headers={'Idempotency-Key':'create-round'})
    assert response.status_code==202,response.text
    from app.domain.tasks.execution import run_task
    task=run_task(response.json()['task']['id'],1)
    return url,client.get(url+'/rounds/'+response.json()['id']).json(),body,base

def test_source_number_display_summary_preserves_saved_identity_and_clarify_pending(client):
    root,base,_,_=published(client);project=root.split('/')[4];url=f'/api/v1/projects/{project}/reviews'
    body=dict(title='Multi-source synthetic',project_type='sci',synthetic=True,manuscript=dict(id=root.split('/')[-1],revision=base['revision'],version_id=base['id']),materials=[],sources=[dict(role='reviewer',label=label,format='text',text='1. Please clarify this comment.') for label in ('Reviewer 1','Reviewer 2')])
    response=client.post(url+'/rounds',json=body,headers={'Idempotency-Key':'multi-source'});assert response.status_code==202,response.text
    from app.domain.tasks.execution import run_task
    run_task(response.json()['task']['id'],1)
    path=url+'/rounds/'+response.json()['id'];r=client.get(path).json()
    identities=[(i['id'],i['number'],i['original_number']) for i in r['items']]
    assert [i['number'] for i in r['items']]==[1,1]
    summary=client.get(f'/api/v1/projects/{project}/summary').json()['future']['revision']
    assert [i['source_label'] for i in summary['opinions']]==['Reviewer 1','Reviewer 2']
    decision=dict(expected_revision=r['revision'],decisions=[dict(item_id=i['id'],action='clarify',reason='已澄清：按既有方法补充表述，不新增分析。') for i in r['items']])
    response=client.post(path+'/decisions',json=decision,headers={'Idempotency-Key':'clarified'});assert response.status_code==201,response.text
    r=response.json();assert all(i['status']=='pending' for i in r['items'])
    assert [(i['id'],i['number'],i['original_number']) for i in r['items']]==identities
    assert [(i['id'],i['number'],i['original_number']) for i in r['versions'][0]['items']]==identities

def edits(round):
    from app.domain.review_rounds.contracts import ItemEdit
    return [{k:i[k] for k in ItemEdit.model_fields} for i in round['items']]

def test_plans_translation_decision_invalidation_common_resume_and_history(client):
    url,r,body,base=make_round(client,'1. Change correlation to causation.')
    assert r['task']['status']=='waiting_confirmation' and r['synthetic']
    item=r['items'][0];original=item['original_text'];identity=r['id'];path=url+'/rounds/'+identity
    workspace=client.get('/api/v1/tasks/'+r['task']['id']+'/workspace')
    assert workspace.status_code==200,workspace.text
    assert workspace.json()['source']['file_id'] is None and workspace.json()['source']['filename']==r['title']
    assert workspace.json()['allowed_actions']==['resume']
    d=dict(expected_revision=r['revision'],decisions=[dict(item_id=item['id'],action='accept',reason='')])
    assert client.post(path+'/decisions',json=d,headers={'Idempotency-Key':'no-reason'}).status_code==422
    d['decisions'][0]['reason']='Explicitly retain association; no scientific causal claim authorized.'
    response=client.post(path+'/decisions',json=d,headers={'Idempotency-Key':'decision'})
    assert response.status_code==201,response.text
    r=response.json();decision=r['items'][0]['decision'];assert decision
    e=edits(r);e[0]['translation']=dict(language='zh-CN',text='把相关改为因果（人工译文）');e[0]['plan']='保留相关表述，列明因果证据缺口。'
    request=dict(expected_revision=r['revision'],items=e)
    response=client.post(path+'/plans',json=request,headers={'Idempotency-Key':'save'})
    assert response.status_code==201,response.text
    r=response.json();assert r['items'][0]['decision'] is None
    assert r['items'][0]['original_text']==original and r['items'][0]['translation']['text'].startswith('把相关')
    assert r['decisions'][0]['id']==decision['id'] and len(r['versions'])==2
    assert r['versions'][0]['items'][0]['translation'] is None
    assert client.post(path+'/plans',json=request,headers={'Idempotency-Key':'save'}).status_code==200
    assert client.post(path+'/plans',json=request,headers={'Idempotency-Key':'stale'}).status_code==409
    d['expected_revision']=r['revision'];d['decisions'][0].update(action='not_adopt',reason='缺少因果依据，明确不采纳。')
    r=client.post(path+'/decisions',json=d,headers={'Idempotency-Key':'final-decision'}).json()
    assert r['items'][0]['status']=='not_adopted'
    task=r['task'];resume=dict(operation='resume',wait_id=r['wait']['id'],expected_task_revision=task['revision'],input_version=task['input_version'])
    response=client.post('/api/v1/tasks/'+task['id']+'/resume',json=resume,headers={'Idempotency-Key':'resume'})
    assert response.status_code==202,response.text
    from app.domain.tasks.execution import run_task
    assert run_task(task['id'],response.json()['revision'])['status']=='succeeded'
    assert client.post('/api/v1/tasks/'+task['id']+'/resume',json=resume,headers={'Idempotency-Key':'resume'}).json()==response.json()
    assert run_task(task['id'],response.json()['revision']) is None
    r=client.get(path).json();assert len(r['items'])==1 and len(r['versions'])==2
    assert r['items'][0]['status']=='not_adopted'

def test_material_wait_cannot_fake_completion_and_foreign_scopes(client):
    url,r,body,base=make_round(client,'1. Add an independent validation experiment.')
    task=r['task'];resume=dict(operation='resume',wait_id=r['wait']['id'],expected_task_revision=task['revision'],input_version=task['input_version'])
    response=client.post('/api/v1/tasks/'+task['id']+'/resume',json=resume,headers={'Idempotency-Key':'missing-resume'})
    assert response.status_code==202,response.text
    from app.domain.tasks.execution import run_task
    assert run_task(task['id'],response.json()['revision'])['status']=='waiting_input'
    fresh=client.get(url+'/rounds/'+r['id']).json()
    assert len(fresh['items'])==1 and fresh['items'][0]['status']=='waiting_material' and len(fresh['versions'])==1
    from tests.api.test_projects import create_project
    other=create_project(client);foreign=f"/api/v1/projects/{other['id']}/reviews"
    assert client.get(foreign+'/rounds/'+r['id']).status_code==404
    assert client.post(foreign+'/rounds',json=body,headers={'Idempotency-Key':'foreign'}).status_code==404
    e=edits(fresh);e[0]['locations']=[dict(kind='paragraph',id='invented')]
    assert client.post(url+'/rounds/'+r['id']+'/plans',json=dict(expected_revision=fresh['revision'],items=e),headers={'Idempotency-Key':'fake-location'}).status_code==422

def test_lossless_segments_and_finite_risks(client):
    from app.domain.review_rounds.planning import split_text,propose,detect_conflicts,update_status
    text='1. Keep limitations.\n2. Remove limitations.\n3. Change correlation to causation.'
    spans=split_text(text)
    assert ''.join(s['text'] for s in spans)==text
    items=[propose(dict(id='source'),n,s) for n,s in enumerate(spans,1)]
    detect_conflicts(items);update_status(items)
    assert items[0]['conflicting_item_ids']==[items[1]['id']]
    assert 'conclusion_change' in items[2]['flags']
    from app.core.exceptions import StorageError
    import pytest
    with pytest.raises(StorageError):split_text(text,[dict(start=0,end=5,original_number='')])

def upload_parsed(client,project,raw,name):
    from tests.api.test_file_batches import register
    from app.domain.tasks.execution import run_task
    base,batch,item,content=register(client,dict(id=project),raw,name,key='upload-'+name)
    response=client.post(base+'/'+batch['id']+'/items/'+item['client_id'],files={'file':(name,content)})
    assert response.status_code==202,response.text
    saved=response.json()
    from app.domain.tasks.file_parse_execution import run_parse_file_task
    if name.endswith('.docx'):
        from app.adapters.docx_parser import parse_docx
        parser=parse_docx
    else:
        from app.adapters.pdf_parser import parse_pdf
        parser=parse_pdf
    result=run_parse_file_task(saved['task']['id'],1,parser=parser);assert result['status']=='succeeded',result
    version=client.get(f"/api/v1/projects/{project}/files/{saved['file_id']}/document-versions").json()['items'][0]
    return dict(role='reviewer',label='Reviewer document',format='document',file_id=saved['file_id'],document_version_id=version['id'])

def test_docx_pdf_fixed_sources_and_blocked_unsupported_content(client):
    from tests.api.test_projects import create_project
    from docx import Document
    from io import BytesIO
    from tests.helpers.pdf_samples import pdf_bytes
    from zipfile import ZipFile
    project=create_project(client)['id'];url=f'/api/v1/projects/{project}/reviews/sources/preview'
    doc=Document();doc.add_paragraph('1. Clarify the method.');doc.add_table(rows=1,cols=1).cell(0,0).text='2. Add validation experiment.'
    buf=BytesIO();doc.save(buf)
    selection=upload_parsed(client,project,buf.getvalue(),'opinions.docx')
    response=client.post(url,json=selection)
    assert response.status_code==200,response.text
    assert response.json()['supported'] and 'Clarify the method' in response.json()['original_text'] and 'validation experiment' in response.json()['original_text']
    assert response.json()['segments'][0]['location']
    selection_pdf=upload_parsed(client,project,pdf_bytes([b'BT /F1 12 Tf 20 350 Td (1. Clarify methods.) Tj ET']),'opinions.pdf')
    result=client.post(url,json=selection_pdf);assert result.status_code==200,result.text
    assert result.json()['supported'] and result.json()['segments'][0]['location']['page_number']==1
    scanned=upload_parsed(client,project,pdf_bytes([b'']),'scanned.pdf')
    result=client.post(url,json=scanned);assert result.status_code==200,result.text
    assert not result.json()['supported'] and result.json()['warnings']
    original=BytesIO(buf.getvalue());modified=BytesIO()
    with ZipFile(original) as old,ZipFile(modified,'w') as new:
        for name in old.namelist():
            content=old.read(name)
            if name=='word/document.xml':content=content.replace(b'</w:body>',b'<w:p><w:commentRangeStart w:id="0"/></w:p></w:body>')
            new.writestr(name,content)
    unsupported=upload_parsed(client,project,modified.getvalue(),'commented.docx')
    result=client.post(url,json=unsupported);assert result.status_code==200,result.text
    assert not result.json()['supported'] and result.json()['warnings']
    foreign=create_project(client)['id']
    assert client.post(f'/api/v1/projects/{foreign}/reviews/sources/preview',json=selection).status_code in (404,409)

def test_thesis_fixed_school_confirmation_roles_and_conflicts(client):
    from tests.api.test_school_exports import published as school_published
    path,base,_,root=school_published(client);version=root['versions'][0];url=path+'/reviews'
    confirmation=version['snapshot']['thesis_context']['confirmation_id']
    from app.domain.tasks.execution import run_task
    newer=dict(version['input'],expected_revision=1,title='较新学校稿件')
    response=client.post(base+'/versions',json=newer,headers={'Idempotency-Key':'newer-school'})
    assert response.status_code==201,response.text
    assert run_task(response.json()['versions'][0]['task_id'],1)['status']=='succeeded'
    available=client.get(url+'/inputs').json()['manuscripts'][0]
    assert available['current_revision']==2 and all(v['status']=='published' for v in available['versions'])
    body=dict(title='合成学校返修',project_type='thesis',manuscript=dict(id=root['id'],revision=1,version_id=version['id']),synthetic=True,school_confirmation_id=confirmation,sources=[dict(role='advisor',label='导师',format='text',text='请保留现有局限性说明。'),dict(role='evaluator',label='评阅',format='text',text='请删除局限性描述。'),dict(role='defense',label='答辩',format='text',text='请把相关改为导致。')],materials=[])
    response=client.post(url+'/rounds',json=body,headers={'Idempotency-Key':'school'})
    assert response.status_code==202,response.text
    from app.domain.tasks.execution import run_task
    assert run_task(response.json()['task']['id'],1)['status']=='waiting_confirmation'
    round=client.get(url+'/rounds/'+response.json()['id']).json()
    assert round['manuscript']['revision']==1 and round['manuscript']['version_id']==version['id']
    assert round['constraints']['school_confirmation_id']==confirmation and round['constraints']['journal_target_id'] is None
    assert 'conflict' in round['items'][0]['flags'] and 'conclusion_change' in round['items'][2]['flags']
    assert client.post(url+'/rounds',json=body|dict(school_confirmation_id=None),headers={'Idempotency-Key':'missing-school'}).status_code==422
    assert client.post(url+'/rounds',json=body|dict(journal_target_id='invented'),headers={'Idempotency-Key':'bad-journal'}).status_code==422
    bad=dict(body);bad['sources']=[dict(role='reviewer',label='bad role',format='text',text='bad')]
    assert client.post(url+'/rounds',json=bad,headers={'Idempotency-Key':'bad-role'}).status_code==422

def test_completed_plan_reopens_wait_and_cannot_hide_frozen_risk(client):
    url,r,base_body,_=make_round(client,'1. Clarify the method.')
    assert r['task']['status']=='succeeded'
    path=url+'/rounds/'+r['id'];e=edits(r);e[0]['flags']=['conflict'];e[0]['plan']='需要用户在两种处理之间明确选择。'
    response=client.post(path+'/plans',json=dict(expected_revision=r['revision'],items=e),headers={'Idempotency-Key':'reopen'})
    assert response.status_code==201,response.text
    r=response.json();assert r['items'][0]['status']=='waiting_confirmation'
    assert r['task']['status']=='waiting_confirmation' and r['wait']['kind']=='review_round_confirmation'
    # A completed initial plan must not make a later important decision unreachable.
    assert client.get('/api/v1/tasks/'+r['task']['id']+'/workspace').json()['allowed_actions']==['resume']
    from app.domain.review_rounds.planning import propose,split_text
    # Original risks are enforced in the API save, not just proposal output.
    url2=url;body=dict(base_body,sources=[dict(role='reviewer',label='Reviewer 2',format='text',text='1. Change correlation to causation.')])
    response=client.post(url+'/rounds',json=body,headers={'Idempotency-Key':'risk-round'});assert response.status_code==202,response.text
    from app.domain.tasks.execution import run_task
    run_task(response.json()['task']['id'],1)
    risk=client.get(url+'/rounds/'+response.json()['id']).json()
    edit=edits(risk);edit[0]['flags']=[]
    response=client.post(url2+'/rounds/'+risk['id']+'/plans',json=dict(expected_revision=risk['revision'],items=edit),headers={'Idempotency-Key':'hide-risk'})
    assert response.status_code==201,response.text
    assert 'conclusion_change' in response.json()['items'][0]['flags']
    assert response.json()['items'][0]['status']=='waiting_confirmation'


def test_cross_owner_round_task_and_original_download_are_hidden(client):
    url,r,body,base=make_round(client,'1. Clarify the method.')
    project=r['project_id'];file=base['snapshot']['sources'][0]['file_id']
    from docx import Document
    from io import BytesIO
    stream=BytesIO();document=Document();document.add_paragraph('1. Clarify the method.');document.save(stream)
    selection=upload_parsed(client,project,stream.getvalue(),'owner-opinions.docx')
    from app.auth.service import create_user
    create_user('other-review@paperassist.local','other-isolated-password')
    response=client.post('/api/v1/auth/login',json=dict(email='other-review@paperassist.local',password='other-isolated-password'))
    assert response.status_code==200,response.text
    client.headers['X-CSRF-Token']=response.json()['csrf_token']
    assert client.get(url+'/rounds/'+r['id']).status_code==404
    assert client.post(url+'/sources/preview',json=selection).status_code==404
    assert client.get('/api/v1/tasks/'+r['task']['id']+'/workspace').status_code==404
    assert client.get(f'/api/v1/projects/{project}/files/{file}/download').status_code==404
    assert client.post(url+'/rounds',json=body,headers={'Idempotency-Key':'foreign-owner'}).status_code==404

def test_dispatcher_selects_queued_review_round_outbox(client):
    url,r,body,_=make_round(client,'1. Clarify methods.')
    from app.db.database import database_connection
    with database_connection(write=True) as db:db.execute('UPDATE task_outbox SET published_at=clock_timestamp() WHERE published_at IS NULL')
    response=client.post(url+'/rounds',json=body,headers={'Idempotency-Key':'queued-dispatch-round'})
    assert response.status_code==202,response.text
    task=response.json()['task'];assert task['status']=='queued'
    from app.workers.dispatcher import dispatch_once
    messages=[]
    assert dispatch_once(publish=messages.append)==1
    assert len(messages)==1 and messages[0]['task_id']==task['id'] and messages[0]['task_revision']==1
    assert set(messages[0])=={'id','task_id','task_revision'}
    from app.domain.tasks.execution import run_task
    assert run_task(messages[0]['task_id'],messages[0]['task_revision'])['status']=='succeeded'
