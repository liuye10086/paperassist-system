"""C: real isolated tasks and synthetic documents, no external transports."""
from copy import deepcopy
from io import BytesIO

import pytest
from docx import Document

from app.db.database import database_connection
from tests.api.test_admin_operations import accounts, admin
from tests.api.test_admin_configurations import baseline, saved, validated, switch, ROOT
from tests.api.test_journals import setup, profile, request, post


def activate(client, version, revision=0, key='first', operation='activate'):
    validated(client, version)
    response = switch(client, version, revision, key, operation)
    assert response.status_code == 200, response.text


def binding(task_id, kind='task'):
    with database_connection() as db:
        row = db.execute('SELECT version_id FROM admin_configuration_bindings WHERE kind=%s AND object_id=%s', (kind, task_id)).fetchone()
        return row['version_id'] if row else None


def preserve_seeded_budgets(client,task):
    from app.domain.model_usage.service import ModelUsageService
    from tests.helpers.tasks import task_owner
    owner=task_owner(client);service=ModelUsageService(owner)
    for scope,key in [('user',owner),('project',task['project_id']),('task',task['id'])]:
        with database_connection() as db:
            old=db.execute('SELECT id FROM model_budgets WHERE user_id=%s AND scope_type=%s AND scope_key=%s',(owner,scope,key)).fetchone()
        if not old:service.set_budget(scope,key,5000000,0)


def test_journal_defaults_intent_replay_retry_and_future_rollback(admin):
    from app.domain.journals.execution import run_journal_task
    one = baseline(admin); activate(admin, one)
    base = setup(admin)
    for i in range(5):
        assert post(admin, base+'/profiles', profile(name='synthetic journal '+str(i)), 'p'+str(i)).status_code == 201
    body = request()
    old = post(admin, base+'/recommendations', body, 'old').json()
    assert binding(old['task']['id']) == one['id']
    with database_connection() as db:
        frozen = dict(db.execute('SELECT input_snapshot,input_digest FROM tasks WHERE id=%s', (old['task']['id'],)).fetchone())
    two = saved(admin, one, 'journals', 'default_count', 2)
    two = saved(admin, two, 'journals', 'max_count', 2)
    activate(admin, two, 1, 'two')
    config = admin.get(base+'/config').json()
    assert (config['default_limit'], config['max_limit']) == (2, 2)
    assert config['configuration']['version_id'] == two['id']
    assert post(admin, base+'/recommendations', body, 'old').json() == old
    assert post(admin, base+'/recommendations', body|{'limit': 3}, 'old').status_code == 409
    assert post(admin, base+'/recommendations', body|{'limit': 3}, 'too-many').status_code == 422
    new = post(admin, base+'/recommendations', body, 'new').json()
    assert new['input_snapshot']['research']['limit'] == 2
    assert binding(new['task']['id']) == two['id']
    def fail(_): raise RuntimeError('synthetic failure')
    failed = run_journal_task(old['task']['id'], 1, recommender=fail)
    retry = post(admin, '/api/v1/tasks/'+failed['id']+'/resume', dict(operation='retry',wait_id=None,expected_task_revision=failed['revision'],input_version=failed['input_version']), 'retry')
    assert retry.status_code == 202, retry.text
    assert run_journal_task(old['task']['id'], retry.json()['revision'])['status'] == 'succeeded'
    assert run_journal_task(new['task']['id'], 1)['status'] == 'succeeded'
    assert len(admin.get(base+'/recommendations/'+old['id']).json()['result']['candidates']) == 3
    assert len(admin.get(base+'/recommendations/'+new['id']).json()['result']['candidates']) == 2
    activate(admin, one, 2, 'back', 'rollback')
    future = post(admin, base+'/recommendations', body, 'future').json()
    assert future['input_snapshot']['research']['limit'] == 3 and binding(future['task']['id']) == one['id']
    with database_connection() as db:
        assert dict(db.execute('SELECT input_snapshot,input_digest FROM tasks WHERE id=%s', (old['task']['id'],)).fetchone()) == frozen


@pytest.mark.parametrize('text', ['Add an independent experiment.', 'Statistic experiment and conclusion.', 'Please supply a reference.', 'Unclear new phrase.'])
def test_editable_categories_cannot_remove_installed_safety(text):
    from app.domain.review_rounds.planning import propose
    from app.domain.admin_configuration.catalog import defaults
    rules = deepcopy(defaults()[0]['review'])
    rules['keywords'] = {k: [] for k in rules['keywords']}
    rules['keywords']['expression'] = [text.lower()]
    item = propose({'id':'source'}, 1, {'text':text,'original_number':'1','location':{}}, rules=rules)
    assert item['category'] == 'expression'
    if 'experiment' in text.lower(): assert any(n['kind']=='experiment' for n in item['needs'])
    if 'reference' in text.lower(): assert any(n['kind']=='literature' for n in item['needs'])
    if 'conclusion' in text.lower(): assert 'conclusion_change' in item['flags']
    if text.startswith('Unclear'): assert 'ambiguous' in item['flags']


def synthetic_alias_docx(heading=False):
    document = Document()
    p = document.add_paragraph('Synthetic Chapter')
    if heading: p.style = 'Heading 2'
    p.paragraph_format.space_after = 0
    document.add_paragraph('Synthetic body example with enough characters for an ordinary paragraph.')
    result = BytesIO(); document.save(result)
    return result.getvalue()


def test_school_alias_uses_actual_xml_properties_and_conflicts_survive_supplements():
    from app.adapters.school_template_parser import inspect_school_template, merge_template_notes
    raw = synthetic_alias_docx()
    before = inspect_school_template(raw, 'synthetic.docx')
    after = inspect_school_template(raw, 'synthetic.docx', heading_aliases=[{'text':'Synthetic Chapter','role':'heading1'}])
    assert not before['chapters'] and after['chapters'][0]['level'] == 1
    assert after['chapters'][0]['source']['text'] == 'Synthetic Chapter'
    assert after['chapters'][0]['source']['part'] == 'word/document.xml'
    other = inspect_school_template(synthetic_alias_docx(True), 'synthetic.docx', heading_aliases=[{'text':'Synthetic Chapter','role':'heading1'}])
    assert any(c.get('kind') == 'heading_mapping' for c in other['conflicts'])
    assert merge_template_notes(other, [])['conflicts'] == other['conflicts']


def test_review_round_defaults_documents_and_children_keep_ancestor(admin,monkeypatch):
    from tests.api.test_review_documents import completed_revision
    from app.domain.tasks.execution import run_task
    monkeypatch.setattr('tests.api.test_manuscript_review_v2_bridge.budgets',preserve_seeded_budgets)
    one = baseline(admin); activate(admin, one)
    url, run = completed_revision(admin)
    assert binding(run['task']['id']) == one['id']
    assert binding(run['result']['output']['task_id']) == one['id']
    two = saved(admin, one, 'review', 'default_language', 'en'); activate(admin, two, 1, 'en')
    round = admin.get(url+'/rounds/'+run['round_id']).json()
    assert round['document_defaults'] == {'language':'zh-CN','document_version':'review-documents-v1'}
    endpoint = url+'/revisions/'+run['id']+'/documents'
    first = post(admin, endpoint, {'kind':'comparison'}, 'default-doc')
    assert first.status_code == 201, first.text
    assert first.json()['language'] == 'zh-CN'
    assert post(admin, endpoint, {'kind':'comparison'}, 'default-doc').json() == first.json()
    assert post(admin, endpoint, {'kind':'comparison','language':'en'}, 'default-doc').status_code == 409
    explicit = post(admin, endpoint, {'kind':'comparison','language':'en'}, 'explicit-doc')
    assert explicit.status_code == 201 and explicit.json()['language'] == 'en'
    exported = post(admin, url+'/documents/'+first.json()['id']+'/exports', {}, 'export').json()
    assert binding(exported['task_id']) == one['id']
    assert run_task(exported['task_id'], 1)['status'] == 'succeeded'


def simple_round(client, *, key='round', previous=None):
    from tests.api.test_journal_format import manuscript
    path,m=manuscript(client);version=m['versions'][0]
    url=path.removesuffix('/manuscripts')+'/reviews'
    body=dict(title='Synthetic rule review',project_type='sci',synthetic=True,
        manuscript=dict(id=m['id'],revision=version['revision'],version_id=version['id']),
        sources=[dict(role='reviewer',label='Synthetic reviewer',format='text',text='1. Custom-marker experiment conclusion.')],materials=[])
    result=post(client,url+'/rounds',body,key)
    assert result.status_code==202,result.text
    return url,result.json(),body


def test_round_new_current_waiting_safety_notes_and_legacy_default(admin):
    from app.domain.tasks.execution import run_task
    from tests.api.test_review_rounds import edits
    one=baseline(admin)
    bundle=deepcopy(one['bundle']);bundle['review']['category_order'].remove('expression');bundle['review']['category_order'].insert(0,'expression')
    bundle['review']['keywords']['expression']=['custom-marker'];bundle['review']['handling_notes']['expression']='Synthetic v1 instruction'
    one=admin.post(ROOT+'/'+one['id']+'/save',json={'bundle':bundle}).json();activate(admin,one)
    url,old,body=simple_round(admin)
    bundle=deepcopy(one['bundle']);bundle['review']['keywords']['expression']=[];bundle['review']['keywords']['format']=['custom-marker']
    bundle['review']['category_order'].remove('format');bundle['review']['category_order'].insert(0,'format')
    bundle['review']['handling_notes']['format']='Synthetic v2 instruction';bundle['review']['default_language']='en'
    two=admin.post(ROOT+'/'+one['id']+'/save',json={'bundle':bundle}).json();activate(admin,two,1,'v2')
    assert run_task(old['task']['id'],1)['status']=='waiting_input'
    old=admin.get(url+'/rounds/'+old['id']).json()
    assert old['items'][0]['category']=='expression' and 'Synthetic v1 instruction' in old['items'][0]['plan']
    weakened=edits(old);weakened[0].update(category='format',needs=[],flags=[],conflicting_item_ids=[])
    changed=post(admin,url+'/rounds/'+old['id']+'/plans',dict(expected_revision=old['revision'],items=weakened),'weaken')
    assert changed.status_code==201,changed.text
    assert changed.json()['items'][0]['status']=='waiting_material'
    assert 'conclusion_change' in changed.json()['items'][0]['flags']
    body['previous_round_id']=old['id']
    new=post(admin,url+'/rounds',body,'new-round').json()
    assert binding(new['id'],'review_round')==two['id']
    assert run_task(new['task']['id'],1)['status']=='waiting_input'
    new=admin.get(url+'/rounds/'+new['id']).json()
    assert new['items'][0]['category']=='format' and 'Synthetic v2 instruction' in new['items'][0]['plan']
    assert new['document_defaults']['language']=='en' and old['document_defaults']['language']=='zh-CN'


def test_legacy_complete_journal_and_school_are_executable_after_activation(admin):
    from app.domain.journals.execution import run_journal_task
    from app.domain.school_templates.execution import run_school_template_task
    from tests.api.test_school_templates import setup as school_setup,analysis
    base=setup(admin);post(admin,base+'/profiles',profile(),'p')
    old=post(admin,base+'/recommendations',request(),'old').json()
    school_path,body=school_setup(admin);school=post(admin,school_path+'/school-templates',body,'school').json()
    one=baseline(admin);one=saved(admin,one,'journals','default_count',1)
    one=saved(admin,one,'schools','heading_aliases',[{'text':'Synthetic Chapter','role':'heading1'}])
    activate(admin,one)
    assert binding(old['task']['id']) is None and binding(school['task']['id']) is None
    assert run_journal_task(old['task']['id'],1)['status']=='succeeded'
    assert run_school_template_task(school['task']['id'],1,parser=analysis)['status']=='succeeded'
    assert admin.get(base+'/recommendations/'+old['id']).json()['input_snapshot']['research']['limit']==3


def test_school_bound_parser_retry_and_confirmation_inherit(admin):
    from app.domain.school_templates.execution import run_school_template_task
    from app.core.exceptions import StorageError
    from app.adapters.school_template_parser import inspect_school_template
    from tests.api.test_file_batches import register
    project=admin.post('/api/v1/projects',json=dict(name='synthetic alias',research_topic='synthetic',project_type='thesis')).json()
    raw=synthetic_alias_docx(True)
    upload,batch,item,_=register(admin,project,raw,'synthetic.docx','docx')
    uploaded=admin.post(upload+'/'+batch['id']+'/items/'+item['client_id'],files={'file':('synthetic.docx',raw)}).json()
    endpoint=f"/api/v1/projects/{project['id']}/school-templates"
    one=baseline(admin);one=saved(admin,one,'schools','heading_aliases',[{'text':'Synthetic Chapter','role':'heading1'}]);activate(admin,one)
    body=dict(name='synthetic template',school='synthetic school',template_file_id=uploaded['file_id'])
    old=post(admin,endpoint,body,'old').json()
    assert binding(old['task']['id'])==one['id']
    def fail(*args,**kwargs):raise StorageError('parse_timeout','synthetic',422)
    assert run_school_template_task(old['task']['id'],1,parser=fail)['status']=='failed'
    two=saved(admin,one,'schools','heading_aliases',[]);activate(admin,two,1,'no-alias')
    retry=post(admin,endpoint+'/'+old['id']+'/retry',{},'retry').json()
    assert run_school_template_task(old['task']['id'],retry['task']['revision'],parser=inspect_school_template)['status']=='succeeded'
    parsed=admin.get(endpoint).json()['items'][0];conflicts=parsed['analysis']['conflicts']
    assert any(c.get('kind')=='heading_mapping' for c in conflicts)
    confirm=dict(analysis_sha256=parsed['analysis_sha256'],chapters=[{k:c[k] for k in ('id','title','level')} for c in parsed['analysis']['chapters']],decisions=[])
    assert post(admin,endpoint+'/'+old['id']+'/confirm',confirm,'missing').status_code==409
    confirm['decisions']=[dict(conflict_id=c['id'],option_id='alias' if c.get('kind')=='heading_mapping' else c['options'][0]['id']) for c in conflicts]
    confirmed=post(admin,endpoint+'/'+old['id']+'/confirm',confirm,'confirm')
    assert confirmed.status_code==201,confirmed.text
    assert binding(confirmed.json()['id'],'school_confirmation')==one['id']
    assert admin.get(ROOT+'/catalog').json()['installed']['supported_schools']==[]


def test_rejected_heading_alias_does_not_become_confirmed_style_rules(admin):
    from app.domain.school_templates.execution import run_school_template_task
    from app.adapters.school_template_parser import inspect_school_template
    from tests.api.test_file_batches import register
    one=baseline(admin);one=saved(admin,one,'schools','heading_aliases',[dict(text='Synthetic Chapter',role='heading1')]);activate(admin,one)
    project=admin.post('/api/v1/projects',json=dict(name='synthetic original role',research_topic='synthetic',project_type='thesis')).json()
    raw=synthetic_alias_docx(True)
    upload,batch,item,_=register(admin,project,raw,'synthetic.docx')
    file=admin.post(upload+'/'+batch['id']+'/items/'+item['client_id'],files={'file':('synthetic.docx',raw)}).json()
    path=f"/api/v1/projects/{project['id']}/school-templates"
    template=post(admin,path,dict(name='Synthetic role conflict',school='Synthetic school',template_file_id=file['file_id']),'parse').json()
    assert run_school_template_task(template['task']['id'],1,parser=inspect_school_template)['status']=='succeeded'
    analysis=admin.get(path).json()['items'][0]
    role=next(r for r in analysis['analysis']['rules'] if r['category']=='heading_mapping')
    assert any(r.get('alias_rule_id')==role['id'] for r in analysis['analysis']['rules'])
    chapters=[{k:c[k] for k in ('id','title','level')} for c in analysis['analysis']['chapters']]
    choices=[dict(conflict_id=c['id'],option_id='original' if c.get('kind')=='heading_mapping' else c['options'][0]['id']) for c in analysis['analysis']['conflicts']]
    body=dict(analysis_sha256=analysis['analysis_sha256'],chapters=chapters,decisions=choices)
    assert post(admin,path+'/'+template['id']+'/confirm',body,'mismatch').status_code==422
    next(c for c in chapters if c['id']==role['mapping']['chapter_id'])['level']=2
    confirmed=post(admin,path+'/'+template['id']+'/confirm',body,'original')
    assert confirmed.status_code==201,confirmed.text
    assert not any(r.get('alias_rule_id')==role['id'] for r in confirmed.json()['rules'])
    assert admin.get(path).json()['items'][0]['analysis']==analysis['analysis']


def test_journal_safe_source_updates_do_not_claim_research_verification(admin):
    one=baseline(admin);activate(admin,one)
    base=setup(admin)
    two=saved(admin,one,'journals','update_result','failed')
    two=saved(admin,two,'journals','update_note','Synthetic manual source review failed')
    activate(admin,two,1,'record-failure')
    config=admin.get(base+'/config').json()
    assert config['updates']['update_result']=='failed'
    assert config['source_policy']['require_verification'] is True
    assert config['source_directory'] and all(s['real_connectivity']=='not_tested' for s in config['source_directory'])
    assert 'source_text' not in str(config['source_directory'])


def test_binding_failure_rolls_back_task_object_and_outbox(admin,monkeypatch):
    one=baseline(admin);activate(admin,one)
    base=setup(admin)
    def counts():
        with database_connection() as db:
            return {table:db.execute('SELECT count(*) AS n FROM '+table).fetchone()['n'] for table in
                    ('journal_recommendations','tasks','task_outbox','admin_configuration_bindings')}
    before=counts()
    def fail(*args,**kwargs):
        from app.core.exceptions import StorageError
        raise StorageError('storage_unavailable','Synthetic atomic binding failure',503)
    monkeypatch.setattr('app.domain.journals.service.bind',fail)
    assert post(admin,base+'/recommendations',request(),'atomic').status_code==503
    assert counts()==before


def test_target_check_adaptation_and_public_exports_use_selected_ancestor(admin,monkeypatch):
    from app.domain.tasks.execution import run_task
    from tests.api.test_journal_adaptation import adaptation_ready
    monkeypatch.setattr('tests.api.test_manuscript_review_v2_bridge.budgets',preserve_seeded_budgets)
    one=baseline(admin);activate(admin,one)
    root,base,version,body,_,_=adaptation_ready(admin)
    old_check=admin.get(base+'/checks/'+body['check_id']).json()
    assert binding(body['target_id'],'journal_confirmation')==one['id']
    two=saved(admin,one,'journals','update_note','Synthetic v2 manual update');activate(admin,two,1,'v2')
    check=post(admin,base+'/checks',dict(target_id=body['target_id'],manuscript=dict(id=body['manuscript']['id'],revision=body['manuscript']['revision'])),'new-check').json()
    assert binding(check['task']['id'])==two['id']
    assert check['input_snapshot']['target']==old_check['input_snapshot']['target']
    assert run_task(check['task']['id'],1)['status']=='succeeded'
    body['check_id']=check['id']
    preview=admin.post(base+'/adaptations/preview',json=body).json()
    three=saved(admin,two,'journals','update_note','Synthetic v3 manual update');activate(admin,three,2,'v3')
    adopted=post(admin,base+'/adaptations',body|dict(expected_current_revision=preview['expected_current_revision'],preview_sha256=preview['preview_sha256'],selected_action_ids=[a['id'] for a in preview['actions']]),'adopt')
    assert adopted.status_code==202,adopted.text
    adopted=adopted.json();assert binding(adopted['task']['id'])==two['id']
    assert run_task(adopted['task']['id'],1)['status']=='succeeded'
    result=admin.get(base+'/adaptations/'+adopted['id']).json()['result']['output']
    assert binding(result['task_id'])==two['id']
    assert run_task(result['task_id'],1)['status']=='succeeded'
    exported=post(admin,root+'/versions/'+str(result['revision'])+'/exports',{},'docx').json()
    assert binding(exported['task_id'])==two['id']


def test_safety_is_idempotent_and_retains_conflicts_and_pending_material():
    from app.domain.review_rounds.planning import propose,enforce_safety,update_status,installed_rules
    rules=installed_rules();rules['keywords']={k:[] for k in rules['keywords']};rules['keywords']['expression']=['limitation','experiment']
    texts=['Delete limitation.','Keep limitation.','Experiment is needed.']
    items=[propose({'id':str(i)},1,dict(text=text,original_number='',location={}),rules=rules) for i,text in enumerate(texts)]
    enforce_safety(items);before=deepcopy(items);enforce_safety(items)
    assert items==before
    assert all('conflict' in item['flags'] for item in items[:2])
    items[-1]['decision']=dict(action='accept',reason='Synthetic material supplied')
    update_status(items)
    assert items[-1]['status']=='waiting_material'


def test_school_mapping_original_choice_matches_confirmed_chapters(admin):
    # The direct parser case makes source attributes and mapping decisions visible.
    from app.adapters.school_template_parser import inspect_school_template
    analysis=inspect_school_template(synthetic_alias_docx(True),'synthetic.docx',heading_aliases=[{'text':'Synthetic Chapter','role':'body'}])
    conflict=next(c for c in analysis['conflicts'] if c.get('kind')=='heading_mapping')
    rule=next(r for r in analysis['rules'] if r['id']==conflict['rule_id'])
    assert rule['mapping']['observed_level']==2
    assert rule['mapping']['observed_role']=='heading2'
    assert not any(c['id']==rule['mapping']['chapter_id'] for c in analysis['chapters'])


@pytest.mark.skipif(__import__('sys').platform!='linux',reason='Existing isolated parser requires Linux')
def test_school_alias_crosses_actual_isolated_transport():
    from app.adapters.school_template_process import inspect_school_template_isolated
    analysis=inspect_school_template_isolated(synthetic_alias_docx(),'synthetic.docx',heading_aliases=[{'text':'Synthetic Chapter','role':'heading1'}])
    assert analysis['chapters'][0]['title']=='Synthetic Chapter' and analysis['chapters'][0]['level']==1


def test_legacy_missing_school_contract_blocks_only_related_fields(admin,accounts):
    from uuid import uuid4
    from datetime import datetime,timezone
    from psycopg.types.json import Jsonb
    from app.domain.literature.service import digest
    one=baseline(admin);activate(admin,one)
    project=admin.post('/api/v1/projects',json=dict(name='synthetic legacy',research_topic='synthetic',project_type='thesis')).json()
    identity=str(uuid4());now=datetime.now(timezone.utc)
    with database_connection(write=True) as db:
        db.execute('''INSERT INTO tasks(id,project_id,user_id,task_type,idempotency_key,input_digest,input_snapshot,
            workflow_version,status,phase,revision,current_attempt,created_at,updated_at)
            VALUES(%s,%s,%s,'school_template_parse','legacy-missing',%s,%s,'school_template_v1','queued','parse',1,0,%s,%s)''',
            (identity,project['id'],accounts[0]['id'],digest({}),Jsonb({}),now,now))
    two=saved(admin,one,'schools','heading_aliases',[dict(text='Synthetic',role='heading1')]);validated(admin,two)
    blocked=switch(admin,two,1,'blocked')
    assert blocked.status_code==409 and blocked.json()['detail']['code']=='configuration_legacy_inflight'
    assert blocked.json()['detail']['params']=={'task_id':identity}
    unrelated=saved(admin,one,'journals','update_note','Synthetic manual metadata')
    activate(admin,unrelated,1,'unrelated')


def existing_sci_fixture(client,project_id,*,prefix='synthetic-c'):
    """Reuse an existing SCI project; only local facts, no model/provider setup."""
    from tests.api.test_file_batches import register
    from tests.api.test_manuscripts import body as manuscript_body
    from app.domain.tasks.execution import run_task
    raw=synthetic_alias_docx()
    upload,batch,item,_=register(client,{'id':project_id},raw,prefix+'.docx',key=prefix+'-upload')
    uploaded=client.post(upload+'/'+batch['id']+'/items/'+item['client_id'],files={'file':(prefix+'.docx',raw)}).json()
    material=dict(kind='original',id=uploaded['file_id'],file_id=uploaded['file_id'])
    path=f'/api/v1/projects/{project_id}'
    manuscript=post(client,path+'/manuscripts',manuscript_body(sources=[material|dict(role='methods')]),prefix+'-draft').json()
    assert run_task(manuscript['versions'][0]['task_id'],1)['status']=='succeeded'
    version=manuscript['versions'][0];url=path+'/reviews'
    response=post(client,url+'/rounds',dict(title='Synthetic local C review',project_type='sci',synthetic=True,
        manuscript=dict(id=manuscript['id'],revision=version['revision'],version_id=version['id']),materials=[],
        sources=[dict(role='reviewer',label='Synthetic reviewer',format='text',text='1. Clarify the method.\n2. Add an independent experiment.')]),prefix+'-round')
    assert response.status_code==202,response.text
    round=response.json()
    run_task(round['task']['id'],1)
    round=client.get(url+'/rounds/'+round['id']).json()
    context=client.get(url+'/rounds/'+round['id']+'/revision-context').json()
    location=next(loc for loc in context['locations'] if loc['kind']=='paragraph' and loc['section_id']=='materials_and_methods')
    request=dict(plan_revision=round['revision'],expected_current_revision=context['expected_current_revision'],operations=[dict(
        item_id=round['items'][0]['id'],kind='replace_paragraph',location_id=location['id'],
        text=location['text']+' The existing method is retained.',evidence_ids=[material['id']],reason='Synthetic clarification of the existing method.')])
    preview=client.post(url+'/rounds/'+round['id']+'/revisions/preview',json=request)
    assert preview.status_code==200 and preview.json()['can_adopt'],preview.text
    adopted=post(client,url+'/rounds/'+round['id']+'/revisions',request|dict(preview_sha256=preview.json()['preview_sha256']),prefix+'-adopt').json()
    assert run_task(adopted['task']['id'],1)['status']=='succeeded'
    run=client.get(url+'/revisions/'+adopted['id']).json()
    assert run_task(run['result']['output']['task_id'],1)['status']=='succeeded'
    return url,round,run


def test_existing_project_fixture_needs_no_model_calls_or_extra_projects(admin):
    project=admin.post('/api/v1/projects',json=dict(name='synthetic D reuse',research_topic='synthetic',project_type='sci')).json()
    one=baseline(admin);activate(admin,one)
    url,round,run=existing_sci_fixture(admin,project['id'])
    document=post(admin,url+'/revisions/'+run['id']+'/documents',{'kind':'comparison'},'comparison')
    assert document.status_code==201,document.text
    with database_connection() as db:
        assert db.execute('SELECT count(*) AS n FROM projects').fetchone()['n']==1
        assert db.execute('SELECT count(*) AS n FROM model_calls').fetchone()['n']==0
    assert any(item['status']=='waiting_material' for item in run['result']['items'])


def existing_thesis_fixture(client,project_id,*,prefix='synthetic-school'):
    """Existing project, real synthetic parser evidence, explicit school confirmation."""
    from tests.api.test_file_batches import register
    from app.adapters.school_template_parser import inspect_school_template
    from app.domain.school_templates.execution import run_school_template_task
    from app.domain.tasks.execution import run_task
    raw=synthetic_alias_docx()
    upload,batch,item,_=register(client,{'id':project_id},raw,prefix+'.docx',key=prefix+'-upload')
    uploaded=client.post(upload+'/'+batch['id']+'/items/'+item['client_id'],files={'file':(prefix+'.docx',raw)}).json()
    path=f'/api/v1/projects/{project_id}'
    created=post(client,path+'/school-templates',dict(name='Synthetic school template',school='Synthetic school',template_file_id=uploaded['file_id']),prefix+'-parse')
    assert created.status_code==202,created.text
    created=created.json()
    assert run_school_template_task(created['task']['id'],1,parser=inspect_school_template)['status']=='succeeded'
    parsed=next(item for item in client.get(path+'/school-templates').json()['items'] if item['id']==created['id'])
    chapters=[{k:c[k] for k in ('id','title','level')} for c in parsed['analysis']['chapters']] or [dict(id='synthetic-method',title='合成研究方法',level=1)]
    decisions=[dict(conflict_id=c['id'],option_id='alias' if c.get('kind')=='heading_mapping' else c['options'][0]['id']) for c in parsed['analysis']['conflicts']]
    confirmed=post(client,path+'/school-templates/'+created['id']+'/confirm',dict(analysis_sha256=parsed['analysis_sha256'],chapters=chapters,decisions=decisions),prefix+'-confirm')
    assert confirmed.status_code==201,confirmed.text
    confirmation=confirmed.json()
    profile=client.get(path+'/thesis-profile').json()
    bound=post(client,path+'/thesis-template-binding',dict(confirmation_id=confirmation['id'],expected_revision=profile['revision']),prefix+'-bind')
    assert bound.status_code==200,bound.text
    body=dict(title='合成学校配置验收',output_language='zh-CN',sources=[],thesis=dict(confirmation_id=confirmation['id'],profile_revision=bound.json()['revision'],chapter_mappings=[dict(chapter_id=c['id'],purpose='materials_and_methods',output_language='project',result_group_ids=[]) for c in chapters]))
    manuscript=post(client,path+'/manuscripts',body,prefix+'-draft')
    assert manuscript.status_code==201,manuscript.text
    manuscript=manuscript.json()
    assert run_task(manuscript['versions'][0]['task_id'],1)['status']=='succeeded'
    return path,parsed,confirmation,manuscript


def test_existing_thesis_fixture_keeps_confirmation_binding_and_no_official_claim(admin):
    from app.domain.tasks.execution import run_task
    project=admin.post('/api/v1/projects',json=dict(name='synthetic D thesis',research_topic='synthetic',project_type='thesis')).json()
    one=baseline(admin);one=saved(admin,one,'schools','heading_aliases',[dict(text='Synthetic Chapter',role='heading1')]);activate(admin,one)
    path,parsed,confirmation,manuscript=existing_thesis_fixture(admin,project['id'])
    assert binding(confirmation['id'],'school_confirmation')==one['id']
    assert binding(manuscript['versions'][0]['task_id'])==one['id']
    exported=post(admin,path+'/manuscripts/'+manuscript['id']+'/versions/1/exports',{'school':{'include_toc':False}},'school-export')
    assert exported.status_code==201,exported.text
    assert binding(exported.json()['task_id'])==one['id']
    exported_task=run_task(exported.json()['task_id'],1)
    assert exported_task['status']=='succeeded',exported_task
    assert admin.get(ROOT+'/catalog').json()['installed']['supported_schools']==[]
    with database_connection() as db:
        assert db.execute('SELECT count(*) AS n FROM projects').fetchone()['n']==1
        assert db.execute('SELECT count(*) AS n FROM model_calls').fetchone()['n']==0


def test_fixed_literature_gate_accepts_only_existing_exact_verified_binding():
    from app.domain.review_rounds.planning import propose,installed_rules
    from app.domain.review_rounds.grounded import apply_operations
    from app.domain.manuscript_writing.contracts import SECTION_IDS
    sections=[dict(id=key,paragraphs=[]) for key in SECTION_IDS]
    paragraph=dict(id='claim',text='Reviewed claim.',origin='user_provided',source_ids=['material'],citations=[],numeric_bindings=[])
    next(s for s in sections if s['id']=='discussion')['paragraphs']=[paragraph]
    support=dict(project_reference_id='ref',review_id='review',claim_id='claim',metadata_version_id='metadata')
    snapshot=dict(output_language='en',sections=sections,sources=[dict(id='material',kind='document_version',content={'text':'Synthetic material'})],results_groups=[],references=[],
        citation_snapshot=dict(paragraphs=[dict(id='c1',fragments=[dict(kind='claim',bindings=[support])])],output=dict(paragraphs=[dict(id='c1',text='Reviewed claim.')])))
    item=propose({'id':'source'},1,dict(text='Add a verified reference.',original_number='',location={}),rules=installed_rules())
    operation=dict(item_id=item['id'],kind='bind_citation',location_id='claim',text='c1',evidence_ids=['material'],reason='Adopt exactly the existing verified statement.')
    _,changes,unexecuted,_,_=apply_operations(snapshot,[item],[],[operation],fixed_safety=True)
    assert changes and not unexecuted
    unverified=deepcopy(snapshot);unverified['citation_snapshot']=None
    _,changes,unexecuted,_,_=apply_operations(unverified,[item],[],[operation],fixed_safety=True)
    assert not changes and unexecuted


def test_new_round_language_default_is_effective_and_replay_ignores_rollback(admin):
    one=baseline(admin);activate(admin,one)
    two=saved(admin,one,'review','default_language','en');activate(admin,two,1,'english')
    project=admin.post('/api/v1/projects',json=dict(name='synthetic English default',research_topic='synthetic',project_type='sci')).json()
    url,round,run=existing_sci_fixture(admin,project['id'],prefix='english-default')
    endpoint=url+'/revisions/'+run['id']+'/documents'
    english=post(admin,endpoint,{'kind':'comparison'},'default-language')
    assert english.status_code==201 and english.json()['language']=='en',english.text
    activate(admin,one,2,'rollback','rollback')
    assert post(admin,endpoint,{'kind':'comparison'},'default-language').json()==english.json()
    assert post(admin,endpoint,{'kind':'comparison'},'same-round-new-document').json()['language']=='en'
    assert post(admin,endpoint,{'kind':'comparison','language':'zh-CN'},'explicit-chinese').json()['language']=='zh-CN'
    document=english.json()
    edited=post(admin,url+'/documents/'+document['id']+'/versions',dict(expected_version=document['version'],items=[dict(item_id=i['item_id'],text=i['text']) for i in document['snapshot']['items']]),'version')
    # A later explicit document is latest, so editing the old version is rejected.
    assert edited.status_code==409
    latest=post(admin,endpoint,{'kind':'reply'},'reply').json()
    version=post(admin,url+'/documents/'+latest['id']+'/versions',dict(expected_version=latest['version'],items=[dict(item_id=i['item_id'],text=i['text']) for i in latest['snapshot']['items']]),'reply-version')
    assert version.status_code==201 and version.json()['language']=='en',version.text


NAMED_CHAPTERS=('学位论文独创性声明','学位论文版权使用授权书','附录A：')


def named_chapter_docx():
    document=Document()
    for title in NAMED_CHAPTERS:document.add_paragraph(title)
    document.add_heading('Synthetic regular chapter',level=1)
    result=BytesIO();document.save(result)
    return result.getvalue()


@pytest.mark.parametrize('alias_role',['body','heading2'])
def test_named_chapter_without_heading_style_requires_alias_conflict(alias_role):
    from app.adapters.school_template_parser import inspect_school_template,merge_template_notes
    raw=named_chapter_docx();before=inspect_school_template(raw,'synthetic-named.docx')
    aliases=[dict(text=title,role=alias_role) for title in NAMED_CHAPTERS]
    after=inspect_school_template(raw,'synthetic-named.docx',heading_aliases=aliases)
    original={chapter['title']:chapter for chapter in before['chapters']}
    conflicts=[c for c in after['conflicts'] if c.get('kind')=='heading_mapping']
    assert len(conflicts)==len(NAMED_CHAPTERS)
    for title in NAMED_CHAPTERS:
        rule=next(r for r in after['rules'] if r['category']=='heading_mapping' and r['mapping']['title']==title)
        conflict=next(c for c in conflicts if c['rule_id']==rule['id'])
        assert rule['mapping']['observed_level']==original[title]['level']==1
        assert rule['source']==original[title]['source']
        assert next(o for o in conflict['options'] if o['id']=='original')['value']=='heading1'
        assert next(o for o in conflict['options'] if o['id']=='alias')['value']==alias_role
    assert merge_template_notes(after,[])['conflicts']==after['conflicts']


@pytest.mark.parametrize('alias_role',['body','heading2'])
def test_named_chapter_original_confirmation_restores_levels_and_preserves_analysis(admin,alias_role):
    from app.adapters.school_template_parser import inspect_school_template
    from app.domain.school_templates.execution import run_school_template_task
    from tests.api.test_file_batches import register
    raw=named_chapter_docx();before=inspect_school_template(raw,'synthetic-named.docx')
    one=baseline(admin);one=saved(admin,one,'schools','heading_aliases',[dict(text=title,role=alias_role) for title in NAMED_CHAPTERS]);activate(admin,one)
    project=admin.post('/api/v1/projects',json=dict(name='synthetic named parts',research_topic='synthetic',project_type='thesis')).json()
    upload,batch,item,_=register(admin,project,raw,'synthetic-named.docx')
    file=admin.post(upload+'/'+batch['id']+'/items/'+item['client_id'],files={'file':('synthetic-named.docx',raw)}).json()
    path=f"/api/v1/projects/{project['id']}/school-templates"
    template=post(admin,path,dict(name='Synthetic named parts',school='Synthetic school',template_file_id=file['file_id']),'parse').json()
    assert run_school_template_task(template['task']['id'],1,parser=inspect_school_template)['status']=='succeeded'
    parsed=admin.get(path).json()['items'][0]
    chapters=[{k:c[k] for k in ('id','title','level')} for c in parsed['analysis']['chapters']]
    confirm=dict(analysis_sha256=parsed['analysis_sha256'],chapters=chapters,decisions=[])
    assert post(admin,path+'/'+template['id']+'/confirm',confirm,'missing-decisions').status_code==409
    confirm['decisions']=[dict(conflict_id=c['id'],option_id='original' if c.get('kind')=='heading_mapping' else c['options'][0]['id']) for c in parsed['analysis']['conflicts']]
    assert post(admin,path+'/'+template['id']+'/confirm',confirm,'wrong-chapters').status_code==422
    confirm['chapters']=[{k:c[k] for k in ('id','title','level')} for c in before['chapters']]
    accepted=post(admin,path+'/'+template['id']+'/confirm',confirm,'restore-original')
    assert accepted.status_code==201,accepted.text
    assert accepted.json()['chapters']==confirm['chapters']
    assert not any(r.get('alias_rule_id') for r in accepted.json()['rules'])
    fresh=admin.get(path).json()['items'][0]
    assert fresh['analysis']==parsed['analysis'] and fresh['analysis_sha256']==parsed['analysis_sha256']
    assert fresh['sources']==parsed['sources']
