"""Synthetic Word inputs shared by opt-in queue/browser acceptance tests."""
from io import BytesIO

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw


def picture_bytes(format='PNG'):
    picture = Image.new('RGB', (160, 80), 'white')
    drawing = ImageDraw.Draw(picture)
    drawing.rectangle((12, 12, 148, 68), outline='#54268c', width=3)
    drawing.line((24, 55, 62, 35, 96, 48, 134, 22), fill='#54268c', width=3)
    stream = BytesIO()
    picture.save(stream, format=format)
    return stream.getvalue()


def document_bytes(*, warnings=False, oversized=False):
    document = Document()
    document.core_properties.title = 'PaperAssist Word结构合成样本'
    document.core_properties.author = 'Synthetic fixture'
    style = document.styles.add_style('Research Note', WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = document.styles['Normal']
    style.font.name = 'Arial'
    style.font.size = Pt(11)
    style.font.color.rgb = RGBColor.from_string('54268C')
    document.add_heading('合成研究资料 / Synthetic research notes', level=1)
    paragraph = document.add_paragraph(style=style)
    paragraph.add_run('本文件仅用于解析验收。').bold = True
    paragraph.add_run('Preserve this English text.').italic = True
    document.add_heading('实验数据说明', level=2)
    document.add_paragraph('表格仅展示原文，不自动进入统计。')
    table = document.add_table(rows=4, cols=3)
    table.style = 'Table Grid'
    for row, values in zip(table.rows, [
        ['组别', '测量值', '说明'], ['A', '1.50 mg', '合成数据'],
        ['B', '2.75 mg', ''], ['B', '3.25 mg', '合成数据'],
    ]):
        for cell, text in zip(row.cells, values):
            cell.text = text
    table.cell(2, 0).merge(table.cell(3, 0))
    nested = table.cell(2, 2).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = '嵌套说明'
    nested.cell(0, 1).text = '原坐标保留'
    document.add_paragraph('图1：合成折线；不据图片推断数据。')
    document.add_picture(BytesIO(picture_bytes()), width=Inches(2))
    document.add_paragraph('图2：JPEG图片字节应保留。')
    document.add_picture(BytesIO(picture_bytes('JPEG')), width=Inches(2))
    document.add_paragraph('图后正文：This paragraph follows both pictures.')
    if warnings:
        document.sections[0].header.paragraphs[0].text = '页眉保留于原件，不是正文'
        paragraph = document.add_paragraph('修订示例：')
        inserted = OxmlElement('w:ins')
        inserted.set(qn('w:id'), '1')
        inserted.set(qn('w:author'), 'Synthetic fixture')
        run = OxmlElement('w:r')
        text = OxmlElement('w:t')
        text.text = '尚未接受的修订内容'
        run.append(text)
        inserted.append(run)
        paragraph._p.append(inserted)
        equation = OxmlElement('m:oMath')
        paragraph._p.append(equation)
    if oversized:
        for _ in range(600):
            document.add_paragraph('SYNTHETIC-OUTPUT-LIMIT-' * 450)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()
