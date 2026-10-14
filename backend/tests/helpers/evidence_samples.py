"""Local synthetic samples; no models or production storage."""
from io import BytesIO
from docx import Document
from docx.shared import Inches
from tests.helpers.document_samples import picture_bytes


def evidence_docx():
    document=Document()
    document.add_paragraph('Figure 1: Synthetic chart; original caption only.')
    document.add_picture(BytesIO(picture_bytes()),width=Inches(2))
    document.add_paragraph('Separate section text.')
    document.add_paragraph('表1：原始测量值 / Original measurements')
    table=document.add_table(rows=2,cols=2)
    table.cell(0,0).text='Group'; table.cell(0,1).text='Mass (mg)'
    table.cell(1,0).text='A'; table.cell(1,1).text='1.50 mg'
    document.add_paragraph('Independent ending text; no inferred n or P.')
    output=BytesIO();document.save(output);return output.getvalue()


def evidence_pdf(*, image_only=False, image_count=1):
    # One uncompressed 1x1 RGB image, plus optional caption/table. Geometry is real PDF content.
    from tests.helpers.pdf_samples import TABLE
    stream=b' '.join([b'q 120 0 0 60 20 300 cm /Im1 Do Q']*image_count)
    if not image_only:
        stream+=b' BT /F1 12 Tf 20 285 Td (Figure 1: Synthetic chart) Tj ET'
        stream+=TABLE+b' BT /F1 12 Tf 20 185 Td (Table 1: Original values) Tj ET'
    objects=[b'<< /Type /Catalog /Pages 2 0 R >>', b'<< /Type /Pages /Kids [4 0 R] /Count 1 >>',
        b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
        b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 400] /Resources << /Font << /F1 3 0 R >> /XObject << /Im1 6 0 R >> >> /Contents 5 0 R >>',
        b'<< /Length '+str(len(stream)).encode()+b' >>\nstream\n'+stream+b'\nendstream',
        b'<< /Type /XObject /Subtype /Image /Width 1 /Height 1 /ColorSpace /DeviceRGB /BitsPerComponent 8 /Length 3 >>\nstream\n\x80\x00\x80\nendstream']
    output=bytearray(b'%PDF-1.4\n'); offsets=[0]
    for i,obj in enumerate(objects,1):
        offsets.append(len(output)); output.extend(f'{i} 0 obj\n'.encode()+obj+b'\nendobj\n')
    start=len(output); output.extend(b'xref\n0 7\n0000000000 65535 f \n')
    for offset in offsets[1:]: output.extend(f'{offset:010d} 00000 n \n'.encode())
    output.extend(f'trailer\n<< /Size 7 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF'.encode())
    return bytes(output)


def evidence_xlsx():
    from tests.adapters.test_excel_structure import sample
    def patch(xml):
        xml=xml.replace('<is><t>name</t></is>','<is><t>Group</t></is>').replace('<is><t>value</t></is>','<is><t>Mass (mg)</t></is>')
        xml=xml.replace('<c r="B30" t="n"><v>30</v></c>','<c r="B30" t="n"><f>1+1</f><v>2</v></c>')
        return xml.replace('<c r="B31" t="n"><v>31</v></c>','<c r="B31" t="inlineStr"><is><t/></is></c><c r="C31"/><c r="D31" t="str"><v/></c>')
    return sample(patch)
