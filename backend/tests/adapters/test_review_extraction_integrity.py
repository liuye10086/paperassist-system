"""Real DOCX parser evidence must survive the review preview boundary."""
from io import BytesIO
from types import SimpleNamespace
from zipfile import ZipFile

import pytest
from docx import Document
from app.adapters.docx_parser import DOCX_SETTINGS, parse_docx
from app.domain.review_rounds.extraction import DocumentService, preview


def parsed_preview(monkeypatch, extra_xml):
    doc = Document()
    doc.add_paragraph('1. Check the following requirement.')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Table comment remains visible.'
    buf = BytesIO(); doc.save(buf)
    modified = BytesIO()
    with ZipFile(BytesIO(buf.getvalue())) as old, ZipFile(modified, 'w') as new:
        for name in old.namelist():
            raw = old.read(name)
            if name == 'word/document.xml':
                raw = raw.replace(b'</w:body>', extra_xml + b'</w:body>')
            new.writestr(name, raw)
    raw = modified.getvalue()
    document = parse_docx(raw, 'opinions.docx', DOCX_SETTINGS).document
    row = {'input_snapshot': {'file': {}}}
    monkeypatch.setattr(DocumentService, '__init__', lambda *a: None)
    monkeypatch.setattr(DocumentService, 'load', lambda *a: row)
    monkeypatch.setattr(DocumentService, 'read', lambda *a: (None, document))
    def texts(blocks):
        for block in blocks:
            if block.kind == 'paragraph':
                yield {'text': block.text, 'location': block.source.model_dump()}
            else:
                for cell in block.cells:
                    yield from texts(cell.blocks)
    service = SimpleNamespace(assets=SimpleNamespace(original=lambda _: raw),
        material=lambda *a: {'content': {'items': list(texts(document.blocks))}})
    result = preview(service, 'project', dict(format='document', file_id='file', document_version_id='version'))
    return document, result


@pytest.mark.parametrize('xml,code,reason', [
    (b'<w:sdt><w:sdtContent><w:p><w:r><w:t>2. Do not change the conclusion.</w:t></w:r></w:p></w:sdtContent></w:sdt>', 'unsupported_blocks', '内容容器'),
    (b'<w:p xmlns:m="http://schemas.openxmlformats.org/officeDocument/2006/math"><m:oMath><m:r><m:t>x=2</m:t></m:r></m:oMath></w:p>', 'equations', '公式'),
])
def test_review_rejects_known_body_omissions(monkeypatch, xml, code, reason):
    document, result = parsed_preview(monkeypatch, xml)
    assert code in [w.code for w in document.warnings]
    assert '2. Do not change the conclusion.' not in result['original_text']
    assert not result['supported']
    assert any(reason in w for w in result['warnings'])


@pytest.mark.parametrize('xml', [b'', b'<w:p><w:fldSimple w:instr="PAGE"><w:r><w:t>1</w:t></w:r></w:fldSimple></w:p>'])
def test_review_keeps_clean_body_tables_and_nonblocking_field_notice(monkeypatch, xml):
    _, result = parsed_preview(monkeypatch, xml)
    assert result['supported']
    assert 'Check the following requirement.' in result['original_text']
    assert 'Table comment remains visible.' in result['original_text']
