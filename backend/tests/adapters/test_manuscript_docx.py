from copy import deepcopy
from io import BytesIO
import hashlib
from zipfile import ZipFile

from docx import Document
from PIL import Image
import pytest


def sample(language='en'):
    image = BytesIO()
    Image.new('RGB', (800, 400), 'white').save(image, format='PNG')
    raw = image.getvalue()
    en = language == 'en'
    return {
        'schema_version': 1, 'workflow_version': 'manuscript_fact_draft_v1',
        'project_id': 'private-project-id', 'title': 'Synthetic cohort' if en else '合成队列样例',
        'output_language': language,
        'notice': 'Offline partial fact draft. Researcher review required.' if en else '离线阶段性事实稿，需要研究人员审核。',
        'methods': [{'id': 'methods-study_design', 'text': 'METHODS-SENTINEL (1)',
                     'origin': 'user_provided', 'source_ids': ['source-1'], 'citations': [1]}],
        'results_groups': [{'id': 'group-1', 'title': 'Measurements' if en else '测量结果',
            'paragraphs': [{'id': 'group-1-result', 'text': 'RESULT-SENTINEL 0.0123',
                            'origin': 'computed_result', 'source_ids': ['source-1'], 'citations': []}],
            'tables': [{'id': 'table-1', 'title': 'Recorded values' if en else '已记录数值',
                        'columns': ['Metric', 'Value'], 'rows': [['n', '12'], ['Mean', '5.24']],
                        'source_ids': ['source-1']}],
            'figures': [{'id': 'figure-1', 'number': 1, 'title': 'Recorded figure' if en else '已保存图',
                         'caption': 'CAPTION-SENTINEL', 'sha256': hashlib.sha256(raw).hexdigest(),
                         'source_ids': ['source-1'], 'panels': [{'label': 'A', 'description': 'PANEL-SENTINEL'}]}]}],
        'references': [{'number': 1, 'text': 'Author A. Saved reference. Journal. 2020;1:1.', 'project_reference_id': 'private-reference'}],
        'issues': [{'id': 'missing-ethics', 'kind': 'missing', 'field': 'ethics',
                    'message': 'ETHICS-MISSING-SENTINEL', 'values': []},
                   {'id': 'conflict-samples', 'kind': 'conflict', 'field': 'samples',
                    'message': 'SAMPLES-CONFLICT-SENTINEL', 'values': ['12', '14']}],
        'sources': [{'kind': 'analysis_run', 'id': 'source-1', 'file_id': 'private-file-id',
                     'role': 'computed_results', 'label': 'Selected analysis', 'sha256': 'a'*64,
                     'download_url': '/private/source', 'content': {'private': 'DO-NOT-DUMP'}}],
        'citation_snapshot': None,
    }, {'figure-1': raw}


@pytest.mark.parametrize('language', ['zh-CN', 'en'])
def test_fact_draft_docx_contains_editable_content_exact_images_and_pending_issues(language):
    from app.adapters.manuscript_docx import build_manuscript_docx
    snapshot, images = sample(language)
    original = deepcopy(snapshot)
    content = build_manuscript_docx(snapshot, images)
    document = Document(BytesIO(content))
    text = '\n'.join(p.text for p in document.paragraphs)
    en = language == 'en'
    assert text.index('Results' if en else '结果') < text.index('Materials and methods' if en else '材料与方法')
    assert text.count('METHODS-SENTINEL (1)') == 1
    assert all(value in text for value in ('RESULT-SENTINEL 0.0123', 'CAPTION-SENTINEL',
        'PANEL-SENTINEL', 'ETHICS-MISSING-SENTINEL', 'SAMPLES-CONFLICT-SENTINEL'))
    assert '(1) Author A.' in text
    assert 'DO-NOT-DUMP' not in text and 'private-project-id' not in text and 'a'*64 not in text
    assert document.paragraphs[0].style.name == 'Title'
    assert any(cell.text == '5.24' for table in document.tables for row in table.rows for cell in row.cells)
    with ZipFile(BytesIO(content)) as archive:
        media = [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
        assert media == [images['figure-1']]
        assert b'bookmarkStart' in archive.read('word/document.xml')
    assert snapshot == original


@pytest.mark.parametrize('change', ['missing', 'tampered'])
def test_fact_draft_docx_rejects_missing_or_changed_figure(change):
    from app.adapters.manuscript_docx import build_manuscript_docx
    from app.core.exceptions import StorageError
    snapshot, images = sample()
    if change == 'missing': images.clear()
    else: images['figure-1'] += b'changed'
    with pytest.raises(StorageError) as error:
        build_manuscript_docx(snapshot, images)
    assert error.value.code in ('figure_missing', 'figure_changed')


def test_fact_draft_docx_repeated_render_is_identical_and_preserves_portrait_ratio():
    from app.adapters.manuscript_docx import build_manuscript_docx
    snapshot, images = sample()
    raw = BytesIO()
    Image.new('RGB', (400, 2000), 'white').save(raw, format='PNG')
    images['figure-1'] = raw.getvalue()
    snapshot['results_groups'][0]['figures'][0]['sha256'] = hashlib.sha256(raw.getvalue()).hexdigest()
    first = build_manuscript_docx(snapshot, images)
    assert first == build_manuscript_docx(snapshot, images)
    shape = Document(BytesIO(first)).inline_shapes[0]
    assert shape.height.inches <= 6.0
    assert shape.width / shape.height == pytest.approx(.2)


def test_new_snapshot_refuses_mismatched_body_reference_while_legacy_snapshot_renders():
    from app.adapters.manuscript_docx import build_manuscript_docx
    from app.core.exceptions import StorageError
    snapshot, images = sample()
    legacy = build_manuscript_docx(snapshot, images)
    assert legacy.startswith(b'PK')
    snapshot['illustration_check_version'] = 1
    snapshot['results_groups'][0]['paragraphs'][0]['cross_references'] = [
        dict(kind='figure', target_id='figure-1', number=99, panel_label='A')]
    with pytest.raises(StorageError) as error:
        build_manuscript_docx(snapshot, images)
    assert error.value.code == 'report_render_failed'
