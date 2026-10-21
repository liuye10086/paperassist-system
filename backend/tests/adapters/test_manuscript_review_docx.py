"""Frozen review caveats must survive Word export without modifying the snapshot."""
from copy import deepcopy
from io import BytesIO

from docx import Document
import pytest

from tests.adapters.test_manuscript_docx import sample
from tests.adapters.test_manuscript_full_docx import full_sample


@pytest.fixture(autouse=True)
def postgres_schema():
    yield None


def reviewed_sample(language, complete):
    snapshot, images = (full_sample if complete else sample)(language)
    snapshot['review'] = dict(
        decisions=[], metadata=[
            dict(target_id='analysis_run:old:/overall/mean', field='unit', value='AU',
                 evidence='Original laboratory record', stale=True),
            dict(target_id='methods:study_design', field='acronym', value='RCT',
                 evidence='Current protocol', stale=False),
        ], addendum=[], relationships=dict(associations=[]),
        report=dict(can_publish=True, scientifically_verified=False, impacts=[
            dict(source_id='old', locations=['results'], pending_review=True)
        ], issues=[
            dict(category='material', severity='human_review', location='analysis_run:old:/overall/mean',
                 message='Original confirmation applies only to the old fact; new material requires renewed confirmation.'),
            dict(category='uncertainty', severity='human_review', location='specific:limitation',
                 message='The supplied protocol leaves the observation period unresolved.'),
        ]))
    return snapshot, images


@pytest.mark.parametrize('language', ['zh-CN', 'en'])
@pytest.mark.parametrize('complete', [False, True])
def test_docx_marks_stale_metadata_at_its_record(language, complete):
    from app.adapters.manuscript_docx import build_manuscript_docx
    snapshot, images = reviewed_sample(language, complete)
    original = deepcopy(snapshot)
    doc = Document(BytesIO(build_manuscript_docx(snapshot, images)))
    paragraph = next(p.text for p in doc.paragraphs if 'unit: AU' in p.text)
    expected = ('Stale confirmation: applies only to the original frozen fact; new material requires renewed confirmation.'
                if language == 'en' else '确认已过期：仅适用于原冻结事实；新资料须重新确认。')
    assert expected in paragraph
    assert 'Original laboratory record' in paragraph
    current = next(p.text for p in doc.paragraphs if 'acronym: RCT' in p.text)
    assert 'Stale confirmation' not in current and '确认已过期' not in current
    assert snapshot == original


@pytest.mark.parametrize('category', ['material', 'uncertainty'])
def test_docx_keeps_each_frozen_unresolved_review_message(category):
    from app.adapters.manuscript_docx import build_manuscript_docx
    snapshot, images = reviewed_sample('en', True)
    item = next(i for i in snapshot['review']['report']['issues'] if i['category'] == category)
    snapshot['review']['report']['issues'].append(deepcopy(item))
    original = deepcopy(snapshot)
    doc = Document(BytesIO(build_manuscript_docx(snapshot, images)))
    text = '\n'.join(p.text for p in doc.paragraphs)
    assert text.count(item['message']) == 1
    assert item['location'] in text
    assert snapshot == original


@pytest.mark.parametrize('language', ['zh-CN', 'en'])
def test_composed_source_refresh_retains_specific_notice_without_repeating_original_issues(language):
    from app.adapters.manuscript_docx import build_manuscript_docx
    from app.domain.manuscript_review.composition import compose
    from tests.domain.test_manuscript_review import sample as factual_sample, review_input
    base = factual_sample()
    base['output_language'] = language
    base['sources'][0]['label'] = 'Frozen original analysis'
    base['issues'] = [dict(id='missing:ethics', kind='missing', field='ethics',
                          message='Original ethics record is still missing.', values=[])]
    metadata = dict(target_id='analysis_run:run:/overall/mean', field='unit',
                    value='AU', evidence='Original protocol unit')
    first, _ = compose(review_input(base, dict(base_revision=1, metadata=[metadata])))
    frozen = review_input(first, dict(base_revision=2))
    frozen['source_state']['changes'] = [
        dict(id='run', kind='analysis_run', reasons=['newer_analysis_run'])]
    candidate, report = compose(frozen)
    original = deepcopy(candidate)
    assert report['can_publish'] and candidate['review']['metadata'][0]['stale']
    notice = next(i['message'] for i in report['issues']
                  if i['category'] == 'material' and i['location'] == metadata['target_id'])
    document = Document(BytesIO(build_manuscript_docx(candidate, {})))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert notice in text
    assert ('Stale confirmation' if language == 'en' else '确认已过期') in text
    assert text.count(base['issues'][0]['message']) == 1
    assert candidate == original
