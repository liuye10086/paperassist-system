"""Standalone Linux tests: no database fixtures or development configuration."""
import importlib.util
import io
import json
import os
from pathlib import Path
import subprocess
import sys
import time
import unittest
from unittest.mock import patch
from zipfile import ZipFile

sys.path.insert(0, '/app' if Path('/app/app').is_dir() else str(Path(__file__).resolve().parents[2]))

from openpyxl import Workbook
from app.adapters.excel import parse_workbook
from app.core.config import ExcelSettings
from app.adapters.file_objects import StorageError


class ExcelProcessTests(unittest.TestCase):
    def settings(self, excel=None):
        return {'excel': (excel or ExcelSettings()).model_dump(), 'process': {
            'timeout_seconds': 60, 'cpu_seconds': 30, 'memory_bytes': 536870912,
            'max_output_bytes': 4194304}}
    def adapter(self):
        self.assertIsNotNone(importlib.util.find_spec('app.adapters.excel_process'),
                             'The isolated parser must exist before parsing untrusted bytes')
        from app.adapters import excel_process
        return excel_process

    def content(self):
        book = Workbook()
        book.active.append(['same', 'same', None, 'formula'])
        book.active.append([1, 2, None, '=A2+B2'])
        book.active.merge_cells('A3:B3')
        book.active['A3'] = 'merged'
        for index in range(24):
            book.active.append([index, index + 1, None, '=A2+B2'])
        book.create_sheet('empty')
        target = io.BytesIO()
        book.save(target)
        return target.getvalue()

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_preview_matches_existing_parser(self):
        adapter = self.adapter()
        data = self.content()
        settings = ExcelSettings()
        start = time.monotonic()
        actual = adapter.parse_excel_isolated(data, 'sample.xlsx', self.settings(settings))
        self.assertEqual(actual.model_dump(), parse_workbook(data, 'sample.xlsx', settings).model_dump())
        self.assertEqual(len(actual.sheets[0].preview_rows), 20)
        print('synthetic preview elapsed_seconds=', round(time.monotonic() - start, 3))

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real native imports')
    def test_both_excel_bootstraps_keep_native_imports_single_threaded(self):
        # The OCR Worker also includes NumPy. openpyxl imports it if installed;
        # native BLAS threads must be constrained before the trusted imports.
        source = '''import json,os,runpy,sys
runpy.run_path(sys.argv[1],run_name='bootstrap_probe')
status={line.split(':',1)[0]:line.split(':',1)[1].strip() for line in open('/proc/self/status') if line.startswith(('Threads:','VmSize:'))}
print(json.dumps(status))
'''
        for name in ('excel_process_child.py','excel_structure_child.py'):
            with self.subTest(bootstrap=name):
                child=Path(self.adapter().__file__).with_name(name)
                probe=subprocess.run([sys.executable,'-I','-B','-c',source,str(child)],
                    env={'OPENBLAS_NUM_THREADS':'8','OMP_NUM_THREADS':'8'},capture_output=True,timeout=15)
                self.assertEqual(probe.returncode,0,probe.stderr.decode())
                status=json.loads(probe.stdout)
                print('bootstrap_native_evidence=',json.dumps(dict(bootstrap=name,**status)))
                self.assertEqual(status['Threads'],'1',status)
                self.assertLess(int(status['VmSize'].split()[0])*1024,536870912,status)

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_invalid_macro_and_expansion_are_safe(self):
        adapter = self.adapter()
        cases = [(b'broken secret contents', ExcelSettings(), 'invalid_workbook')]
        target = io.BytesIO()
        with ZipFile(target, 'w') as archive:
            archive.writestr('[Content_Types].xml', '<Types/>')
            archive.writestr('xl/workbook.xml', '<workbook/>')
            archive.writestr('xl/vbaProject.bin', 'macro')
        cases.append((target.getvalue(), ExcelSettings(), 'invalid_workbook'))
        cases.append((self.content(), ExcelSettings(max_uncompressed_bytes=10), 'workbook_too_large'))
        for data, settings, code in cases:
            with self.subTest(code=code):
                with self.assertRaises(StorageError) as caught:
                    adapter.parse_excel_isolated(data, 'safe.xlsx', self.settings(settings))
                self.assertEqual(caught.exception.code, code)
                self.assertNotIn('secret', str(caught.exception))

    def test_windows_fails_closed(self):
        if sys.platform == 'linux':
            self.skipTest('Windows-only platform rejection')
        adapter = self.adapter()
        with self.assertRaises(StorageError) as caught:
            adapter.parse_excel_isolated(self.content(), 'safe.xlsx', self.settings())
        self.assertEqual(caught.exception.code, 'parse_environment_unavailable')

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_real_sandbox_probe(self):
        self.adapter()
        probe = Path(__file__).with_name('excel_sandbox_probe.py')
        result = subprocess.run([sys.executable, str(probe), 'permissions'],
                                capture_output=True, env={'PYTHONPATH': '/app', 'SYNTHETIC_SECRET': 'must-not-inherit'}, timeout=10)
        self.assertEqual(result.returncode, 0, result.stderr.decode())
        data = json.loads(result.stdout)
        self.assertEqual(data['denied'], ['read', 'write', 'network', 'fork', 'exec', 'proc_environment'])
        self.assertEqual(data['limits'], [30, 536870912, 64, 0])

    def test_memory_error_is_not_disguised_as_input_failure(self):
        with patch('app.adapters.excel.inspect_archive', side_effect=MemoryError):
            with self.assertRaises(MemoryError):
                parse_workbook(self.content(), 'safe.xlsx', ExcelSettings())

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_cpu_and_address_space_are_enforced(self):
        self.adapter()
        probe = str(Path(__file__).with_name('excel_sandbox_probe.py'))
        memory = subprocess.run([sys.executable, probe, 'memory'], env={'PYTHONPATH': '/app'},
                                capture_output=True, timeout=10)
        self.assertEqual(memory.returncode, 0, memory.stderr.decode())
        self.assertEqual(memory.stdout.strip(), b'memory_denied')
        cpu = subprocess.run([sys.executable, probe, 'cpu'], env={'PYTHONPATH': '/app'},
                             capture_output=True, timeout=10)
        self.assertIn(cpu.returncode, [-9, -24])

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_transport_bounded_output_timeout_exit_and_clean_environment(self):
        adapter = self.adapter()
        cases = [
            ("import time; time.sleep(5)", 'parse_timeout'),
            ("import os; os.write(1,b'x'*4096)", 'parse_output_invalid'),
            ("raise SystemExit(7)", 'parse_process_failed'),
            ("print('not-json')", 'parse_output_invalid'),
            ("print('READY',flush=True); print('{\"x\":1,\"x\":2}')", 'parse_output_invalid'),
            ("print('READY',flush=True); print('{\"x\":NaN}')", 'parse_output_invalid'),
        ]
        for script, expected in cases:
            with self.subTest(expected=expected):
                with self.assertRaises(StorageError) as caught:
                    adapter._run_child([sys.executable, '-c', script], b'', timeout_seconds=.3,
                                       max_output_bytes=1024)
                self.assertEqual(caught.exception.code, expected)
        os.environ['SYNTHETIC_SECRET'] = 'synthetic-only'
        try:
            result = adapter._run_child([sys.executable, '-c',
                "import os,json; print('READY',flush=True); print(json.dumps({'has_secret':'SYNTHETIC_SECRET' in os.environ}))"],
                b'', timeout_seconds=3, max_output_bytes=1024)
            self.assertFalse(result['has_secret'])
        finally:
            os.environ.pop('SYNTHETIC_SECRET', None)

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_parent_death_kills_child(self):
        self.adapter()
        import ctypes
        libc = ctypes.CDLL(None)
        self.assertEqual(libc.prctl(36, 1, 0, 0, 0), 0)  # become child subreaper, reap test orphan
        probe = str(Path(__file__).with_name('excel_sandbox_probe.py'))
        launcher = "import subprocess,sys; p=subprocess.Popen([sys.executable,sys.argv[1],'sleep']); p.wait()"
        parent = subprocess.Popen([sys.executable, '-c', launcher, probe], stdout=subprocess.PIPE,
                                  stderr=subprocess.PIPE, env={'PYTHONPATH': '/app'})
        child_pid = int(parent.stdout.readline())
        parent.kill()
        parent.wait(timeout=3)
        deadline = time.monotonic() + 3
        while time.monotonic() < deadline:
            reaped, status = os.waitpid(child_pid, os.WNOHANG)
            if reaped:
                self.assertEqual(os.waitstatus_to_exitcode(status), -9)
                break
            time.sleep(.02)
        else:
            os.kill(child_pid, 9)
            os.waitpid(child_pid, 0)
            self.fail('Child survived parent death')
        parent.stdout.close()
        parent.stderr.close()

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_missing_linux_capability_fails_closed(self):
        self.adapter()
        probe = str(Path(__file__).with_name('excel_sandbox_probe.py'))
        result = subprocess.run([sys.executable, probe, 'missing_capability'],
                                env={'PYTHONPATH': '/app'}, capture_output=True, timeout=5)
        self.assertEqual(result.returncode, 0, result.stderr.decode())
        self.assertEqual(result.stdout.strip(), b'fail_closed')

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_limits_and_settings_cannot_be_weakened(self):
        adapter = self.adapter()
        for mutate in [lambda s: s['process'].update(memory_bytes=1073741824),
                       lambda s: s['process'].update(extra=1),
                       lambda s: s['excel'].update(max_rows=200000)]:
            settings = self.settings()
            mutate(settings)
            with self.assertRaises(StorageError) as caught:
                adapter.parse_excel_isolated(self.content(), 'safe.xlsx', settings)
            self.assertEqual(caught.exception.code, 'parse_environment_unavailable')

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_callback_exception_reaps_child(self):
        adapter = self.adapter()
        children = []
        def fail(pid):
            children.append(pid)
            raise RuntimeError('synthetic callback interruption')
        with self.assertRaises(RuntimeError):
            adapter.parse_excel_isolated(self.content(), 'safe.xlsx', self.settings(), on_started=fail)
        with self.assertRaises(ProcessLookupError):
            os.kill(children[0], 0)

    @unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
    def test_sandbox_cannot_read_or_change_parent_resource_limits(self):
        self.adapter()
        probe = str(Path(__file__).with_name('excel_sandbox_probe.py'))
        # This expendable parent, not the test process/Worker, has its limits
        # potentially changed by the reproducer. No development state is used.
        launcher = '''import json,resource,subprocess,sys
before = resource.getrlimit(resource.RLIMIT_NOFILE)
child = subprocess.run([sys.executable,sys.argv[1],'prlimit'],capture_output=True)
after = resource.getrlimit(resource.RLIMIT_NOFILE)
print(json.dumps({'before':before,'after':after,'code':child.returncode,
                  'probe':json.loads(child.stdout),'stderr':child.stderr.decode()}))
'''
        result = subprocess.run([sys.executable, '-c', launcher, probe],
                                env={'PYTHONPATH': '/app'}, capture_output=True, timeout=5)
        self.assertEqual(result.returncode, 0, result.stderr.decode())
        data = json.loads(result.stdout)
        self.assertEqual(data['code'], 0, data['stderr'])
        print('prlimit_same_uid_parent_evidence=', json.dumps(data))
        self.assertEqual(data['probe']['denied'], ['parent_read', 'parent_write', 'self_write',
                                                 'pid_high', 'new_limit_high'])
        self.assertEqual(data['before'], data['after'], 'Sandbox child changed its same-UID parent limits')
        self.assertEqual(data['probe']['self_read'], [64, 64])
        self.assertEqual(data['probe']['self_read_raw'], [64, 64])


if __name__ == '__main__':
    unittest.main(verbosity=2)
