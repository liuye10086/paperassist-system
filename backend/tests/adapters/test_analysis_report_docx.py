from io import BytesIO
from zipfile import ZipFile
from copy import deepcopy
from docx import Document
import pytest
from tests.domain.test_correlation_artifacts import result
from tests.adapters.test_analysis_plot_response import png


def snapshot(language='en',mode='pair'):
    value=result(language,mode)
    value.update(started_at='2026-10-14T00:00:00+00:00',completed_at='2026-10-14T00:01:00+00:00',
        engine={'id':'correlation_statistics_dataset_v1','python_version':'3.11','parser_version':'test','scipy_version':'1.17.1','numpy_version':'2.4.6'})
    return {'project':{'id':'project','name':'Research','research_topic':'Association','project_type':'sci','created_at':'2026'},
        'result':value,'figure':{'id':'figure','sha256':'c'*64,'title':'Figure 1','caption':'Complete records.',
            'created_at':'2026','engine':{'id':'openai_analysis_plot_v2'},'provenance':{'model':'synthetic','response_id':'resp'}},
        'explanation':{'id':'explanation','created_at':'2026','language':language,'engine':{'id':'analysis_explanation_v2','model':'synthetic'},
            'provenance':{'input_sha256':'d'*64,'response_id':'resp'},'sections':[],
            'limitations':['Review required' if language=='en' else '需要人工审阅'],
            'verification':{'note':'Checked' if language=='en' else '已核对'}}}


@pytest.mark.parametrize('language',['zh-CN','en'])
@pytest.mark.parametrize('mode',['pair','matrix'])
def test_modern_word_includes_local_statistics_and_original_png_bytes(language,mode):
    from app.adapters.report_docx import build_report
    from app.adapters.report_store import ReportStore
    saved=snapshot(language,mode)
    metadata=ReportStore.metadata(saved)
    assert metadata['language']==language and metadata['engine']['id']=='word_analysis_report_v2'
    content=build_report(saved,metadata,png())
    document=Document(BytesIO(content))
    text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert ('Correlation analysis report' if language=='en' else '相关分析报告') in text
    assert '0.8' in text and '4' in text
    if mode=='pair': assert '0.2' in text and '-0.5' in text
    else:
        assert '0.2' not in text
        assert ('exploratory' if language=='en' else '探索') in text
    if language=='en': assert '原始文件' not in text and '描述统计分析报告' not in text
    with ZipFile(BytesIO(content)) as archive:
        media=[archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]
    assert media==[png()]


def test_unavailable_spearman_interval_has_explicit_reason_in_word():
    from app.adapters.report_docx import build_report
    from app.adapters.report_store import ReportStore
    saved=snapshot()
    saved['result']['statistics']['method']='spearman'
    saved['result']['statistics']['confidence_interval'].update(low=None,high=None,method=None,reason='spearman_interval_not_implemented')
    text='\n'.join(p.text for p in Document(BytesIO(build_report(saved,ReportStore.metadata(saved),png()))).paragraphs)
    assert 'not implemented' in text


@pytest.mark.parametrize('language',['zh-CN','en'])
def test_modern_report_retains_local_research_declarations_and_assumptions(language):
    from app.adapters.report_docx import build_report
    from app.adapters.report_store import ReportStore
    saved=snapshot(language)
    plan=saved['result']['analysis_context']['plan']
    plan['research_request']={'question':'LOCAL-QUESTION-SENTINEL','population':'LOCAL-POPULATION',
        'study_design':'LOCAL-DESIGN','observation_unit':'LOCAL-OBSERVATION',
        'observation_structure':'independent','relationship_description':'LOCAL-RELATIONSHIP',
        'output_requirements':'LOCAL-OUTPUT','output_language':language}
    saved['result']['selection']['assumptions']={'independent_observations':True,'relationship_form':'linear','inference_assumptions_confirmed':True}
    text='\n'.join(p.text for p in Document(BytesIO(build_report(saved,ReportStore.metadata(saved),png()))).paragraphs)
    assert all(sentinel in text for sentinel in ['LOCAL-QUESTION-SENTINEL','LOCAL-POPULATION','LOCAL-DESIGN','LOCAL-OBSERVATION','LOCAL-RELATIONSHIP','LOCAL-OUTPUT'])
    assert ('用户研究声明' if language=='zh-CN' else 'User research declarations') in text
    assert ('独立观察声明' if language=='zh-CN' else 'Independent-observation declaration') in text
