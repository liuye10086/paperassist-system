"""Standalone real Linux structure protocol checks; no database or app config."""
import io
from pathlib import Path
import sys
import unittest
from zipfile import ZipFile, ZIP_DEFLATED

sys.path.insert(0, '/app' if Path('/app/app').is_dir() else str(Path(__file__).resolve().parents[2]))

from openpyxl import Workbook
from app.adapters.excel_structure import extract_excel_structure
from app.adapters.excel_structure_process import parse_excel_structure_isolated, PROCESS_SETTINGS
from app.core.config import ExcelSettings
from app.core.exceptions import StorageError


@unittest.skipUnless(sys.platform == 'linux', 'Requires real Linux sandbox')
class StructureProcessTests(unittest.TestCase):
    def settings(self):
        return {'excel': ExcelSettings().model_dump(), 'process': dict(PROCESS_SETTINGS)}

    def content(self):
        book = Workbook()
        sheet = book.active
        sheet.title = '测量'
        sheet.append(['值', '组', '公式'])
        for value in range(1, 31):
            sheet.append([value, ' A ', '=A2*2'])
        book.create_sheet('隐藏').sheet_state = 'hidden'
        book.create_sheet('空白')
        stream = io.BytesIO()
        book.save(stream)
        result = io.BytesIO()
        with ZipFile(io.BytesIO(stream.getvalue())) as original, ZipFile(result, 'w', ZIP_DEFLATED) as output:
            for entry in original.infolist():
                data = original.read(entry)
                if entry.filename == 'xl/worksheets/sheet1.xml':
                    data = data.replace(b'<v>1</v>', b'<v>9007199254740993</v>', 1)
                    data = data.replace(b'<f>A2*2</f><v></v>', b'<f>A2*2</f><v>0</v>', 1)
                output.writestr(entry, data)
        return result.getvalue()

    def test_actual_child_keeps_all_rows_raw_precision_and_formula_evidence(self):
        content = self.content()
        actual = parse_excel_structure_isolated(content, '合成.xlsx', self.settings())
        self.assertEqual(actual.model_dump(), extract_excel_structure(content, '合成.xlsx', ExcelSettings()).model_dump())
        self.assertEqual(actual.sheets[0].row_count, 31)
        cells = {cell.coordinate: cell for cell in actual.sheets[0].cells}
        self.assertEqual(cells['A2'].value, '9007199254740993')
        self.assertEqual(cells['A31'].value, '30')
        self.assertEqual(cells['B2'].value, ' A ')
        self.assertEqual(cells['C2'].formula.cached_value, '0')
        self.assertEqual(cells['C2'].formula.cache_validity, 'unverified')
        self.assertEqual(actual.sheets[1].state, 'hidden')
        self.assertEqual(actual.sheets[2].cells, [])

    def test_actual_child_rejects_complete_output_over_limit(self):
        book = Workbook()
        book.active.append(['值'])
        for _ in range(145):
            book.active.append(['x' * 32000])
        stream = io.BytesIO()
        book.save(stream)
        with self.assertRaises(StorageError) as caught:
            parse_excel_structure_isolated(stream.getvalue(), '超限.xlsx', self.settings())
        self.assertEqual(caught.exception.code, 'parse_output_invalid')

    def test_weakened_structure_settings_are_rejected_before_child(self):
        for key, value in [('timeout_seconds', 61), ('memory_bytes', 1073741824)]:
            settings = self.settings()
            settings['process'][key] = value
            with self.assertRaises(StorageError) as caught:
                parse_excel_structure_isolated(self.content(), '合成.xlsx', settings)
            self.assertEqual(caught.exception.code, 'parse_environment_unavailable')


if __name__ == '__main__':
    unittest.main()
