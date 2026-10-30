"""Synthetic source-bound front matter; no real school/identity acceptance."""
from copy import deepcopy
from io import BytesIO
import sys

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from app.core.exceptions import StorageError
from app.domain.literature.service import digest
from tests.adapters.test_school_docx import fixture, plan, parts, text, NS


@pytest.fixture(autouse=True)
def postgres_schema():
    yield None


def synthetic(label='培养类型', size=18, alignment=WD_ALIGN_PARAGRAPH.CENTER):
    doc=Document();p=doc.add_paragraph(label+'：示例值不能作为本人信息')
    p.alignment=alignment;p.paragraph_format.space_after=Pt(8)
    p.add_run().font.size=Pt(size)
    p.runs[0].font.size=Pt(size)
    doc.add_heading('学位论文独创性声明',1)
    doc.add_paragraph('作者签名：示例签名')
    doc.add_heading('致谢',1)
    doc.add_paragraph('封面')
    stream=BytesIO();doc.save(stream)
    from app.adapters.school_template_parser import inspect_school_template
    return inspect_school_template(stream.getvalue(),'synthetic.docx')


def configured():
    value,analysis=fixture();parsed=synthetic()
    analysis['front_matter_sources']=parsed['front_matter_sources']
    source=analysis['front_matter_sources'][0]['id']
    signature=analysis['front_matter_sources'][2]['id']
    fm=dict(schema_version=1,cover_section_id='section-1',fields=[
        dict(id='degree',label='培养类型',kind='text',required=True,source_id=source,style_source_id=source,section_id='section-1',chapter_id=None),
        dict(id='signature',label='作者签名',kind='signature',required=True,source_id=signature,style_source_id=None,section_id='section-1',chapter_id='declaration')],
        content_requirements=[dict(id='declare',kind='declaration',required=True,chapter_id='declaration',source_id=signature)])
    value['thesis_context']['confirmation']['front_matter']=fm
    value['thesis_context']['confirmation']['analysis_sha256']=digest(analysis)
    value['thesis_context']['confirmation_sha256']=digest(value['thesis_context']['confirmation'])
    return value,analysis,fm


def test_two_synthetic_sources_retain_different_labels_and_styles_and_cover_is_incomplete():
    a=synthetic();b=synthetic('论文类别',14,WD_ALIGN_PARAGRAPH.RIGHT)
    assert a['front_matter_sources'][0]['suggested_label']=='培养类型'
    assert b['front_matter_sources'][0]['suggested_label']=='论文类别'
    assert a['front_matter_sources'][0]['style']['font_size_pt']==18
    assert b['front_matter_sources'][0]['style']['alignment']=='right'
    assert a['front_matter_sources'][0]['style_sources']['font_size_pt']['path']
    assert any(i['id']=='cover_rules' for i in a['missing'])


@pytest.mark.skipif(sys.platform!='linux',reason='Existing fenced parser requires Linux')
def test_existing_linux_isolation_returns_valid_front_source_contract():
    from app.adapters.school_template_process import inspect_school_template_isolated
    from app.adapters.school_template_parser import RESULT_KEYS
    doc=Document();doc.add_paragraph('Synthetic field: never copy example identity');doc.add_heading('Introduction',1)
    stream=BytesIO();doc.save(stream)
    result=inspect_school_template_isolated(stream.getvalue(),'synthetic.docx')
    assert set(result)==RESULT_KEYS
    assert result['front_matter_sources'][0]['suggested_label']=='Synthetic field'


def test_dynamic_required_values_and_signature_remain_missing_with_nonempty_declaration():
    value,analysis,_=configured();before=deepcopy((value,analysis))
    p=plan(value,analysis,manual_chapters=[dict(chapter_id='declaration',text='本人提供的声明正文')],field_values={'degree':'合成学位'})
    ids={i['id'] for i in p['report']['issues']}
    assert p['schema_version']==2
    assert 'school:field:signature' in ids and 'school:field:degree' not in ids
    assert not any(i.startswith('school:manual:') or i.startswith('school:keywords:') for i in ids)
    assert next(i for i in p['report']['issues'] if i['id']=='school:field:signature')['source_evidence']
    assert (value,analysis)==before
    from app.adapters.manuscript_docx import build_manuscript_docx
    from lxml import etree as E
    raw=build_manuscript_docx(value,{},school_plan=p)
    assert '培养类型：合成学位' in text(raw) and '作者签名：待填写' in text(raw)
    assert '示例签名' not in text(raw) and '示例值不能作为本人信息' not in text(raw)
    root=E.fromstring(parts(raw)['word/document.xml'])
    para=root.xpath('//w:body/w:p[w:r/w:t="培养类型：合成学位"]',namespaces=NS)[0]
    assert para.xpath('./w:pPr/w:jc/@w:val',namespaces=NS)==['center']
    assert para.xpath('./w:r/w:rPr/w:sz/@w:val',namespaces=NS)==['36']


@pytest.mark.parametrize('mutation',['source','style','section','chapter','duplicate','unknown_value'])
def test_unfrozen_references_and_unknown_values_are_rejected(mutation):
    value,analysis,fm=configured();options={}
    if mutation=='unknown_value':options['field_values']={'alien':'not accepted'}
    else:
        if mutation=='source':fm['fields'][0]['source_id']='invented'
        if mutation=='style':fm['fields'][0]['style_source_id']='invented'
        if mutation=='section':fm['fields'][0]['section_id']='invented'
        if mutation=='chapter':fm['content_requirements'][0]['chapter_id']='invented'
        if mutation=='duplicate':fm['fields'].append(deepcopy(fm['fields'][0]))
        value['thesis_context']['confirmation_sha256']=digest(value['thesis_context']['confirmation'])
    with pytest.raises((StorageError,ValueError)):plan(value,analysis,**options)


def test_empty_field_contract_does_not_invent_manual_fields_or_keyword_limits():
    value,analysis,fm=configured();fm['fields']=[];fm['content_requirements']=[]
    value['thesis_context']['confirmation_sha256']=digest(value['thesis_context']['confirmation'])
    p=plan(value,analysis)
    assert p['front_matter']['fields']==[]
    assert not any(i['id'].startswith(('school:manual:','school:keywords:','school:manual-chapter:')) for i in p['report']['issues'])


def test_second_template_changes_order_field_names_styles_and_cover_section():
    from app.adapters.manuscript_docx import build_manuscript_docx
    from lxml import etree as E
    value,analysis,fm=configured()
    other=synthetic('论文类别',14,WD_ALIGN_PARAGRAPH.RIGHT)
    analysis['front_matter_sources']=other['front_matter_sources']
    fm['cover_section_id']='section-2'
    fm['fields'][0].update(id='category',label='论文类别')
    fm['fields'].insert(0,dict(id='optional',label='可选联系字段',kind='text',required=False,source_id=fm['fields'][0]['source_id'],style_source_id=None,section_id='section-1',chapter_id=None))
    value['thesis_context']['confirmation']['analysis_sha256']=digest(analysis)
    value['thesis_context']['confirmation_sha256']=digest(value['thesis_context']['confirmation'])
    before=deepcopy(value)
    p=plan(value,analysis,field_values={'category':'合成类别','signature':'用户签名文字'})
    assert p['sections'][0]['id']=='section-2'
    assert not any(i['id']=='school:field:optional' for i in p['report']['issues'])
    raw=build_manuscript_docx(value,{},school_plan=p);content=text(raw)
    assert content.index('可选联系字段：待填写')<content.index('论文类别：合成类别')<content.index('作者签名：用户签名文字')
    root=E.fromstring(parts(raw)['word/document.xml'])
    para=root.xpath('//w:body/w:p[w:r/w:t="论文类别：合成类别"]',namespaces=NS)[0]
    assert para.xpath('./w:pPr/w:jc/@w:val',namespaces=NS)==['right']
    assert para.xpath('./w:r/w:rPr/w:sz/@w:val',namespaces=NS)==['28']
    assert value==before


def test_required_content_missing_and_signature_values_are_separate_checks():
    value,analysis,fm=configured()
    p=plan(value,analysis,field_values={'signature':'手工文字'})
    ids={i['id'] for i in p['report']['issues']}
    assert 'school:required-content:declare' in ids
    assert 'school:signature-review:signature' in ids and 'school:field:signature' not in ids
    fm['content_requirements'][0].update(kind='acknowledgements')
    value['thesis_context']['confirmation_sha256']=digest(value['thesis_context']['confirmation'])
    with pytest.raises(StorageError):plan(value,analysis)


def test_sealed_v2_plan_rejects_page_source_forgery():
    from app.domain.manuscripts.school_export import validate_school_export
    value,analysis,_=configured();p=plan(value,analysis)
    p['front_matter_sections']['section-1']['page']['width_cm']=99
    with pytest.raises(StorageError):validate_school_export(p,value,analysis)


def test_field_target_in_reference_chapter_is_rendered_after_references():
    from app.adapters.manuscript_docx import build_manuscript_docx
    value,analysis,fm=configured();fm['fields'][0]['chapter_id']='bibliography'
    value['thesis_context']['confirmation_sha256']=digest(value['thesis_context']['confirmation'])
    raw=build_manuscript_docx(value,{},school_plan=plan(value,analysis,field_values={'degree':'用户补充'}))
    assert '培养类型：用户补充' in text(raw)


def test_legacy_options_normalization_removes_only_new_empty_default():
    from app.domain.manuscripts.service import ManuscriptService
    value,_=fixture();version={'snapshot_json':value}
    old={'manual_fields':{'author':'','student_id':'','department':'','major':'','class_name':'','supervisor':'','submission_date':''},'manual_chapters':[],'english_title':'','keywords':[],'include_toc':True,'toc_depth':3}
    normalized=ManuscriptService.school_options(None,version,{'school':{}})
    assert normalized==old and 'field_values' not in normalized
    assert normalized['include_toc'] is True and normalized['toc_depth']==3 and 'author' in normalized['manual_fields']
    with pytest.raises(StorageError):ManuscriptService.school_options(None,version,{'school':{'field_values':{'x':'y'}}})
