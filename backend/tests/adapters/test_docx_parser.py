from io import BytesIO
import base64
import pytest
from docx import Document
from app.adapters.docx_parser import parse_docx, DOCX_SETTINGS
from app.core.exceptions import StorageError


def sample():
    doc = Document()
    doc.add_heading('中文标题', 1)
    t = doc.add_table(rows=2, cols=2)
    t.cell(0, 0).merge(t.cell(0, 1)).text = '合并'
    t.cell(1, 0).add_table(rows=1, cols=1).cell(0, 0).text = '内表'
    doc.add_paragraph('末尾')
    b = BytesIO(); doc.save(b)
    return b.getvalue()


def test_order_merge_nested_and_sources():
    raw = sample(); result = parse_docx(raw, '样本.docx', DOCX_SETTINGS)
    assert [b.kind for b in result.document.blocks] == ['paragraph', 'table', 'paragraph']
    assert result.document.blocks[0].heading_level == 1
    table = result.document.blocks[1]
    assert table.cells[0].column_span == 2
    assert table.cells[0].source.part == 'word/document.xml'
    assert any(b.kind == 'table' for b in table.cells[1].blocks)
    assert raw.startswith(b'PK')


def test_reject_wrong_office_and_limit():
    with pytest.raises(StorageError) as e: parse_docx(b'PKbad', 'x.docx', DOCX_SETTINGS)
    assert e.value.code == 'invalid_document'
    settings = {**DOCX_SETTINGS, 'docx': {**DOCX_SETTINGS['docx'], 'max_blocks': 1}}
    with pytest.raises(StorageError) as e: parse_docx(sample(), 'x.docx', settings)
    assert e.value.code == 'document_too_large'

@pytest.fixture(autouse=True)
def postgres_schema():
    """Pure parser tests require no database."""
    yield None

def rewrite(raw, changes):
    import zipfile
    src=zipfile.ZipFile(BytesIO(raw)); out=BytesIO()
    with zipfile.ZipFile(out,'w') as target:
        for name in src.namelist():target.writestr(name,changes.get(name,src.read(name)))
    return out.getvalue()

def test_inline_anchor_images_bytes_and_revisions():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1sAAAAASUVORK5CYII=')
    doc=Document();doc.add_picture(BytesIO(png));p=doc.add_paragraph('保留')
    ins=OxmlElement('w:ins');r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='修订';r.append(t);ins.append(r);p._p.append(ins)
    out=BytesIO();doc.save(out);raw=out.getvalue();original=bytes(raw)
    result=parse_docx(raw,'x.docx',DOCX_SETTINGS)
    assert base64.b64decode(result.image_payloads['image-1'])==png
    assert result.document.images[0].inline_preview
    assert result.document.blocks[-1].text=='保留'
    assert any(w.code=='revisions' for w in result.document.warnings)
    assert raw==original
    import zipfile
    xml=zipfile.ZipFile(BytesIO(raw)).read('word/document.xml').replace(b'wp:inline',b'wp:anchor')
    anchored=parse_docx(rewrite(raw,{'word/document.xml':xml}),'x.docx',DOCX_SETTINGS)
    assert anchored.document.blocks[0].image_ids==['image-1']

@pytest.mark.parametrize('mutation',['doctype','xlsx','macro','duplicate','path','ziplimit'])
def test_package_security(mutation):
    import zipfile
    raw=sample();src=zipfile.ZipFile(BytesIO(raw));changes={};settings=DOCX_SETTINGS
    if mutation=='doctype':changes['word/document.xml']=b'<!DOCTYPE x [<!ENTITY x "boom">]><x/>'
    if mutation=='xlsx':changes['[Content_Types].xml']=src.read('[Content_Types].xml').replace(b'wordprocessingml.document.main+xml',b'spreadsheetml.sheet.main+xml')
    if mutation=='macro':changes['[Content_Types].xml']=src.read('[Content_Types].xml').replace(b'wordprocessingml.document.main+xml',b'wordprocessingml.document.macroEnabled.main+xml')
    raw=rewrite(raw,changes)
    if mutation in ('duplicate','path'):
        out=BytesIO(raw)
        with zipfile.ZipFile(out,'a') as z:
            if mutation=='duplicate':
                with pytest.warns(UserWarning,match='Duplicate name'):z.writestr('word/document.xml',b'bad')
            else:z.writestr('../evil',b'bad')
        raw=out.getvalue()
    if mutation=='ziplimit':settings={**settings,'docx':{**settings['docx'],'max_uncompressed_bytes':1}}
    with pytest.raises(StorageError):parse_docx(raw,'x.docx',settings)

def test_vml_image_and_merged_continuation_are_reported():
    import zipfile
    raw=sample();src=zipfile.ZipFile(BytesIO(raw));xml=src.read('word/document.xml')
    xml=xml.replace(b'</w:body>',b'<w:p><w:r><w:pict><v:imagedata xmlns:v="urn:schemas-microsoft-com:vml"/></w:pict></w:r></w:p></w:body>')
    result=parse_docx(rewrite(raw,{'word/document.xml':xml}),'x.docx',DOCX_SETTINGS)
    assert any(w.code=='legacy_images' for w in result.document.warnings)

def test_vertical_merge_styles_external_and_other_parts():
    import zipfile
    doc=Document(); t=doc.add_table(rows=2,cols=1);t.cell(0,0).merge(t.cell(1,0)).text='纵向'
    doc.sections[0].header.paragraphs[0].text='页眉';doc.add_paragraph('样式',style='Title')
    out=BytesIO();doc.save(out);raw=out.getvalue();src=zipfile.ZipFile(BytesIO(raw))
    xml=src.read('word/document.xml').replace(b'</w:body>',b'<w:p><w:r><w:drawing><a:blip xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" r:link="external-test"/></w:drawing></w:r></w:p></w:body>')
    rel=src.read('word/_rels/document.xml.rels').replace(b'</Relationships>',b'<Relationship Id="external-test" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" Target="https://example.invalid/image.png" TargetMode="External"/></Relationships>')
    result=parse_docx(rewrite(raw,{'word/document.xml':xml,'word/_rels/document.xml.rels':rel}),'x.docx',DOCX_SETTINGS)
    assert len(result.document.blocks[0].cells)==1
    assert result.document.blocks[0].cells[0].row_span==2
    assert any(s.id=='Title' and s.based_on for s in result.document.styles)
    assert {w.code for w in result.document.warnings}>={'external_images','other_parts'}

def test_nested_limit_and_invalid_table_geometry():
    import zipfile
    raw=sample();src=zipfile.ZipFile(BytesIO(raw));xml=src.read('word/document.xml').replace(b'w:val="2"',b'w:val="-2"')
    with pytest.raises(StorageError):parse_docx(rewrite(raw,{'word/document.xml':xml}),'x.docx',DOCX_SETTINGS)

def doc_bytes(doc):
    out=BytesIO();doc.save(out);return out.getvalue()

def test_review_move_revisions_excluded():
    from docx.oxml import OxmlElement
    doc=Document();p=doc.add_paragraph('base')
    for tag in ('moveFrom','moveTo'):
        e=OxmlElement('w:'+tag);r=OxmlElement('w:r');t=OxmlElement('w:t');t.text='MOVED';r.append(t);e.append(r);p._p.append(e)
    result=parse_docx(doc_bytes(doc),'x.docx',DOCX_SETTINGS)
    assert result.document.blocks[0].text=='base'

def test_review_explicit_styles_keep_boolean_and_container_path():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();run=doc.add_paragraph().add_run('bold italic');run.bold=True;run.italic=True
    style=doc.styles['Normal']
    ppr=style.element.get_or_add_pPr();element=OxmlElement('w:color');element.set(qn('w:val'),'paragraph');ppr.append(element)
    rpr=style.element.get_or_add_rPr();element=OxmlElement('w:color');element.set(qn('w:val'),'run');rpr.append(element)
    result=parse_docx(doc_bytes(doc),'x.docx',DOCX_SETTINGS)
    assert result.document.blocks[0].runs[0].properties=={'b.val':'true','i.val':'true'}
    props=next(s.properties for s in result.document.styles if s.id=='Normal')
    assert props['pPr.color.val']=='paragraph'
    assert props['rPr.color.val']=='run'

def test_review_hmerge_anchor():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='merged'
    for col,value in ((0,'restart'),(1,'continue')):
        element=OxmlElement('w:hMerge');element.set(qn('w:val'),value);table.cell(0,col)._tc.get_or_add_tcPr().append(element)
    result=parse_docx(doc_bytes(doc),'x.docx',DOCX_SETTINGS)
    assert [(c.column,c.column_span) for c in result.document.blocks[0].cells]==[(0,2)]

def test_review_absolute_package_image_target():
    import zipfile
    png=base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAQAAAC1HAwCAAAAC0lEQVR42mP8/x8AAwMCAO+aX1sAAAAASUVORK5CYII=')
    doc=Document();doc.add_picture(BytesIO(png));raw=doc_bytes(doc)
    rels=zipfile.ZipFile(BytesIO(raw)).read('word/_rels/document.xml.rels')
    result=parse_docx(rewrite(raw,{'word/_rels/document.xml.rels':rels.replace(b'Target="media/image1.png"',b'Target="/word/media/image1.png"')}),'x.docx',DOCX_SETTINGS)
    assert len(result.document.images)==1
    assert base64.b64decode(result.image_payloads['image-1'])==png

def test_hmerge_with_vertical_merge_keeps_single_rectangle():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document();table=doc.add_table(rows=2,cols=2)
    for row in range(2):
        for col in range(2):
            tc=table.cell(row,col)._tc
            for tag,value in (('hMerge','restart' if col==0 else 'continue'),('vMerge','restart' if row==0 else 'continue')):
                element=OxmlElement('w:'+tag);element.set(qn('w:val'),value);tc.get_or_add_tcPr().append(element)
    result=parse_docx(doc_bytes(doc),'x.docx',DOCX_SETTINGS)
    assert [(c.row_span,c.column_span) for c in result.document.blocks[0].cells]==[(2,2)]

def test_outline_level_nine_is_body_and_retains_explicit_property():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc=Document()
    for value in (9,0,8):
        p=doc.add_paragraph('正文或标题')
        outline=OxmlElement('w:outlineLvl');outline.set(qn('w:val'),str(value));p._p.get_or_add_pPr().append(outline)
    result=parse_docx(doc_bytes(doc),'outline.docx',DOCX_SETTINGS)
    assert [b.heading_level for b in result.document.blocks]==[None,1,9]
    assert result.document.blocks[0].properties['outlineLvl.val']=='9'
