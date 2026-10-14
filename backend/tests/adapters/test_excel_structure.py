from io import BytesIO
from zipfile import ZipFile, ZIP_DEFLATED
import pytest
from openpyxl import Workbook
from app.core.config import ExcelSettings


def sample(patch=None):
    book = Workbook()
    sheet = book.active
    sheet.append(['name', 'value'])
    for row in range(2, 32):
        sheet.append([f'row{row}', row])
    data = BytesIO()
    book.save(data)
    if patch:
        out = BytesIO()
        with ZipFile(data) as source, ZipFile(out, 'w', ZIP_DEFLATED) as target:
            for name in source.namelist():
                value = source.read(name)
                if name == 'xl/worksheets/sheet1.xml':
                    value = patch(value.decode()).encode()
                target.writestr(name, value)
        return out.getvalue()
    return data.getvalue()


def test_full_structure_and_lexical_precision():
    from app.adapters.excel_structure import extract_excel_structure
    content = sample(lambda xml: xml.replace('<v>31</v>', '<v>9007199254740993</v>'))
    result = extract_excel_structure(content, 'test.xlsx', ExcelSettings())
    sheet = result.sheets[0]
    assert sheet.cells[-1].row == 31
    assert sheet.cells[-1].value == '9007199254740993'
    assert sheet.row_count == 31
    assert sheet.nonempty_data_row_count == 30


@pytest.mark.parametrize('text_xml,expected', [
    ('<t>base</t>', 'base'),
    ('<r><t>rich</t></r><r><t> text</t></r>', 'rich text'),
    ('<t/>', ''),
])
def test_inline_phonetic_annotations_are_not_cell_text(text_xml, expected):
    from app.adapters.excel_structure import extract_excel_structure
    content = sample(lambda xml: xml.replace('<is><t>name</t></is>',
        '<is>' + text_xml + '<rPh sb="0" eb="1"><t>phonetic</t></rPh></is>'))
    sheet = extract_excel_structure(content, 'test.xlsx', ExcelSettings()).sheets[0]
    assert sheet.cells[0].value == expected
    assert sheet.fields[0].header_value == expected


@pytest.mark.parametrize('kind,value,raw_type,raw_value', [
    ('boolean', True, 'b', '0'), ('boolean', False, 'b', '1'),
    ('boolean', True, 'b', 'true'), ('boolean', True, 'n', '1'),
    ('text', 'x', 'b', '0'), ('text', 'x', 'str', 'y'),
    ('error', '#DIV/0!', 'e', '#VALUE!'), ('error', '#DIV/0!', 'n', '1'),
    ('blank', None, 'n', '1'), ('number', '1', 'b', '1'),
    ('text', 'x', 'unknown', None),
])
def test_contract_rejects_contradictory_raw_evidence(kind, value, raw_type, raw_value):
    from app.domain.datasets.contracts import StructureCell
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        StructureCell(row=1, column=1, coordinate='A1', kind=kind, value=value,
                      raw_type=raw_type, raw_value=raw_value, number_format='General', formula=None)


@pytest.mark.parametrize('kind,value,raw_type,raw_value', [
    ('boolean', True, 'b', '1'), ('boolean', False, 'b', '0'),
    ('text', 'resolved shared text', 's', '12'),
    ('text', 'inline text', 'inlineStr', None), ('text', '', 'str', None),
    ('error', '#DIV/0!', 'e', '#DIV/0!'), ('blank', None, 'n', None),
    ('datetime', '2026-10-10T00:00:00', 'n', '46305'),
    ('datetime', '2026-10-10T00:00:00', 'd', '2026-10-10T00:00:00'),
])
def test_contract_accepts_coherent_raw_evidence(kind, value, raw_type, raw_value):
    from app.domain.datasets.contracts import StructureCell
    cell = StructureCell(row=1, column=1, coordinate='A1', kind=kind, value=value,
                         raw_type=raw_type, raw_value=raw_value, number_format='General', formula=None)
    assert cell.value == value
    assert cell.raw_value == raw_value


def test_iso_duration_xml_preserves_duration_and_raw_evidence():
    from app.adapters.excel_structure import extract_excel_structure
    content = sample(lambda xml: xml.replace('<c r="B31" t="n"><v>31</v></c>',
        '<c r="B31" t="d"><v>PT1H</v></c>'))
    cell = extract_excel_structure(content, 'test.xlsx', ExcelSettings()).sheets[0].cells[-1]
    assert cell.kind == 'duration'
    assert cell.value == '1:00:00'
    assert cell.raw_type == 'd'
    assert cell.raw_value == 'PT1H'


@pytest.mark.parametrize('cache,state,value', [('', 'missing', None), ('<v/>','empty',None), ('<v>0</v>','present','0')])
def test_formula_xml_cache(cache, state, value):
    from app.adapters.excel_structure import extract_excel_structure
    content = sample(lambda xml: xml.replace('<v>31</v>', '<f t="shared" si="0">SUM(B2:B30)</f>' + cache))
    cell = extract_excel_structure(content, 'test.xlsx', ExcelSettings()).sheets[0].cells[-1]
    assert cell.kind == 'formula'
    assert cell.formula.cache_state == state
    assert cell.formula.cached_value == value
    assert cell.formula.cache_validity == 'unverified'


def test_contract_rejects_boolean_coordinate():
    from app.domain.datasets.contracts import StructureCell
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        StructureCell(row=True, column=1, coordinate='A1', kind='blank', value=None,
                      raw_type=None, raw_value=None, number_format='General', formula=None)


@pytest.mark.parametrize('kind,value', [('number','NaN'),('number','oops'),('boolean','true'),('blank',''),('text',True)])
def test_contract_rejects_inconsistent_value(kind,value):
    from app.domain.datasets.contracts import StructureCell
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        StructureCell(row=1,column=1,coordinate='A1',kind=kind,value=value,raw_type=None,
                      raw_value=None,number_format='General',formula=None)


def test_empty_strings_blanks_missing_and_wrong_dimension():
    from app.adapters.excel_structure import extract_excel_structure
    def patch(xml):
        xml = xml.replace('ref="A1:B31"','ref="A1:A1"')
        return xml.replace('<c r="B31" t="n"><v>31</v></c>', '<c r="B31" t="inlineStr"><is><t> </t></is></c><c r="C31" t="inlineStr"><is><t/></is></c><c r="D31"/>')
    sheet = extract_excel_structure(sample(patch),'test.xlsx',ExcelSettings()).sheets[0]
    assert sheet.row_count == 31
    assert [(c.kind,c.value) for c in sheet.cells[-3:]] == [('text',' '),('text',''),('blank',None)]
    assert all(c.coordinate != 'E31' for c in sheet.cells)


@pytest.mark.parametrize('replacement', ['r="B30"','r="A31"','r="B0"','r="B031"'])
def test_rejects_duplicate_or_contradictory_xml_coordinate(replacement):
    from app.adapters.excel_structure import extract_excel_structure
    from app.core.exceptions import StorageError
    with pytest.raises(StorageError):
        extract_excel_structure(sample(lambda xml:xml.replace('r="B31"',replacement)),'test.xlsx',ExcelSettings())


def test_hidden_merge_date_and_empty_sheet():
    from datetime import datetime
    from app.adapters.excel_structure import extract_excel_structure
    book=Workbook()
    sheet=book.active
    sheet['A1']='repeat'
    sheet['B1']='repeat'
    sheet['A2']=datetime(2026,10,10)
    sheet.merge_cells('A3:C5')
    sheet['A3']='anchor'
    hidden=book.create_sheet('hidden'); hidden.sheet_state='hidden'
    data=BytesIO(); book.save(data)
    result=extract_excel_structure(data.getvalue(),'test.xlsx',ExcelSettings())
    sheet=result.sheets[0]
    assert (sheet.row_count,sheet.column_count)==(5,3)
    assert sheet.merged_ranges==['A3:C5']
    assert next(c for c in sheet.cells if c.coordinate=='A2').kind=='datetime'
    assert not any(c.coordinate=='B3' for c in sheet.cells)
    assert result.sheets[1].row_count==0
    assert 'structure_hidden_sheet' in result.sheets[1].warnings


def test_output_limit_fails_instead_of_truncating():
    from app.adapters.excel_structure import extract_excel_structure
    from app.core.exceptions import StorageError
    book=Workbook(); sheet=book.active
    for row in range(1,151): sheet.cell(row,1,'x'*30000)
    data=BytesIO(); book.save(data)
    with pytest.raises(StorageError) as exc:
        extract_excel_structure(data.getvalue(),'test.xlsx',ExcelSettings())
    assert exc.value.code=='parse_output_invalid'


def test_contract_rejects_forged_header_and_extent():
    from app.adapters.excel_structure import extract_excel_structure
    from app.domain.datasets.contracts import ExcelStructure
    from pydantic import ValidationError
    data=extract_excel_structure(sample(),'test.xlsx',ExcelSettings()).model_dump()
    data['sheets'][0]['fields'][0]['header_value']='forged'
    with pytest.raises(ValidationError): ExcelStructure.model_validate(data)


def test_contract_schema_version_is_not_boolean():
    from app.adapters.excel_structure import extract_excel_structure
    from app.domain.datasets.contracts import ExcelStructure
    from pydantic import ValidationError
    data=extract_excel_structure(sample(),'test.xlsx',ExcelSettings()).model_dump()
    data['schema_version']=True
    with pytest.raises(ValidationError): ExcelStructure.model_validate(data)


@pytest.mark.parametrize('raw_type,cache', [('b','1'),('e','#DIV/0!')])
def test_array_formula_preserves_cache_type(raw_type,cache):
    from app.adapters.excel_structure import extract_excel_structure
    def patch(xml):
        return xml.replace('<c r="B31" t="n"><v>31</v></c>',f'<c r="B31" t="{raw_type}"><f t="array" ref="B31:B31">A1</f><v>{cache}</v></c>')
    cell=extract_excel_structure(sample(patch),'test.xlsx',ExcelSettings()).sheets[0].cells[-1]
    assert cell.formula.cached_type==raw_type
    assert cell.formula.cached_value==cache
    assert cell.formula.attributes=={'t':'array','ref':'B31:B31'}
