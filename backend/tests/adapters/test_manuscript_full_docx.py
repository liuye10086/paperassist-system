"""Complete chapters use their frozen writing text, retaining factual table/image assets."""
from copy import deepcopy
from io import BytesIO

from docx import Document
import pytest

from tests.adapters.test_manuscript_docx import sample


ORDER = ('title', 'abstract', 'introduction', 'results', 'discussion',
         'materials_and_methods', 'references')


def full_sample(language='en'):
    snapshot, images = sample(language)
    snapshot['title'] = 'Generated manuscript title' if language == 'en' else '由正文提炼的稿件标题'
    snapshot['sections'] = [
        dict(id=section, paragraphs=[
            dict(id=f'writing:{section}:1', text=snapshot['title'] if section == 'title'
                 else 'NEW-' + section.upper(), origin='model_generated',
                 source_ids=[], citations=[], numeric_bindings=[])
        ] if section != 'references' else [])
        for section in ORDER
    ]
    snapshot['keywords'] = ['user keyword', 'explicit keyword']
    snapshot['acknowledgements'] = 'USER-ACKNOWLEDGEMENT'
    snapshot['writing_provenance'] = dict(
        provider_provenance='synthetic_fixture', verification_status='not_verified',
        request_id='private-writing-request', prompt_version='internal-prompt-id')
    return snapshot, images


@pytest.mark.parametrize('language', ['zh-CN', 'en'])
def test_full_docx_orders_chapters_and_renders_rewritten_text_once(language):
    from app.adapters.manuscript_docx import build_manuscript_docx
    snapshot, images = full_sample(language)
    original = deepcopy(snapshot)
    document = Document(BytesIO(build_manuscript_docx(snapshot, images)))
    text = '\n'.join(p.text for p in document.paragraphs)
    headings = [p.text for p in document.paragraphs if p.style.name == 'Heading 1']
    expected = (['Abstract', 'Introduction', 'Results', 'Discussion', 'Materials and methods', 'References']
                if language == 'en' else ['摘要', '引言', '结果', '讨论', '材料与方法', '参考文献'])
    assert headings[:6] == expected
    assert document.paragraphs[0].text == snapshot['title']
    assert document.paragraphs[0].style.name == 'Title'
    for section in ORDER[1:-1]:
        assert text.count('NEW-' + section.upper()) == 1
    assert 'METHODS-SENTINEL' not in text and 'RESULT-SENTINEL' not in text
    assert all(item in text for item in ('CAPTION-SENTINEL', 'PANEL-SENTINEL',
                                         'ETHICS-MISSING-SENTINEL', 'SAMPLES-CONFLICT-SENTINEL'))
    assert text.count('USER-ACKNOWLEDGEMENT') == 1
    assert 'user keyword' in text and 'explicit keyword' in text
    assert '(1) Author A.' in text
    assert 'private-writing-request' not in text and 'internal-prompt-id' not in text
    assert len(document.inline_shapes) == 1
    assert any(cell.text == '5.24' for table in document.tables for row in table.rows for cell in row.cells)
    assert snapshot == original


def test_full_docx_empty_references_remain_seventh_part_and_no_default_extras():
    from app.adapters.manuscript_docx import build_manuscript_docx
    snapshot, images = full_sample()
    snapshot.update(references=[], keywords=[], acknowledgements='')
    document = Document(BytesIO(build_manuscript_docx(snapshot, images)))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert 'References' in text
    assert 'No verified references' in text
    assert 'Keywords' not in text and 'Acknowledgements' not in text


@pytest.mark.parametrize('invalid', ['missing', 'duplicate', 'reordered', 'title_mismatch'])
def test_full_docx_does_not_silently_fallback_or_export_ambiguous_chapters(invalid):
    from app.adapters.manuscript_docx import build_manuscript_docx
    from app.core.exceptions import StorageError
    snapshot, images = full_sample()
    if invalid == 'missing':
        snapshot['sections'].pop()
    elif invalid == 'duplicate':
        snapshot['sections'][2]['id'] = 'abstract'
    elif invalid == 'reordered':
        snapshot['sections'][2], snapshot['sections'][3] = snapshot['sections'][3], snapshot['sections'][2]
    else:
        snapshot['sections'][0]['paragraphs'][0]['text'] = 'Different title'
    with pytest.raises(StorageError) as caught:
        build_manuscript_docx(snapshot, images)
    assert caught.value.code == 'report_render_failed'


def test_full_docx_rejects_unrecorded_body_reference_to_missing_figure():
    from app.adapters.manuscript_docx import build_manuscript_docx
    from app.core.exceptions import StorageError
    snapshot,images=full_sample()
    snapshot['sections'][3]['paragraphs'][0]['text']='See Figure 12 for the saved image.'
    with pytest.raises(StorageError) as caught:
        build_manuscript_docx(snapshot,images)
    assert caught.value.code=='report_render_failed'
