from io import BytesIO
from hashlib import sha256
from zipfile import ZipFile, ZIP_DEFLATED

import pytest
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt


@pytest.fixture(autouse=True)
def postgres_schema():
    yield None


def sample(*, conflict=False):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.left_margin = Cm(2.5)
    section.bottom_margin = section.right_margin = Cm(2)
    section.header.paragraphs[0].text = '学校论文页眉'
    doc.styles['Normal'].font.name = 'Times New Roman'
    doc.styles['Normal'].font.size = Pt(12)
    doc.add_paragraph('页面设置为上方和左侧分别留边2.5cm，下方和右侧分别留2.0cm')
    doc.add_paragraph('中文论文题目')
    doc.add_paragraph('摘 要')
    doc.add_paragraph('摘要字数300字左右。摘要的字数控制在300-500。')
    doc.add_paragraph('关键词：工程；设计')
    doc.add_heading('1 前言', 1)
    doc.add_paragraph('这是一段从默认样式继承字体和字号的正文，保留其属性来源并进行人工核对。')
    doc.add_heading('2 工艺设计', 1)
    doc.add_paragraph('图2.1 结构图', 'Caption')
    doc.add_paragraph('学位论文作者签名：由本人提供。')
    doc.add_paragraph('参考文献')
    doc.add_paragraph('凡正文引用他人观点，要在当页进行脚注标注。')
    if conflict:
        doc.add_section(WD_SECTION_START.NEW_PAGE).left_margin = Cm(3)
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def rewrite(raw, changes):
    out = BytesIO()
    with ZipFile(BytesIO(raw)) as source, ZipFile(out, 'w', ZIP_DEFLATED) as target:
        for name in source.namelist():
            target.writestr(name, changes.get(name, source.read(name)))
        for name, value in changes.items():
            if name not in source.namelist():
                target.writestr(name, value)
    return out.getvalue()


def test_reads_real_properties_and_school_chapter_order_without_mutation():
    from app.adapters.school_template_parser import inspect_school_template
    raw = sample()
    original = sha256(raw).hexdigest()
    result = inspect_school_template(raw, '学校.docx')
    assert sha256(raw).hexdigest() == original
    assert result['parser_version'] == 'school-template-openxml-v1'
    assert [c['title'] for c in result['chapters'] if c['title'][0].isdigit()] == ['1 前言', '2 工艺设计']
    assert '摘 要' in [c['title'] for c in result['chapters']]
    assert result['sections'][0]['page']['margin_left_cm'] == 2.5
    body = next(r for r in result['rules'] if r['id'] == 'body.font_size_pt')
    assert body['value'] == 12
    assert body['source']['part'] == 'word/styles.xml'
    assert body['origin'] == 'style_inheritance'
    assert result['raw_styles'] and result['defaults']


def test_margin_and_text_conflicts_have_concrete_source_options():
    from app.adapters.school_template_parser import inspect_school_template
    result = inspect_school_template(sample(conflict=True), '学校.docx')
    conflict = next(c for c in result['conflicts'] if c['rule_id'] == 'section-2.margin_left_cm')
    assert [o['value'] for o in conflict['options']] == [3.0, 2.5]
    assert all(o['source']['path'] for o in conflict['options'])
    assert any(c['rule_id'] == 'abstract.word_count' for c in result['conflicts'])


def test_header_link_and_page_fields_are_inspected_not_updated():
    from app.adapters.school_template_parser import inspect_school_template
    doc = Document(BytesIO(sample()))
    p = doc.sections[0].footer.paragraphs[0]
    field = OxmlElement('w:fldSimple')
    field.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}instr', 'PAGE')
    p._p.append(field)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    out = BytesIO(); doc.save(out)
    result = inspect_school_template(out.getvalue(), '学校.docx')
    assert result['sections'][1]['headers'][0]['linked_to_previous'] is True
    assert result['sections'][1]['headers'][0]['text'] == '学校论文页眉'
    assert 'PAGE' in result['sections'][0]['footers'][0]['fields']
    assert any('签名' in field for field in result['identity_fields'])


def test_textbox_norms_quoted_once_and_never_used_as_chapter_body():
    from app.adapters.school_template_parser import inspect_school_template
    raw = sample()
    with ZipFile(BytesIO(raw)) as package:
        xml = package.read('word/document.xml')
    box = '<w:p><w:r><w:pict><w:txbxContent><w:p><w:r><w:t>一级标题：黑体，三号字</w:t></w:r></w:p></w:txbxContent></w:pict></w:r></w:p>'.encode()
    result = inspect_school_template(rewrite(raw, {'word/document.xml': xml.replace(b'</w:body>', box + box + b'</w:body>')}), '学校.docx')
    assert sum('一级标题：黑体' in n['text'] for n in result['notes']) == 1
    assert all('一级标题：黑体' not in c['title'] for c in result['chapters'])


def test_ole_is_preserved_as_unsupported_and_unknown_template_is_honest():
    from app.adapters.school_template_parser import inspect_school_template
    doc = Document(); doc.add_paragraph('任意用户文字')
    out = BytesIO(); doc.save(out)
    raw = rewrite(out.getvalue(), {'word/embeddings/oleObject1.bin': b'passive binary'})
    result = inspect_school_template(raw, '未知.docx')
    assert not result['chapters']
    assert result['missing']
    assert any(x['id'] == 'embedded_objects' for x in result['unsupported'])
    assert any(x['id'] == 'text_rules_manual' for x in result['unsupported'])


def test_invalid_docx_uses_existing_bounded_validation():
    from app.adapters.school_template_parser import inspect_school_template
    from app.core.exceptions import StorageError
    with pytest.raises(StorageError):
        inspect_school_template(b'bad', 'bad.docx')


def test_direct_outline_levels_and_cached_toc_are_distinguished():
    from app.adapters.school_template_parser import inspect_school_template
    doc = Document()
    p = doc.add_paragraph('1 实际学校章节')
    outline = OxmlElement('w:outlineLvl')
    outline.set('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}val', '0')
    p._p.get_or_add_pPr().append(outline)
    cached = doc.add_paragraph('1 实际学校章节1', 'Heading 1')
    instruction = OxmlElement('w:instrText'); instruction.text = 'PAGEREF _Toc123 \\h'
    cached.add_run()._r.append(instruction)
    out = BytesIO(); doc.save(out)
    result = inspect_school_template(out.getvalue(), '学校.docx')
    assert [c['title'] for c in result['chapters']] == ['1 实际学校章节']


def test_supplemental_notes_join_same_conflict_check_with_their_source():
    from app.adapters.school_template_parser import inspect_school_template, merge_template_notes
    base = inspect_school_template(sample(), '学校.docx')
    note = {'id': 'supplement-1', 'text': '页面设置为上方和左侧分别留边3.0cm，下方和右侧分别留2.0cm',
            'source': {'part': 'supplement.pdf', 'path': 'page:2', 'text': '学院补充规范'}}
    result = merge_template_notes(base, [note])
    conflict = next(c for c in result['conflicts'] if c['rule_id'] == 'section-1.margin_left_cm')
    assert {o['value'] for o in conflict['options']} == {2.5, 3.0}
    assert any(o['source']['path'] == 'page:2' for o in conflict['options'])
    assert len(base['notes']) < len(result['notes'])


def test_instruction_paragraphs_with_heading_attributes_are_not_chapters():
    from app.adapters.school_template_parser import inspect_school_template
    doc = Document()
    doc.add_heading('1 实际章节', 1)
    doc.add_paragraph('自绘图应去掉底色、网格线等，只能用黑色的字体和线条。', 'Heading 2')
    out = BytesIO(); doc.save(out)
    result = inspect_school_template(out.getvalue(), 'x.docx')
    assert [c['title'] for c in result['chapters']] == ['1 实际章节']
    assert any('自绘图应去掉' in n['text'] for n in result['notes'])


def test_nonabstract_word_counts_remain_quotes_instead_of_false_abstract_rules():
    from app.adapters.school_template_parser import inspect_school_template, merge_template_notes
    base = inspect_school_template(sample(), '学校.docx')
    note = {'id': 'full-text', 'text': '正文字数要求8000-12000字。全文字数要求10000字左右。',
            'source': {'part': 'supplement_text', 'path': 'user', 'text': '全文要求'}}
    result = merge_template_notes(base, [note])
    conflict = next(c for c in result['conflicts'] if c['rule_id'] == 'abstract.word_count')
    assert all('10000' not in str(o['value']) and '8000' not in str(o['value']) for o in conflict['options'])
    assert note['text'] in [n['text'] for n in result['notes']]


def test_heading_font_notes_do_not_read_the_next_heading_role():
    from app.adapters.school_template_parser import inspect_school_template, merge_template_notes
    base = inspect_school_template(sample(), '学校.docx')
    note = {'id': 'headings', 'text': '一级标题：三号宋体；二级标题：四号黑体。',
            'source': {'part': 'supplement_text', 'path': 'user', 'text': '标题要求'}}
    result = merge_template_notes(base, [note])
    first = next(r for r in result['rules'] if r['id'] == 'heading1.font_east_asia')
    assert first['value'] == '宋体'


def test_template_process_never_falls_back_on_windows(monkeypatch):
    from app.adapters import school_template_process
    from app.core.exceptions import StorageError
    monkeypatch.setattr(school_template_process.sys, 'platform', 'win32')
    with pytest.raises(StorageError) as caught:
        school_template_process.inspect_school_template_isolated(sample(), 'x.docx')
    assert caught.value.code == 'parse_environment_unavailable'
