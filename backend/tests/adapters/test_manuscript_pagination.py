"""SCI table pagination must not suppress rows in the pinned PDF converter."""
from io import BytesIO

from docx import Document
from docx.oxml.ns import qn
import pytest

from tests.adapters.test_manuscript_docx import sample


@pytest.mark.parametrize('language', ['en', 'zh-CN'])
def test_table_rows_allow_page_continuation_and_retain_repeated_header(language):
    from app.adapters.manuscript_docx import build_manuscript_docx

    snapshot, images = sample(language)
    table = snapshot['results_groups'][0]['tables'][0]
    table['rows'] = [[f'ROW-{i:02d}', str(i)] for i in range(37)]
    document = Document(BytesIO(build_manuscript_docx(snapshot, images)))
    rendered = document.tables[0]
    assert [[cell.text for cell in row.cells] for row in rendered.rows] == [table['columns'], *table['rows']]
    assert rendered.rows[0]._tr.trPr.find(qn('w:tblHeader')) is not None
    # The pinned LibreOffice converter can discard the continuation of a
    # bottom-of-page table when its data rows prohibit splitting. The default
    # row policy permits continuation, including rows taller than one page.
    for row in rendered.rows[1:]:
        policy = row._tr.find(qn('w:trPr') + '/' + qn('w:cantSplit'))
        assert policy is None or policy.get(qn('w:val')) in ('0', 'false', 'off')

@pytest.mark.parametrize('language', ['en', 'zh-CN'])
def test_table_cells_do_not_inherit_body_trailing_space_across_page_boundary(language):
    from app.adapters.manuscript_docx import build_manuscript_docx

    snapshot, images = sample(language)
    table = snapshot['results_groups'][0]['tables'][0]
    table['rows'] = [[f'ROW-{i:02d}', str(i)] for i in range(37)]
    document = Document(BytesIO(build_manuscript_docx(snapshot, images)))
    # Normal's 6pt body spacing can continue alone after a page break in the
    # pinned converter, leaving an empty grid row below the repeated header.
    # Cells need their own zero trailing space; ordinary prose keeps its gap.
    assert document.styles['Normal'].paragraph_format.space_after.pt == 6
    for row in document.tables[0].rows:
        for cell in row.cells:
            assert cell.paragraphs[0].paragraph_format.space_after is not None
            assert cell.paragraphs[0].paragraph_format.space_after.pt == 0
