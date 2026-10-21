"""Scientific content must survive additive layout without image regeneration."""
import hashlib
import importlib
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from lxml import etree
import pytest

from app.core.exceptions import StorageError

NS = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
      'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}


@pytest.fixture(autouse=True)
def postgres_schema():
    yield None


@pytest.fixture
def source_docx():
    from PIL import Image
    doc = Document()
    doc.add_paragraph('User-authored scientific text; p=0.012; 中文。')
    doc.add_heading('Methods', level=1)
    doc.add_paragraph('Unmodified formula source anchor.')
    picture = BytesIO()
    Image.new('RGB', (30, 20), 'blue').save(picture, format='PNG')
    doc.add_picture(BytesIO(picture.getvalue()))
    doc.add_paragraph('Figure 1. Original image.', style='Caption')
    table = doc.add_table(rows=2, cols=2)
    for row in table.rows:
        for cell in row.cells:
            cell.text = 'Original table'
    doc.add_paragraph('Table 1. Original table.', style='Caption')
    doc.add_heading('Results', level=1)
    doc.add_paragraph('Final original paragraph.')
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


def api():
    return importlib.import_module('app.adapters.document_layout')


def options(value):
    from app.domain.document_layout import LayoutOptions
    return LayoutOptions.model_validate(value)


def parts(data):
    with ZipFile(BytesIO(data)) as z:
        return {n: z.read(n) for n in z.namelist()}


def xml(data):
    return etree.fromstring(parts(data)['word/document.xml'])


def test_catalog_and_layout_api_exist():
    assert importlib.util.find_spec('app.adapters.document_layout') is not None


def test_catalog_is_bound_to_bytes_and_original_body_positions(source_docx):
    catalog = api().catalog_docx(source_docx)
    assert catalog['source_sha256'] == hashlib.sha256(source_docx).hexdigest()
    assert catalog['source_size'] == len(source_docx)
    assert {'paragraph', 'heading', 'figure', 'table'} <= {a['type'] for a in catalog['anchors']}
    assert catalog == api().catalog_docx(source_docx)
    assert len({a['id'] for a in catalog['anchors']}) == len(catalog['anchors'])


def test_unknown_anchor_does_not_silently_append(source_docx):
    with pytest.raises(StorageError) as error:
        api().apply_layout(source_docx, options({'formulas': [
            {'id': 'eq1', 'anchor_id': 'missing', 'expression': r'\frac{a}{b}', 'placement': 'after'}
        ]}), language='zh-CN')
    assert error.value.code == 'layout_target_invalid'
    assert error.value.params['path'] == 'formulas.0.anchor_id'


def test_source_mismatch_and_unsupported_input_rejected(source_docx):
    with pytest.raises(StorageError, match='来源'):
        api().apply_layout(source_docx, options({'source_sha256': 'a' * 64}), language='en')
    with pytest.raises(StorageError) as error:
        api().catalog_docx(b'not zip')
    assert error.value.code == 'layout_source_invalid'


def test_omml_is_native_ordered_and_original_text_images_preserved(source_docx):
    anchors = api().catalog_docx(source_docx)['anchors']
    values = {'formulas': [
        {'id': 'later', 'anchor_id': anchors[-1]['id'], 'expression': r'\matrix{a & b \\ c & d}'},
        {'id': 'earlier', 'anchor_id': anchors[1]['id'], 'expression': r'\frac{x_i^2}{\sqrt{a}}+\alpha'},
        {'id': 'inline', 'anchor_id': anchors[2]['id'], 'expression': r'\int_0^1{x dx}', 'placement': 'inline'},
    ]}
    result = api().apply_layout(source_docx, options(values), language='en')
    root = xml(result)
    assert len(root.xpath('//m:oMath', namespaces=NS)) == 3
    for tag in ('f', 'rad', 'sSubSup', 'm', 'nary'):
        assert root.xpath(f'//m:{tag}', namespaces=NS)
    assert root.xpath('//m:naryPr/m:subHide/@m:val', namespaces=NS) == ['0']
    assert root.xpath('//m:naryPr/m:supHide/@m:val', namespaces=NS) == ['0']
    original_text = xml(source_docx).xpath('//w:t/text()', namespaces=NS)
    new_text = iter(root.xpath('//w:t/text()', namespaces=NS))
    assert all(any(v == text for v in new_text) for text in original_text)
    for name, content in parts(source_docx).items():
        if name.startswith('word/media/'):
            assert parts(result)[name] == content
    assert b'user_input' in parts(result)['customXml/paperassist-layout.xml']
    assert root.xpath('//w:instrText/text()', namespaces=NS).count(' SEQ PAEquation \\* ARABIC ') == 2


def test_toc_cross_references_sections_and_repeat_header(source_docx):
    anchors = api().catalog_docx(source_docx)['anchors']
    heading = anchors[1]['id']
    table = next(a['id'] for a in anchors if a['type'] == 'table')
    result = api().apply_layout(source_docx, options({
        'toc': {'depth': 3}, 'pagination': {'mode': 'roman_front', 'body_anchor_id': heading},
        'landscape_anchor_ids': [table],
        'references': [{'anchor_id': anchors[-1]['id'], 'target_id': heading, 'kind': 'page'}],
    }), language='en')
    root = xml(result)
    fields = root.xpath('//w:instrText/text()', namespaces=NS)
    assert any('TOC ' in f for f in fields)
    assert any('PAGEREF PA' in f for f in fields)
    assert root.xpath('//w:pgSz[@w:orient="landscape"]', namespaces=NS)
    assert root.xpath('//w:pgNumType[@w:fmt="lowerRoman"]', namespaces=NS)
    assert root.xpath('//w:pgNumType[@w:fmt="decimal"][@w:start="1"]', namespaces=NS)
    assert root.xpath('//w:tblHeader', namespaces=NS)
    assert root.xpath('//w:sectPr', namespaces=NS)[-1].find('w:pgSz', NS).get(f'{{{NS["w"]}}}orient') != 'landscape'


def test_reference_to_paragraph_or_inline_formula_rejected(source_docx):
    anchors = api().catalog_docx(source_docx)['anchors']
    for target in (anchors[0]['id'], 'formula:in'):
        with pytest.raises(StorageError) as error:
            api().apply_layout(source_docx, options({
                'formulas': [{'id': 'in', 'anchor_id': anchors[0]['id'], 'expression': 'x', 'placement': 'inline'}],
                'references': [{'anchor_id': anchors[2]['id'], 'target_id': target}],
            }), language='en')
        assert error.value.code == 'layout_target_invalid'


def test_cache_merge_requires_complete_results_and_preserves_media_math(source_docx):
    import json
    anchors = api().catalog_docx(source_docx)['anchors']
    pending = api().apply_layout(source_docx, options({
        'toc': {'depth': 2},
        'formulas': [{'id': 'eq', 'anchor_id': anchors[1]['id'], 'expression': r'\frac{a}{b}'}],
        'references': [{'anchor_id': anchors[-1]['id'], 'target_id': 'formula:eq', 'kind': 'number'}],
    }), language='en')
    manifest = json.loads(etree.fromstring(parts(pending)['customXml/paperassist-layout.xml']).text)
    cache = {'source_sha256': hashlib.sha256(pending).hexdigest(), 'fields': {}, 'toc_lines': ['Methods\t1', 'Results\t2']}
    for field in manifest['fields']:
        if field['role'] != 'toc':
            cache['fields'][field['bookmark']] = '1'
    finished = api().merge_field_cache(pending, cache)
    assert len(xml(finished).xpath('//m:oMath', namespaces=NS)) == 1
    assert 'Methods1' in ''.join(xml(finished).xpath('//w:t/text()', namespaces=NS))
    assert xml(finished).xpath('//w:hyperlink[@w:anchor]/w:r/w:tab', namespaces=NS)
    assert b'complete' in parts(finished)['customXml/paperassist-layout.xml']
    assert all(parts(finished)[n] == b for n, b in parts(pending).items() if n.startswith('word/media/'))
    cache['fields'] = {}
    with pytest.raises(StorageError) as error:
        api().merge_field_cache(pending, cache)
    assert error.value.code == 'layout_field_update_failed'


def test_unbreakable_wide_expression_rejected(source_docx):
    anchor = api().catalog_docx(source_docx)['anchors'][0]['id']
    with pytest.raises(StorageError) as error:
        api().apply_layout(source_docx, options({'formulas': [{'id': 'wide', 'anchor_id': anchor, 'expression': 'x+' * 40 + 'x'}]}), language='en')
    assert error.value.code == 'layout_formula_too_wide'


def test_figure_text_reference_uses_original_caption_not_image(source_docx):
    anchors = api().catalog_docx(source_docx)['anchors']
    figure = next(a for a in anchors if a['type'] == 'figure')
    result = api().apply_layout(source_docx, options({
        'references': [{'anchor_id': anchors[-1]['id'], 'target_id': figure['id'], 'kind': 'text'}],
    }), language='en')
    root = xml(result)
    instruction = next(x for x in root.xpath('//w:instrText/text()', namespaces=NS) if 'REF ' in x)
    name = instruction.split()[1]
    mark = root.xpath('//w:bookmarkStart[@w:name=$name]', namespaces=NS, name=name)[0]
    assert ''.join(mark.getparent().xpath('.//w:t/text()', namespaces=NS)) == 'Figure 1. Original image.'
    assert not mark.getparent().xpath('.//w:drawing', namespaces=NS)


def test_existing_footer_text_survives_added_page_number():
    doc = Document()
    doc.add_heading('Body', 1)
    doc.sections[0].footer.paragraphs[0].text = 'Confidential footer'
    stream = BytesIO()
    doc.save(stream)
    result = api().apply_layout(stream.getvalue(), options({'pagination': {'mode': 'continuous'}}), language='en')
    footer_parts = {n: b for n, b in parts(result).items() if n.startswith('word/footer')}
    assert any(b'Confidential footer' in b and b'PAGE' in b for b in footer_parts.values())


def test_landscape_keeps_preceding_table_caption_with_table():
    doc = Document()
    doc.add_heading('Body', 1)
    doc.add_paragraph('Table 1. Original pre-table caption.', style='Caption')
    doc.add_table(rows=2, cols=2).cell(0, 0).text = 'Original table cell'
    doc.add_paragraph('After table')
    stream = BytesIO()
    doc.save(stream)
    content = stream.getvalue()
    table = next(a for a in api().catalog_docx(content)['anchors'] if a['type'] == 'table')
    result = api().apply_layout(content, options({'landscape_anchor_ids': [table['id']]}), language='en')
    root = xml(result)
    caption = root.xpath('//w:p[w:r/w:t="Table 1. Original pre-table caption."]', namespaces=NS)[0]
    assert caption.getnext().tag == f'{{{NS["w"]}}}tbl'
    assert caption.find('w:pPr/w:keepNext', NS) is not None
    assert caption.getprevious().find('w:pPr/w:pStyle', NS).get(f'{{{NS["w"]}}}val') == 'Heading1'
    assert caption.getprevious().getprevious() is None


@pytest.mark.parametrize('expression', [r'x^\sum', r'x_\int', r'\sum_\sum{x}', r'x^\sum{y}'])
def test_bare_nary_scripts_fail_with_public_location(source_docx, expression):
    anchor = api().catalog_docx(source_docx)['anchors'][0]['id']
    with pytest.raises(StorageError) as error:
        api().apply_layout(source_docx, options({'formulas': [
            {'id': 'bad', 'anchor_id': anchor, 'expression': expression}]}), language='en')
    assert error.value.code == 'layout_formula_invalid'
    assert error.value.params['path'] == 'formulas.0.expression'
    assert error.value.params['formula_id'] == 'bad'
    assert isinstance(error.value.params['position'], int)


def test_footer_inheritance_preserves_default_first_even_and_independent_sections():
    from docx.enum.section import WD_SECTION_START
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    first = doc.sections[0]
    first.different_first_page_header_footer = True
    first.footer.paragraphs[0].text = 'Confidential default'
    first.first_page_footer.paragraphs[0].text = 'Confidential first'
    first.even_page_footer.paragraphs[0].text = 'Confidential even'
    doc.add_paragraph('First section')
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph('All footer types inherited')
    third = doc.add_section(WD_SECTION_START.NEW_PAGE)
    third.footer.is_linked_to_previous = False
    third.footer.paragraphs[0].text = 'Independent default'
    doc.add_paragraph('Independent default; inherited first/even')
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    result = api().apply_layout(source, options({'pagination': {'mode': 'continuous'}}), language='en')
    ns = NS | {'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    refs = lambda data: [
        [(n.get(f'{{{NS["w"]}}}type'), n.get(f'{{{ns["r"]}}}id')) for n in s.findall('w:footerReference', ns)]
        for s in xml(data).xpath('//w:sectPr', namespaces=ns)]
    assert refs(result) == refs(source)
    for name, value in parts(result).items():
        if name.startswith('word/footer') and name.endswith('.xml'):
            assert b'PAGE' in value


@pytest.mark.parametrize('start', [1, 5])
def test_landscape_split_keeps_source_page_start_only_on_first_fragment(start):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.enum.section import WD_SECTION_START
    doc = Document()
    doc.add_heading('First source section', 1)
    for _ in range(2):
        doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Landscape object'
        doc.add_paragraph('Continue original section')
    # Explicit intentional restart on the next original source section.
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading('Second source section', 1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Third landscape object'
    doc.add_paragraph('Final source paragraph')
    for section, number in zip(doc.sections, (start, 11)):
        pg = OxmlElement('w:pgNumType')
        pg.set(qn('w:start'), str(number))
        section._sectPr.append(pg)
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    tables = [a['id'] for a in api().catalog_docx(source)['anchors'] if a['type'] == 'table']
    result = api().apply_layout(source, options({'landscape_anchor_ids': tables}), language='en')
    sections = xml(result).xpath('//w:sectPr', namespaces=NS)
    starts = [s.xpath('./w:pgNumType/@w:start', namespaces=NS) for s in sections]
    # Headings now travel with their tables. When the block is first in an
    # original section, it owns that section's restart without an empty page.
    assert starts == [[str(start)], [], [], [], ['11'], []]


def test_toc_owned_bookmarks_are_not_shared_with_cross_references(source_docx):
    import json
    anchors = api().catalog_docx(source_docx)['anchors']
    pending = api().apply_layout(source_docx, options({
        'toc': {'depth': 2},
        'references': [{'anchor_id': anchors[-1]['id'], 'target_id': anchors[1]['id'], 'kind': 'page'}],
    }), language='en')
    meta = json.loads(etree.fromstring(parts(pending)['customXml/paperassist-layout.xml']).text)
    toc_names = {target['bookmark'] for target in meta['toc_targets']}
    ref_names = {field['instruction'].split()[1] for field in meta['fields'] if field['role'] == 'reference'}
    assert toc_names.isdisjoint(ref_names), 'Word deletes old TOC link targets while rebuilding its TOC'
    cache = {'source_sha256': hashlib.sha256(pending).hexdigest(),
             'fields': {field['bookmark']: '1' for field in meta['fields'] if field['role'] != 'toc'},
             'toc_lines': ['Methods\t1', 'Results\t1']}
    ready = api().merge_field_cache(pending, cache)
    assert set(xml(ready).xpath('//w:hyperlink/@w:anchor', namespaces=NS)) == toc_names


def _figure_block_source(*, styled_caption):
    from PIL import Image
    from docx.shared import Inches
    doc = Document()
    doc.add_paragraph('Earlier source content')
    doc.add_heading('Labelled figures', 1)
    image = BytesIO()
    Image.new('RGB', (600, 400), 'blue').save(image, format='PNG')
    doc.add_picture(BytesIO(image.getvalue()), width=Inches(6))
    doc.add_paragraph('Figure 1  Original plot', style='Caption' if styled_caption else None)
    doc.add_paragraph('Complete original caption with all explanatory words.', style='Caption' if styled_caption else None)
    if styled_caption:
        doc.add_paragraph('A  Original panel A', style='Caption')
        doc.add_paragraph('B  Original panel B', style='Caption')
    doc.add_heading('Following section', 1)
    doc.add_paragraph('Following scientific text')
    out = BytesIO()
    doc.save(out)
    return out.getvalue()


@pytest.mark.parametrize('styled_caption', [True, False])
def test_landscape_figure_contains_heading_and_complete_caption_block(styled_caption):
    source = _figure_block_source(styled_caption=styled_caption)
    figure = next(a for a in api().catalog_docx(source)['anchors'] if a['type'] == 'figure')
    result = api().apply_layout(source, options({'landscape_anchor_ids': [figure['id']]}), language='en')
    root = xml(result)
    heading = root.xpath('//w:p[w:r/w:t="Labelled figures"]', namespaces=NS)[0]
    assert heading.getprevious().find('w:pPr/w:sectPr', NS) is not None
    last_text = 'B  Original panel B' if styled_caption else 'Complete original caption with all explanatory words.'
    last = root.xpath('//w:p[w:r/w:t=$text]', namespaces=NS, text=last_text)[0]
    assert last.getnext().find('w:pPr/w:sectPr', NS) is not None
    current = heading
    while current is not last:
        assert current.find('w:pPr/w:sectPr', NS) is None
        assert current.find('w:pPr/w:keepNext', NS) is not None
        current = current.getnext()
    assert last.find('w:pPr/w:keepLines', NS) is not None
    assert all(parts(result)[n] == b for n, b in parts(source).items() if n.startswith('word/media/'))


def test_portrait_figure_keeps_full_plain_caption_with_image():
    source = _figure_block_source(styled_caption=False)
    result = api().apply_layout(source, options({}), language='en')
    root = xml(result)
    title = root.xpath('//w:p[w:r/w:t="Figure 1  Original plot"]', namespaces=NS)[0]
    caption = title.getnext()
    assert title.find('w:pPr/w:keepNext', NS) is not None
    assert caption.find('w:pPr/w:keepLines', NS) is not None
    assert caption.find('w:pPr/w:keepNext', NS).get(f'{{{NS["w"]}}}val') == 'false'


def test_landscape_table_keeps_immediately_preceding_heading():
    doc = Document()
    doc.add_paragraph('Earlier source content')
    doc.add_heading('Statistical results', 1)
    doc.add_paragraph('Table 1. Original result', style='Caption')
    doc.add_table(rows=2, cols=2).cell(0, 0).text = 'Header'
    doc.add_paragraph('Following original text')
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    table = next(a for a in api().catalog_docx(source)['anchors'] if a['type'] == 'table')
    result = api().apply_layout(source, options({'landscape_anchor_ids': [table['id']]}), language='en')
    heading = xml(result).xpath('//w:p[w:r/w:t="Statistical results"]', namespaces=NS)[0]
    assert heading.getprevious().find('w:pPr/w:sectPr', NS) is not None
    assert heading.getnext().getnext().tag == f'{{{NS["w"]}}}tbl'


def test_first_object_heading_does_not_create_empty_portrait_section():
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    doc.add_heading('First content', 1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'First table'
    doc.add_paragraph('Following source')
    number = OxmlElement('w:pgNumType')
    number.set(qn('w:start'), '5')
    doc.sections[0]._sectPr.append(number)
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    target = next(a['id'] for a in api().catalog_docx(source)['anchors'] if a['type'] == 'table')
    root = xml(api().apply_layout(source, options({'landscape_anchor_ids': [target]}), language='en'))
    body = root.find('w:body', NS)
    assert body[0].find('w:pPr/w:sectPr', NS) is None
    sections = root.findall('.//w:sectPr', NS)
    assert len(sections) == 2
    assert sections[0].find('w:pgNumType', NS).get(qn('w:start')) == '5'
    assert sections[1].find('w:pgNumType', NS).get(qn('w:start')) is None


@pytest.mark.parametrize('boundary', ['heading', 'section', 'empty'])
def test_caption_detection_does_not_cross_source_boundary(boundary):
    from app.adapters.document_layout.engine import Editor
    from docx.oxml import OxmlElement
    from PIL import Image
    doc = Document()
    image = BytesIO()
    Image.new('RGB', (20, 20), 'blue').save(image, format='PNG')
    doc.add_picture(BytesIO(image.getvalue()))
    paragraph = doc.add_paragraph(
        '' if boundary == 'empty' else 'Figure 1  Following source',
        style='Heading 1' if boundary == 'heading' else 'Caption')
    if boundary == 'section':
        paragraph._p.get_or_add_pPr().append(OxmlElement('w:sectPr'))
    doc.add_paragraph('Unrelated source prose')
    stream = BytesIO()
    doc.save(stream)
    editor = Editor(stream.getvalue(), options({}), 'en')
    assert editor.figure_captions(editor.body[0]) == []


def test_roman_body_and_landscape_start_share_one_boundary():
    doc = Document()
    doc.add_heading('Body begins', 1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Original table'
    doc.add_paragraph('After table')
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    anchors = api().catalog_docx(source)['anchors']
    result = api().apply_layout(source, options({
        'toc': {'depth': 1},
        'pagination': {'mode': 'roman_front', 'body_anchor_id': anchors[0]['id']},
        'landscape_anchor_ids': [anchors[1]['id']],
    }), language='en')
    root = xml(result)
    heading = root.xpath('//w:p[w:r/w:t="Body begins"]', namespaces=NS)[0]
    assert heading.getprevious().find('w:pPr/w:sectPr', NS) is not None
    assert heading.getprevious().getprevious().find('w:pPr/w:sectPr', NS) is None
    sections = root.findall('.//w:sectPr', NS)
    assert len(sections) == 3
    assert [s.xpath('./w:pgNumType/@w:fmt', namespaces=NS) for s in sections] == [
        ['lowerRoman'], ['decimal'], ['decimal']]
    assert [s.xpath('./w:pgNumType/@w:start', namespaces=NS) for s in sections] == [['1'], ['1'], []]


def test_preceding_caption_section_boundary_preserves_original_section():
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    doc.sections[0].footer.paragraphs[0].text = 'FIRST FOOTER'
    doc.add_paragraph('Previous source section content')
    caption = doc.add_paragraph('Table 1. Original caption', style='Caption')
    second = doc.add_section(WD_SECTION_START.NEW_PAGE)
    second.footer.is_linked_to_previous = False
    second.footer.paragraphs[0].text = 'SECOND FOOTER'
    separator = caption._p.getnext()
    section = separator.find('w:pPr/w:sectPr', NS)
    section.getparent().remove(section)
    caption._p.get_or_add_pPr().append(section)
    separator.getparent().remove(separator)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Second section table'
    doc.add_paragraph('Second section after')
    for section, start in zip(doc.sections, (5, 11)):
        number = OxmlElement('w:pgNumType')
        number.set(qn('w:start'), str(start))
        section._sectPr.append(number)
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    target = next(a['id'] for a in api().catalog_docx(source)['anchors'] if a['type'] == 'table')
    result = api().apply_layout(source, options({'landscape_anchor_ids': [target]}), language='en')
    original_sections = xml(source).findall('.//w:sectPr', NS)
    sections = xml(result).findall('.//w:sectPr', NS)
    assert len(sections) == 3
    assert etree.tostring(sections[0]) == etree.tostring(original_sections[0])
    assert [s.xpath('./w:pgNumType/@w:start', namespaces=NS) for s in sections] == [['5'], ['11'], []]
    assert [s.find('w:footerReference', NS).get(qn('r:id')) for s in sections] == [
        original_sections[0].find('w:footerReference', NS).get(qn('r:id')),
        original_sections[1].find('w:footerReference', NS).get(qn('r:id')),
        original_sections[1].find('w:footerReference', NS).get(qn('r:id')),
    ]


@pytest.mark.parametrize('before', [True, False])
def test_trailing_table_caption_reuses_its_original_section_end(before):
    from pathlib import Path
    import runpy
    from app.adapters.document_layout.engine import Editor
    fixture = runpy.run_path(str(Path(__file__).parents[1] / 'fixtures/layout_section_caption.py'))
    source = fixture['source_docx'](before)
    target = next(a['id'] for a in api().catalog_docx(source)['anchors'] if a['type'] == 'table')
    result = api().apply_layout(source, options({'landscape_anchor_ids': [target]}), language='en')
    root = xml(result)
    table = root.find('w:body/w:tbl', NS)
    caption = table.getnext()
    assert caption.xpath('.//w:t/text()', namespaces=NS) == ['Table 1. Original caption']
    end = caption.find('w:pPr/w:sectPr', NS)
    assert end is not None
    assert end.find('w:pgSz', NS).get(f'{{{NS["w"]}}}orient') == 'landscape'
    assert len(root.findall('.//w:sectPr', NS)) == (3 if before else 2)
    original = xml(source)
    assert etree.tostring(root.find('w:body/w:sectPr', NS)) == etree.tostring(original.find('w:body/w:sectPr', NS))
    # The next original section starts immediately after the original caption.
    assert caption.getnext().xpath('.//w:t/text()', namespaces=NS) == ['Second source content']
    assert all(parts(result)[n] == value for n, value in parts(source).items() if n.startswith('word/footer'))
    editor = Editor(result, options({}), 'en')
    assert editor.effective_section(editor.body.find(f'{{{NS["w"]}}}tbl')).find('w:pgSz', NS).get(f'{{{NS["w"]}}}orient') == 'landscape'
