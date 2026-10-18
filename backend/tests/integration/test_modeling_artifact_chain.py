"""Real isolated persistence with synthetic independent curve computation only."""
from tests.helpers.statistics import statistics_post
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import json
import math
import pytest
from psycopg.types.json import Jsonb
from app.db.database import database_connection
from tests.api.test_analysis import client
from tests.helpers.modeling import sample,selection,prepared_modeling
from tests.integration.test_correlation_plan import prepared_correlations
from tests.integration.test_comparison_plan import make_plan,approve
from tests.integration.test_dataset_artifact_chain import confirmed_task
from tests.integration.test_boxplot_execution import SyntheticPlotProvider
from tests.integration.test_correlation_artifact_chain import ModernExplanationProvider,finish_plot
from tests.integration.test_explanation_execution import run as run_explanation


def result_for(client,method,language='en'):
    from tests.domain.test_correlation_plan_contracts import research
    base,file,prepared=prepared_modeling(client,method)
    declaration=research();declaration.update(output_language=language,question='LOCAL-QUESTION-SENTINEL')
    task,_=make_plan(client,base,prepared,selection(method,prepared['source']['sheet_name']),declaration)
    approve(client,task)
    response=statistics_post(client, base+'/analysis-runs',json={'expected_revision':1});assert response.status_code==200,response.text
    return base,file,response.json()


def artifact_task(client,base,file,result,kind,figure=None,explanation=None):
    task,owner=confirmed_task(client,base,file,result,kind,figure,explanation)
    if kind!='word_report':
        from app.domain.plot_policy import derive_execution_policy as plot_policy
        from app.domain.explanation_policy import derive_execution_policy as explanation_policy
        key='boxplot_policy' if kind=='boxplot' else 'explanation_policy'
        with database_connection(write=True) as db:
            snapshot=db.execute('SELECT input_snapshot FROM tasks WHERE id=%s',(task['id'],)).fetchone()['input_snapshot']
            policy=(plot_policy if kind=='boxplot' else explanation_policy)(snapshot[key],{'schema_version':4})
            db.execute('UPDATE tasks SET input_snapshot=jsonb_set(input_snapshot,%s,%s) WHERE id=%s',([key],Jsonb(policy),task['id']))
    return task,owner


def independent_roc(observations,direction):
    positive=sum(r['outcome'] for r in observations);negative=len(observations)-positive
    points=[dict(threshold_kind='above_all' if direction=='higher_is_positive' else 'below_all',threshold=None,tpr=0.,fpr=0.,sensitivity=0.,specificity=1.)]
    for threshold in sorted({r['score'] for r in observations},reverse=direction=='higher_is_positive'):
        chosen=[r for r in observations if r['score']>=threshold] if direction=='higher_is_positive' else [r for r in observations if r['score']<=threshold]
        tp=sum(r['outcome'] for r in chosen);fp=len(chosen)-tp;tpr=tp/positive;fpr=fp/negative
        points.append(dict(threshold_kind='value',threshold=threshold,tpr=tpr,fpr=fpr,sensitivity=tpr,specificity=1-fpr))
    area=sum((b['fpr']-a['fpr'])*(b['tpr']+a['tpr'])/2 for a,b in zip(points,points[1:]))
    return dict(n=len(observations),positive_n=positive,negative_n=negative,auc=area,points=points)


def independent_km(observations,labels):
    maximum=max(r['time'] for r in observations);ticks=sorted(set(maximum*i/4 for i in range(5)));groups=[]
    for index,label in enumerate(labels):
        chosen=[r for r in observations if r['group']==index];survival=1.;greenwood=0.;points=[]
        for time in sorted({r['time'] for r in chosen}):
            risk=sum(r['time']>=time for r in chosen);events=sum(r['time']==time and r['event'] for r in chosen);censored=sum(r['time']==time and not r['event'] for r in chosen)
            survival*=1-events/risk
            if risk==events: greenwood=None
            elif events: greenwood+=events/(risk*(risk-events))
            if survival in (0,1):variance=lo=hi=None;reason='boundary_survival_interval_unavailable'
            else:
                variance=survival**2*greenwood;se=math.sqrt(greenwood)/abs(math.log(survival));transformed=math.log(-math.log(survival));z=1.959963984540054
                lo=math.exp(-math.exp(transformed+z*se));hi=math.exp(-math.exp(transformed-z*se));reason=None
            points.append(dict(time=time,n_at_risk=risk,n_event=events,n_censored=censored,survival=survival,greenwood_variance=variance,ci_low=lo,ci_high=hi,ci_reason=reason))
        groups.append(dict(label=label,n=len(chosen),event_n=sum(r['event'] for r in chosen),censored_n=sum(not r['event'] for r in chosen),max_time=max(r['time'] for r in chosen),points=points,
            risk_table=[dict(time=t,n_at_risk=sum(r['time']>=t for r in chosen),events_before=sum(r['time']<t and r['event'] for r in chosen),censored_before=sum(r['time']<t and not r['event'] for r in chosen)) for t in ticks]))
    return dict(n=len(observations),groups=groups)


class ModelingPlotProvider(SyntheticPlotProvider):
    """Does not receive local expected/result; implements independent ROC/KM calculations."""
    def __init__(self,corrupt=False):super().__init__(None,pending=False);self.corrupt=corrupt
    def receipt(self):
        value=super().receipt();value['response']['output'][1]['content'][0]['annotations'][0]['filename']='/mnt/data/figure.png';return value
    def download(self,file_id,container_id,limit,*,policy):
        self.downloads.append(file_id)
        if file_id=='cfile_png':return self.png
        p=self.payload;kind=p['figure_kind'];config=p['plot_config']
        numeric=independent_roc(p['observations'],config['direction']) if kind=='roc_curve' else independent_km(p['observations'],p['group_labels']) if kind=='kaplan_meier' else dict(n=p['sample_count'],effects=p['effects'])
        if self.corrupt:numeric['n']+=1
        manifest={k:deepcopy(p[k]) for k in ('schema_version','figure_kind','language','variables','plot_config')}
        manifest.update(p['labels'],statistics=numeric)
        return json.dumps(manifest,ensure_ascii=False,allow_nan=False).encode()


@pytest.mark.parametrize('method,language',[('ols','en'),('logistic','zh-CN'),('roc_auc','en'),('kaplan_meier','zh-CN')])
def test_modeling_real_consumers_preserve_original_png_and_complete_tables(client,monkeypatch,method,language):
    from tests.browser.task_browser_support import block_real_providers
    blocked=block_real_providers(monkeypatch)
    base,file,result=result_for(client,method,language)
    task,_=artifact_task(client,base,file,result,'boxplot');provider=ModelingPlotProvider();done=finish_plot(task,provider)
    assert done and done['status']=='succeeded',done
    outbound=json.dumps(provider.payload)
    assert 'statistics' not in provider.payload and 'LOCAL-QUESTION-SENTINEL' not in outbound and 'PRIVATE' not in outbound
    if method in ('roc_auc','kaplan_meier'):assert all('row' not in r for r in provider.payload['observations'])
    with database_connection() as db:figure=json.loads(db.execute('SELECT figure_json FROM figures WHERE id=%s',(done['result_figure_id'],)).fetchone()['figure_json'])
    assert figure['schema_version']==4 and figure['engine']['id']=='openai_modeling_plot_v4'
    task,_=artifact_task(client,base,file,result,'explanation',figure);explain=ModernExplanationProvider();done=run_explanation(task,explain)
    assert done and done['status']=='succeeded',done
    assert 'LOCAL-QUESTION-SENTINEL' not in json.dumps(explain.posts[0][1])
    with database_connection() as db:explanation=json.loads(db.execute('SELECT explanation_json FROM explanations WHERE id=%s',(done['result_explanation_id'],)).fetchone()['explanation_json'])
    assert explanation['schema_version']==4 and len(explanation['sections'])==6
    task,_=artifact_task(client,base,file,result,'word_report',figure,explanation)
    from app.domain.tasks.execution import run_word_task
    done=run_word_task(task['id'],task['revision']);assert done and done['status']=='succeeded',done
    workspace=client.get('/api/v1/tasks/'+task['id']+'/workspace');assert workspace.status_code==200,workspace.text
    url=workspace.json()['artifacts']['report']['download_url'];response=client.get(url);assert response.status_code==200,response.text if response.status_code!=200 else ''
    from docx import Document
    doc=Document(BytesIO(response.content));assert len(doc.tables)>=3
    assert 'LOCAL-QUESTION-SENTINEL' in '\n'.join(p.text for p in doc.paragraphs)
    with ZipFile(BytesIO(response.content)) as archive:assert [archive.read(n) for n in archive.namelist() if n.startswith('word/media/')]==[provider.png]
    assert blocked==[]
    png_url=base+'/analysis-runs/'+result['id']+'/figures/'+figure['id']+'/download'
    for change in ('count','scale','engine'):
        changed=deepcopy(figure)
        if change=='count':changed['statistics']['n']+=1
        elif change=='scale':changed['plot_config']['verification_scope']='independent_model_refit'
        else:changed['engine']['id']='openai_modeling_plot_v5'
        with database_connection(write=True) as db:db.execute('UPDATE figures SET figure_json=%s WHERE id=%s',(json.dumps(changed),figure['id']))
        assert client.get(png_url).status_code==409,(method,change)
        assert client.get(url).status_code==409,(method,change)
    with database_connection(write=True) as db:db.execute('UPDATE figures SET figure_json=%s WHERE id=%s',(json.dumps(figure),figure['id']))
    from tests.api.test_analysis import selection as descriptive_selection
    revision=client.put(base+'/analysis-setup',json=descriptive_selection(dataset_id=result['source']['dataset_id'],expected_revision=1,
        sheet_name=result['selection']['sheet_name'],numeric_column='B',group_column=None,unit='day' if method=='kaplan_meier' else ''))
    assert revision.status_code==200,revision.text
    assert client.get(url).content==response.content
    assert client.get(png_url).content==provider.png
    from app.maintenance.database_audit import audit_database
    with database_connection() as db:assert audit_database(db)['ok']


@pytest.mark.parametrize('method',['ols','roc_auc','kaplan_meier'])
def test_modeling_corrupt_cloud_count_fails_without_publish(client,method):
    base,file,result=result_for(client,method);task,_=artifact_task(client,base,file,result,'boxplot')
    done=finish_plot(task,ModelingPlotProvider(corrupt=True));assert done and done['status']=='failed' and done['error_code']=='plot_result_mismatch',done
    with database_connection() as db:assert db.execute('SELECT count(*) AS n FROM figures').fetchone()['n']==0
    workspace=client.get('/api/v1/tasks/'+task['id']+'/workspace');assert workspace.status_code==200,workspace.text
    failure=workspace.json()['verification_failure'];assert failure['differences'] and 'observations' not in json.dumps(failure)
