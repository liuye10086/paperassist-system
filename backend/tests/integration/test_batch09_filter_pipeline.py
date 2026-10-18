"""Real isolated PostgreSQL filter publication, provenance and iteration."""
from tests.helpers.legacy_plans import seed_legacy_plan
from tests.helpers.statistics import statistics_post
import copy
import json
from pathlib import Path
from tests.api.test_analysis import client
from tests.integration.test_batch07_quality_pipeline import published,quality_task
from tests.integration.test_dataset_preparation import run,review,decide,decision
from app.db.database import database_connection,migration_connection

ROWS=[['age','group','value'],[17,'A',10],[18,'A',20],[30,'B',30],[65,'A',40],[None,'A',50],[40,'A',60],[45,'B',70]]
RECIPE=dict(conditions=[dict(column=1,kind='number',operator='between',lower='18',upper='65'),dict(column=2,kind='text',operator='eq',value='A')],reason='成人A组',unit_acknowledged=True)

def filtering(client,base,parent,key='filter',recipe=None):
    request=dict(base_dataset_id=parent['id'],recipe=recipe or RECIPE)
    response=client.post(base+'/dataset-population-filters',json=request,headers={'Idempotency-Key':key})
    assert response.status_code==202,response.text
    return response.json(),request

def accept(client,task,key='decision'):
    waiting=run(task)
    assert waiting['status']=='waiting_confirmation',waiting
    view=review(client,task)
    response=decide(client,task,decision(view),key)
    assert response.status_code==202,response.text
    assert run(response.json())['status']=='succeeded'
    return review(client,task)['result'],view,response.json()

def test_complete_population_publish_counts_download_and_replay(client):
    base,parent=published(client,ROWS)
    old=client.get(parent['download_url']).content
    task,request=filtering(client,base,parent)
    result,view,accepted=accept(client,task)
    assert result['source']['schema_version']==3 and result['operation']=='population_filter'
    assert result['summary']['retained_record_count']==3
    assert result['summary']['mismatch_excluded_row_count']==3 and result['summary']['unknown_excluded_row_count']==1
    rowview=client.get(base+'/prepared-dataset-versions/'+result['id']+'/rows').json()
    assert [r['row'] for r in rowview['rows']]==[3,5,7]
    entries=client.get(result['impact_download_url']).json()['entries']
    assert len(entries)==7
    assert client.get(result['impact_url'],params={'operation':'unknown'}).json()['items'][0]['row']==6
    assert client.get(parent['download_url']).content==old
    assert client.post(base+'/dataset-population-filters',json=request,headers={'Idempotency-Key':'filter'}).json()['id']==task['id']
    assert run(accepted) is None
    with database_connection() as db:
        assert db.execute('SELECT count(*) AS n FROM prepared_dataset_derivations_v3').fetchone()['n']==1
    target=Path(__file__).resolve().parents[3]/'script/.runtime/stage04-batch09/filter-api-fixture.json'
    target.write_text(json.dumps(dict(defaults=client.get(base+'/prepared-dataset-versions/'+parent['id']+'/filter-defaults').json(),review=view,published=result,rows=rowview,impacts=client.get(result['impact_url']).json()),ensure_ascii=False,indent=2),encoding='utf-8')
    from tests.api.test_analysis import selection
    from tests.helpers.analysis_plan import run_plan,review_plan,plan_decision,decide_plan
    response=seed_legacy_plan(client, base+'/analysis-plans',json=dict(dataset_id=result['id'],expected_setup_revision=0,
        selection=selection(numeric_column='C',group_column=None)),headers={'Idempotency-Key':'filtered-plan'})
    assert response.status_code==202,response.text
    planned=response.json()
    assert run_plan(planned)['status']=='waiting_confirmation'
    assert client.get(base+'/analysis-setup').json() is None
    approved=decide_plan(client,planned,plan_decision(review_plan(client,planned)))
    assert approved.status_code==202,approved.text
    assert run_plan(approved.json())['status']=='succeeded'
    computed=statistics_post(client, base+'/analysis-runs',json={'expected_revision':1})
    assert computed.status_code==200,computed.text
    assert computed.json()['overall']['n']==3 and computed.json()['overall']['mean']==40
    assert computed.json()['source']==result['source']
    from app.domain.dataset_profile.service import DatasetProfileService
    from app.adapters.storage import ProjectStore
    from app.core.config import get_data_dir
    from tests.helpers.tasks import task_owner
    store=ProjectStore(get_data_dir(),owner_id=task_owner(client))
    profile=DatasetProfileService(store).profile(base.split('/')[4],base.split('/')[6],result['id'])
    assert profile['record_count']==3
    # Real verified data feed plotting material and Word; bitmap and prose are explicit synthetic fixtures.
    from app.domain.dataset_analysis import load_dataset
    from app.domain.boxplot_material import build_material
    from app.domain.external_processing import disclosure,freeze_confirmation
    from app.adapters.task_store import TaskStore
    from app.domain.tasks.contracts import TaskCreateRequest
    from app.core.config import get_excel_settings
    scoped=TaskStore(task_owner(client),store.db_config)
    with scoped.connection() as db:
        frozen=scoped.source_snapshot(db,scoped.require_project(db,base.split('/')[4]),TaskCreateRequest(task_type='boxplot',file_id=base.split('/')[6],analysis_run_id=computed.json()['id'],expected_revision=1))
    disclosure_view=disclosure(frozen,'boxplot',task_owner(client))
    external=freeze_confirmation(frozen,'boxplot',task_owner(client),dict(version=disclosure_view['version'],confirmed=True,
        source_digest=disclosure_view['source_digest'],labels={x['key']:x['value'] for x in disclosure_view['labels']}))
    record,_,sheet=load_dataset(store,base.split('/')[4],base.split('/')[6],result['id'],expected_source=result['source'])
    payload,expected=build_material(computed.json(),record,None,get_excel_settings(),dataset_sheet=sheet,external=external)
    assert payload['values']==[20.0,40.0,60.0] and expected['overall']['n']==3
    from tests.adapters.test_analysis_report_docx import snapshot
    from tests.adapters.test_analysis_plot_response import png
    from app.domain.dataset_processing import load_processing
    from app.adapters.report_docx import build_report
    from app.adapters.report_store import ReportStore
    from io import BytesIO
    from docx import Document
    from app.domain.artifact_contracts import EXPLANATION_V1
    saved=snapshot('zh-CN');saved['result']=computed.json();saved['explanation']['engine']['id']=EXPLANATION_V1
    saved['processing']=load_processing(store,base.split('/')[4],result['source'])
    document=Document(BytesIO(build_report(saved,ReportStore.metadata(saved),png())))
    text='\n'.join(p.text for p in document.paragraphs)
    assert '当前整理记录数：3' in text and '本轮筛选缺失未决排除：1' in text
    assert 'AND' in text and '成人A组' in text
    table_values=[[cell.text for cell in row.cells] for table in document.tables for row in table.rows]
    assert any(row[:3]==['总体','3','40'] for row in table_values)
    from app.maintenance.dataset_source_audit import audit_dataset_sources
    with database_connection() as db:
        checks=audit_dataset_sources(db)
    assert all(c['invalid']==0 for c in checks),checks

def test_mixed_chain_quality_filter_quality_filter_keeps_exclusions(client):
    base,parent=published(client,ROWS)
    task,_=quality_task(client,base,parent,dict(fields=[dict(column=c) for c in (1,2,3)],manual_exclusions=[dict(row=3,reason='user exclusion')]))
    parent,_,_=accept(client,task)
    assert parent['source']['schema_version']==2
    task,_=filtering(client,base,parent);parent,_,_=accept(client,task)
    assert parent['summary']['retained_record_count']==2
    task,_=quality_task(client,base,parent,dict(fields=[dict(column=c) for c in (1,2,3)],manual_exclusions=[dict(row=5,reason='another exclusion')]),key='v3quality')
    parent,_,_=accept(client,task)
    assert parent['source']['schema_version']==3 and parent['operation']=='quality'
    task,_=filtering(client,base,parent,key='filter2');result,_,_=accept(client,task)
    assert result['summary']['inherited_excluded_row_count']==6
    assert [r['row'] for r in client.get(base+'/prepared-dataset-versions/'+result['id']+'/rows').json()['rows']]==[7]
    with database_connection() as db:
        assert db.execute('SELECT count(*) AS n FROM prepared_dataset_derivations_v3').fetchone()['n']==3
    from app.adapters.task_store import DatasetTaskAssets
    from app.domain.dataset_processing import load_processing,quality_lines
    from app.db.database import get_database_config
    from tests.helpers.tasks import task_owner
    processing=load_processing(DatasetTaskAssets(task_owner(client),get_database_config()),task['project_id'],result['source'])
    assert [p['operation'] for p in processing['ancestors']]==['quality','population_filter','quality']
    assert any('AND' in line for line in quality_lines(processing,'en'))

def test_revise_stale_zero_and_source_tamper(client):
    base,parent=published(client,ROWS)
    task,_=filtering(client,base,parent);run(task);view=review(client,task)
    recipe=copy.deepcopy(RECIPE);recipe['conditions']=[dict(column=1,kind='number',operator='gt',value='1000')]
    revised=decide(client,task,decision(view,'revise',recipe),'revise')
    assert revised.status_code==202,revised.text
    assert decide(client,task,decision(view),'stale').status_code==409
    assert client.get(view['impact_download_url']).status_code==409
    run(revised.json());new=review(client,task)
    assert new['preview']['retained_record_count']==0
    assert decide(client,task,decision(new),'empty').status_code==422
    assert new['task']['status']=='waiting_confirmation'
    other,_=filtering(client,base,parent,key='other');result,_,_=accept(client,other)
    with migration_connection(write=True) as db:
        db.execute('UPDATE prepared_dataset_derivations_v3 SET manifest_sha256=%s WHERE prepared_dataset_id=%s',('0'*64,result['id']))
    assert client.get(result['download_url']).status_code==409

def test_full_pagination_permissions_and_parallel_single_publication(client):
    from concurrent.futures import ThreadPoolExecutor
    from tests.api.test_analysis import add_file
    from app.adapters.task_store import DatasetTaskAssets
    from app.domain.dataset_preparation.derivation_service import DerivationPreparationService
    from app.db.database import get_database_config
    from app.core.exceptions import StorageError
    import pytest
    base,parent=published(client,[['age','group','value'],*[[20,'A',n] for n in range(45)]])
    task,request=filtering(client,base,parent);run(task);view=review(client,task)
    assert [len(client.get(view['impact_url'],params={'page':p}).json()['items']) for p in (1,2,3)]==[20,20,5]
    assert len(client.get(view['impact_download_url']).json()['entries'])==45
    other,_=add_file(client)
    assert client.post(other+'/dataset-population-filters',json=request,headers={'Idempotency-Key':'cross'}).status_code==404
    with pytest.raises(StorageError):DerivationPreparationService(DatasetTaskAssets('other-owner',get_database_config())).manifest_for_task(task['id'])
    body=decision(view)
    accepted=decide(client,task,body).json()
    assert decide(client,task,body).json()==accepted
    with ThreadPoolExecutor(max_workers=2) as pool:results=list(pool.map(run,[accepted,accepted]))
    assert sum(r is not None and r['status']=='succeeded' for r in results)==1
    with database_connection() as db:
        assert db.execute('SELECT count(*) AS n FROM prepared_dataset_versions WHERE task_id=%s',(task['id'],)).fetchone()['n']==1
        assert db.execute('SELECT count(*) AS n FROM prepared_dataset_derivations_v3').fetchone()['n']==1
