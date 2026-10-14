"""Fixed multi-round retention oracle, independent of scenario implementation."""
from io import BytesIO
from copy import deepcopy
import pytest
from docx import Document
from openpyxl import load_workbook
from tests.api.test_analysis import client  # noqa: F401
from tests.helpers.version_retention import workbook_bytes, create_retention_scenario, _assert_retained


def test_retention_guard_rejects_identity_setup_budget_and_unrelated_changes():
    before = ({'projects': [{'id':'project', 'updated_at':'1', 'name':'sample'},
                            {'id':'other', 'updated_at':'1', 'name':'other'}],
               'files': [{'id':'file', 'sha256':'original', 'size_bytes':10, 'object_key':'original.xlsx'}],
               'analysis_setups': [{'file_id':'file', 'revision':1, 'setup_json':'original-selection'}],
               'model_budgets': [{'id':'budget', 'scope_key':'project', 'limit_micro_usd':100}]},
              {'original.xlsx':b'original'})
    after = deepcopy(before)
    after[0]['projects'][0]['updated_at'] = '2'
    _assert_retained(before, after, project_id='project')
    for table, index, field, value in [
        ('files',0,'sha256','altered'), ('files',0,'object_key','elsewhere.xlsx'),
        ('analysis_setups',0,'revision',2), ('analysis_setups',0,'setup_json','altered-selection'),
        ('model_budgets',0,'scope_key','other'), ('model_budgets',0,'limit_micro_usd',999),
        ('projects',0,'name','altered'), ('projects',1,'updated_at','2')]:
        damaged = deepcopy(after)
        damaged[0][table][index][field] = value
        with pytest.raises(AssertionError):
            _assert_retained(before, damaged, project_id='project')
    expected = dict(file_id='file', revision=2, setup_json='{"selection":"explicit-new-selection"}')
    after[0]['analysis_setups'][0] = expected.copy()
    _assert_retained(before, after, project_id='project', expected_setup=expected)
    after[0]['analysis_setups'][0]['setup_json'] = '{"selection":"unapproved-selection"}'
    with pytest.raises(AssertionError):
        _assert_retained(before, after, project_id='project', expected_setup=expected)


def test_multiround_versions_keep_frozen_inputs_assets_and_historical_artifacts(client, monkeypatch):
    content = workbook_bytes()
    assert workbook_bytes() == content
    book = load_workbook(BytesIO(content))
    assert book.sheetnames == ['研究 & 数据']
    assert (book.active.max_row, book.active.max_column) == (11, 4)
    assert book.active['D4'].value == 900
    assert all(cell.data_type != 'f' for row in book.active for cell in row)
    book.close()
    scenario = create_retention_scenario(client, monkeypatch)
    assert [r['version'] for r in scenario['raws']] == [1, 2, 3]
    assert [p['version'] for p in scenario['prepared']] == [1, 1]
    assert scenario['prepared'][0]['id'] != scenario['prepared'][1]['id']
    for raw, prepared in zip(scenario['raws'], scenario['prepared']):
        assert prepared['parent_dataset_id'] == raw['id']
    old, new = scenario['old_result'], scenario['new_result']
    assert (old['overall']['n'], old['overall']['mean']) == (4, 5.5)
    assert (new['overall']['n'], new['overall']['mean']) == (2, 3)
    for result, valid, excluded in [(old,4,2),(new,2,4)]:
        assert (result['check']['row_count'],result['check']['valid_count'],result['check']['excluded_count']) == (6,valid,excluded)
    assert old['source'] == scenario['prepared'][0]['source']
    assert new['source'] == scenario['prepared'][1]['source']
    assert [e['stage'] for e in scenario['stage_evidence']] == ['R2','P2','R3','select-P2']
    assert scenario['provider_attempts'] == []
    assert scenario['synthetic_model_calls'] > 0
    report = client.get(scenario['historical_urls']['report'])
    document = Document(BytesIO(report.content))
    text = '\n'.join(p.text for p in document.paragraphs)
    assert '原始数据版本 1 · 整理数据版本 1' in text
    assert '所选整理区域 6 行' in text and '体重（列 D）' in text
    main_text = text.split('附录 A 解释依据')[0]
    for forbidden in (old['source']['dataset_id'], old['source']['dataset_sha256'], scenario['file']['id'], scenario['file']['sha256']):
        assert forbidden not in main_text
    assert old['source']['dataset_id'] not in text and old['source']['dataset_sha256'] not in text
