"""DB-free source whose trailing table caption terminates an original section."""
from io import BytesIO
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def source_docx(before=True):
    doc = Document()
    doc.sections[0].footer.paragraphs[0].text = 'FIRST FOOTER '
    if before:
        doc.add_paragraph('First source content')
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Target wide table'
    caption = doc.add_paragraph('Table 1. Original caption', style='Caption')
    second = doc.add_section(WD_SECTION_START.NEW_PAGE)
    second.footer.is_linked_to_previous = False
    second.footer.paragraphs[0].text = 'SECOND FOOTER '
    separator = caption._p.getnext()
    section = separator.find(qn('w:pPr')).find(qn('w:sectPr'))
    section.getparent().remove(section)
    caption._p.get_or_add_pPr().append(section)
    separator.getparent().remove(separator)
    doc.add_paragraph('Second source content')
    for section, start in zip(doc.sections, (5, 11)):
        number = OxmlElement('w:pgNumType')
        number.set(qn('w:start'), str(start))
        section._sectPr.append(number)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        section.footer.paragraphs[0]._p.append(field)
    stream = BytesIO()
    doc.save(stream)
    return stream.getvalue()
