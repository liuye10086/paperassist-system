from datetime import datetime
from io import BytesIO
import json
import os
from pathlib import Path
import sqlite3
import subprocess
import sys
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.main import app


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv('PAPERASSIST_DATA_DIR', str(tmp_path / 'data'))
    with TestClient(app) as instance:
        yield instance


def add_file(client, rows=None):
    project = client.post('/api/v1/projects', json={
        'name': '分析配置测试', 'research_topic': '合成数据', 'project_type': 'sci',
    }).json()
    book = Workbook()
    # A legal sheet name containing a URL-sensitive character.
    book.active.title = '研究 & 数据'
    for row in rows or [['指标', '组别'], [1, 'A'], [2, 'A'], [None, 'B'], [3, None], [0, 'B']]:
        book.active.append(row)
    book.create_sheet('空表')
    book.create_sheet('仅表头').append(['值'])
    content = BytesIO()
    book.save(content)
    book.close()
    record = client.post(f'/api/v1/projects/{project["id"]}/files',
                         files={'file': ('data.xlsx', content.getvalue())}).json()['file']
    return f'/api/v1/projects/{project["id"]}/files/{record["id"]}', record


def selection(**changes):
    return {'task_type': 'descriptive_boxplot', 'sheet_name': '研究 & 数据', 'numeric_column': 'A',
            'group_column': 'B', 'unit': 'mg/L', 'missing_policy': 'exclude_selected_missing', **changes}


def test_profile_reads_beyond_preview_and_distinguishes_column_positions(client):
    rows = [['指标', '指标', None]] + [[index, 'A', ''] for index in range(25)] + [['文本数字', None, True]]
    base, _ = add_file(client, rows)
    response = client.get(base + '/analysis-profile', params={'sheet': '研究 & 数据'})
    assert response.status_code == 200, response.text
    profile = response.json()
    assert profile['row_count'] == 26
    assert [column['id'] for column in profile['columns']] == ['A', 'B', 'C']
    first, second, third = profile['columns']
    assert first['name'] == second['name'] == '指标'
    assert first['type'] == 'mixed' and first['numeric_count'] == 25
    assert not first['can_be_numeric']
    assert first['issue_cells'] == ['A27']
    assert second['missing_count'] == 1
    assert third['type'] == 'boolean' and third['missing_count'] == 25


def test_pairwise_missing_counts_and_group_counts(client):
    base, _ = add_file(client)
    response = client.post(base + '/analysis-check', json=selection())
    assert response.status_code == 200, response.text
    check = response.json()
    assert check['ready'] and check['row_count'] == 5
    assert (check['valid_count'], check['excluded_count']) == (3, 2)
    assert check['groups'] == [{'label': 'A', 'count': 2}, {'label': 'B', 'count': 1}]
    assert check['warnings']
    ungrouped = client.post(base + '/analysis-check', json=selection(group_column=None)).json()
    assert ungrouped['valid_count'] == 4 and ungrouped['excluded_count'] == 1


@pytest.mark.parametrize('value', ['12', True, datetime(2026, 9, 29), '=1+1', '#DIV/0!'])
def test_non_numeric_values_are_not_silently_converted_or_dropped(client, value):
    base, _ = add_file(client, [['指标', '组别'], [1, 'A'], [value, 'A']])
    check = client.post(base + '/analysis-check', json=selection()).json()
    assert not check['ready'] and check['issues']
    response = client.put(base + '/analysis-setup', json=selection(expected_revision=0))
    assert response.status_code == 422
    assert client.get(base + '/analysis-setup').json() is None


def test_whitespace_is_missing_but_na_remains_text(client):
    base, _ = add_file(client, [['数值', '空列', '标签'], [0, '  ', 'NA'], [1, None, 'A']])
    columns = client.get(base + '/analysis-profile', params={'sheet': '研究 & 数据'}).json()['columns']
    assert columns[0]['numeric_count'] == 2
    assert columns[1]['type'] == 'empty' and columns[1]['missing_count'] == 2
    assert columns[2]['type'] == 'text' and columns[2]['missing_count'] == 0


@pytest.mark.parametrize('changes', [
    {'numeric_column': 'Z'}, {'group_column': 'A'}, {'sheet_name': 'missing'},
    {'task_type': 'arbitrary_code'}, {'missing_policy': 'fill_zero'}, {'numeric_column': True},
])
def test_invalid_selection_cannot_be_saved(client, changes):
    base, _ = add_file(client)
    assert client.put(base + '/analysis-setup', json=selection(**changes, expected_revision=0)).status_code == 422


@pytest.mark.parametrize('sheet', ['空表', '仅表头'])
def test_empty_sheets_have_no_usable_selection(client, sheet):
    base, _ = add_file(client)
    profile = client.get(base + '/analysis-profile', params={'sheet': sheet}).json()
    assert profile['row_count'] == 0
    response = client.put(base + '/analysis-setup', json=selection(sheet_name=sheet, group_column=None, expected_revision=0))
    assert response.status_code == 422


def test_no_complete_pairs_and_excessive_groups_are_rejected(client):
    base, _ = add_file(client, [['值', '组'], [1, None], [None, 'A']])
    assert not client.post(base + '/analysis-check', json=selection()).json()['ready']
    base, _ = add_file(client, [['值', '组']] + [[i, str(i)] for i in range(21)])
    check = client.post(base + '/analysis-check', json=selection()).json()
    assert not check['ready'] and any('20' in issue for issue in check['issues'])


def test_save_reload_revision_conflict_and_new_process(client):
    base, record = add_file(client)
    assert client.get(base + '/analysis-setup').json() is None
    response = client.put(base + '/analysis-setup', json=selection(expected_revision=0))
    assert response.status_code == 200, response.text
    saved = response.json()
    assert saved['revision'] == 1 and saved['status'] == 'configured'
    assert saved['source_sha256'] == record['sha256']
    assert saved['selection'] == selection()
    assert saved['check']['valid_count'] == 3
    assert client.get(base + '/analysis-setup').json() == saved
    assert client.put(base + '/analysis-setup', json=selection(expected_revision=0)).status_code == 409
    updated = client.put(base + '/analysis-setup', json=selection(group_column=None, expected_revision=1)).json()
    assert updated['revision'] == 2 and updated['check']['valid_count'] == 4
    code = "from fastapi.testclient import TestClient; from app.main import app; import sys; print(TestClient(app).get(sys.argv[1]).text)"
    result = subprocess.run([sys.executable, '-c', code, base + '/analysis-setup'],
        cwd=Path(__file__).resolve().parents[1], capture_output=True, text=True, check=True, timeout=30)
    assert json.loads(result.stdout) == updated


def test_original_integrity_and_project_ownership_are_checked(client):
    base, record = add_file(client)
    other_base, _ = add_file(client)
    wrong = other_base.rsplit('/', 1)[0] + '/' + record['id']
    assert client.get(wrong + '/analysis-profile', params={'sheet': '研究 & 数据'}).status_code == 404
    assert client.put(wrong + '/analysis-setup', json=selection(expected_revision=0)).status_code == 404
    path = Path(os.environ['PAPERASSIST_DATA_DIR']) / 'files' / (record['id'] + '.xlsx')
    path.write_bytes(b'changed')
    assert client.put(base + '/analysis-setup', json=selection(expected_revision=0)).status_code == 409


def test_existing_schema_one_data_survives_upgrade(client):
    base, record = add_file(client)
    database = Path(os.environ['PAPERASSIST_DATA_DIR']) / 'paperassist.sqlite3'
    with sqlite3.connect(database) as db:
        db.execute('DROP TABLE IF EXISTS analysis_setups')
        db.execute('PRAGMA user_version = 1')
    response = client.put(base + '/analysis-setup', json=selection(expected_revision=0))
    assert response.status_code == 200, response.text
    assert client.get(base + '/preview').status_code == 200
    assert client.get(base.rsplit('/', 1)[0]).json()[0]['sha256'] == record['sha256']
    with sqlite3.connect(database) as db:
        assert db.execute('PRAGMA user_version').fetchone()[0] == 2


def test_two_simultaneous_saves_cannot_overwrite_each_other(client):
    base, _ = add_file(client)
    barrier = Barrier(2)

    def save(unit):
        barrier.wait(timeout=5)
        return client.put(base + '/analysis-setup', json=selection(unit=unit, expected_revision=0))

    with ThreadPoolExecutor(max_workers=2) as executor:
        results = list(executor.map(save, ['mg/L', 'mmol/L']))
    assert sorted(result.status_code for result in results) == [200, 409]
    winning = next(result.json() for result in results if result.status_code == 200)
    assert client.get(base + '/analysis-setup').json() == winning


def test_database_failure_keeps_existing_configuration(client):
    base, _ = add_file(client)
    original = client.put(base + '/analysis-setup', json=selection(expected_revision=0)).json()
    database = Path(os.environ['PAPERASSIST_DATA_DIR']) / 'paperassist.sqlite3'
    with sqlite3.connect(database) as db:
        db.execute("CREATE TRIGGER reject_setup BEFORE UPDATE ON analysis_setups BEGIN SELECT RAISE(ABORT, 'test'); END;")
    response = client.put(base + '/analysis-setup', json=selection(expected_revision=1, unit='new'))
    assert response.status_code == 503
    assert client.get(base + '/analysis-setup').json() == original


def test_profile_obeys_current_parse_limits(client, monkeypatch):
    base, _ = add_file(client)
    monkeypatch.setenv('EXCEL_MAX_ROWS', '3')
    response = client.get(base + '/analysis-profile', params={'sheet': '研究 & 数据'})
    assert response.status_code == 422
    assert response.json()['detail']['code'] == 'workbook_too_large'
