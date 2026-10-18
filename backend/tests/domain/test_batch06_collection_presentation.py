from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from docx import Document
import pytest
from tests.domain.test_batch05_collection_contracts import bundle
from tests.adapters.test_analysis_plot_response import png
from app.domain.report_collections.contracts import freeze_collection
from app.adapters.collection_docx import build_collection_report


@pytest.mark.parametrize('language',['zh-CN','en'])
def test_presentation_is_same_in_frozen_preview_and_word_and_never_rewrites_png(language):
    item=bundle(language=language)
    presentation={'schema_version':1,'axes':{'x':{'label':'Time (day)','unit':'day'},'y':{'label':'Within-group proportion (0–1)','unit':'0–1'}},'legend':{'title':'Group','entries':[{'label':'Treatment','color':'#0072B2'}]},'color_scale':None,'abbreviations':[{'term':'SD','definition':'Sample standard deviation'}],'statistical_notes':['Missing SD is not zero.']}
    item['figure']['presentation']=presentation
    item['explanation']['sections'][0]['evidence']=[{'key':'internal.method.fact','label':'Method','value':'Counts are records, not participants.'}]
    item['explanation']['limitations']=['Design not verified.']
    before=deepcopy(item)
    snapshot=freeze_collection([item],language=language,title='Contract report')
    section=next(s for s in snapshot['sections'] if s['key']=='figures')
    assert section['content'][0]['presentation']==presentation
    raw=build_collection_report(snapshot,{'id':'report','export_revision':1},{'figure1':png()})
    document=Document(BytesIO(raw));text='\n'.join(p.text for p in document.paragraphs)
    assert all(v in text for v in ['Time (day)','Within-group proportion (0–1)','Treatment','Sample standard deviation','Missing SD is not zero.','Counts are records, not participants.','Design not verified.'])
    assert 'internal.method.fact' not in text
    with ZipFile(BytesIO(raw)) as archive:
        assert all(archive.read(n)==png() for n in archive.namelist() if n.startswith('word/media/'))
    assert item==before


def test_old_snapshot_remains_without_optional_presentation_fields():
    snapshot=freeze_collection([bundle()],language='zh-CN',title='Legacy')
    assert all('presentation' not in row for s in snapshot['sections'] if s['key']=='figures' for row in s['content'])


@pytest.mark.parametrize('language',['zh-CN','en'])
def test_word_color_scale_retains_range_center_ticks_and_palette_order(language):
    item=bundle(language=language)
    scale={'label':'Correlation coefficient','min':-1,'max':1,'center':0,'ticks':[-1,0,1],
        'palette':['#2166AC','#F7F7F7','#B2182B']}
    item['figure']['presentation']={'schema_version':1,
        'axes':{'x':{'label':'Variables','unit':''},'y':{'label':'Variables','unit':''}},
        'legend':{'title':'','entries':[]},'color_scale':scale,
        'abbreviations':[],'statistical_notes':[]}
    snapshot=freeze_collection([item],language=language,title='Heatmap report')
    raw=build_collection_report(snapshot,{'id':'report','export_revision':1},{'figure1':png()})
    document=Document(BytesIO(raw))
    prefix='Color scale: ' if language=='en' else '色标: '
    line=next(p.text for p in document.paragraphs if p.text.startswith(prefix))
    assert 'Correlation coefficient; -1 – 1' in line
    assert ('center 0; ticks -1, 0, 1' if language=='en' else '中心 0; 刻度 -1, 0, 1') in line
    assert ('palette ' if language=='en' else '配色 ')+'#2166AC → #F7F7F7 → #B2182B' in line
    assert snapshot['items'][0]['figure']['presentation']['color_scale']==scale
