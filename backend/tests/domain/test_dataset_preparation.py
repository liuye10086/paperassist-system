"""The approved rectangle keeps Excel coordinates and only explicit text rules."""
import importlib
from copy import deepcopy
import pytest
from pydantic import ValidationError
from openpyxl.utils.cell import get_column_letter
from app.core.exceptions import StorageError
from app.domain.analysis_sources import DatasetAnalysisSource
from app.domain.datasets.contracts import SheetStructure, STRUCTURE_PARSER_VERSION


@pytest.fixture(autouse=True)
def postgres_schema():
    # This pure transform has no database dependency.
    yield None


def raw_sheet(values, merged_ranges=None):
    cells = []
    for (row, column), (kind, value) in sorted(values.items()):
        raw_type = {'blank': None, 'text': 'inlineStr', 'number': 'n', 'boolean': 'b',
                    'date': 'd', 'formula': 'n', 'error': 'e'}[kind]
        raw_value = None if kind in ('blank', 'text', 'formula') else ('1' if value else '0') if kind == 'boolean' else value
        formula = dict(text=value[1:], attributes={}, cache_state='missing', cached_type='n', cached_value=None) if kind == 'formula' else None
        cells.append(dict(row=row, column=column, coordinate=f'{get_column_letter(column)}{row}',
                          kind=kind, value=value, raw_type=raw_type, raw_value=raw_value,
                          number_format='General', formula=formula))
    def extent(points):
        if not points:
            return None
        return f'{get_column_letter(min(c for r, c in points))}{min(r for r, c in points)}:{get_column_letter(max(c for r, c in points))}{max(r for r, c in points)}'
    positions = list(values)
    merged_ranges = merged_ranges or []
    from app.domain.datasets.contracts import rectangle
    span = positions + [point for area in merged_ranges for point in rectangle(area)]
    rows = max((r for r, c in span), default=0)
    columns = max((c for r, c in span), default=0)
    nonempty = {r for (r, c), (kind, value) in values.items() if value not in (None, '')}
    return SheetStructure(index=0, name='实验', state='visible', row_count=rows, column_count=columns,
        nonempty_row_count=len(nonempty), nonempty_data_row_count=len(nonempty - {1}),
        stored_range=extent(positions), content_range=extent([p for p, (k, v) in values.items() if v not in (None, '')]),
        merged_ranges=merged_ranges, cells=cells, warnings=[], fields=[dict(column=c, letter=get_column_letter(c),
            header_value=values.get((1, c), ('blank', None))[1], observed_types=sorted({k for (r, col), (k, v) in values.items() if r > 1 and col == c})) for c in range(1, columns + 1)])


def parent_source():
    return DatasetAnalysisSource(kind='dataset_version', schema_version=1, dataset_id='raw-id',
        structure_id='structure-id', dataset_version=3, file_id='file-id', source_sha256='a' * 64,
        dataset_sha256='b' * 64, dataset_size_bytes=100, dataset_schema_version=1,
        sheet_index=0, sheet_name='实验', parser_version=STRUCTURE_PARSER_VERSION)


def preparation():
    return importlib.import_module('app.domain.dataset_preparation.transform')


def contracts():
    return importlib.import_module('app.domain.dataset_preparation.contracts')


def test_selected_rectangle_keeps_blank_rows_coordinates_and_parent():
    sheet = raw_sheet({(1, 1): ('text', '说明'), (3, 3): ('text', '组别'), (3, 4): ('text', '值'),
                       (4, 3): ('text', '第二条说明'), (5, 3): ('text', ' A '), (5, 4): ('number', '1.00'),
                       (7, 4): ('number', '2'), (35, 4): ('blank', None)})
    before = deepcopy(sheet.model_dump())
    recipe = preparation().build_recipe(sheet, header_row=3, data_range='C5:D35')
    obj = preparation().prepare_dataset(sheet, parent_source(), recipe)
    assert obj.sheet.row_count == obj.summary.record_count == 31
    assert obj.sheet.column_count == 2
    assert [f.column for f in obj.sheet.fields] == [3, 4]
    assert [f.name for f in recipe.fields] == ['组别', '值']
    assert {c.coordinate for c in obj.sheet.cells} == {'C5', 'D5', 'D7', 'D35'}
    assert obj.summary.columns[0].missing_count == 30
    assert obj.summary.columns[1].missing_count == 29
    assert obj.summary.changed_cell_count == 0
    assert obj.sheet.cells[0].value == ' A '
    assert sheet.model_dump() == before
    assert obj.parent_source == parent_source()


def test_default_fields_keep_duplicate_blank_and_whitespace_labels():
    sheet = raw_sheet({(3, 3): ('text', '重复'), (3, 4): ('text', '重复'),
                       (3, 6): ('text', ' '), (5, 3): ('number', '1'), (5, 6): ('number', '2')})
    recipe = preparation().build_recipe(sheet, header_row=3)
    assert recipe.data_range == 'C4:F5'
    assert [f.name for f in recipe.fields] == ['重复', '重复', '未命名列 E', ' ']
    assert all(not f.trim_text and f.missing_markers == [] and f.unit.status == 'unconfirmed' for f in recipe.fields)


def test_long_header_still_offers_editable_bounded_name_without_changing_raw():
    header = '长字段' * 100
    sheet = raw_sheet({(1, 1): ('text', header), (2, 1): ('number', '1')})
    recipe = preparation().build_recipe(sheet)
    assert recipe.fields[0].name == header[:160]
    assert sheet.cells[0].value == header
    recipe.fields[0].name = '研究者修改后的名称'
    obj = preparation().prepare_dataset(sheet, parent_source(), recipe)
    assert obj.sheet.fields[0].name == '研究者修改后的名称'


def test_trim_then_case_sensitive_exact_markers_counts_one_changed_cell():
    sheet = raw_sheet({(1, 3): ('text', '字段'), (2, 3): ('text', ' NA '),
                       (3, 3): ('text', 'na'), (4, 3): ('text', ' N/A '), (5, 3): ('text', ' '),
                       (6, 3): ('number', '10')})
    recipe = preparation().build_recipe(sheet, data_range='C2:C6')
    recipe.fields[0].trim_text = True
    recipe.fields[0].missing_markers = ['NA']
    obj = preparation().prepare_dataset(sheet, parent_source(), recipe)
    assert [(c.kind, c.value) for c in obj.sheet.cells] == [('blank', None), ('text', 'na'), ('text', 'N/A'), ('text', ''), ('number', '10')]
    assert obj.summary.changed_cell_count == obj.summary.changed_row_count == 3
    assert obj.summary.trimmed_cell_count == 3
    assert obj.summary.missing_marker_cell_count == 1
    assert obj.summary.columns[0].missing_count == 2
    assert obj.changes[0].operations == ['trim_text', 'missing_marker']
    assert obj.changes[0].before == {'kind': 'text', 'value': ' NA '}
    assert obj.changes[0].after == {'kind': 'blank', 'value': None}


def test_non_text_values_and_formula_display_are_unchanged():
    sheet = raw_sheet({(1, 1): ('text', '类型'), (2, 1): ('number', '1.00'),
                       (3, 1): ('formula', '=SUM(B3:C3)'), (4, 1): ('date', '2026-10-12'),
                       (5, 1): ('boolean', True), (6, 1): ('error', '#VALUE!')})
    recipe = preparation().build_recipe(sheet)
    recipe.fields[0].trim_text = True
    recipe.fields[0].missing_markers = ['1.00', '=SUM(B3:C3)', '2026-10-12', '#VALUE!']
    obj = preparation().prepare_dataset(sheet, parent_source(), recipe)
    assert [(c.kind, c.value) for c in obj.sheet.cells] == [(c.kind, c.value) for c in sheet.cells[1:]]
    assert obj.summary.changed_cell_count == 0
    issues = {i.code: i.count for i in obj.summary.unresolved}
    assert issues == {'formula_unverified': 1, 'error_values': 1, 'mixed_types': 5, 'unconfirmed_unit': 1}


def test_unit_confirmation_and_intersecting_merges_are_metadata_only():
    sheet = raw_sheet({(1, 3): ('text', '浓度'), (2, 3): ('number', '1'), (3, 4): ('number', '2')},
                      ['C2:D2', 'A1:B1'])
    recipe = preparation().build_recipe(sheet, data_range='C2:D3')
    recipe.fields[0].unit = contracts().PreparationUnit(status='confirmed', value='mg/L')
    recipe.fields[1].unit = contracts().PreparationUnit(status='confirmed', value='')
    obj = preparation().prepare_dataset(sheet, parent_source(), recipe)
    assert obj.sheet.merged_ranges == ['C2:D2']
    assert [(i.code, i.column, i.count) for i in obj.summary.unresolved] == [('merged_cells', None, 1)]
    assert obj.summary.columns[0].unit.value == 'mg/L'
    assert [(c.coordinate, c.value) for c in obj.sheet.cells] == [('C2', '1'), ('D3', '2')]


def test_changes_preview_is_bounded_but_all_counts_and_cells_are_complete():
    sheet = raw_sheet({(1, 1): ('text', '组'), **{(r, 1): ('text', ' x ') for r in range(2, 27)}})
    recipe = preparation().build_recipe(sheet)
    recipe.fields[0].trim_text = True
    obj = preparation().prepare_dataset(sheet, parent_source(), recipe)
    assert len(obj.changes) == 20
    assert obj.summary.changes_truncated is True
    assert obj.summary.changed_cell_count == obj.summary.changed_row_count == 25
    assert len(obj.sheet.cells) == 25
    assert all(c.value == 'x' for c in obj.sheet.cells)


@pytest.mark.parametrize('values,header,area', [({}, 1, None), ({(1, 1): ('text', 'h')}, 1, None),
    ({(1, 1): ('text', 'h'), (2, 1): ('number', '1')}, 2, 'A2:A2'),
    ({(1, 1): ('text', 'h'), (2, 1): ('number', '1')}, 1, 'A2:B3')])
def test_invalid_selection_is_actionable(values, header, area):
    with pytest.raises(StorageError) as error:
        preparation().build_recipe(raw_sheet(values), header_row=header, data_range=area)
    assert error.value.code == 'preparation_invalid'


def test_single_cell_rectangle_is_canonical_and_fields_cannot_skip_columns():
    sheet = raw_sheet({(1, 3): ('text', 'h'), (2, 3): ('number', '1')})
    recipe = preparation().build_recipe(sheet, data_range='C2')
    assert recipe.data_range == 'C2:C2'
    with pytest.raises(ValidationError):
        contracts().PreparationRecipe(header_row=1, data_range='C2:D2', fields=recipe.fields)


@pytest.mark.parametrize('unit', [{'status': 'unconfirmed', 'value': 'mg'}, {'status': 'confirmed', 'value': None}])
def test_unit_requires_explicit_coherent_confirmation(unit):
    with pytest.raises(ValidationError):
        contracts().PreparationUnit.model_validate(unit)


def test_recipe_does_not_normalize_markers_and_rejects_duplicates():
    model = contracts().PreparationField
    field = model(column=1, name=' x ', unit={'status': 'unconfirmed', 'value': None}, missing_markers=[' NA ', 'NA', 'na'])
    assert field.name == ' x ' and field.missing_markers == [' NA ', 'NA', 'na']
    with pytest.raises(ValidationError):
        model(column=1, name='x', unit={'status': 'unconfirmed', 'value': None}, missing_markers=['NA', 'NA'])


def test_object_boundary_rejects_mismatched_identity_and_effective_range():
    sheet = raw_sheet({(1, 3): ('text', 'h'), (2, 3): ('number', '1')})
    recipe = preparation().build_recipe(sheet)
    obj = preparation().prepare_dataset(sheet, parent_source(), recipe)
    data = obj.model_dump()
    data['sheet']['fields'][0]['name'] = 'other'
    with pytest.raises(ValidationError):
        contracts().PreparedDatasetObject.model_validate(data)
    wrong_parent = parent_source().model_copy(update={'sheet_name': 'other'})
    with pytest.raises(StorageError) as error:
        preparation().prepare_dataset(sheet, wrong_parent, recipe)
    assert error.value.code == 'preparation_invalid'


def test_validate_recipe_rejects_mutated_and_out_of_sheet_rules_without_transform():
    sheet = raw_sheet({(1, 3): ('text', 'h'), (2, 3): ('number', '1')})
    recipe = preparation().build_recipe(sheet)
    validated = preparation().validate_recipe(sheet, recipe)
    assert validated == recipe and validated is not recipe
    recipe.fields[0].column = 2
    with pytest.raises(StorageError) as error:
        preparation().validate_recipe(sheet, recipe)
    assert error.value.code == 'preparation_invalid'
    recipe = preparation().build_recipe(sheet)
    recipe.data_range = 'C2:C3'
    with pytest.raises(StorageError):
        preparation().validate_recipe(sheet, recipe)


def test_object_boundary_rejects_summary_that_disagrees_with_effective_cells():
    sheet = raw_sheet({(1, 3): ('text', 'h'), (2, 3): ('number', '1'), (3, 3): ('blank', None)})
    obj = preparation().prepare_dataset(sheet, parent_source(), preparation().build_recipe(sheet, data_range='C2:C3'))
    data = obj.model_dump()
    data['summary']['columns'][0]['missing_count'] = 0
    data['summary']['columns'][0]['non_missing_count'] = 2
    with pytest.raises(ValidationError):
        contracts().PreparedDatasetObject.model_validate(data)


def test_two_changed_cells_in_one_row_and_recipe_mutation_do_not_alter_candidate():
    sheet = raw_sheet({(1, 3): ('text', 'g'), (1, 4): ('text', 'h'),
                       (2, 3): ('text', ' NA '), (2, 4): ('text', ' NA ')})
    recipe = preparation().build_recipe(sheet)
    for field in recipe.fields:
        field.trim_text = True
        field.missing_markers = ['NA']
    obj = preparation().prepare_dataset(sheet, parent_source(), recipe)
    assert obj.summary.changed_cell_count == 2
    assert obj.summary.changed_row_count == 1
    assert obj.summary.trimmed_cell_count == obj.summary.missing_marker_cell_count == 2
    recipe.fields[0].name = 'changed upstream'
    recipe.fields[0].missing_markers.append('new')
    assert obj.recipe.fields[0].name == 'g'
    assert obj.recipe.fields[0].missing_markers == ['NA']
