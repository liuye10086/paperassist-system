from copy import deepcopy
from io import BytesIO
from docx import Document
import pytest


@pytest.fixture(autouse=True)
def postgres_schema():
    yield None


def processing():
    before=dict(record_count=10,missing_cell_count=1,missing_row_count=1,blank_row_count=0,duplicate_group_count=2,duplicate_member_count=4,duplicate_extra_row_count=2,outlier_candidate_count=2)
    after={**before,'record_count':7,'missing_cell_count':2,'missing_row_count':2,'duplicate_group_count':0,'duplicate_member_count':0,'duplicate_extra_row_count':0,'outlier_candidate_count':1}
    return dict(schema_version=2,recipe=dict(fields=[],manual_exclusions=[{'row':9,'reason':'LOCAL-REASON'},{'row':11,'reason':'LOCAL-REASON'}],duplicate_policy='exact'),
        receipt=dict(root_record_count=10,base_record_count=10,retained_record_count=7,inherited_excluded_row_count=0,manual_excluded_row_count=2,duplicate_excluded_row_count=1,cumulative_excluded_row_count=3,changed_cell_count=1,retained_changed_cell_count=1,excluded_changed_cell_count=0,changed_row_count=1,trimmed_cell_count=1,missing_marker_cell_count=1,impact_count=4,before_quality=before,after_quality=after),
        manifest=dict(object_key='quality-manifests/local/manifest.json',sha256='f'*64,size_bytes=100,entry_count=4),base_source={'dataset_version':1})


@pytest.mark.parametrize('language',['zh-CN','en'])
def test_single_word_reports_quality_before_after_cumulative_and_separate_analysis_n(language):
    from tests.adapters.test_analysis_report_docx import snapshot
    from tests.adapters.test_analysis_plot_response import png
    from app.adapters.report_docx import build_report
    from app.adapters.report_store import ReportStore
    saved=snapshot(language)
    saved['result']['source']={'kind':'prepared_dataset_version','schema_version':2,'dataset_version':2,'parent_dataset_version':1}
    saved['processing']=processing()
    saved['result']['check'].update(row_count=7,valid_count=5,excluded_count=2)
    saved['result']['statistics']['n']=5
    document=Document(BytesIO(build_report(saved,ReportStore.metadata(saved),png())))
    text='\n'.join(p.text for p in document.paragraphs)
    assert ('原始选定区域记录数：10' if language=='zh-CN' else 'Original selected records: 10') in text
    assert 'LOCAL-REASON' in text
    assert ('累计整理排除：3' if language=='zh-CN' else 'Cumulative preparation exclusions: 3') in text
    assert ('未插补、去重或删除离群值' if language=='zh-CN' else 'No imputation, deduplication or outlier removal was performed.') not in text
    assert 'quality-manifests/' not in text


def test_v1_collection_snapshot_is_preserved_and_mixed_quality_versions_rejected():
    from tests.domain.test_batch05_collection_contracts import bundle
    from app.domain.report_collections.contracts import freeze_collection
    from app.core.exceptions import StorageError
    first,second=bundle(),bundle(2)
    saved=freeze_collection([first,second],language='zh-CN',title='Old')
    assert saved['schema_version']==1 and all('processing' not in x for x in saved['items'])
    second['result']['source']['schema_version']=2
    with pytest.raises(StorageError):
        freeze_collection([first,second],language='zh-CN',title='Mixed')


def test_collection_v2_processing_is_bounded_and_preserves_separate_n():
    from tests.domain.test_batch05_collection_contracts import bundle
    from tests.adapters.test_analysis_plot_response import png
    from app.domain.report_collections.contracts import freeze_collection
    from app.adapters.collection_docx import build_collection_report
    first,second=bundle(),bundle(2)
    for item,n in [(first,6),(second,5)]:
        item['result']['source']['schema_version']=2
        item['processing']=processing()
        item['result']['check'].update(row_count=7,valid_count=n,excluded_count=7-n)
    saved=freeze_collection([first,second],language='zh-CN',title='Quality')
    assert [x['result']['check']['valid_count'] for x in saved['items']]==[6,5]
    text='\n'.join(p.text for p in Document(BytesIO(build_collection_report(saved,{'id':'report','export_revision':1},{'figure1':png(),'figure2':png()}))).paragraphs)
    assert '原始选定区域记录数：10' in text and 'LOCAL-REASON' in text
    document=Document(BytesIO(build_collection_report(saved,{'id':'report','export_revision':1},{'figure1':png(),'figure2':png()})))
    labels=' '.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert '原始记录数' not in labels
    assert '整理后保留记录数' in labels


def test_quality_collection_localizes_analysis_stage_limitations():
    from app.adapters.collection_labels import rows
    limitations=[
        '统计阶段采用当前整理保留行的共同完整记录，不新增填补、去重或自动剔除。',
        '统计阶段采用当前整理保留行的完整记录或完整对，不新增插补、去重或自动剔除。',
        '当前整理保留行的共同完整记录可能产生选择偏差；统计阶段不新增插补或自动删异常。']
    rendered=list(rows({'limitations':limitations},True,prepared_quality=True))
    assert all(not any('\u4e00'<=c<='\u9fff' for c in value) for _,value in rendered)


def test_word_task_snapshot_rejects_changed_frozen_quality_receipt():
    from app.adapters.task_store import snapshots_match
    original={key:{} for key in ['file','result','figure','versions','output_language','explanation']}
    original['processing']=processing()
    current=deepcopy(original)
    current['processing']['receipt']['manual_excluded_row_count']=0
    assert not snapshots_match(current,original)


def test_cloud_facts_only_correct_stage_labels_without_expanding_material():
    from tests.domain.test_correlation_artifacts import result
    from app.domain.correlation_facts import build_payload
    value=result('en')
    figure={'id':'figure','figure_number':1,'title':'Figure','sha256':'a'*64}
    old=build_payload(value,figure)
    value['source']={'schema_version':2}
    current=build_payload(value,figure)
    assert set(current['facts'])==set(old['facts'])
    assert current['facts']['data.total']['label']=='Records retained after preparation'
    assert 'statistical analysis' in current['facts']['method.summary']['value']
    assert 'LOCAL-REASON' not in str(current)
