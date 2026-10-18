"""Independent expectations: quality exclusions never become analysis missingness."""
from types import SimpleNamespace
import pytest
from app.domain.dataset_preparation.quality_contracts import PreparedSheetV2 as SparsePreparedSheet


@pytest.fixture(autouse=True)
def postgres_schema():
    yield None


def sheet_for(rows, active=None, area=None):
    from openpyxl.utils.cell import get_column_letter
    active = list(rows) if active is None else active
    width = len(next(iter(rows.values())))
    return SparsePreparedSheet(index=0, name='Synthetic', data_range=area or f'A2:{get_column_letter(width)}{max(rows)}',
        active_source_rows=active, merged_ranges=[],
        fields=[dict(column=c, name=f'Field {c}', unit=dict(status='confirmed', value='')) for c in range(1,width+1)],
        cells=[dict(row=r,column=c,coordinate=f'{get_column_letter(c)}{r}',kind='number' if isinstance(v,(int,float)) else 'text',value=str(v))
               for r,values in rows.items() if r in active for c,v in enumerate(values,1) if v is not None])


def main_sheet():
    return sheet_for({2:[10,1,'A'],3:[20,2,'A'],4:[20,2,'A'],5:[None,3,'A'],6:[40,None,'B'],
                      7:[1000,10,'B'],8:[50,5,'B'],9:[1000,10,'B'],10:[60,6,'B'],11:[70,7,'B']},[2,3,5,6,7,8,10])


def test_main_sample_descriptive_n_six_sum_1180_and_missing_one():
    from app.domain.analysis import AnalysisSelection
    from app.domain.dataset_analysis import scan_dataset
    profile,scan=scan_dataset(main_sheet(),AnalysisSelection(task_type='descriptive_boxplot',sheet_name='Synthetic',numeric_column='A',group_column=None,unit='',missing_policy='exclude_selected_missing'),True)
    assert profile.row_count==7
    assert scan.check(profile).valid_count==6 and scan.check(profile).excluded_count==1
    assert sum(scan.values)==1180


@pytest.mark.parametrize('method',['pearson','spearman'])
def test_main_sample_correlation_n_five_excludes_only_active_missing(method):
    from app.domain.correlation.checks import selected_records
    rows,_,audit=selected_records(main_sheet(),SimpleNamespace(sheet_name='Synthetic',columns=['A','B'],units={'A':'','B':''},method=method))
    assert len(rows)==5
    assert audit['included_rows']==[2,3,7,8,10]
    assert [r['row'] for r in audit['excluded_rows']]==[5,6]


def test_quality_profile_missing_pages_never_return_excluded_rows():
    from app.domain.dataset_profile.scan import ProfileScan
    profile=ProfileScan(main_sheet())
    assert profile.record_count==7 and profile.blank_row_count==0
    assert profile.duplicate_extra_rows==0
    count,items=profile.issues('missing',1,1)
    assert count==1 and [i['coordinate'] for i in items]==['A5']


@pytest.mark.parametrize('method',['welch_t','mann_whitney_u','paired_t','wilcoxon','welch_anova','kruskal_wallis'])
def test_comparison_projects_only_active_rows_and_checks_active_ids(method):
    from app.domain.comparison.checks import selected_records
    paired=method in ('paired_t','wilcoxon')
    sheet=sheet_for({2:[1,2,'A'],3:[8,9,'A'],4:[3,7,'B'],5:[5,9,'C']},[2,4,5])
    selection=SimpleNamespace(sheet_name='Synthetic',design='paired' if paired else 'multiple',unit='',numeric_column='A',group_column='C',
        pair_id_column='C',before_column='A',after_column='B',groups=[] if paired else [SimpleNamespace(kind='text',value=g) for g in 'ABC'],missing_policy='complete_pairs' if paired else 'complete_cases')
    _,_,audit,_=selected_records(sheet,selection)
    assert audit['included_rows']==[2,4,5] and audit['excluded_rows']==[]


@pytest.mark.parametrize('method',['ols','logistic','roc_auc','kaplan_meier'])
def test_modeling_projects_only_active_rows_and_preserves_rectangle_limit(method):
    from app.domain.regression.checks import selected_records
    from app.core.exceptions import StorageError
    from tests.helpers.modeling import selection,sample
    from app.domain.correlation.contracts import parse_selection
    data=sample(method)[1:]
    rows={r:row for r,row in enumerate(data,2)}
    active=[r for r in rows if r!=3]
    sheet=sheet_for(rows,active)
    chosen=selection(method,'Synthetic')
    if method=='kaplan_meier':
        # Prepared units must remain inherited; projection itself does not alter them.
        chosen['time_unit']='day'
    _,_,audit=selected_records(sheet,parse_selection(chosen))
    assert audit['included_rows']==active and audit['excluded_rows']==[]
    oversized=sheet.model_copy(update={'data_range':f'A2:{chr(64+len(data[0]))}10002'})
    with pytest.raises(StorageError) as error:
        selected_records(oversized,parse_selection(chosen))
    assert error.value.code=='numeric_row_limit'


@pytest.mark.parametrize('kind,projection',[('bar',{'category_column':'C','bar_stat':'count'}),('line',{'x_column':'A','y_column':'B'}),('violin',{'y_column':'A'})])
def test_descriptive_figures_cannot_restore_quality_exclusions(kind,projection):
    from app.domain.figure_workbench.sources import population_rows,projection_rows
    sheet=main_sheet()
    population,excluded=population_rows(sheet,{'selection':{}})
    assert population=={2,3,5,6,7,8,10} and excluded==[]
    assert [r['row'] for r in projection_rows(sheet,{'figure_kind':kind,'projection':projection},included_rows=set(range(2,12)))]==[2,3,5,6,7,8,10]


def test_evidence_uses_active_source_rows_and_does_not_invent_v2_header_coordinate():
    from app.domain.parsed_evidence.datasets import project_sheet
    from app.domain.parsed_evidence.contracts import Origin
    sheet=main_sheet()
    prepared=SimpleNamespace(schema_version=2,recipe=SimpleNamespace(),parent_source=SimpleNamespace(dataset_id='raw'))
    origin=Origin(project_id='p',file_id='f',version_id='v',version=2,kind='prepared_dataset_version',filename='synthetic.xlsx',sha256='a'*64,source_sha256='b'*64,original_download_url='/original',object_download_url='/object')
    raw=SimpleNamespace(cells=sheet.cells,stored_range='A1:C11',content_range='A1:C11')
    result=project_sheet(sheet,origin,'excel_prepare_v2',1,prepared=prepared,raw_sheet=raw)
    assert [x['location']['row'] for x in result['items'] if x['kind']=='row']==[2,3,5,6,7,8,10]
    assert all(x['location']['coordinate'] is None for x in result['items'] if x['kind']=='field')


def test_confirmed_plan_snapshot_accepts_v2_source_without_rewriting_v1():
    from tests.domain.test_prepared_analysis import raw_source
    from app.domain.analysis_sources import prepared_dataset_source
    from app.domain.analysis_plan.contracts import AnalysisPlanSnapshot
    source=prepared_dataset_source(dict(id='prepared',version=2,parent_dataset_id='raw',sha256='c'*64,size_bytes=200),raw_source(),schema_version=2)
    value=dict(schema_version=1,kind='analysis_plan',workflow_version='analysis_plan_v1',rules_version='analysis_plan_rules_v1',project_id='p',filename='synthetic.xlsx',source=source,expected_setup_revision=0,
        selection=dict(task_type='descriptive_boxplot',sheet_name='研究数据',numeric_column='A',group_column=None,unit='',missing_policy='exclude_selected_missing'))
    assert AnalysisPlanSnapshot.model_validate(value).model_dump()['source']==source
