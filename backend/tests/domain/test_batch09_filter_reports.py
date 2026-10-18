"""V3 single and collection reports retain population conditions and distinct n."""
from io import BytesIO
from docx import Document
import pytest

@pytest.fixture(autouse=True)
def postgres_schema():yield None

def processing():
    from tests.domain.test_batch09_filter_transform import recipe
    return dict(schema_version=3,operation='population_filter',recipe=recipe(),receipt=dict(root_record_count=7,base_record_count=7,
        retained_record_count=3,inherited_excluded_row_count=0,mismatch_excluded_row_count=3,unknown_excluded_row_count=1),
        manifest=dict(sha256='0123456789abcdef'*4),ancestors=[])

@pytest.mark.parametrize('language',['zh-CN','en'])
def test_v3_word_includes_population_and_separate_analysis_n(language):
    from tests.adapters.test_analysis_report_docx import snapshot
    from tests.adapters.test_analysis_plot_response import png
    from app.adapters.report_docx import build_report
    from app.adapters.report_store import ReportStore
    saved=snapshot(language)
    saved['result']['source']={'kind':'prepared_dataset_version','schema_version':3,'dataset_version':2,'parent_dataset_version':1}
    saved['processing']=processing();saved['result']['check'].update(row_count=3,valid_count=2,excluded_count=1)
    saved['result']['statistics']['n']=2
    text='\n'.join(p.text for p in Document(BytesIO(build_report(saved,ReportStore.metadata(saved),png()))).paragraphs)
    assert 'AND' in text and '18' in text and '65' in text
    assert ('当前整理记录数：3' if language=='zh-CN' else 'Current prepared records: 3') in text
    assert ('本轮筛选缺失未决排除：1' if language=='zh-CN' else 'Excluded as UNKNOWN in this filter: 1') in text
    assert '"column"' not in text and '"operator"' not in text and '0123456789abcdef'*4 not in text

def test_v3_collection_renders_population_receipt_as_prose():
    from tests.domain.test_batch05_collection_contracts import bundle
    from tests.adapters.test_analysis_plot_response import png
    from app.domain.report_collections.contracts import freeze_collection
    from app.adapters.collection_docx import build_collection_report
    first,second=bundle(language='en'),bundle(2,language='en')
    for item in (first,second):
        item['result']['source']['schema_version']=3;item['processing']=processing()
    saved=freeze_collection([first,second],language='en',title='Population')
    document=Document(BytesIO(build_collection_report(saved,{'id':'report','export_revision':1},{'figure1':png(),'figure2':png()})))
    text='\n'.join(p.text for p in document.paragraphs)
    assert 'Current prepared records: 3' in text and 'Approved population condition (AND):' in text

def test_v3_side_relation_helper_does_not_relax_v2():
    from app.domain.file_lifecycle.quality_references import valid_derivation,valid_derivation_v3
    child=dict(id='child',file_id='f',created_by='u',parent_dataset_id='root',version=2)
    parent={**child,'id':'parent','version':1}
    indexes={'prepared_dataset_versions':{'child':child,'parent':parent}}
    row=dict(prepared_dataset_id='child',base_prepared_dataset_id='parent',operation='population_filter',object_schema_version=3,
        transform_version='excel_prepare_v3',workflow_version='dataset_preparation_v3',rules_version='population_filter_rules_v1',
        manifest_object_key='derivation-manifests-v3/child/'+'a'*64+'.json',manifest_sha256='a'*64,manifest_size_bytes=100,manifest_entry_count=7)
    assert valid_derivation_v3(row,indexes) and not valid_derivation(row,indexes)
    assert not valid_derivation_v3({**row,'rules_version':'quality_rules_v1'},indexes)
