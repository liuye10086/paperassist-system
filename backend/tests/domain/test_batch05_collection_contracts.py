from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import hashlib
import pytest
from docx import Document


def bundle(i=1, language='zh-CN'):
    from tests.adapters.test_analysis_plot_response import png
    image = png()
    source = {'kind':'prepared_dataset_version','dataset_id':'prepared','sheet_name':'Sheet1','dataset_sha256':'a'*64}
    result = {'id':f'run{i}','file_id':'file','setup_revision':i,'source':source,'source_sha256':'b'*64,
        'analysis_plan':{'plan_id':f'plan{i}','sha256':str(i)*64},'output_language':language,
        'statistics':{'n':3,'mean':i},'check':{'valid_count':3,'excluded_count':1},'selection':{'sheet_name':'Sheet1'}}
    figure = {'id':f'figure{i}','analysis_run_id':result['id'],'file_id':'file','source':source,
        'source_sha256':'b'*64,'setup_revision':i,'schema_version':5,'sha256':hashlib.sha256(image).hexdigest(),
        'language':language,'figure_config_version_id':f'config{i}','figure_number':None,'numbering_mode':'external',
        'title':f'Title {i}','caption':'Frozen caption','verification':{'status':'matched'}}
    explanation = {'id':f'exp{i}','analysis_run_id':result['id'],'figure_id':figure['id'],'figure_sha256':figure['sha256'],
        'language':language,'verification':{'status':'references_checked'},'engine':{'id':'figure_explanation_v1'},
        'sections':[{'key':key,'title':key,'text':f'{key} preserved','evidence':[]} for key in ('purpose','data','methods','results','interpretation','paper_text')]}
    return {'result':result,'figure':figure,'explanation':explanation,'configuration':{'id':f'config{i}','sha256':'c'*64},'projection':None}


def test_collection_freezes_multiple_historical_plans_and_deduplicates_source():
    from app.domain.report_collections.contracts import freeze_collection
    saved=freeze_collection([bundle(),bundle(2)],language='zh-CN',title='Report')
    assert len(saved['sources']) == 1
    assert [x['figure_number'] for x in saved['items']] == [1,2]
    assert [x['result']['setup_revision'] for x in saved['items']] == [1,2]
    assert all(x['hashes']['explanation'] for x in saved['items'])
    assert len(saved['sections']) == 9
    assert saved['missing'] == []
    assert 'total_n' not in saved


@pytest.mark.parametrize('change,code',[
    ('source','collection_source_conflict'),('language','collection_source_conflict'),
    ('duplicate','collection_duplicate_figure'),('legacy','collection_legacy_numbering'),
    ('plan','collection_plan_required'),('binding','collection_explanation_conflict')])
def test_collection_rejects_invalid_sources(change,code):
    from app.domain.report_collections.contracts import freeze_collection
    from app.core.exceptions import StorageError
    a,b=bundle(),bundle(2)
    if change=='source': b['result']['source']['dataset_id']='other'
    if change=='language': b['figure']['language']='en'
    if change=='duplicate': b=a
    if change=='legacy': b['figure']['numbering_mode']='embedded'; b['figure']['figure_number']=1
    if change=='plan': b['result'].pop('analysis_plan')
    if change=='binding': b['explanation']['figure_sha256']='e'*64
    with pytest.raises(StorageError) as error: freeze_collection([a,b],language='zh-CN',title='Report')
    assert error.value.code==code


def test_missing_explanation_persists_dependency_instead_of_silently_omitting():
    from app.domain.report_collections.contracts import freeze_collection
    item=bundle(); item['explanation']=None
    saved=freeze_collection([item],language='zh-CN',title='Report')
    assert saved['missing']==[{'figure_id':'figure1','kind':'explanation'}]
    assert not saved['ready']


@pytest.mark.parametrize('language',['zh-CN','en'])
def test_collection_word_embeds_exact_png_and_nine_sections_and_number_map(language):
    from app.domain.report_collections.contracts import freeze_collection
    from app.adapters.collection_docx import build_collection_report
    from tests.adapters.test_analysis_plot_response import png
    saved=freeze_collection([bundle(language=language),bundle(2,language)],language=language,title='Report')
    content=build_collection_report(saved,{'id':'report','export_revision':1}, {'figure1':png(),'figure2':png()})
    document=Document(BytesIO(content))
    text='\n'.join(p.text for p in document.paragraphs)
    assert ('图 1' if language=='zh-CN' else 'Figure 1') in text
    assert ('图 2' if language=='zh-CN' else 'Figure 2') in text
    assert all(s['title'] in text for s in saved['sections'])
    assert 'results preserved' in text
    with ZipFile(BytesIO(content)) as archive:
        assert all(archive.read(n)==png() for n in archive.namelist() if n.startswith('word/media/'))


def test_schema_defines_five_independent_tables_and_unique_export_intent():
    import sqlalchemy as sa
    from app.db.report_collection_schema import define_report_collection_tables
    metadata=sa.MetaData()
    tables=define_report_collection_tables(metadata)
    assert len(tables)==5
    request=metadata.tables['report_build_requests']
    unique={tuple(c.name for c in x.columns) for x in request.constraints if isinstance(x,sa.UniqueConstraint)}
    assert ('version_id','request_key') in unique
    assert ('version_id','renderer_version','export_revision') in unique
    assert 'collection_reports' in metadata.tables and 'reports' not in metadata.tables



def test_word_archive_is_stable_for_fenced_crash_replay():
    from app.domain.report_collections.contracts import freeze_collection
    from app.adapters.collection_docx import build_collection_report
    from tests.adapters.test_analysis_plot_response import png
    snapshot=freeze_collection([bundle()],language='zh-CN',title='Report')
    content=build_collection_report(snapshot,{'id':'report','export_revision':1},{'figure1':png()})
    with ZipFile(BytesIO(content)) as archive:
        assert all(entry.date_time==(1980,1,1,0,0,0) for entry in archive.infolist())


def test_freeze_accepts_persisted_configuration_timestamps_and_records_caption_author():
    from datetime import datetime,timezone
    from app.domain.report_collections.contracts import freeze_collection
    value=bundle(); value['configuration']['created_at']=datetime(2026,1,1,tzinfo=timezone.utc)
    value['caption_provenance']={'origin':'user','created_by':'user1'}
    saved=freeze_collection([value],language='zh-CN',title='Report')
    assert saved['items'][0]['configuration']['created_at']=='2026-01-01T00:00:00+00:00'
    assert saved['items'][0]['caption_provenance']['created_by']=='user1'


def test_pending_config_freezes_figure_dependency_without_model_work():
    from app.domain.report_collections.contracts import freeze_collection,CollectionRequest
    item=bundle(); item['figure']=None; item['explanation']=None
    request=CollectionRequest(title='Report',language='zh-CN',items=[{'figure_config_version_id':'config1'}])
    assert request.items[0].figure_id is None
    saved=freeze_collection([item],language='zh-CN',title='Report')
    assert saved['missing']==[{'figure_config_version_id':'config1','kind':'figure'}]
    assert saved['items'][0]['configuration']['id']=='config1'
    assert saved['sections'][5]['content'][0]['missing'] is True


@pytest.mark.parametrize('language',['zh-CN','en'])
def test_word_primary_tables_use_human_labels(language):
    from app.domain.report_collections.contracts import freeze_collection
    from app.adapters.collection_docx import build_collection_report
    from tests.adapters.test_analysis_plot_response import png
    saved=freeze_collection([bundle(language=language)],language=language,title='Report')
    content=build_collection_report(saved,{'id':'report','export_revision':1},{'figure1':png()})
    document=Document(BytesIO(content))
    text='\n'.join(c.text for t in document.tables for r in t.rows for c in r.cells)
    assert ('有效记录数' if language=='zh-CN' else 'Valid records') in text
    assert 'analysis_run_id' not in text
    assert '{"' not in text


def test_collection_renderer_fits_tall_original_png_without_reencoding():
    from PIL import Image
    from app.domain.report_collections.contracts import freeze_collection
    from app.adapters.collection_docx import build_collection_report
    image=BytesIO(); Image.new('RGB',(1600,4000),'white').save(image,format='PNG'); raw=image.getvalue()
    item=bundle(); item['figure']['sha256']=hashlib.sha256(raw).hexdigest(); item['explanation']['figure_sha256']=item['figure']['sha256']
    saved=freeze_collection([item],language='zh-CN',title='Report')
    content=build_collection_report(saved,{'id':'report','export_revision':1},{'figure1':raw})
    document=Document(BytesIO(content))
    assert document.inline_shapes[0].height.inches <= 7
    with ZipFile(BytesIO(content)) as archive: assert archive.read('word/media/image1.png')==raw


def test_user_supplement_never_replaces_program_caption():
    from app.domain.report_collections.contracts import freeze_collection
    value=bundle(); value['caption']='User annotation'
    saved=freeze_collection([value],language='zh-CN',title='Report')
    assert 'Frozen caption' in saved['items'][0]['caption']
    assert 'User annotation' in saved['items'][0]['caption']


def test_pending_figure_caption_is_frozen_from_confirmed_material():
    from app.domain.report_collections.contracts import freeze_collection
    item=bundle();item['figure']=None;item['explanation']=None
    item['planned_caption']='Confirmed rules and denominator';item['planned_title']='Line plot'
    saved=freeze_collection([item],language='zh-CN',title='Report')
    assert saved['items'][0]['caption']=='Confirmed rules and denominator'


def test_same_configuration_cannot_enter_twice_as_figure_and_pending_dependency():
    from app.domain.report_collections.contracts import freeze_collection
    from app.core.exceptions import StorageError
    ready=bundle();pending=deepcopy(ready);pending['figure']=None;pending['explanation']=None
    with pytest.raises(StorageError) as error: freeze_collection([ready,pending],language='zh-CN',title='Report')
    assert error.value.code=='collection_duplicate_figure'
