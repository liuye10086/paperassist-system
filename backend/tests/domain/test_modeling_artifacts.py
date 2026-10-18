from copy import deepcopy
import pytest
from app.core.exceptions import StorageError, PlotError
from app.domain.artifact_contracts import figure_kind, figure_renderer, explanation_version, report_version
from app.domain.external_processing import disclosure_version, default_labels, digest


def roc_result():
    return {'schema_version':5,'result_type':'roc','id':'run','file_id':'file','source_sha256':'a'*64,'setup_revision':1,
        'analysis_plan':{'sha256':'b'*64},'selection':{'sheet_name':'secret','direction':'higher_is_positive','evaluation_role':'held_out_validation'},
        'variables':[{'column':'A','name':'Outcome','unit':''},{'column':'B','name':'Score','unit':''}],
        'check':{'row_count':4,'valid_count':4,'excluded_count':0},
        'statistics':{'method':'roc_auc','n':4,'positive_n':2,'negative_n':2,'auc':.75,
          'confidence_interval':{'low':.1,'high':1.,'level':.95,'method':'delong_normal_clipped','reason':None},
          'direction':'higher_is_positive','evaluation_role':'held_out_validation','warnings':[],
          'points':[{'threshold_kind':kind,'threshold':threshold,'tpr':tpr,'fpr':fpr,'sensitivity':tpr,'specificity':1-fpr}
            for kind,threshold,tpr,fpr in [('above_all',None,0.,0.),('value',.9,.5,0.),('value',.8,.5,.5),('value',.7,1.,.5),('value',.1,1.,1.)]]}}


def external(result,task_type='boxplot',figure=None):
    labels=default_labels({'result':result,'figure':figure},task_type)
    confirmation={'version':6,'confirmed':True,'source_digest':'c'*64,'labels':labels}
    return {'version':6,'provider':'openai','task_type':task_type,'source_digest':'c'*64,'labels':labels,'confirmation_digest':digest(confirmation)}


@pytest.mark.parametrize('kind,figure',[('regression','forest'),('roc','roc_curve'),('survival','kaplan_meier')])
def test_new_artifacts_are_explicit_v4(kind,figure):
    result={'schema_version':5,'result_type':kind}
    assert figure_kind(result)==figure
    assert figure_renderer(result)=='openai_modeling_plot_v4'
    assert explanation_version(result)=='analysis_explanation_v4'
    assert report_version(result)=='word_analysis_report_v4'
    assert disclosure_version(result)==6
    with pytest.raises(StorageError): figure_renderer({**result,'schema_version':6})


def test_roc_material_sends_anonymous_rows_never_expected_curve():
    from app.domain.modeling_material import build_material
    result=roc_result()
    rows=[{'outcome':y,'score':s,'row':i} for i,(y,s) in enumerate([(1,.9),(0,.8),(1,.7),(0,.1)],2)]
    payload,expected=build_material(result,rows,external=external(result))
    assert payload['observations']==[{'outcome':r['outcome'],'score':r['score']} for r in rows]
    assert 'statistics' not in payload and 'points' not in payload and 'auc' not in payload
    assert expected['statistics']['auc']==.75
    assert expected['statistics']['points']==result['statistics']['points']
    assert payload['plot_config']['verification_scope']=='independent_curve_calculation'
    with pytest.raises(StorageError): build_material(result,rows[:3],external=external(result))


def test_new_six_section_facts_reject_fabricated_auc_or_hazard():
    from app.domain.explanation_content import build_payload,render_explanation
    result=roc_result(); figure={'id':'fig','sha256':'d'*64,'figure_number':1,'title':'ROC'}
    payload=build_payload(result,figure)
    draft={key:' '.join('{{'+ref+'}}' for ref in refs) for key,refs in payload['required_references'].items()}
    assert len(render_explanation(draft,payload))==6
    for text in [' AUC is excellent',' 风险比提示获益',' odds ratio proves benefit']:
        with pytest.raises(PlotError): render_explanation({**draft,'interpretation':draft['interpretation']+text},payload)


@pytest.mark.parametrize('language,unavailable',[('zh-CN','本次不提供AUC区间'),('en','No AUC interval is supplied')])
def test_roc_method_summary_does_not_claim_an_unavailable_interval(language,unavailable):
    from app.domain.modeling_facts import method_summary
    result=roc_result()
    assert 'DeLong' in method_summary(result,language)
    result['statistics']['confidence_interval'].update(low=None,high=None,reason='evaluation_role_not_independent')
    summary=method_summary(result,language)
    assert unavailable in summary
    assert 'DeLong' not in summary


def test_v4_word_has_actual_roc_table_and_original_png():
    from io import BytesIO
    from zipfile import ZipFile
    from docx import Document
    from PIL import Image
    from app.adapters.report_docx import build_report
    result=roc_result();result.update(filename='data.csv',row_audit={'included_rows':[2,3,4,5],'excluded_rows':[]},engine={'id':'roc_statistics_dataset_v1'})
    figure={'id':'fig','sha256':'d'*64,'figure_number':1,'title':'ROC','caption':'independent curve','engine':{'id':'openai_modeling_plot_v4'}}
    explanation={'id':'exp','language':'en','sections':[],'limitations':[],'engine':{'id':'analysis_explanation_v4'},'provenance':{'input_sha256':'e'*64}}
    report={'id':'rep','created_at':'2026-10-14','input_sha256':'f'*64,'engine':{'id':'word_analysis_report_v4'}}
    snapshot={'project':{'id':'p','name':'Study','research_topic':''},'result':result,'figure':figure,'explanation':explanation}
    output=BytesIO();Image.new('RGB',(1600,1000),'white').save(output,format='PNG');png=output.getvalue()
    content=build_report(snapshot,report,png);document=Document(BytesIO(content))
    assert len(document.tables)>=2
    assert 'AUC' in '\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    with ZipFile(BytesIO(content)) as archive: assert [archive.read(n) for n in archive.namelist() if n.startswith('word/media/')]==[png]


def test_v4_disclosure_and_material_exclude_id_and_private_declarations():
    from app.domain.external_material import explanation_material
    result=roc_result();result['selection'].update(id_column='C',score_source='PRIVATE-SOURCE',formation='PRIVATE-FORMATION',outcome_time='PRIVATE-TIME')
    result['variables'].append({'column':'C','name':'PRIVATE-ID','unit':''})
    figure={'id':'fig','sha256':'d'*64,'figure_number':1,'title':'ROC','engine':{'id':'openai_modeling_plot_v4'}}
    confirmation=external(result,'explanation',figure)
    assert 'field:C' not in confirmation['labels']
    local,outbound=explanation_material(result,figure,confirmation)
    assert local['schema_version']==4
    assert 'PRIVATE' not in str(outbound)


def test_unknown_saved_artifact_engines_are_rejected():
    from app.domain.artifact_contracts import saved_explanation_version,saved_report_version
    with pytest.raises(StorageError):saved_explanation_version({'engine':{'id':'openai_modeling_plot_v5'}})
    with pytest.raises(StorageError):saved_report_version({'engine':{'id':'analysis_explanation_v5'}})


def test_km_external_time_unit_uses_confirmed_label_everywhere():
    from app.domain.modeling_material import build_material
    from app.domain.external_material import explanation_material
    from tests.helpers.modeling import selection
    result={'schema_version':5,'result_type':'survival','id':'run','source_sha256':'a'*64,
        'selection':selection('kaplan_meier','Sheet'),
        'variables':[{'column':'A','name':'ID','unit':''},{'column':'B','name':'Time','unit':'PRIVATE-UNIT'},
                     {'column':'C','name':'Status','unit':''}],
        'check':{'row_count':1,'valid_count':1,'excluded_count':0},
        'statistics':{'method':'kaplan_meier','n':1,'time_unit':'PRIVATE-UNIT','warnings':[],
            'groups':[{'identity':None,'label':'Overall','n':1,'event_n':0,'censored_n':1,'max_time':1,
                       'median':None,'median_reason':'not_reached','points':[],'risk_table':[]}]}}
    result['selection']['time_unit']='PRIVATE-UNIT'
    original=deepcopy(result)
    def confirmed(kind,figure=None):
        value=external(result,kind,figure)
        value['labels']['unit:B']='day'
        value['confirmation_digest']=digest({'version':6,'confirmed':True,
            'source_digest':value['source_digest'],'labels':value['labels']})
        return value
    payload,expected=build_material(result,[{'time':1.,'event':False,'group':0}],external=confirmed('boxplot'))
    assert payload['plot_config']['time_unit']==payload['labels']['x_label']=='day'
    assert 'PRIVATE-UNIT' not in str(payload) and 'PRIVATE-UNIT' not in str(expected)
    figure={'id':'fig','sha256':'d'*64,'figure_number':1,'title':'KM','engine':{'id':'openai_modeling_plot_v4'}}
    local,outbound=explanation_material(result,figure,confirmed('explanation',figure))
    assert outbound['facts']['modeling.time_unit']['raw']=='day'
    assert 'PRIVATE-UNIT' not in str(outbound) and 'PRIVATE-UNIT' not in str(local)
    assert result==original


def test_v4_manifest_thresholds_and_endpoints_are_exact_not_tolerant():
    from app.adapters.openai_plot import plot_differences
    expected={'schema_version':4,'statistics':{'auc':.75,'points':[{'threshold':.9,'tpr':0.,'fpr':0.}]}}
    for key in ('threshold','tpr'):
        changed=deepcopy(expected);changed['statistics']['points'][0][key]+=1e-15
        assert plot_differences(changed,expected)
