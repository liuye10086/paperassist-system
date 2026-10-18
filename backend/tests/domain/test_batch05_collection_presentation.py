"""Report presentation must explain saved facts without translating user data."""
from copy import deepcopy
from io import BytesIO
import pytest
from docx import Document
from docx.oxml.ns import qn
from app.adapters.collection_labels import rows
from app.adapters.collection_docx import build_collection_report
from app.domain.report_collections.contracts import freeze_collection
from tests.domain.test_batch05_collection_contracts import bundle
from tests.adapters.test_analysis_plot_response import png


@pytest.mark.parametrize('en',[False,True])
def test_saved_method_warning_and_status_are_readable_without_engine_metadata(en):
    value={'status':'confirmed','warnings':['correlation_not_causation'],
        'p_value':{'method':'pearson_beta_two_sided','value':0.01},
        'engine':{'id':'private-engine','version':'private-version'},'workflow_version':'analysis_plan_v3',
        'message':'样本量较小，结果可能不稳定。',
        'research_request':{'question':'样本量较小，结果可能不稳定。'},
        'selected_fields':[{'name':'confirmed','unit':{'status':'unknown','value':'unit-name'}}]}
    before=deepcopy(value);presented=list(rows(value,en));text='\n'.join(a+': '+b for a,b in presented)
    assert ('Confirmed' if en else '已确认') in text
    assert ('Correlation does not establish causation' if en else '相关不等于因果') in text
    assert ('Pearson two-sided beta test' if en else 'Pearson双侧Beta检验') in text
    assert 'private-engine' not in text and 'analysis_plan_v3' not in text
    assert any(label.endswith('Name' if en else '名称') and val=='confirmed' for label,val in presented)
    assert value==before
    assert ('The sample is small and results may be unstable.' if en else '样本量较小，结果可能不稳定。') in text
    assert any(label.endswith('Research question' if en else '研究问题') and val=='样本量较小，结果可能不稳定。' for label,val in presented)


@pytest.mark.parametrize('language',['zh-CN','en'])
def test_word_provenance_labels_identify_each_frozen_artifact(language):
    saved=freeze_collection([bundle(language=language)],language=language,title='Report')
    before=deepcopy(saved)
    content=build_collection_report(saved,{'id':'report','export_revision':1},{'figure1':png()})
    document=Document(BytesIO(content));text='\n'.join(p.text for p in document.paragraphs)
    assert ('Report ID: report' if language=='en' else '报告记录ID：report') in text
    assert ('Statistical result ID: run1' if language=='en' else '统计结果ID：run1') in text
    assert ('Figure ID: figure1' if language=='en' else '图表记录ID：figure1') in text
    assert ('Explanation ID: exp1' if language=='en' else '解释记录ID：exp1') in text
    assert 'PNG SHA-256' in text
    assert saved==before


def test_word_headings_use_explicit_bilingual_fonts_without_theme_override():
    saved=freeze_collection([bundle(language='zh-CN')],language='zh-CN',title='Report')
    document=Document(BytesIO(build_collection_report(saved,{'id':'report','export_revision':1},{'figure1':png()})))
    for name in ['Title','Heading 1','Heading 2','Heading 3']:
        fonts=document.styles[name].element.get_or_add_rPr().get_or_add_rFonts()
        assert fonts.get(qn('w:eastAsia'))=='Microsoft YaHei'
        assert fonts.get(qn('w:ascii'))=='Calibri'
        assert fonts.get(qn('w:hAnsi'))=='Calibri'
        assert all(fonts.get(qn('w:'+key)) is None for key in ['asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'])
