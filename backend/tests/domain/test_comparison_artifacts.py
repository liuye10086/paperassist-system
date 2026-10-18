"""Comparison artifacts use fixed local inference and independent cloud plot summaries."""
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
import json

import pytest
from docx import Document


def comparison_result(language='en', design='independent', method=None, posthoc='none'):
    method = method or {'independent':'welch_t','paired':'paired_t','multiple':'welch_anova'}[design]
    groups = [[1., 3., 5.], [2., 5., 8.]]
    if design == 'multiple': groups.append([4., 8., 12.])
    from app.domain.descriptive import summarize
    variables = ([{'column':'B','name':'PRIVATE BEFORE','unit':'cm'}, {'column':'C','name':'PRIVATE AFTER','unit':'cm'}]
                 if design == 'paired' else [{'column':'B','name':'PRIVATE OUTCOME','unit':'cm'}])
    statistic = {'name':'F' if design=='multiple' else 't','value':1.25,
                 'df_numerator':2. if design=='multiple' else None, 'df_denominator':4.}
    estimate = {'name':None if design=='multiple' else 'mean_difference','value':None if design=='multiple' else -2.,
                'unit':None if design=='multiple' else 'outcome','reason':'omnibus_no_single_effect' if design=='multiple' else None}
    ci = {'low':None if design=='multiple' else -6.,'high':None if design=='multiple' else 2., 'level':.95,
          'method':None if design=='multiple' else method,'reason':'omnibus_no_interval' if design=='multiple' else None,'simultaneous':False}
    inference = {'alternative':'omnibus' if design=='multiple' else 'two-sided','rank_method':None,'zero_method':None,
                 'continuity_correction':False,'exact':False,'zero_count':0}
    stats = {'method':method,'design':design,'sample_sizes':[3]*len(groups),'statistic':statistic,'estimate':estimate,
             'p_value':{'value':.25,'method':'welch_anova_f' if design=='multiple' else method+'_two_sided'},
             'confidence_interval':ci,'inference':inference,'family':{'policy':posthoc,'size':3 if posthoc!='none' else 0,'alpha':.05},
             'pairwise':[],'warnings':['observational_not_causal']}
    if posthoc!='none':
        for left,right,p,adjusted in [(0,1,.02,.06),(0,2,.01,.03),(1,2,.2,.2)]:
            stats['pairwise'].append({'left_index':left,'right_index':right,'method':'welch_t','statistic':{**statistic,'name':'t','df_numerator':None},
                'estimate':{**estimate,'name':'mean_difference','value':-2.,'unit':'outcome','reason':None},
                'p_value':{'value':p,'method':'welch_t_two_sided'},'confidence_interval':{**ci,'low':-6.,'high':2.,'method':'welch_t','reason':None},
                'inference':{**inference,'alternative':'two-sided'},'adjusted_p_value':adjusted,'adjustment':'holm'})
    value = {'id':'run','file_id':'file','filename':'source.xlsx','source_sha256':'a'*64,'setup_revision':1,
        'schema_version':4,'result_type':'group_comparison','analysis_plan':{'id':'plan','sha256':'b'*64},
        'engine':{'id':'group_comparison_statistics_dataset_v1'},
        'analysis_context':{'plan':{'research_request':{'output_language':language,'question':'PRIVATE RESEARCH',
            'population':'PRIVATE POPULATION','study_design':'USER DESIGN','observation_unit':'PERSON',
            'observation_structure':'paired' if design=='paired' else 'independent','relationship_description':'LOCAL RELATIONSHIP','output_requirements':'LOCAL OUTPUT'}}},
        'selection':{'task_type':'group_comparison','sheet_name':'PRIVATE SHEET','design':design,'method':method,
            'numeric_column':None if design=='paired' else 'B','group_column':None if design=='paired' else 'D',
            'groups':[] if design=='paired' else [{'kind':'text','value':f'PRIVATE GROUP {i}'} for i in range(len(groups))],
            'pair_id_column':'A' if design=='paired' else None,'before_column':'B' if design=='paired' else None,'after_column':'C' if design=='paired' else None,
            'unit':'cm','missing_policy':'complete_pairs' if design=='paired' else 'complete_cases','posthoc':posthoc,
            'assumptions':{'independent_units':True,'distribution':'normal','inference_confirmed':True}},
        'variables':variables,'groups':[{'label':f'PRIVATE GROUP {i}','statistics':summarize(g)} for i,g in enumerate(groups)],
        'statistics':stats,'check':{'row_count':10,'valid_count':3 if design=='paired' else 3*len(groups),
            'excluded_count':7 if design=='paired' else 10-3*len(groups),'groups':[],'issues':[],'warnings':[]},
        'row_audit':{'included_rows':[2,3,4],'excluded_rows':[{'row':5,'coordinates':['B5'],'reason':'missing_measurement'}],
            'missing_policy':'complete_pairs' if design=='paired' else 'complete_cases','group_rows':[[2,3,4]]*len(groups),'pair_id_column':'A' if design=='paired' else None}}
    return value, groups


def confirmation(value, task_type='boxplot', figure=None):
    from app.domain.external_processing import disclosure, freeze_confirmation
    snapshot={'result':value,'file':{'id':'file','filename':'source.xlsx'},'output_language':value['analysis_context']['plan']['research_request']['output_language']}
    if figure: snapshot['figure']=figure
    preview=disclosure(snapshot,task_type,'owner')
    labels={item['key']:item['value'] for item in preview['labels']}
    for index,variable in enumerate(value['variables']): labels['field:'+variable['column']]='Outcome' if len(value['variables'])==1 else ('Before' if index==0 else 'After')
    for index in range(len(value['groups'])): labels['group:'+str(index)]=('Before' if index==0 else 'After') if value['selection']['design']=='paired' else f'Group {index+1}'
    if task_type=='explanation': labels['research_summary']='Approved summary'
    return freeze_confirmation(snapshot,task_type,'owner',{'version':preview['version'],'confirmed':True,'source_digest':preview['source_digest'],'labels':labels})


@pytest.mark.parametrize('language',['zh-CN','en'])
@pytest.mark.parametrize('design,kind',[('independent','group_boxplot'),('multiple','group_boxplot'),('paired','paired_connected')])
def test_comparison_routes_are_explicit_and_versioned(language,design,kind):
    from app.domain.artifact_contracts import figure_renderer,figure_kind,explanation_version,report_version
    value,_=comparison_result(language,design)
    assert figure_renderer(value)=='openai_comparison_plot_v3'
    assert figure_kind(value)==kind
    assert explanation_version(value)=='analysis_explanation_v3'
    assert report_version(value)=='word_analysis_report_v3'


def test_unknown_result_cannot_fall_back_to_descriptive_artifact():
    from app.domain.artifact_contracts import figure_kind
    from app.core.exceptions import StorageError
    with pytest.raises(StorageError): figure_kind({'schema_version':4,'result_type':'unknown'})


@pytest.mark.parametrize('design',['independent','paired','multiple'])
def test_comparison_cloud_plot_only_receives_reviewed_numeric_material(design):
    from app.domain.comparison_material import build_material
    value,groups=comparison_result(design=design)
    external=confirmation(value)
    payload,expected=build_material(value,groups,external=external)
    assert payload['schema_version']==3 and expected['statistics']['sample_sizes']==[3]*len(groups)
    assert 'PRIVATE' not in repr(payload)
    assert not any(key in payload for key in ('statistics','analysis_run_id','source_sha256','selection','row_audit'))
    assert all('statistics' not in group for group in payload['groups'])
    assert expected['plot_config']['inference_displayed'] is False
    if design=='paired':
        assert payload['paired_values']==[[1.,2.],[3.,5.],[5.,8.]]
        assert expected['paired_summary']['n_pairs']==3
        assert expected['paired_summary']['difference']['mean']==-2.
        assert 'ID' not in repr(payload) and payload['plot_config']['paired_direction']=='before_minus_after'
    else:
        assert len(expected['series'])==len(groups)


def test_comparison_disclosure_has_no_pair_identifier_and_requires_unique_group_labels():
    from app.domain.external_processing import disclosure,freeze_confirmation
    from app.core.exceptions import StorageError
    value,_=comparison_result(design='paired')
    snapshot={'result':value,'file':{'filename':'source.xlsx'}}
    preview=disclosure(snapshot,'boxplot','owner')
    assert preview['version']==5 and preview['summary']['design']=='paired'
    assert preview['summary']['valid_count']==3 and preview['summary']['paired_direction']=='before_minus_after'
    assert not any('pair_id' in key for key in preview)
    labels={item['key']:item['value'] for item in preview['labels']}
    labels['group:1']=labels['group:0']
    with pytest.raises(StorageError): freeze_confirmation(snapshot,'boxplot','owner',{'version':5,'confirmed':True,'source_digest':preview['source_digest'],'labels':labels})


@pytest.mark.parametrize('language',['zh-CN','en'])
def test_comparison_six_sections_bind_all_local_inference_and_pairwise_family(language):
    from app.domain.explanation_content import build_payload,render_explanation,limitations_for
    value,_=comparison_result(language,'multiple',posthoc='all_pairs_holm')
    payload=build_payload(value,{'id':'figure','title':'Figure 1','figure_number':1,'sha256':'c'*64})
    assert payload['schema_version']==3
    assert payload['facts']['comparison.family.size']['raw']==3
    assert payload['facts']['pairwise.0.adjusted_p_value']['raw']==.06
    assert payload['facts']['pairwise.1.p_value']['raw']==.01
    assert payload['facts']['comparison.interval_caveat']['raw']
    assert 'relationship_form' not in repr(payload)
    assert 'comparison.p_value' in payload['required_references']['results']
    assert all(f'pairwise.{i}.adjusted_p_value' in payload['required_references']['results'] for i in range(3))
    prefix='依据 ' if language=='zh-CN' else 'Based on '
    draft={key:prefix+', '.join('{{'+ref+'}}' for ref in refs)+'.' for key,refs in payload['required_references'].items()}
    sections=render_explanation(draft,payload)
    assert len(sections)==6 and '{{' not in sections[3]['text']
    assert ('Holm' in sections[3]['text'])
    assert any(('simultaneous' in item if language=='en' else '同时' in item) for item in limitations_for(value,language))


@pytest.mark.parametrize('bad',['The test statistic is large.','The effect size is substantial.','总体检验显著。','P value was low.','The degrees of freedom are limited.'])
def test_comparison_guard_rejects_unbound_statistical_claims(bad):
    from app.domain.explanation_content import build_payload,render_explanation
    from app.core.exceptions import PlotError
    value,_=comparison_result()
    payload=build_payload(value,{'id':'figure','title':'Figure 1','figure_number':1,'sha256':'c'*64})
    draft={key:'Based on '+', '.join('{{'+ref+'}}' for ref in refs)+'.' for key,refs in payload['required_references'].items()}
    draft['interpretation']+=bad
    with pytest.raises(PlotError): render_explanation(draft,payload)


def test_comparison_explanation_sends_confirmed_summary_and_no_raw_audit():
    from app.domain.external_material import explanation_material
    value,_=comparison_result(design='paired')
    figure={'id':'figure','title':'Figure 1','figure_number':1,'sha256':'c'*64,'engine':{'id':'openai_comparison_plot_v3'},'language':'en'}
    external=confirmation(value,'explanation',figure)
    local,outbound=explanation_material(value,figure,external)
    assert outbound['research_summary']=='Approved summary'
    assert 'PRIVATE' not in repr(outbound)
    assert not any(key in outbound for key in ('row_audit','paired_values','values','pair_id_column','source_sha256','analysis_run_id'))
    assert local['facts']['comparison.direction']['raw']=='Before - After'


@pytest.mark.parametrize('language',['zh-CN','en'])
@pytest.mark.parametrize('design',['paired','multiple'])
def test_comparison_word_uses_correct_assumptions_local_family_and_original_png(language,design):
    from app.adapters.report_docx import build_report
    from tests.adapters.test_analysis_plot_response import png
    value,_=comparison_result(language,design,posthoc='all_pairs_holm' if design=='multiple' else 'none')
    figure={'id':'figure','sha256':'c'*64,'title':'Figure 1','caption':'Checked observations','engine':{'id':'openai_comparison_plot_v3'}}
    explanation={'id':'explanation','language':language,'engine':{'id':'analysis_explanation_v3'},'sections':[],
        'limitations':[],'verification':{'note':'Checked'},'provenance':{'input_sha256':'d'*64}}
    saved={'project':{'id':'project','name':'Research','research_topic':'LOCAL TOPIC'},'result':value,'figure':figure,'explanation':explanation}
    report={'id':'report','created_at':'2026','engine':{'id':'word_analysis_report_v3'},'input_sha256':'e'*64}
    content=build_report(saved,report,png())
    document=Document(BytesIO(content))
    text='\n'.join(p.text for p in document.paragraphs)+'\n'+'\n'.join(cell.text for table in document.tables for row in table.rows for cell in row.cells)
    assert ('Group comparison report' if language=='en' else '组间比较分析报告') in text
    assert 'PRIVATE RESEARCH' in text and 'LOCAL RELATIONSHIP' in text and 'USER DESIGN' in text
    assert ('distribution' in text.lower() if language=='en' else '分布' in text)
    assert ('Relationship form' not in text and '关系形态' not in text)
    if design=='multiple': assert 'Holm' in text and '0.06' in text and ('simultaneous' in text if language=='en' else '同时' in text)
    else: assert ('3 complete pairs' in text if language=='en' else '3 个完整对' in text)
    with ZipFile(BytesIO(content)) as archive: assert [archive.read(name) for name in archive.namelist() if name.startswith('word/media/')]==[png()]


@pytest.mark.parametrize('design',['independent','paired'])
def test_v3_plot_receipt_accepts_figure_file_and_verifies_independent_counts(design):
    from app.domain.comparison_material import build_material
    from app.domain.plot_response import parse_plot_response,download_plot,PlotVerificationError
    from tests.adapters.test_analysis_plot_response import png,parsed
    value,groups=comparison_result(design=design)
    _,expected=build_material(value,groups,external=confirmation(value))
    response={'id':'resp_test','status':'completed','output':[
        {'type':'code_interpreter_call','container_id':'cntr_test','status':'completed'},
        {'type':'message','content':[{'annotations':[{'type':'container_file_citation','container_id':'cntr_test','filename':name,'file_id':file} for name,file in parsed()['files'].items()]}]}]}
    receipt=parse_plot_response(response,'cntr_test',expected)
    assert receipt['provenance']['prompt_version']=='openai_comparison_plot_v3'
    actual=deepcopy(expected)
    class Provider:
        def download(self,file_id,*args,**kwargs): return png() if file_id=='cfile_png' else json.dumps(actual).encode()
    assert download_plot(receipt,Provider(),None,expected)['manifest']==expected
    actual['statistics']['sample_sizes'][0]=True
    with pytest.raises(PlotVerificationError) as error: download_plot(receipt,Provider(),None,expected)
    assert any(item['path']=='statistics.sample_sizes.0' for item in error.value.differences)
    for injection in ('p_value','raw_rows'):
        actual=deepcopy(expected)
        actual[injection]=.01 if injection=='p_value' else ['PRIVATE RAW ID']
        with pytest.raises(PlotVerificationError) as error: download_plot(receipt,Provider(),None,expected)
        assert 'PRIVATE' not in repr(error.value.differences)


def test_legacy_chinese_plot_download_carries_bounded_differences():
    from app.domain.plot_response import download_plot,PlotVerificationError
    from tests.adapters.test_analysis_plot_response import png
    expected={'title':'旧中文图','statistics':{'n':3},'series':[{'fliers':[1.,2.]}]}
    actual={**expected,'statistics':{'n':2},'values':['PRIVATE ID VALUE'],'series':[{'fliers':[3.]}]}
    parsed={'files':{'boxplot.png':'png','result.json':'manifest'},'provenance':{'response_id':'resp_old','container_id':'cntr_old'}}
    class Provider:
        def download(self,file_id,*args,**kwargs): return png() if file_id=='png' else json.dumps(actual).encode()
    with pytest.raises(PlotVerificationError) as error: download_plot(parsed,Provider(),None,expected)
    assert any(item['path']=='statistics.n' for item in error.value.differences)
    assert 'PRIVATE' not in repr(error.value.manifest)
    assert 'observation values omitted' in repr(error.value.differences)


def test_v3_admission_keeps_amounts_and_checks_material_before_provider():
    from app.domain import plot_policy,explanation_policy
    from app.core.exceptions import StorageError
    plot={'policy':{'model':'synthetic','prompt_version':'openai_boxplot_v1','max_output_tokens':10000,'timeout_seconds':15,
        'tools':['code_interpreter'],'output_format':'text','price':{'version':'test','currency':'USD','model':'synthetic',
        'input_micro_usd_per_million':1,'cached_input_micro_usd_per_million':1,'output_micro_usd_per_million':1,'cache_write_micro_usd_per_million':1}},
        'input_token_allowance':100000,'task_limit_micro_usd':1000000,'tool_reserve_micro_usd':1000,'tool_reserve_version':'test'}
    value,groups=comparison_result()
    from app.domain.comparison_material import build_material
    payload,_=build_material(value,groups,external=confirmation(value))
    derived=plot_policy.derive_execution_policy(plot,payload)
    assert derived['policy']['prompt_version']=='openai_comparison_plot_v3'
    assert derived['policy']['price']==plot['policy']['price'] and derived['tool_reserve_micro_usd']==plot['tool_reserve_micro_usd']
    with pytest.raises(StorageError) as error: plot_policy.validate_execution_policy({**derived,'input_token_allowance':1},payload)
    assert error.value.code=='plot_input_limit'
    explanation={key:value for key,value in plot.items() if not key.startswith('tool_')}
    explanation['policy']={**plot['policy'],'tools':[],'prompt_version':'analysis_explanation_v1','output_format':'analysis_explanation_v1'}
    derived=explanation_policy.derive_execution_policy(explanation,{'schema_version':3})
    assert derived['policy']['prompt_version']==derived['policy']['output_format']=='analysis_explanation_v3'
    assert derived['policy']['price']==plot['policy']['price']


def test_comparison_material_rejects_changed_group_values_even_with_same_n():
    from app.domain.comparison_material import build_material
    from app.core.exceptions import StorageError
    value,groups=comparison_result()
    groups[0][0]=100.
    with pytest.raises(StorageError) as error: build_material(value,groups,external=confirmation(value))
    assert error.value.code=='result_mismatch'


def test_v3_report_metadata_and_source_validation_use_frozen_comparison_contract():
    from app.adapters.report_store import ReportStore,digest
    from app.domain.artifact_contracts import figure_metadata,explanation_metadata
    from app.domain.comparison_material import build_material
    from app.domain.external_material import explanation_material
    from app.domain.explanation_content import render_explanation
    from app.domain.reports import validate_sources
    from app.core.exceptions import StorageError
    value,groups=comparison_result()
    external=confirmation(value)
    _,expected=build_material(value,groups,external=external)
    figure={**figure_metadata(value,expected,{'model':'synthetic'},external=external),'id':'figure','sha256':'c'*64}
    explanation_external=confirmation(value,'explanation',figure)
    payload,_=explanation_material(value,figure,explanation_external)
    draft={key:'Based on '+', '.join('{{'+ref+'}}' for ref in refs)+'.' for key,refs in payload['required_references'].items()}
    explanation={**explanation_metadata(value,figure,render_explanation(draft,payload),payload,{'model':'synthetic','external_processing':explanation_external}),'id':'explanation'}
    snapshot={'result':value,'figure':figure,'explanation':explanation}
    report=ReportStore.metadata(snapshot)
    assert report['engine']['id']=='word_analysis_report_v3' and report['language']=='en' and report['schema_version']==3
    validate_sources(value,figure,explanation)
    figure['groups'][0]['statistics']['mean']=999.
    with pytest.raises(StorageError): validate_sources(value,figure,explanation)


def test_all_six_groups_and_fifteen_pairs_fit_existing_six_section_output_bounds():
    from app.domain.descriptive import summarize
    from app.domain.explanation_content import build_payload,render_explanation
    value,_=comparison_result(design='multiple',posthoc='all_pairs_holm')
    value['groups']=[{'label':f'Group {i}','statistics':summarize([float(i),float(i+1),float(i+3)])} for i in range(6)]
    seed=value['statistics']['pairwise'][0]
    value['statistics']['pairwise']=[{**deepcopy(seed),'left_index':i,'right_index':j} for i in range(6) for j in range(i+1,6)]
    value['statistics']['sample_sizes']=[3]*6;value['statistics']['family']['size']=15
    payload=build_payload(value,{'id':'f','title':'Figure 1','figure_number':1,'sha256':'c'*64})
    draft={key:'Based on '+', '.join('{{'+ref+'}}' for ref in refs)+'.' for key,refs in payload['required_references'].items()}
    assert max(map(len,draft.values()))<=4000
    assert len(render_explanation(draft,payload))==6


def test_malformed_cloud_evidence_never_retains_raw_arrays_in_differences():
    from app.adapters.openai_plot import plot_differences
    differences=plot_differences({'statistics':['PRIVATE RAW ID',123]}, {'statistics':{'n':3}})
    assert 'PRIVATE' not in repr(differences) and '123' not in repr(differences)


def test_cloud_material_accepts_platform_summary_roundoff_and_retains_frozen_facts():
    import math
    from app.domain.comparison_material import build_material
    from app.domain.comparison.service import statistics_match
    from app.core.exceptions import StorageError
    value,groups=comparison_result()
    original=deepcopy(value['groups'])
    value['groups'][0]['statistics']['std']=math.nextafter(value['groups'][0]['statistics']['std'],math.inf)
    assert statistics_match(value['groups'],original,summaries=True)
    frozen=deepcopy(value)
    _,expected=build_material(value,groups,external=confirmation(value))
    assert expected['groups'][0]['statistics']['std']==value['groups'][0]['statistics']['std']
    assert value==frozen
    for _ in range(64): value['groups'][0]['statistics']['std']=math.nextafter(value['groups'][0]['statistics']['std'],math.inf)
    with pytest.raises(StorageError): build_material(value,groups,external=confirmation(value))


@pytest.mark.parametrize('language',['zh-CN','en'])
@pytest.mark.parametrize('method,design,label',[('welch_t','independent',('均值差','Mean difference')),
    ('mann_whitney_u','independent',('秩双列相关','Rank-biserial correlation')),
    ('paired_t','paired',('均值差','Mean difference')),('wilcoxon','paired',('秩双列相关','Rank-biserial correlation')),
    ('welch_anova','multiple',('单一效应量','Single effect')),('kruskal_wallis','multiple',('单一效应量','Single effect'))])
def test_comparison_visible_facts_and_word_use_localized_algorithm_effect_and_boolean_names(language,method,design,label):
    from app.domain.comparison.kernel import compute_comparison
    from app.domain.comparison_facts import build_payload
    from app.adapters.report_docx import build_report
    from app.domain.descriptive import summarize
    from tests.adapters.test_analysis_plot_response import png
    value,_=comparison_result(language,design)
    groups=([[8.,9.,7.,10.,11.],[2.,3.,4.,6.,8.]] if design=='paired' else
        [[1.,2.,4.,5.,7.],[3.,5.,6.,8.,11.]]+([[6.,7.,9.,12.,13.]] if design=='multiple' else []))
    value['statistics']=compute_comparison(groups,method=method,posthoc='all_pairs_holm' if design=='multiple' else 'none')
    value['selection']['method']=method
    value['groups']=[{'label':f'Group {i}','statistics':summarize(group)} for i,group in enumerate(groups)]
    figure={'id':'figure','sha256':'c'*64,'title':'Figure 1','caption':'Checked observations','figure_number':1,'engine':{'id':'openai_comparison_plot_v3'}}
    payload=build_payload(value,figure)
    assert payload['facts']['comparison.estimate']['label']==label[0 if language=='zh-CN' else 1]
    visible='\n'.join(fact['label']+': '+fact['value'] for fact in payload['facts'].values())
    forbidden=['mean_difference','rank_biserial','welch_t_two_sided','paired_t_two_sided','mann_whitney_',
        'wilcoxon_sign_exact','wilcoxon_asymptotic','welch_anova_f','kruskal_chi_square','zero_method','wilcox;', 'True','False']
    assert not any(token in visible for token in forbidden)
    algorithm=payload['facts']['comparison.inference_summary']['value']
    assert ('解析' in algorithm if language=='zh-CN' else 'analytical' in algorithm) if method in ('welch_t','paired_t') else ('检验' in algorithm if language=='zh-CN' else 'test' in algorithm or 'enumeration' in algorithm or 'approximation' in algorithm)
    assert ('连续性校正：否' in algorithm or '连续性校正：是' in algorithm) if language=='zh-CN' else ('continuity correction: no' in algorithm or 'continuity correction: yes' in algorithm)
    if method=='wilcoxon': assert ('符号秩中排除零差值' in algorithm if language=='zh-CN' else 'exclude zero differences from signed ranks' in algorithm)
    explanation={'id':'explanation','language':language,'engine':{'id':'analysis_explanation_v3'},'sections':[],
        'limitations':[],'verification':{'note':'Checked'},'provenance':{'input_sha256':'d'*64}}
    snapshot={'project':{'id':'project','name':'Research','research_topic':'LOCAL TOPIC'},'result':value,'figure':figure,'explanation':explanation}
    report={'id':'report','created_at':'2026','engine':{'id':'word_analysis_report_v3'},'input_sha256':'e'*64}
    document=Document(BytesIO(build_report(snapshot,report,png())))
    text='\n'.join(p.text for p in document.paragraphs).split('Appendix B' if language=='en' else '附录 B')[0]
    assert label[0 if language=='zh-CN' else 1] in text
    assert not any(token in text for token in forbidden)
