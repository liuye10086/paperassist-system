from io import BytesIO
import pytest
from openpyxl import Workbook
from app.adapters.excel_structure import extract_excel_structure
from app.core.config import ExcelSettings
from app.domain.analysis import AnalysisSelection


def sheet(rows, *, tail=False):
    book = Workbook()
    book.active.title = 'Data'
    for row in rows:
        book.active.append(row)
    if tail:
        book.active['D100'].number_format = '0.00'
        book.active.merge_cells('C80:D81')
    stream = BytesIO()
    book.save(stream)
    return extract_excel_structure(stream.getvalue(), 'data.xlsx', ExcelSettings()).sheets[0]


def selection(**changes):
    return AnalysisSelection(task_type='descriptive_boxplot', sheet_name='Data', numeric_column='A',
        group_column='B', unit='', missing_policy='exclude_selected_missing', **changes)


def test_full_sparse_range_preserves_intermediate_missing_and_ignores_format_tail():
    from app.domain import dataset_analysis
    value = sheet([['Value', 'Group'], [1, 'A'], [None, None]] + [[n, 'A'] for n in range(2, 32)], tail=True)
    profile, scan = dataset_analysis.scan_dataset(value, selection(), collect_values=True)
    assert profile.row_count == 32
    assert len(profile.columns) == 2
    assert scan.check(profile).valid_count == 31
    assert scan.values == list(map(float, range(1, 32)))


@pytest.mark.parametrize('lexical', ['9007199254740992', '9007199254740993', '1e309', '1e-400', '-1e-999'])
def test_unsafe_binary64_numbers_are_rejected(lexical):
    from app.domain.dataset_analysis import scan_dataset
    value = sheet([['Value', 'Group'], [1, 'A']])
    cell = value.cells[-2].model_copy(update={'value': lexical, 'raw_value': lexical})
    value = value.model_copy(update={'cells': [*value.cells[:-2], cell, value.cells[-1]]})
    with pytest.raises(Exception, match='精度|范围'):
        scan_dataset(value, selection(), collect_values=True)


def test_decimal_group_identity_is_exact_and_labels_stable():
    from app.domain.dataset_analysis import scan_dataset
    value = sheet([['Value', 'Group']] + [[1, n] for n in range(6)])
    lexical = ['1', '1.0', '1e0', '-0', '9007199254740992', '9007199254740993']
    cells = [c.model_copy(update={'value': lexical[c.row-2], 'raw_value': lexical[c.row-2]})
        if c.row > 1 and c.column == 2 else c for c in value.cells]
    profile, scan = scan_dataset(value.model_copy(update={'cells': cells}), selection(), collect_values=True)
    assert [(g.label, g.count) for g in scan.check(profile).groups] == [
        ('1', 3), ('0', 1), ('9007199254740992', 1), ('9007199254740993', 1)]


@pytest.mark.parametrize('bad', ['12', '=1+1', True])
def test_invalid_numeric_kind_is_not_dropped(bad):
    from app.domain.dataset_analysis import scan_dataset
    profile, scan = scan_dataset(sheet([['Value','Group'], [1,'A'], [bad,'A']]), selection())
    assert not scan.check(profile).ready


def test_strict_source_rejects_boolean_version_and_extra_path():
    from app.domain.analysis_sources import DatasetAnalysisSource
    source = dict(kind='dataset_version', schema_version=1, dataset_id='d', structure_id='s',
        dataset_version=1,file_id='f',source_sha256='a'*64,dataset_sha256='b'*64,
        dataset_size_bytes=1,dataset_schema_version=1,sheet_index=0,sheet_name='Data',parser_version='openpyxl-3.1.5-structure-v1')
    for field in ('schema_version','dataset_schema_version'):
        with pytest.raises(ValueError):
            DatasetAnalysisSource.model_validate({**source,field:True})
    with pytest.raises(ValueError):
        DatasetAnalysisSource.model_validate({**source,'object_key':'secret'})
    with pytest.raises(ValueError):
        DatasetAnalysisSource.model_validate({**source,'parser_version':'arbitrary'})


@pytest.mark.parametrize('cache_state,cached', [('missing',None),('empty',''),('present','2')])
def test_formula_cache_never_becomes_numeric(cache_state,cached):
    from app.domain.dataset_analysis import scan_dataset
    from app.domain.datasets.contracts import StructureFormula
    value=sheet([['Value','Group'],[1,'A'],['=1+1','A']])
    cells=[c.model_copy(update={'formula':StructureFormula(text='1+1',attributes={},cache_state=cache_state,
        cached_type='n',cached_value=cached),'raw_type':'n','raw_value':cached}) if c.kind=='formula' else c for c in value.cells]
    profile,scan=scan_dataset(value.model_copy(update={'cells':cells}),selection())
    assert not scan.check(profile).ready
    assert profile.columns[0].type=='mixed'


@pytest.mark.parametrize('bad_kind', ['date','datetime','time','duration','error'])
def test_temporal_and_error_cells_are_never_numeric(bad_kind):
    from app.domain.dataset_analysis import scan_dataset
    value=sheet([['Value','Group'],[1,'A'],[2,'A']])
    cells=[c.model_copy(update={'kind':bad_kind,'value':'2'}) if c.coordinate=='A3' else c for c in value.cells]
    profile,scan=scan_dataset(value.model_copy(update={'cells':cells}),selection())
    assert not scan.check(profile).ready


def test_headers_whitespace_and_boolean_labels():
    from app.domain.dataset_analysis import scan_dataset
    value=sheet([['Value','Value',None],[1,True,'  '],[2,False,''],[3,None,'NA']])
    profile,scan=scan_dataset(value,selection(),collect_values=True)
    assert [c.name for c in profile.columns]==['Value','Value','未命名列 C']
    assert profile.columns[2].missing_count==2
    assert profile.columns[2].type=='text'
    assert [(g.label,g.count) for g in scan.check(profile).groups]==[('TRUE',1),('FALSE',1)]
    assert scan.values==[1.0,2.0]


def test_decimal_label_does_not_round_through_decimal_context():
    from app.domain.dataset_analysis import decimal_label
    from decimal import Decimal,localcontext
    with localcontext() as context:
        context.prec=3
        assert decimal_label(Decimal('123456789012345678901234567890.00'))=='1.2345678901234567890123456789E+29'


def test_extreme_explicit_float_lexical_remains_supported():
    from app.domain.dataset_analysis import scan_dataset
    value=sheet([['Value','Group'],[1,'A']])
    cells=[c.model_copy(update={'value':'1e308','raw_value':'1e308'}) if c.coordinate=='A2' else c for c in value.cells]
    profile,scan=scan_dataset(value.model_copy(update={'cells':cells}),selection(),collect_values=True)
    assert scan.check(profile).ready
    assert scan.values==[1e308]


def test_dataset_source_must_match_owning_result_fields():
    from app.domain.analysis_sources import statistics_engine
    source=dict(kind='dataset_version',schema_version=1,dataset_id='d',structure_id='s',dataset_version=1,
        file_id='f',source_sha256='a'*64,dataset_sha256='b'*64,dataset_size_bytes=1,
        dataset_schema_version=1,sheet_index=0,sheet_name='Data',parser_version='openpyxl-3.1.5-structure-v1')
    for value in ({'file_id':'other'}, {'source_sha256':'c'*64}, {'selection':{'sheet_name':'Other'}}):
        with pytest.raises(Exception):
            statistics_engine({**value,'source':source})


def test_sparse_far_coordinate_never_expands_missing_rows(monkeypatch):
    from app.domain.dataset_analysis import scan_dataset,DatasetScan
    from app.domain.datasets.contracts import SheetStructure, StructureField
    from openpyxl.utils.cell import get_column_letter
    value=sheet([['Value','Group'],[1,'A']])
    cells=[*value.cells,
        value.cells[-1].model_copy(update={'column':256,'coordinate':'IV2','value':'Far column'}),
        value.cells[-2].model_copy(update={'row':100000,'coordinate':'A100000','value':'2','raw_value':'2'}),
        value.cells[-1].model_copy(update={'row':100000,'coordinate':'B100000'})]
    fields=[StructureField(column=c,letter=get_column_letter(c),
        header_value=next((cell.value for cell in cells if cell.row==1 and cell.column==c),None),
        observed_types=list(dict.fromkeys(cell.kind for cell in cells if cell.row>1 and cell.column==c)))
        for c in range(1,257)]
    value=SheetStructure(index=0,name='Data',state='visible',row_count=100000,column_count=256,
        nonempty_row_count=3,nonempty_data_row_count=2,stored_range='A1:IV100000',
        content_range='A1:IV100000',merged_ranges=[],fields=fields,cells=cells,warnings=[])
    observed=[]
    original=DatasetScan.observe_dataset
    def observe(scan,row,cells):
        assert row in (2,100000),'Do not expand absent rows'
        assert len(cells)<=3,'Do not expand absent columns'
        observed.append(row)
        return original(scan,row,cells)
    monkeypatch.setattr(DatasetScan,'observe_dataset',observe)
    profile,scan=scan_dataset(value,selection(),collect_values=True)
    assert observed==[2,100000]
    assert profile.row_count==99999 and len(profile.columns)==256
    assert profile.columns[0].missing_count==99997
    assert scan.check(profile).excluded_count==99997
    assert scan.values==[1.0,2.0]


def test_merged_group_coverage_is_not_filled():
    from app.domain.dataset_analysis import scan_dataset
    book=Workbook();book.active.title='Data'
    for row in [['Value','Group'],[1,'A'],[2,None]]:book.active.append(row)
    book.active.merge_cells('B2:B3')
    stream=BytesIO();book.save(stream)
    value=extract_excel_structure(stream.getvalue(),'data.xlsx',ExcelSettings()).sheets[0]
    profile,scan=scan_dataset(value,selection(),collect_values=True)
    assert profile.row_count==2
    assert scan.check(profile).valid_count==1
    assert scan.check(profile).excluded_count==1
    assert scan.values==[1.0]
    assert profile.warnings


def test_group_mixed_types_and_category_limit_reject_selection():
    from app.domain.dataset_analysis import scan_dataset
    for rows in ([['Value','Group'],[1,'A'],[2,True]],
                 [['Value','Group']]+[[n,str(n)] for n in range(21)]):
        profile,scan=scan_dataset(sheet(rows),selection())
        assert not scan.check(profile).ready
