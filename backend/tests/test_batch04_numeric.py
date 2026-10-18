import pytest

def test_roc_hand_count_and_delong():
    from app.domain.roc.kernel import compute_roc
    result=compute_roc([{'outcome':1,'score':3.0,'row':1},{'outcome':1,'score':1.0,'row':2},{'outcome':0,'score':2.0,'row':3},{'outcome':0,'score':0.0,'row':4}], {'direction':'higher_is_positive','evaluation_role':'external_validation'})
    assert result['auc']==.75
    assert result['variance']==.125

def test_km_tied_event_censor():
    from app.domain.survival.kernel import compute_survival
    rows=[dict(time=t,event=e,group=0,row=i) for i,(t,e) in enumerate([(1.,True),(1.,False),(2.,True),(3.,False)])]
    result=compute_survival(rows,{'groups':[],'time_unit':'day'})
    assert [p['survival'] for p in result['groups'][0]['points']]==[.75,.375,.375]
    assert result['groups'][0]['median']==2

def test_ols_hand_symmetric_residuals():
    from app.domain.regression.kernel import compute_regression
    rows=[dict(x=[float(x)],y=2+3*x+e,row=i) for i,(x,e) in enumerate(zip(range(-5,5),[1,-1,-1,1,0,0,1,-1,-1,1]))]
    result=compute_regression(rows,{'method':'ols','covariates':[{'column':'A','type':'continuous','unit':'x','levels':[],'reference':None}]})
    assert result['coefficients'][0]['beta']==pytest.approx(2)
    assert result['coefficients'][1]['beta']==pytest.approx(3)
    assert result['residual_sd']==pytest.approx(1)

def test_logistic_balanced_two_cell_hand_mle():
    from app.domain.regression.kernel import compute_regression
    import math
    rows=[]
    for x,events in [(0.,5),(1.,15)]:
        rows += [dict(x=[x],y=float(i<events),row=len(rows)+i) for i in range(20)]
    result=compute_regression(rows,{'method':'logistic','covariates':[{'column':'A','type':'continuous','unit':'x','levels':[],'reference':None}]})
    assert result['coefficients'][0]['beta']==pytest.approx(-math.log(3))
    assert result['coefficients'][1]['beta']==pytest.approx(math.log(9))
    assert result['coefficients'][1]['se']==pytest.approx((8/15)**.5)

def test_logistic_quasi_separation_is_blocked():
    from app.domain.regression.kernel import compute_regression
    from app.domain.regression.common import NumericError
    rows=[dict(x=[float(x)],y=float(y),row=i) for i,(x,y) in enumerate([(-1,0)]*5+[(0,0)]*5+[(0,1)]*5+[(1,1)]*5)]
    with pytest.raises(NumericError,match='分离'):compute_regression(rows,{'method':'logistic','covariates':[{'column':'A','type':'continuous','unit':'','levels':[],'reference':None}]})

def test_roc_direction_ties_and_no_false_degenerate_interval():
    from app.domain.roc.kernel import compute_roc
    rows=[dict(outcome=i%2,score=1.,row=i) for i in range(8)]
    result=compute_roc(rows,{'direction':'lower_is_positive','evaluation_role':'held_out_validation'})
    assert result['auc']==.5
    assert result['points'][0]['threshold_kind']=='below_all'
    assert result['confidence_interval']['reason']=='degenerate_delong_variance'
    assert result['confidence_interval']['low'] is None

def test_survival_all_censored_risk_identity():
    from app.domain.survival.kernel import compute_survival
    rows=[dict(time=float(i),event=False,group=0,row=i) for i in range(4)]
    g=compute_survival(rows,{'groups':[],'time_unit':'d'})['groups'][0]
    assert g['median'] is None and g['median_reason']=='not_reached'
    assert all(p['ci_low'] is None and p['survival']==1 for p in g['points'])
    assert all(p['n_at_risk']+p['events_before']+p['censored_before']==4 for p in g['risk_table'])

def test_fixed_wrapper_timeout_and_whitelist():
    from app.domain.regression.bounded import compute_bounded
    from app.domain.regression.common import NumericError
    with pytest.raises(NumericError,match='硬超时'):compute_bounded('roc',[dict(outcome=1,score=1.,row=1)],{'direction':'higher_is_positive','evaluation_role':'unknown'},timeout_seconds=.00001)
    with pytest.raises(NumericError,match='未知'):compute_bounded('os.system',[],{})

def prepared(values):
    from app.domain.dataset_preparation.contracts import PreparedSheet
    from openpyxl.utils.cell import get_column_letter
    cells=[]
    for row,items in enumerate(values,2):
        for col,value in enumerate(items,1):
            kind='blank' if value is None else 'boolean' if type(value) is bool else 'number' if isinstance(value,(int,float)) else 'text'
            cells.append(dict(row=row,column=col,coordinate=f'{get_column_letter(col)}{row}',kind=kind,value=value if kind in ('blank','boolean') else str(value)))
    return PreparedSheet(index=0,name='Data',data_range=f'A2:{get_column_letter(len(values[0]))}{len(values)+1}',fields=[dict(column=i,name=f'field{i}',unit=dict(status='unconfirmed',value=None)) for i in range(1,len(values[0])+1)],cells=cells,merged_ranges=[])

def test_boolean_encoding_and_source_audit():
    from app.domain.roc.contracts import RocSelection
    from app.domain.roc.checks import selected_records
    from tests.helpers.modeling import selection
    s=selection('roc_auc','Data');s['positive']=dict(kind='boolean',value='true');s['negative']=dict(kind='boolean',value='false')
    rows,variables,audit=selected_records(prepared([[True,.8],[False,.2],[None,.5]]),RocSelection.model_validate(s))
    assert [r['outcome'] for r in rows]==[1,0]
    assert audit['excluded_rows']==[dict(row=4,coordinates=['A4'],reason='selected_missing')]

def test_statistics_reconstruction_never_tolerates_threshold_change():
    from app.domain.regression.service import statistics_match
    assert not statistics_match({'points':[{'threshold':.5}]},{'points':[{'threshold':.5+1e-13}]})

def test_binary_categories_strict_result_internal_counts():
    from app.domain.roc.contracts import RocStatistics
    from app.domain.roc.kernel import compute_roc
    from pydantic import ValidationError
    result=compute_roc([dict(outcome=i%2,score=float(i),row=i) for i in range(8)],{'direction':'higher_is_positive','evaluation_role':'unknown'})
    result['positive_n']=99
    with pytest.raises(ValidationError):RocStatistics.model_validate(result)

def test_engine_and_audit_have_strict_bounded_metadata():
    from app.domain.regression.contracts import StatisticalEngine,RowAudit
    from pydantic import ValidationError
    engine=dict(id='regression_statistics_dataset_v1',python_version='3.11.4',numpy_version='2.4.6',scipy_version='1.17.1',parser_version='excel_parser_v1',encoding_version='typed_reference_dummy_v1',timeout_seconds=30,statsmodels_version='0.15.0',solver='newton',max_iterations=100,tolerance=1e-10,scale='center_scale_restored_original_units',separation_solver='highs',separation_timeout_seconds=5)
    StatisticalEngine.model_validate(engine)
    with pytest.raises(ValidationError):StatisticalEngine.model_validate(dict(engine,solver='arbitrary'))
    with pytest.raises(ValidationError):RowAudit.model_validate(dict(included_rows=[2,2],excluded_rows=[],missing_policy='complete_cases',encoding_counts={}))

def test_expanded_columns_limit_checked_before_data_allocation():
    from app.domain.regression.contracts import RegressionSelection
    from app.domain.regression.checks import selected_records
    from tests.helpers.modeling import selection,identity
    from app.core.exceptions import StorageError
    s=selection('ols','Data');s['covariates']=[dict(column=c,type='categorical',unit='',levels=[identity(i) for i in range(6)],reference=identity(0)) for c in ['B','C','D','E','F']]
    sheet=prepared([[1,0,0,0,0,0]])
    with pytest.raises(StorageError,match='20'):selected_records(sheet,RegressionSelection.model_validate(s))

def test_invalid_numeric_even_in_missing_record_is_blocked():
    from app.domain.roc.contracts import RocSelection
    from app.domain.roc.checks import selected_records
    from tests.helpers.modeling import selection
    from app.core.exceptions import StorageError
    with pytest.raises(StorageError):selected_records(prepared([[None,'0.7']]),RocSelection.model_validate(selection('roc_auc','Data')))

def test_duplicate_subject_ids_rejected_before_exclusion():
    from app.domain.survival.contracts import SurvivalSelection
    from app.domain.survival.checks import selected_records
    from tests.helpers.modeling import selection
    from app.core.exceptions import StorageError
    with pytest.raises(StorageError):selected_records(prepared([['a',1,1],['a',None,0]]),SurvivalSelection.model_validate(selection('kaplan_meier','Data')))

def test_regression_schema_rejects_covariance_and_dimensions():
    from app.domain.regression.kernel import compute_regression
    from app.domain.regression.contracts import RegressionStatistics
    from pydantic import ValidationError
    rows=[dict(x=[float(x)],y=2+3*x+e,row=i) for i,(x,e) in enumerate(zip(range(-5,5),[1,-1,-1,1,0,0,1,-1,-1,1]))]
    result=compute_regression(rows,{'method':'ols','covariates':[{'column':'A','type':'continuous','unit':'x','levels':[],'reference':None}]})
    result['covariance']=[[1.]]
    with pytest.raises(ValidationError):RegressionStatistics.model_validate(result)

def test_km_schema_rejects_broken_risk_identity():
    from app.domain.survival.kernel import compute_survival
    from app.domain.survival.contracts import SurvivalStatistics
    from pydantic import ValidationError
    result=compute_survival([dict(time=1.,event=True,group=0,row=2)],{'groups':[],'time_unit':'d'})
    result['groups'][0]['risk_table'][0]['n_at_risk']=100
    with pytest.raises(ValidationError):SurvivalStatistics.model_validate(result)

def test_bounded_cache_input_identity_deepcopy_and_failures(monkeypatch):
    import app.domain.regression.bounded as bounded
    from app.domain.regression.common import NumericError
    calls=[]
    def fake(kind,rows,selection,timeout_seconds=30):
        calls.append((kind,rows,selection))
        if selection.get('fail'):raise NumericError('numeric_failure','failure')
        return {'value':[rows[0]['score']], 'direction':selection['direction']}
    monkeypatch.setattr(bounded,'_compute_uncached',fake)
    bounded.clear_compute_cache()
    rows=[dict(outcome=1,score=.73,row=2)];s={'direction':'higher_is_positive','evaluation_role':'unknown'}
    first=bounded.compute_bounded('roc',rows,s);first['value'][0]=999
    assert bounded.compute_bounded('roc',rows,s)['value']==[.73]
    assert len(calls)==1
    bounded.compute_bounded('roc',[dict(outcome=1,score=.74,row=2)],s)
    bounded.compute_bounded('roc',rows,dict(s,direction='lower_is_positive'))
    assert len(calls)==3
    for i in range(2):
        with pytest.raises(NumericError):bounded.compute_bounded('roc',rows,dict(s,fail=True))
    assert len(calls)==5
    bounded.clear_compute_cache()

def test_km_hand_greenwood_loglog_and_ols_standard_errors():
    import math
    from app.domain.survival.kernel import compute_survival
    from app.domain.regression.kernel import compute_regression
    from scipy.stats import t
    rows=[dict(time=tt,event=e,group=0,row=i) for i,(tt,e) in enumerate([(1.,True),(1.,False),(2.,True),(3.,False)])]
    point=compute_survival(rows,{'groups':[],'time_unit':'day'})['groups'][0]['points'][0]
    greenwood=1/12
    assert point['greenwood_variance']==pytest.approx(.75**2*greenwood)
    a=math.log(-math.log(.75));b=1.959963984540054*math.sqrt(greenwood)/abs(math.log(.75))
    assert point['ci_low']==pytest.approx(math.exp(-math.exp(a+b)))
    assert point['ci_high']==pytest.approx(math.exp(-math.exp(a-b)))
    rows=[dict(x=[float(x)],y=2+3*x+e,row=i) for i,(x,e) in enumerate(zip(range(-5,5),[1,-1,-1,1,0,0,1,-1,-1,1]))]
    result=compute_regression(rows,{'method':'ols','covariates':[{'column':'A','type':'continuous','unit':'x','levels':[],'reference':None}]})
    intercept,slope=result['coefficients']
    assert slope['se']==pytest.approx(1/math.sqrt(82.5))
    assert intercept['se']==pytest.approx(math.sqrt(.1+.25/82.5))
    assert slope['ci_low']==pytest.approx(3-t.ppf(.975,8)/math.sqrt(82.5))

@pytest.mark.parametrize('bad', ['perfect','rank','constant','limit'])
def test_regression_boundaries_stop_inference(bad):
    from app.domain.regression.kernel import compute_regression
    from app.domain.regression.common import NumericError
    count=10001 if bad=='limit' else 20
    rows=[dict(x=[float(i)] if bad!='rank' else [1.],y=3.*i if bad=='perfect' else 1. if bad=='constant' else float(i%3),row=i) for i in range(count)]
    with pytest.raises(NumericError):compute_regression(rows,{'method':'ols','covariates':[{'column':'A','type':'continuous','unit':'','levels':[],'reference':None}]})

def test_reference_disappears_after_complete_case_selection():
    from app.domain.regression.checks import selected_records
    from app.domain.regression.contracts import RegressionSelection
    from tests.helpers.modeling import selection,identity
    from app.core.exceptions import StorageError
    s=selection('ols','Data');s['covariates']=[dict(column='B',type='categorical',unit='',levels=[identity('a'),identity('b')],reference=identity('a'))]
    with pytest.raises(StorageError):selected_records(prepared([[None,'a'],[1,'b']]),RegressionSelection.model_validate(s))

def test_original_unit_scaling_is_recorded():
    from app.domain.regression.kernel import compute_regression
    rows=[dict(x=[1000.+10*x],y=2+3*x+e,row=i) for i,(x,e) in enumerate(zip(range(-5,5),[1,-1,-1,1,0,0,1,-1,-1,1]))]
    result=compute_regression(rows,{'method':'ols','covariates':[{'column':'A','type':'continuous','unit':'unit','levels':[],'reference':None}]})
    assert result['diagnostics']['centers']==[995.]
    assert result['diagnostics']['scales'][0]==pytest.approx((8250/10)**.5)
    assert result['coefficients'][1]['beta']==pytest.approx(.3)

def test_bounded_lock_wait_counts_toward_wall_clock_budget(monkeypatch):
    import threading,time
    import app.domain.regression.bounded as bounded
    from app.domain.regression.common import NumericError
    entered=threading.Event();release=threading.Event()
    def hold():
        with bounded._CACHE_LOCK:entered.set();release.wait(1)
    holder=threading.Thread(target=hold);holder.start();assert entered.wait(1)
    monkeypatch.setattr(bounded,'_compute_uncached',lambda *args: {'unused':True})
    started=time.monotonic()
    try:
        with pytest.raises(NumericError,match='超时'):bounded.compute_bounded('roc',[dict(outcome=1,score=.5,row=2)],{'direction':'higher_is_positive','evaluation_role':'unknown'},timeout_seconds=.05)
        assert time.monotonic()-started<.5
    finally:release.set();holder.join()

def test_check_unit_issues_have_unique_field_identity(monkeypatch):
    import app.domain.regression.bounded as bounded
    from app.domain.regression.checks import build_check
    from app.domain.regression.contracts import RegressionSelection
    from tests.helpers.modeling import selection
    monkeypatch.setattr(bounded,'compute_bounded',lambda *args,**kwargs:{})
    research=dict(question='Synthetic',output_language='zh-CN',population='Synthetic',study_design='Synthetic',observation_unit='subject',observation_structure='independent',relationship_description='',output_requirements='')
    check=build_check(prepared([[1,1],[2,2]]),RegressionSelection.model_validate(selection('ols','Data')),research)
    issues=[i for i in check['issues'] if i['code']=='unit_unconfirmed']
    assert len(issues)==2 and len({i['id'] for i in issues})==2
    assert {i['field'] for i in issues}=={'A','B'}

def _daemon_fixed_numeric(connection):
    from app.domain.regression.bounded import compute_bounded
    try:
        result=compute_bounded('roc',[dict(outcome=i%2,score=float(i),row=i+2) for i in range(8)],dict(direction='higher_is_positive',evaluation_role='unknown'))
        connection.send(dict(ok=True,auc=result['auc']))
    except Exception as exc:connection.send(dict(ok=False,error=repr(exc)))
    finally:connection.close()

def test_fixed_numeric_runs_inside_daemon_worker():
    import multiprocessing as mp
    ctx=mp.get_context('spawn');parent,child=ctx.Pipe(duplex=False)
    process=ctx.Process(target=_daemon_fixed_numeric,args=(child,));process.daemon=True
    process.start();child.close()
    try:
        assert parent.poll(15),'Daemon worker returned no result.'
        message=parent.recv()
        assert message['ok'],message
        assert message['auc']==pytest.approx(.625)
    finally:
        process.join(2)
        if process.is_alive():process.kill();process.join(2)
        parent.close();process.close()

def test_fixed_runner_rejects_unknown_type_and_oversize_input():
    from app.domain.regression.runner import compute_request
    from app.domain.fixed_statistics_contract import MAX_INPUT_BYTES
    from app.domain.regression.bounded import _compute_uncached
    from app.domain.regression.common import NumericError
    with pytest.raises(NumericError):
        compute_request(dict(task_type='os.system',rows=[{}],selection={}))
    # The parent caps the serialized minimal request before any pipe spawn.
    with pytest.raises(NumericError) as caught:
        _compute_uncached('roc',[dict(outcome=1,score=.5,row=2,padding='x'*(MAX_INPUT_BYTES+1))],
            dict(direction='higher_is_positive',evaluation_role='unknown'))
    assert caught.value.code=='numeric_input_bytes_limit'


def test_fixed_subprocess_timeout_reaps_child_without_files(monkeypatch):
    import subprocess
    import app.domain.regression.bounded as bounded
    from app.domain.regression.common import NumericError
    processes=[];original_popen=subprocess.Popen
    def tracked_popen(*args,**kwargs):
        process=original_popen(*args,**kwargs);processes.append(process);return process
    monkeypatch.setattr(subprocess,'Popen',tracked_popen)
    with pytest.raises(NumericError,match='超时'):
        bounded._compute_uncached('regression',[dict(y=float(i),x=[float(i)],row=i+2) for i in range(20)],dict(method='ols',covariates=[dict(column='A',type='continuous',unit='',levels=[],reference=None)]),timeout_seconds=.2)
    assert len(processes)==1 and processes[0].poll() is not None
    assert processes[0].stdin.closed and processes[0].stdout.closed


def test_fixed_subprocess_result_pipe_cap_is_enforced(monkeypatch):
    import app.domain.regression.bounded as bounded
    from app.domain.regression.common import NumericError
    monkeypatch.setattr(bounded,'_MAX_OUTPUT_BYTES',16)
    with pytest.raises(NumericError) as caught:
        bounded._compute_uncached('roc',[dict(outcome=i%2,score=float(i),row=i+2) for i in range(8)],dict(direction='higher_is_positive',evaluation_role='unknown'))
    assert caught.value.code=='numeric_output_limit'
