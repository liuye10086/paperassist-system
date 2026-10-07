"""Synthetic schema-6 import and actual child-process interruption rehearsal."""
import importlib.util
import json
from unittest.mock import MagicMock, patch

import pytest

from app.maintenance import legacy_import
from app.db.database import migration_connection
from scripts._import_drill_fixture import BUSINESS_SENTINEL, CLOUD_SENTINEL, create_complex_source, file_hashes


def test_synthetic_drill_fixture_is_available():
    assert importlib.util.find_spec('scripts._import_drill_fixture') is not None, 'synthetic drill fixture is missing'


@pytest.fixture
def complex_source(tmp_path):
    return create_complex_source(tmp_path / 'synthetic')


def drill_module():
    assert importlib.util.find_spec('scripts.rehearse_import') is not None, 'rehearsal runner is missing'
    from scripts import rehearse_import
    return rehearse_import


def test_complex_fixture_assets_are_valid_formats_and_backup_preserves_every_value(complex_source, tmp_path):
    from openpyxl import load_workbook
    from PIL import Image
    from docx import Document
    source, directory = complex_source
    before = file_hashes(directory)
    snapshot = legacy_import.inspect_source(source, directory)
    assert snapshot.counts == {'projects': 1, 'files': 2, 'analysis_setups': 1, 'analysis_runs': 2,
                               'figures': 2, 'figure_jobs': 3, 'explanations': 2, 'explanation_jobs': 3, 'reports': 2}
    failed = next(row for row in snapshot.rows['files'] if row['parse_status'] == 'failed')
    assert json.loads(failed['error_json'])['code'] == 'invalid_workbook'
    assert snapshot.rows['analysis_setups'][0]['revision'] == 3
    assert [row['setup_revision'] for row in snapshot.rows['analysis_runs']] == [1, 2]
    assert [row['seq'] for row in snapshot.rows['figure_jobs']] == [11, 29, 3]
    assert sorted(row['seq'] for row in snapshot.rows['explanation_jobs']) == [5, 19, 41]
    for asset in directory.glob('files/*.xlsx'):
        workbook = load_workbook(asset, read_only=True)
        assert workbook.active.max_row == 3
        workbook.close()
    for asset in directory.glob('figures/*.png'):
        with Image.open(asset) as image:
            image.verify()
    for asset in directory.glob('reports/*.docx'):
        assert BUSINESS_SENTINEL in [paragraph.text for paragraph in Document(asset).paragraphs]
    backup = tmp_path / 'verified-backup'
    legacy_import.snapshot_source(source, directory, backup)
    copied = legacy_import.inspect_source(backup / 'paperassist.sqlite3', backup)
    assert copied.rows == snapshot.rows and copied.assets == snapshot.assets
    assert file_hashes(directory) == before


def test_actual_process_interruption_retry_idempotency_and_pending_scans(complex_source, tmp_path, postgres_migration_config):
    module = drill_module()
    source, directory = complex_source
    evidence = tmp_path / 'evidence'
    evidence.mkdir()
    before = file_hashes(directory)
    with patch('openai.OpenAI', side_effect=AssertionError('cloud client must never be created')) as cloud:
        result = module.rehearse_on_config(source, directory, evidence, postgres_migration_config)
    cloud.assert_not_called()
    assert result['interruption']['method'] == 'terminate_owned_child_before_reports_insert'
    assert result['interruption']['all_tables_empty'] is True
    assert result['interruption']['sequences_unchanged'] is True
    assert result['first_status'] == 'imported'
    assert result['repeat_status'] == 'already_imported'
    assert result['source_unchanged'] is True
    assert result['next_sequences'] == {'figure_jobs': 30, 'explanation_jobs': 42}
    assert result['pending_jobs'] == {'figure_jobs': ['fj-11', 'fj-29'], 'explanation_jobs': ['ej-19', 'ej-41']}
    assert file_hashes(directory) == before
    snapshot = legacy_import.inspect_source(source, directory)
    assert module.read_target(postgres_migration_config) == snapshot.rows
    with migration_connection(config=postgres_migration_config) as db:
        assert db.execute('SELECT owner_id, default_output_language FROM projects').fetchone() == {'owner_id': None, 'default_output_language': 'zh-CN'}
    verification = json.loads((evidence / 'imported.json').read_text(encoding='utf-8'))['verification']
    assert verification['matches']
    for table, detail in verification['tables'].items():
        assert detail['source_count'] == detail['target_count'] == snapshot.counts[table]
        assert detail['source_sha256'] == detail['target_sha256'] == snapshot.hashes[table]
        assert detail['matches']
        assert all(item['source_id'] == item['target_id'] and item['matches'] for item in verification['identity_mappings'][table])
    for table in legacy_import.JOB_TABLES:
        assert sorted(item['source_rowid'] for item in verification['job_sequence_mappings'][table]) == sorted(row['seq'] for row in snapshot.rows[table])
        assert all(item['source_rowid'] == item['target_seq'] and item['matches'] for item in verification['job_sequence_mappings'][table])
    assert all(item['checked'] == item['valid'] and item['matches'] for item in verification['relationships'].values())
    assert len(verification['assets']) == 7
    assert all(item['expected'] == item['actual'] and item['exists'] and item['matches'] for item in verification['assets'].values())
    assert verification['scope']
    reports = '\n'.join(path.read_text(encoding='utf-8') for path in evidence.glob('*.json'))
    assert BUSINESS_SENTINEL not in reports and CLOUD_SENTINEL not in reports


def test_default_prepare_does_not_connect_to_any_database(monkeypatch, tmp_path):
    module = drill_module()
    monkeypatch.setattr(module, 'BACKUPS', tmp_path / 'backups')
    with patch('sqlalchemy.create_engine', side_effect=AssertionError('prepare must not connect')):
        assert module.main([]) == 0
    evidence, = list((tmp_path / 'backups').iterdir())
    assert json.loads((evidence / 'prepared.json').read_text(encoding='utf-8'))['status'] == 'prepared'
    assert (evidence / 'backup' / 'paperassist.sqlite3').is_file()
    assert not (evidence / 'execution.json').exists()


def test_runner_rejects_development_or_arbitrary_schema_before_connecting(postgres_migration_config):
    module = drill_module()
    from app.db.database import DatabaseConfig
    unsafe = DatabaseConfig('development', 'postgresql://synthetic:synthetic@localhost/paperassist_system', 'public', 'migration')
    with patch('sqlalchemy.create_engine', side_effect=AssertionError('unsafe target must not connect')):
        with pytest.raises(ValueError, match='isolated test'):
            module.validate_target(unsafe)


@pytest.mark.parametrize('create_succeeds, cleanup_database', [(False, 'paperassist_system_test'), (True, 'paperassist_system')])
def test_cleanup_only_drops_schema_created_here_in_actual_test_database(monkeypatch, tmp_path, create_succeeds, cleanup_database):
    module = drill_module()
    monkeypatch.setenv('PAPERASSIST_TEST_DATABASE_URL', 'postgresql://synthetic_runtime:synthetic@localhost/paperassist_system_test')
    monkeypatch.setenv('PAPERASSIST_TEST_MIGRATION_DATABASE_URL', 'postgresql://synthetic_owner:synthetic@localhost/paperassist_system_test')
    engine = MagicMock()
    engine.dialect.identifier_preparer.quote.side_effect = lambda name: '"' + name + '"'
    first, cleanup = MagicMock(), MagicMock()
    first.exec_driver_sql.return_value.scalar_one.return_value = 'paperassist_system_test'
    cleanup.exec_driver_sql.return_value.scalar_one.return_value = cleanup_database
    contexts = []
    for connection in (first, cleanup):
        context = MagicMock()
        context.__enter__.return_value = connection
        contexts.append(context)
    engine.begin.side_effect = contexts
    if not create_succeeds:
        def collide(statement):
            if statement.startswith('CREATE SCHEMA'):
                raise RuntimeError('synthetic pre-existing schema collision')
            return first.exec_driver_sql.return_value
        first.exec_driver_sql.side_effect = collide
    with patch('sqlalchemy.create_engine', return_value=engine), patch('app.db.database.migrate_database', side_effect=RuntimeError('synthetic migration failure')):
        with pytest.raises((RuntimeError, ValueError)):
            module.execute_new_schema(tmp_path / 'source', tmp_path / 'assets', tmp_path)
    assert not any(call.args[0].startswith('DROP SCHEMA') for call in first.exec_driver_sql.call_args_list + cleanup.exec_driver_sql.call_args_list)
    assert engine.begin.call_count == (2 if create_succeeds else 1)
    assert (tmp_path / 'created-schema.json').exists() is create_succeeds
    engine.dispose.assert_called_once()
