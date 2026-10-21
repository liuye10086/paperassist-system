"""Opt-in real isolated renderer regression, no model or business DB access.

PAPERASSIST_LAYOUT_LIVE=1 PAPERASSIST_PDF_RENDERER_URL=<loopback|internal>
python -m pytest infra/pdf/test_layout_live.py -q
"""
import ctypes
from io import BytesIO
import json
import os
from pathlib import Path
import sys
from zipfile import ZipFile

from docx import Document
from lxml import etree as E
import pypdfium2 as pdfium
from pypdfium2 import raw
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[2] / 'backend'))
from app.adapters.document_layout import catalog_docx, prepare_layout_docx
from app.adapters.pdf_renderer import render_pdf, update_layout_fields
from app.domain.document_layout import LayoutOptions

pytestmark = pytest.mark.skipif(os.environ.get('PAPERASSIST_LAYOUT_LIVE') != '1',
                                reason='requires explicitly selected isolated local renderer')


def test_actual_office_toc_math_fields_and_pdf_internal_destinations(tmp_path):
    doc = Document()
    doc.add_paragraph('Front matter')
    doc.add_heading('Exact body heading', 1)
    doc.add_paragraph('Scientific body retained.')
    doc.add_heading('Second heading', 1)
    stream = BytesIO()
    doc.save(stream)
    content = stream.getvalue()
    anchors = catalog_docx(content)['anchors']
    options = LayoutOptions.model_validate({
        'toc': {'depth': 2},
        'pagination': {'mode': 'roman_front', 'body_anchor_id': anchors[1]['id']},
        'formulas': [{'id': 'formula1', 'anchor_id': anchors[1]['id'], 'expression': r'\frac{a}{b}+\sqrt{x}+\alpha'}],
        'references': [{'anchor_id': anchors[-1]['id'], 'target_id': anchors[1]['id'], 'kind': 'page', 'prefix': 'BODY_PAGE='},
                       {'anchor_id': anchors[-1]['id'], 'target_id': 'formula:formula1', 'kind': 'number', 'prefix': 'EQUATION='}],
    })
    result = prepare_layout_docx(content, options, language='en')
    (tmp_path / 'layout.docx').write_bytes(result)
    ns = {'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
          'm': 'http://schemas.openxmlformats.org/officeDocument/2006/math'}
    with ZipFile(BytesIO(result)) as package:
        root = E.fromstring(package.read('word/document.xml'))
        assert len(root.xpath('//m:oMath', namespaces=ns)) == 1
        assert len(root.xpath('//w:instrText[contains(text(),"TOC")]', namespaces=ns)) == 1
        meta = json.loads(E.fromstring(package.read('customXml/paperassist-layout.xml')).text)
        assert meta['cache_status'] == 'complete'
    # A second actual Writer import must not duplicate our TOC.
    cache = update_layout_fields(result)
    assert sum('Exact body heading' in line for line in cache['toc_lines']) == 1
    pdf = render_pdf(result).content
    (tmp_path / 'layout.pdf').write_bytes(pdf)
    document = pdfium.PdfDocument(pdf)
    text, destinations = [], []
    for page in document:
        tp = page.get_textpage()
        text.append(tp.get_text_range())
        tp.close()
        position, link = ctypes.c_int(0), raw.FPDF_LINK()
        while raw.FPDFLink_Enumerate(page.raw, ctypes.byref(position), ctypes.byref(link)):
            destination = raw.FPDFLink_GetDest(document.raw, link)
            if destination:
                destinations.append(raw.FPDFDest_GetDestPageIndex(document.raw, destination))
        page.close()
    joined = '\n'.join(text)
    assert 'α' in joined and 'a' in joined and 'b' in joined, 'math must render, not just equation numbers'
    assert 'BODY_PAGE=1' in joined and 'EQUATION=1' in joined
    assert joined.count('Exact body heading') == 2, 'one TOC entry and one original heading'
    assert len(destinations) >= 4 and all(0 <= d < len(document) for d in destinations)
    document.close()


def test_source_numbering_survives_landscape_split_and_later_explicit_restart(tmp_path):
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    doc.add_paragraph('Original portrait source content')
    doc.add_heading('Before landscape', 1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Landscape table'
    doc.add_paragraph('After landscape')
    footer = doc.sections[0].footer.paragraphs[0]
    footer.text = 'Source confidentiality '
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_heading('Intentional source restart', 1)
    for section, number in zip(doc.sections, (5, 11)):
        pg = OxmlElement('w:pgNumType')
        pg.set(qn('w:start'), str(number))
        section._sectPr.append(pg)
    stream = BytesIO()
    doc.save(stream)
    content = stream.getvalue()
    target = next(a['id'] for a in catalog_docx(content)['anchors'] if a['type'] == 'table')
    result = prepare_layout_docx(content, LayoutOptions(landscape_anchor_ids=[target]), language='en')
    pdf = render_pdf(result).content
    (tmp_path / 'source-numbering.docx').write_bytes(result)
    (tmp_path / 'source-numbering.pdf').write_bytes(pdf)
    document = pdfium.PdfDocument(pdf)
    assert len(document) == 4
    for page, expected in zip(document, (5, 6, 7, 11)):
        tp = page.get_textpage()
        text = tp.get_text_range()
        assert f'Source confidentiality {expected}' in text
        tp.close()
        page.close()
    document.close()


def test_continuous_numbering_preserves_inherited_visible_footer(tmp_path):
    from docx.enum.section import WD_SECTION_START
    doc = Document()
    doc.sections[0].footer.paragraphs[0].text = 'Original confidential footer'
    doc.add_paragraph('First source section')
    doc.add_section(WD_SECTION_START.NEW_PAGE)
    doc.add_paragraph('Second source section inherits footer')
    stream = BytesIO()
    doc.save(stream)
    result = prepare_layout_docx(stream.getvalue(), LayoutOptions(pagination={'mode': 'continuous'}), language='en')
    pdf = render_pdf(result).content
    (tmp_path / 'inherited-footer.pdf').write_bytes(pdf)
    document = pdfium.PdfDocument(pdf)
    assert len(document) == 2
    for page in document:
        tp = page.get_textpage()
        assert 'Original confidential footer' in tp.get_text_range()
        tp.close()
        page.close()
    document.close()


@pytest.mark.parametrize('styled', [False, True])
def test_landscape_figure_heading_and_complete_caption_render_together(tmp_path, styled):
    from PIL import Image
    from docx.shared import Inches
    doc = Document()
    doc.add_paragraph('Earlier content')
    doc.add_heading('Figure section', 1)
    image = BytesIO()
    Image.new('RGB', (600, 400), (20, 70, 190)).save(image, format='PNG')
    doc.add_picture(BytesIO(image.getvalue()), width=Inches(6))
    doc.add_paragraph('Figure 1  Original plot', style='Caption' if styled else None)
    doc.add_paragraph('Complete explanatory caption END-CAPTION', style='Caption' if styled else None)
    if styled:
        doc.add_paragraph('A Panel description', style='Caption')
        doc.add_paragraph('B Final panel END-PANEL', style='Caption')
    doc.add_heading('Following source section', 1)
    doc.add_paragraph('Original prose')
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    target = next(a['id'] for a in catalog_docx(source)['anchors'] if a['type'] == 'figure')
    result = prepare_layout_docx(source, LayoutOptions(landscape_anchor_ids=[target]), language='en')
    pdfbytes = render_pdf(result).content
    (tmp_path / f'figure-{styled}.docx').write_bytes(result)
    (tmp_path / f'figure-{styled}.pdf').write_bytes(pdfbytes)
    pdf = pdfium.PdfDocument(pdfbytes)
    image_pages = []
    for page in pdf:
        if list(page.get_objects(filter=[raw.FPDF_PAGEOBJ_IMAGE])):
            tp = page.get_textpage()
            text = tp.get_text_range()
            tp.close()
            assert 'Figure section' in text and 'END-CAPTION' in text
            if styled:
                assert 'END-PANEL' in text
            assert 'Following source section' not in text
            # Presence of an XObject is insufficient: the rendered center must
            # contain the original colored pixels, not white or a clipped box.
            raster = page.render(scale=1).to_pil().convert('RGB')
            count = sum(b > 130 and b > r * 2 and b > g * 2 for r, g, b in raster.getdata())
            assert count > 100000
            raster.save(tmp_path / f'figure-{styled}.png')
            image_pages.append(text)
        page.close()
    pdf.close()
    assert len(image_pages) == 1


def _section_evidence(tmp_path, source, options):
    """Keep the real source, PDF text/raster and identities even on RED."""
    import hashlib
    (tmp_path / 'source.docx').write_bytes(source)
    result = prepare_layout_docx(source, options, language='en')
    (tmp_path / 'layout.docx').write_bytes(result)
    pdfbytes = render_pdf(result).content
    (tmp_path / 'layout.pdf').write_bytes(pdfbytes)
    pdf = pdfium.PdfDocument(pdfbytes)
    pages = []
    for index, page in enumerate(pdf):
        tp = page.get_textpage()
        pages.append({'text': tp.get_text_range(), 'size': page.get_size()})
        tp.close()
        page.render(scale=1.5).to_pil().save(tmp_path / f'page-{index+1}.png')
        page.close()
    pdf.close()
    evidence = {'source_sha256': hashlib.sha256(source).hexdigest(),
                'docx_sha256': hashlib.sha256(result).hexdigest(),
                'pdf_sha256': hashlib.sha256(pdfbytes).hexdigest(), 'pages': pages}
    (tmp_path / 'evidence.json').write_text(json.dumps(evidence, indent=2), encoding='utf-8')
    return pages


def test_roman_body_landscape_start_has_no_empty_pdf_section(tmp_path):
    doc = Document()
    doc.add_heading('Body begins', 1)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Original table'
    doc.add_paragraph('After table')
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    anchors = catalog_docx(source)['anchors']
    options = LayoutOptions(toc={'depth': 1},
                            pagination={'mode': 'roman_front', 'body_anchor_id': anchors[0]['id']},
                            landscape_anchor_ids=[anchors[1]['id']])
    pages = _section_evidence(tmp_path, source, options)
    assert len(pages) == 3
    for page, content, number, landscape in zip(pages, ('Contents', 'Original table', 'After table'),
                                               ('i', '1', '2'), (False, True, False)):
        assert content in page['text']
        assert page['text'].strip().splitlines()[-1] == number
        assert (page['size'][0] > page['size'][1]) == landscape


def test_caption_at_previous_section_end_keeps_pdf_footer_and_restart(tmp_path):
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document()
    doc.sections[0].footer.paragraphs[0].text = 'FIRST FOOTER '
    doc.add_paragraph('Previous source section content')
    caption = doc.add_paragraph('Table 1. Original caption', style='Caption')
    second = doc.add_section(WD_SECTION_START.NEW_PAGE)
    second.footer.is_linked_to_previous = False
    second.footer.paragraphs[0].text = 'SECOND FOOTER '
    separator = caption._p.getnext()
    section = separator.find(qn('w:pPr')).find(qn('w:sectPr'))
    section.getparent().remove(section)
    caption._p.get_or_add_pPr().append(section)
    separator.getparent().remove(separator)
    doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Second section table'
    doc.add_paragraph('Second section after')
    for section, start in zip(doc.sections, (5, 11)):
        number = OxmlElement('w:pgNumType')
        number.set(qn('w:start'), str(start))
        section._sectPr.append(number)
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        section.footer.paragraphs[0]._p.append(field)
    stream = BytesIO()
    doc.save(stream)
    source = stream.getvalue()
    target = next(a['id'] for a in catalog_docx(source)['anchors'] if a['type'] == 'table')
    pages = _section_evidence(tmp_path, source, LayoutOptions(landscape_anchor_ids=[target]))
    assert len(pages) == 3
    for page, footer, number, content in zip(pages, ('FIRST', 'SECOND', 'SECOND'), (5, 11, 12),
                                            ('Previous source section content', 'Second section table', 'Second section after')):
        assert f'{footer} FOOTER {number}' in page['text']
        assert content in page['text']
    assert 'SECOND FOOTER' not in pages[0]['text']
    assert 'Table 1. Original caption' in pages[0]['text']


@pytest.mark.parametrize('before', [True, False])
def test_trailing_caption_boundary_keeps_table_landscape_and_source_footer(tmp_path, before):
    import runpy
    fixture = runpy.run_path(str(Path(__file__).resolve().parents[2] / 'backend/tests/fixtures/layout_section_caption.py'))
    source = fixture['source_docx'](before)
    target = next(a['id'] for a in catalog_docx(source)['anchors'] if a['type'] == 'table')
    pages = _section_evidence(tmp_path, source, LayoutOptions(landscape_anchor_ids=[target]))
    assert len(pages) == (3 if before else 2)
    table_page = pages[1 if before else 0]
    assert table_page['size'][0] > table_page['size'][1]
    assert 'Target wide table' in table_page['text'] and 'Table 1. Original caption' in table_page['text']
    assert f'FIRST FOOTER {6 if before else 5}' in table_page['text']
    assert 'Second source content' in pages[-1]['text'] and 'SECOND FOOTER 11' in pages[-1]['text']
    assert pages[-1]['size'][0] < pages[-1]['size'][1]
    if before:
        assert 'First source content' in pages[0]['text'] and 'FIRST FOOTER 5' in pages[0]['text']
